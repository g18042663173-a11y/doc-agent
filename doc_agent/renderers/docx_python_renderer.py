from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from doc_agent.config import get_settings, load_style_profile
from doc_agent.ir.schemas import WordIR
from doc_agent.renderers.base import BaseRenderer
from doc_agent.utils.file_utils import ensure_parent
from doc_agent.validators.word_validator import WordValidator


class PythonDocxRenderer(BaseRenderer):
    def __init__(self, style_profile: dict[str, Any] | None = None) -> None:
        self.settings = get_settings()
        self.style = style_profile or load_style_profile()
        self.fonts = self.style.get("fonts", {})
        self.rules = self.style.get("docx_rules", {})

    def render(self, ir: dict[str, Any], output_path: str | Path) -> Path:
        word = WordIR.model_validate(ir)
        word, errors, _ = WordValidator(self.style).validate_and_fix(word)
        if errors:
            raise ValueError("; ".join(errors))

        doc = self._new_document()
        self._apply_base_styles(doc)
        title = doc.add_heading(word.title, level=0)
        self._format_paragraph(title, int(self.rules.get("heading_1_size", 18)) + 4, bold=True)
        if word.subtitle:
            subtitle = doc.add_paragraph(word.subtitle)
            self._format_paragraph(subtitle, int(self.rules.get("body_size", 11)), italic=True)

        for block in word.blocks:
            if block.type == "heading":
                paragraph = doc.add_heading(block.text or "未命名章节", level=block.level or 1)
                size_key = f"heading_{block.level or 1}_size"
                self._format_paragraph(paragraph, int(self.rules.get(size_key, 13)), bold=True)
            elif block.type == "paragraph":
                paragraph = doc.add_paragraph(block.text or "")
                self._format_paragraph(paragraph, int(self.rules.get("body_size", 11)))
            elif block.type == "bullet_list":
                for item in block.items:
                    paragraph = doc.add_paragraph(item, style="List Bullet")
                    self._format_paragraph(paragraph, int(self.rules.get("body_size", 11)))
            elif block.type == "table":
                self._add_table(doc, block.table_headers, block.table_rows)

        output = ensure_parent(output_path)
        doc.save(str(output))
        return output

    def _new_document(self):
        reference = self.settings.reference_docx
        if reference.exists() and reference.stat().st_size > 0:
            try:
                return Document(str(reference))
            except Exception:
                pass
        return Document()

    def _apply_base_styles(self, doc) -> None:
        zh_font = self.fonts.get("zh", "Microsoft YaHei")
        en_font = self.fonts.get("en", "Arial")
        normal = doc.styles["Normal"]
        normal.font.name = en_font
        normal.font.size = Pt(int(self.rules.get("body_size", 11)))
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), zh_font)

    def _format_paragraph(self, paragraph, size: int, bold: bool = False, italic: bool = False) -> None:
        zh_font = self.fonts.get("zh", "Microsoft YaHei")
        en_font = self.fonts.get("en", "Arial")
        for run in paragraph.runs:
            run.font.name = en_font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run._element.rPr.rFonts.set(qn("w:eastAsia"), zh_font)

    def _add_table(self, doc, headers: list[str], rows: list[list[str]]) -> None:
        headers = headers or ["项目", "说明"]
        rows = rows or [["暂无数据", "输入内容未提供表格"]]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row[: len(headers)]):
                cells[index].text = value
