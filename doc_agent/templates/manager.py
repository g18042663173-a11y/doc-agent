from __future__ import annotations

import json
import shutil
import uuid
import zipfile
from collections import Counter
from xml.etree import ElementTree
from pathlib import Path
from typing import Any

from doc_agent.config import get_settings
from doc_agent.templates.models import TemplateCategory, TemplateConfig, TemplateLayout, utc_now


class TemplateManager:
    def __init__(self, templates_dir: str | Path | None = None) -> None:
        self.settings = get_settings()
        self.templates_dir = Path(templates_dir) if templates_dir else self.settings.data_dir / "templates"
        self.templates_dir.mkdir(parents=True, exist_ok=True)
        self._load_system_templates()

    def create_template(self, template_config: TemplateConfig) -> TemplateConfig:
        template_file = self._template_path(template_config.template_id)
        if template_file.exists():
            raise ValueError(f"Template {template_config.template_id} already exists")
        self._save_template(template_config)
        return template_config

    def upsert_template(self, template_config: TemplateConfig) -> TemplateConfig:
        template = template_config.model_copy(deep=True)
        file_path = Path(template.file_path)
        local_candidate = self.templates_dir / "custom" / file_path.name
        if not file_path.exists() and local_candidate.exists():
            template.file_path = str(local_candidate)
        self._save_template(template)
        return template

    def get_template(self, template_id: str) -> TemplateConfig | None:
        template_file = self._template_path(template_id)
        if not template_file.exists():
            return None
        return TemplateConfig.model_validate(json.loads(template_file.read_text(encoding="utf-8")))

    def list_templates(self, category: TemplateCategory | str | None = None) -> list[TemplateConfig]:
        category_value = category.value if isinstance(category, TemplateCategory) else category
        templates: list[TemplateConfig] = []
        for template_file in sorted(self.templates_dir.glob("*.json")):
            template = TemplateConfig.model_validate(json.loads(template_file.read_text(encoding="utf-8")))
            if category_value is None or template.category.value == category_value:
                templates.append(template)
        return templates

    def apply_template(self, deck_ir: dict, template_id: str) -> dict:
        template = self.get_template(template_id)
        if template is None:
            raise ValueError(f"Template {template_id} not found")
        enhanced = dict(deck_ir)
        enhanced["template_id"] = template.template_id
        enhanced["template_name"] = template.name
        enhanced["color_scheme"] = template.color_scheme
        enhanced["font_config"] = template.font_config
        return enhanced

    def import_template(self, file_path: str | Path, name: str, category: TemplateCategory | str) -> TemplateConfig:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-pptx is required to import templates") from exc

        source = Path(file_path)
        try:
            presentation = Presentation(str(source))
        except Exception as exc:
            raise ValueError(f"Invalid PPTX template: {exc}") from exc

        category_value = TemplateCategory(category)
        template_id = uuid.uuid4().hex
        target = self.templates_dir / "custom" / f"{template_id}.pptx"
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(source, target)

        layout_config = self._analyze_template_layout(presentation)
        style_analysis = self._analyze_template_style(source, presentation)
        template = TemplateConfig(
            template_id=template_id,
            name=name,
            category=category_value,
            description=f"Custom template: {name}",
            file_path=str(target),
            layout_config=layout_config,
            font_config=style_analysis.get("font_config", {}),
            custom_settings={"analysis": style_analysis},
            is_system=False,
        )
        self._save_template(template)
        return template

    def style_overrides(self, template_id: str | None) -> dict:
        if not template_id:
            return {}
        template = self.get_template(template_id)
        if template is None:
            raise ValueError(f"Template {template_id} not found")
        overrides: dict[str, Any] = {}
        if template.font_config:
            overrides["fonts"] = template.font_config
        analysis = template.custom_settings.get("analysis", {}) if template.custom_settings else {}
        colors = analysis.get("colors") or {}
        if colors:
            overrides["colors"] = colors
        slide_size = analysis.get("slide_size") or {}
        if slide_size:
            overrides["slide_size"] = slide_size
        return overrides

    def _load_system_templates(self) -> None:
        system_layout = TemplateLayout(
            layout_id="business_layout",
            name="Business Layout",
            description="Standard 16:9 business presentation layout",
            slide_masters=["slide_master_0"],
            default_layouts={"cover": "Title Slide", "content": "Title and Content", "section": "Section Header"},
        )
        system_templates = [
            TemplateConfig(
                template_id="business_report",
                name="商务报告模板",
                category=TemplateCategory.BUSINESS,
                description="专业商务风格，适合内部汇报和项目报告",
                file_path=str(self.settings.template_pptx),
                layout_config=system_layout,
                color_scheme="business_blue",
                font_config={"zh": "Microsoft YaHei", "en": "Arial"},
                is_system=True,
            ),
            TemplateConfig(
                template_id="tech_summary",
                name="技术总结模板",
                category=TemplateCategory.TECH,
                description="清晰克制的技术汇报模板",
                file_path=str(self.settings.template_pptx),
                layout_config=system_layout.model_copy(update={"layout_id": "tech_layout", "name": "Tech Layout"}),
                color_scheme="tech_green",
                font_config={"zh": "Microsoft YaHei", "en": "Arial"},
                is_system=True,
            ),
        ]
        for template in system_templates:
            existing = self.get_template(template.template_id)
            if existing is None or existing.is_system:
                self._save_template(template)

    def _analyze_template_layout(self, presentation) -> TemplateLayout:
        slide_masters = [f"slide_master_{index}" for index, _ in enumerate(presentation.slide_masters)]
        return TemplateLayout(
            layout_id=f"layout_{uuid.uuid4().hex[:8]}",
            name="Imported Layout",
            description="Layout analyzed from imported PPTX template",
            slide_masters=slide_masters,
            default_layouts={},
        )

    def _analyze_template_style(self, source: Path, presentation) -> dict[str, Any]:
        width_inches = round(presentation.slide_width / 914400, 3)
        height_inches = round(presentation.slide_height / 914400, 3)
        fonts = self._extract_theme_fonts(source)
        colors = self._extract_theme_colors(source)
        return {
            "mode": "style_apply",
            "note": "Imported PPTX is used for style cues; complex masters, animations, and placeholders are not replicated.",
            "slide_size": {
                "width_inches": width_inches,
                "height_inches": height_inches,
                "ratio": f"{width_inches:g}:{height_inches:g}",
            },
            "master_count": len(presentation.slide_masters),
            "slide_count": len(presentation.slides),
            "font_config": fonts,
            "colors": colors,
        }

    def _extract_theme_fonts(self, source: Path) -> dict[str, str]:
        try:
            with zipfile.ZipFile(source) as archive:
                theme_xml = archive.read("ppt/theme/theme1.xml")
        except Exception:
            return {}
        root = ElementTree.fromstring(theme_xml)
        namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        major_latin = root.find(".//a:majorFont/a:latin", namespaces)
        minor_latin = root.find(".//a:minorFont/a:latin", namespaces)
        major_ea = root.find(".//a:majorFont/a:ea", namespaces)
        minor_ea = root.find(".//a:minorFont/a:ea", namespaces)
        zh = major_ea if major_ea is not None else minor_ea
        en = major_latin if major_latin is not None else minor_latin
        fonts: dict[str, str] = {}
        if zh is not None and zh.attrib.get("typeface"):
            fonts["zh"] = zh.attrib["typeface"]
        if en is not None and en.attrib.get("typeface"):
            fonts["en"] = en.attrib["typeface"]
        return fonts

    def _extract_theme_colors(self, source: Path) -> dict[str, str]:
        values: list[str] = []
        try:
            with zipfile.ZipFile(source) as archive:
                theme_xml = archive.read("ppt/theme/theme1.xml")
        except Exception:
            return {}
        root = ElementTree.fromstring(theme_xml)
        namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for node in root.iter():
            if node.tag in {f"{namespace}srgbClr", f"{namespace}sysClr"}:
                value = node.attrib.get("val") or node.attrib.get("lastClr")
                if value and len(value) == 6:
                    values.append(value.upper())
        common = [color for color, _ in Counter(values).most_common()]
        if not common:
            return {}
        colors: dict[str, str] = {}
        for key, index in [("primary", 0), ("secondary", 1), ("accent", 2), ("title", 3), ("body", 4)]:
            if index < len(common):
                colors[key] = common[index]
        if "background" not in colors:
            colors["background"] = "FFFFFF"
        if "white" not in colors:
            colors["white"] = "FFFFFF"
        return colors

    def _save_template(self, template: TemplateConfig) -> None:
        template.updated_at = utc_now()
        self._template_path(template.template_id).write_text(template.model_dump_json(indent=2), encoding="utf-8")

    def _template_path(self, template_id: str) -> Path:
        return self.templates_dir / f"{Path(template_id).name}.json"
