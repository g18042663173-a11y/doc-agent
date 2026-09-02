from __future__ import annotations

import zipfile
from collections import Counter
from pathlib import Path
import re
from typing import Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pydantic import BaseModel, Field


Severity = Literal["Error", "Warning", "Info"]


class ComplianceItem(BaseModel):
    severity: Severity
    code: str
    message: str
    slide_index: int | None = None


class ComplianceReport(BaseModel):
    score: int = Field(ge=0, le=100)
    summary: str
    items: list[ComplianceItem] = Field(default_factory=list)

    @property
    def error_count(self) -> int:
        return sum(1 for item in self.items if item.severity == "Error")

    @property
    def warning_count(self) -> int:
        return sum(1 for item in self.items if item.severity == "Warning")

    @property
    def info_count(self) -> int:
        return sum(1 for item in self.items if item.severity == "Info")


class HuaweiPptxComplianceChecker:
    cn_fonts = {"microsoft yahei", "微软雅黑"}
    en_fonts = {"arial"}
    required_footer = "HUAWEI CONFIDENTIAL"
    cjk_pattern = re.compile(r"[\u4e00-\u9fff]")
    ascii_letter_pattern = re.compile(r"[A-Za-z]")

    def check(self, pptx_path: str | Path) -> ComplianceReport:
        path = Path(pptx_path)
        items: list[ComplianceItem] = []
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            return self._report([ComplianceItem(severity="Error", code="pptx.invalid", message=f"PPTX cannot be opened: {exc}")])

        if self._has_animation_or_transition(path):
            items.append(
                ComplianceItem(
                    severity="Error",
                    code="pptx.animation",
                    message="Animations or slide transitions were found; Huawei static deck output should not use them.",
                )
            )

        for slide_index, slide in enumerate(presentation.slides, start=1):
            text = "\n".join(self._shape_text(shape) for shape in slide.shapes)
            if self.required_footer not in text:
                items.append(
                    ComplianceItem(
                        severity="Error",
                        code="footer.missing",
                        message=f"Footer must contain {self.required_footer}.",
                        slide_index=slide_index,
                    )
                )

            sizes = self._slide_font_sizes(slide, slide_index, items)
            if len(sizes) > 3:
                items.append(
                    ComplianceItem(
                        severity="Warning",
                        code="font.size.too_many",
                        message=f"Slide uses {len(sizes)} font sizes; keep no more than 3.",
                        slide_index=slide_index,
                    )
                )

            self._check_shape_lines(slide, slide_index, items)

        return self._report(items)

    def _slide_font_sizes(self, slide, slide_index: int, items: list[ComplianceItem]) -> set[int]:
        sizes: set[int] = set()
        seen_font_issues: set[tuple[str, str]] = set()
        seen_color_issues: set[str] = set()
        for text_frame in self._iter_text_frames(slide.shapes):
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    text = run.text.strip()
                    if not text:
                        continue
                    name = run.font.name or paragraph.font.name
                    normalized_name = name.strip().lower() if name else ""
                    if self.cjk_pattern.search(text) and normalized_name and normalized_name not in self.cn_fonts:
                        issue_key = ("cn", name)
                        if issue_key not in seen_font_issues:
                            seen_font_issues.add(issue_key)
                            items.append(
                                ComplianceItem(
                                    severity="Error",
                                    code="font.cn_family",
                                    message=f"Chinese text should use Microsoft YaHei, found {name}.",
                                    slide_index=slide_index,
                                )
                            )
                    elif self._is_english_text(text) and normalized_name and normalized_name not in self.en_fonts:
                        if not (normalized_name == "impact" and self._is_large_number(text)):
                            issue_key = ("en", name)
                            if issue_key not in seen_font_issues:
                                seen_font_issues.add(issue_key)
                                items.append(
                                    ComplianceItem(
                                        severity="Warning",
                                        code="font.en_family",
                                        message=f"English text should use Arial, found {name}.",
                                        slide_index=slide_index,
                                    )
                                )
                    if run.font.size:
                        size = round(run.font.size.pt)
                        sizes.add(size)
                        if size < 6:
                            items.append(
                                ComplianceItem(
                                    severity="Error",
                                    code="font.size.too_small",
                                    message=f"Font size {size} pt is below the 6 pt minimum.",
                                    slide_index=slide_index,
                                )
                            )
                    color = self._font_rgb(run.font.color)
                    if color and not self._is_allowed_color(color) and color not in seen_color_issues:
                        seen_color_issues.add(color)
                        items.append(
                            ComplianceItem(
                                severity="Info",
                                code="color.unexpected",
                                message=f"Text color #{color} is outside the black/red/white/gray palette.",
                                slide_index=slide_index,
                            )
                        )
        return sizes

    def _check_shape_lines(self, slide, slide_index: int, items: list[ComplianceItem]) -> None:
        widths: Counter[float] = Counter()
        for shape in self._iter_shapes(slide.shapes):
            for width in self._shape_line_widths(shape):
                widths[width] += 1
        for width, count in widths.items():
            if abs(width - 0.5) > 0.05:
                items.append(
                    ComplianceItem(
                        severity="Info",
                        code="line.width",
                        message=f"{count} shape line(s) use {width} pt; Huawei reference line width is 0.5 pt.",
                        slide_index=slide_index,
                    )
                )

    @staticmethod
    def _shape_text(shape) -> str:
        if not getattr(shape, "has_text_frame", False):
            return ""
        return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)

    @classmethod
    def _iter_shapes(cls, shapes):
        for shape in shapes:
            yield shape
            nested = getattr(shape, "shapes", None)
            if nested is not None:
                yield from cls._iter_shapes(nested)

    @classmethod
    def _iter_text_frames(cls, shapes):
        for shape in cls._iter_shapes(shapes):
            if getattr(shape, "has_text_frame", False):
                yield shape.text_frame
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

    @staticmethod
    def _shape_line_widths(shape) -> list[float]:
        widths: list[float] = []
        element = getattr(shape, "_element", None)
        if element is not None:
            for node in element.iter():
                if not str(node.tag).endswith("}ln"):
                    continue
                raw = node.get("w")
                if raw and raw.isdigit():
                    widths.append(round(int(raw) / 12700, 2))
        if widths:
            return [width for width in widths if width > 0]
        line = getattr(shape, "line", None)
        if line is not None and line.width is not None:
            width = round(float(line.width.pt), 2)
            if width > 0:
                return [width]
        return []

    @staticmethod
    def _font_rgb(color) -> str | None:
        try:
            rgb = color.rgb
        except Exception:
            return None
        if not isinstance(rgb, RGBColor):
            return None
        return str(rgb).upper()

    @classmethod
    def _is_english_text(cls, text: str) -> bool:
        return bool(cls.ascii_letter_pattern.search(text)) and not cls.cjk_pattern.search(text)

    @staticmethod
    def _is_large_number(text: str) -> bool:
        compact = text.replace(",", "").replace(".", "")
        return compact.isdigit() and len(compact) >= 3

    @staticmethod
    def _is_allowed_color(color: str) -> bool:
        try:
            red = int(color[0:2], 16)
            green = int(color[2:4], 16)
            blue = int(color[4:6], 16)
        except ValueError:
            return False
        is_grayscale = red == green == blue
        is_huawei_red = red >= 150 and green <= 60 and blue <= 60
        return is_grayscale or is_huawei_red

    @staticmethod
    def _has_animation_or_transition(path: Path) -> bool:
        markers = (b"<p:timing", b"<p:anim", b"<p:animEffect", b"<p:transition")
        try:
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                        continue
                    payload = archive.read(name)
                    if any(marker in payload for marker in markers):
                        return True
        except zipfile.BadZipFile:
            return True
        return False

    @staticmethod
    def _report(items: list[ComplianceItem]) -> ComplianceReport:
        errors = sum(1 for item in items if item.severity == "Error")
        warnings = sum(1 for item in items if item.severity == "Warning")
        info = sum(1 for item in items if item.severity == "Info")
        score = max(0, 100 - errors * 20 - warnings * 6 - info * 2)
        if items:
            summary = f"{errors} Error, {warnings} Warning, {info} Info"
        else:
            summary = "No Huawei compliance issues found"
        return ComplianceReport(score=score, summary=summary, items=items)
