from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from doc_agent.compliance.pptx import HuaweiPptxComplianceChecker
from doc_agent.config import Settings, get_settings
from doc_agent.exporters import preview
from doc_agent.llm import NGAClient, get_llm_client
from doc_agent.workflow import run_generate


def test_provider_aliases_keep_old_env_values_runnable(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    monkeypatch.setenv("PPT_RENDERER", "python_pptx")

    settings = get_settings()

    assert settings.llm_provider == "stub"
    assert settings.ppt_renderer == "stub"


def test_nga_client_is_explicit_internal_placeholder() -> None:
    client = get_llm_client(Settings(llm_provider="nga"))

    assert isinstance(client, NGAClient)
    with pytest.raises(RuntimeError, match="internal-network adapter placeholder"):
        client.generate_json("return {}")


def test_hw_skill_renderer_is_explicit_internal_placeholder(tmp_path: Path) -> None:
    source = tmp_path / "input.md"
    source.write_text("# Demo\n\n- one\n", encoding="utf-8")

    with pytest.raises(RuntimeError, match="HuaweiSkillRenderer"):
        run_generate(
            source,
            "pptx",
            tmp_path / "out.pptx",
            settings=Settings(llm_provider="stub", ppt_renderer="hw_skill"),
        )


def test_presenton_renderer_is_no_longer_supported(tmp_path: Path) -> None:
    source = tmp_path / "input.md"
    source.write_text("# Demo\n\n- one\n", encoding="utf-8")

    with pytest.raises(RuntimeError, match="presenton is no longer supported"):
        run_generate(
            source,
            "pptx",
            tmp_path / "out.pptx",
            settings=Settings(llm_provider="stub", ppt_renderer="presenton"),
        )


def test_huawei_compliance_checker_reports_footer_and_size_issues(tmp_path: Path) -> None:
    pptx_path = tmp_path / "bad.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text(slide, "中文标题", 5, "SimSun", RGBColor(0, 0, 255))
    _add_text(slide, "English Title", 18, "Calibri", RGBColor(0, 0, 0), top=1.5)
    prs.save(str(pptx_path))

    report = HuaweiPptxComplianceChecker().check(pptx_path)

    assert report.score < 100
    assert any(item.code == "footer.missing" and item.severity == "Error" for item in report.items)
    assert any(item.code == "font.cn_family" and item.severity == "Error" for item in report.items)
    assert any(item.code == "font.en_family" and item.severity == "Warning" for item in report.items)
    assert any(item.code == "font.size.too_small" and item.severity == "Error" for item in report.items)
    assert any(item.code == "color.unexpected" and item.severity == "Info" for item in report.items)


def test_huawei_compliance_checker_accepts_minimal_static_deck(tmp_path: Path) -> None:
    pptx_path = tmp_path / "good.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text(slide, "Quarterly Review", 14, "Arial", RGBColor(0, 0, 0))
    _add_text(slide, "HUAWEI CONFIDENTIAL", 9, "Arial", RGBColor(128, 128, 128), top=6.8)
    prs.save(str(pptx_path))

    report = HuaweiPptxComplianceChecker().check(pptx_path)

    assert report.error_count == 0
    assert report.summary == "No Huawei compliance issues found"


def test_preview_export_reports_missing_tools(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(preview.shutil, "which", lambda _binary: None)

    with pytest.raises(RuntimeError, match="LibreOffice"):
        preview.export_pages_to_images(tmp_path / "missing.pptx", tmp_path / "images")


def test_preview_export_uses_libreoffice_and_pdftoppm(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "input.pptx"
    source.write_bytes(b"pptx")
    output_dir = tmp_path / "images"
    calls: list[list[str]] = []

    monkeypatch.setattr(preview.shutil, "which", lambda binary: f"/usr/bin/{binary}")

    def fake_run(command, check, capture_output, text):
        calls.append(command)
        if command[1] == "--headless":
            outdir = Path(command[command.index("--outdir") + 1])
            outdir.mkdir(parents=True, exist_ok=True)
            (outdir / "input.pdf").write_bytes(b"pdf")
        elif command[1] == "-png":
            Path(f"{command[-1]}-1.png").write_bytes(b"png")

    monkeypatch.setattr(preview.subprocess, "run", fake_run)

    images = preview.export_pages_to_images(source, output_dir)

    assert len(images) == 1
    assert images[0].name == "input-1.png"
    assert calls[0][0].endswith("soffice")
    assert calls[1][0].endswith("pdftoppm")


def _add_text(slide, text: str, size: int, font: str, color: RGBColor, top: float = 1.0) -> None:
    shape = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(10), Inches(0.4))
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = font
    run.font.color.rgb = color
