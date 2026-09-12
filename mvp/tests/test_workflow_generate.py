from pathlib import Path

from docx import Document
from pptx import Presentation

from doc_agent.workflow import run_generate


def test_run_generate_creates_pptx_and_docx_from_markdown(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    source = tmp_path / "input.md"
    source.write_text(
        "# 企业文档生成 Agent\n\n这是摘要。\n\n- 解析输入\n- 渲染输出\n",
        encoding="utf-8",
    )

    pptx_path = run_generate(source, "pptx", tmp_path / "demo.pptx", target_slide_count=6)
    docx_path = run_generate(source, "docx", tmp_path / "demo.docx", target_slide_count=6)

    assert len(Presentation(str(pptx_path)).slides) >= 5
    assert len(Document(str(docx_path)).paragraphs) > 0
