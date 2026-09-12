from __future__ import annotations

import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api import app
from doc_agent.parsers.router import ParserRouter
from doc_agent.smartart import SmartArtConfig, SmartArtManager, SmartArtNode, SmartArtType
from doc_agent.validators.output_validator import OutputValidator
from doc_agent.workflow import run_generate


def test_legacy_generate_download_and_error_paths(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "outputs"))
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    client = TestClient(app)

    health_start = time.perf_counter()
    health = client.get("/health")
    assert health.status_code == 200
    assert time.perf_counter() - health_start < 2

    missing_download = client.get("/download/missing.pptx")
    assert missing_download.status_code == 404
    unsafe_download = client.get("/download/%2E%2E%2Fsecret.pptx")
    assert unsafe_download.status_code in {400, 404}

    missing_file = client.post("/generate", data={"target": "pptx"})
    assert missing_file.status_code == 400

    invalid_target = client.post(
        "/generate",
        files={"file": ("input.md", b"# Demo", "text/markdown")},
        data={"target": "pdf"},
    )
    assert invalid_target.status_code == 400
    invalid_slides = client.post(
        "/generate",
        files={"file": ("input.md", b"# Demo", "text/markdown")},
        data={"target": "pptx", "slides": "not-a-number"},
    )
    assert invalid_slides.status_code == 400
    zero_slides = client.post(
        "/generate",
        files={"file": ("input.md", b"# Demo", "text/markdown")},
        data={"target": "pptx", "slides": "0"},
    )
    assert zero_slides.status_code == 400

    invalid_file = client.post(
        "/generate",
        files={"file": ("input.txt", b"Demo", "text/plain")},
        data={"target": "pptx"},
    )
    assert invalid_file.status_code == 400
    invalid_api_slides = client.post(
        "/api/generate/start",
        files={"file": ("input.md", b"# Demo", "text/markdown")},
        data={"target": "pptx", "slides": "-1"},
    )
    assert invalid_api_slides.status_code == 400

    response = client.post(
        "/generate",
        files={"file": ("input.md", b"# Demo\n\n- One\n- Two\n", "text/markdown")},
        data={"target": "docx", "slides": "5"},
    )
    assert response.status_code == 200
    download = client.get(response.json()["download_url"])
    assert download.status_code == 200
    assert download.content


def test_design_api_negative_paths_and_updates(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "outputs"))
    client = TestClient(app)

    assert client.get("/api/templates/missing").status_code == 404
    assert client.post(
        "/api/templates/import",
        files={"file": ("template.txt", b"not pptx", "text/plain")},
        data={"name": "Bad", "category": "business"},
    ).status_code == 400

    assert client.get("/api/colors/missing").status_code == 404
    assert client.post(
        "/api/colors/create",
        json={
            "name": "Bad",
            "category": "business",
            "primary_color": "123456",
            "secondary_color": "#234567",
            "accent_color": "#345678",
            "background_color": "#FFFFFF",
            "text_color": "#111111",
        },
    ).status_code == 400

    bad_connection = client.post(
        "/api/users/test-connection",
        json={"nga_endpoint": "not-a-url", "auth_token": "token"},
    )
    assert bad_connection.status_code == 200
    assert bad_connection.json()["success"] is False

    missing_update = client.put("/api/users/missing", json={"user_name": "none"})
    assert missing_update.status_code == 404
    missing_switch = client.post("/api/users/switch/missing")
    assert missing_switch.status_code == 404

    created = client.post(
        "/api/users/create",
        json={"user_name": "Quality User", "nga_endpoint": "http://example.test/v1", "auth_token": "secret"},
    )
    assert created.status_code == 200
    user_id = created.json()["user_id"]

    updated = client.put(user_path := f"/api/users/{user_id}", json={"user_name": "Renamed", "preferences": {"theme": "dark"}})
    assert updated.status_code == 200
    assert updated.json()["user_name"] == "Renamed"
    assert updated.json()["auth_token"] == "***HIDDEN***"
    assert client.get(user_path).json()["preferences"]["theme"] == "dark"
    assert client.delete(user_path).json()["success"] is True
    assert client.delete(user_path).json()["success"] is False


def test_docx_and_pptx_parsers_are_routed(tmp_path: Path) -> None:
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_heading("Doc Title", level=1)
    doc.add_paragraph("First paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    doc.save(docx_path)

    pptx_path = tmp_path / "input.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Deck Title"
    slide.placeholders[1].text = "Subtitle"
    presentation.save(pptx_path)

    router = ParserRouter()
    docx_ir = router.parse(docx_path)
    pptx_ir = router.parse(pptx_path)

    assert docx_ir.source_type == "docx"
    assert docx_ir.title == "Doc Title"
    assert any(block.type == "table" for block in docx_ir.blocks)
    assert pptx_ir.source_type == "pptx"
    assert pptx_ir.title == "Deck Title"
    assert pptx_ir.meta["slide_count"] == 1


def test_output_validator_and_smartart_edge_layouts(tmp_path: Path) -> None:
    validator = OutputValidator()
    assert validator.validate(tmp_path / "missing.pptx", "pptx")

    wrong_suffix = tmp_path / "wrong.txt"
    wrong_suffix.write_text("not a pptx", encoding="utf-8")
    errors = validator.validate(wrong_suffix, "pptx")
    assert any("suffix" in error for error in errors)
    assert any("cannot be opened" in error for error in errors)
    assert validator.validate(wrong_suffix, "pdf")[-1] == "Unsupported target: pdf"

    manager = SmartArtManager()
    nodes = [SmartArtNode(id="n1", text="Root")]
    layouts = {
        SmartArtType.HIERARCHY: "tree",
        SmartArtType.ORG_CHART: "tree",
        SmartArtType.MATRIX: "grid",
        SmartArtType.PYRAMID: "pyramid",
        SmartArtType.RELATIONSHIP: "radial",
    }
    for smartart_type, layout in layouts.items():
        result = manager.generate_smartart(nodes, SmartArtConfig(smartart_type=smartart_type))
        assert result["layout"] == layout
    assert set(SmartArtManager.supported_types()) >= {item.value for item in layouts}


def test_generation_performance_and_three_concurrent_tasks(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "outputs"))
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    monkeypatch.setenv("PPT_RENDERER", "python_pptx")
    source = tmp_path / "input.md"
    source.write_text(
        "# Performance Demo\n\n"
        "This input verifies the quality gate.\n\n"
        "- Parse content\n- Plan slides\n- Render editable output\n",
        encoding="utf-8",
    )

    started = time.perf_counter()
    ten_slide_output = run_generate(source, "pptx", tmp_path / "ten-slides.pptx", target_slide_count=10)
    assert time.perf_counter() - started < 30
    assert len(Presentation(str(ten_slide_output)).slides) >= 5

    def generate(index: int) -> Path:
        return run_generate(source, "pptx", tmp_path / f"concurrent-{index}.pptx", target_slide_count=5)

    concurrent_started = time.perf_counter()
    with ThreadPoolExecutor(max_workers=3) as executor:
        outputs = list(executor.map(generate, range(3)))
    assert time.perf_counter() - concurrent_started < 30
    assert all(path.exists() and path.stat().st_size > 0 for path in outputs)
