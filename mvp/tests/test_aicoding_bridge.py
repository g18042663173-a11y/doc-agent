from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

openpyxl = pytest.importorskip("openpyxl")

from app.api import app
from doc_agent.parsers.router import ParserRouter


def _sample_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "收入"
    sheet.append(["月份", "收入", "成本"])
    sheet.append(["一月", 120, 80])
    sheet.append(["二月", 150, 95])
    sheet["D1"] = "利润"
    sheet["D2"] = "=B2-C2"
    sheet.merge_cells("A5:B5")
    sheet["A5"] = "备注"
    workbook.save(path)


def test_xlsx_parser_builds_document_ir(tmp_path: Path) -> None:
    source = tmp_path / "finance.xlsx"
    _sample_xlsx(source)

    document_ir = ParserRouter().parse(source)

    assert document_ir.source_type == "xlsx"
    assert document_ir.title == "finance"
    assert document_ir.meta["sheet_count"] == 1
    assert any(block.type == "table" and block.meta["sheet_name"] == "收入" for block in document_ir.blocks)
    assert "数值列统计" in "\n".join(block.text or "" for block in document_ir.blocks)


def test_aicoding_prompt_and_manual_word_render(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "outputs"))
    source = tmp_path / "finance.xlsx"
    _sample_xlsx(source)
    client = TestClient(app)

    prompt = client.post(
        "/api/aicoding/prompt",
        files={"file": ("finance.xlsx", source.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"target": "docx", "slides": "6"},
    )

    assert prompt.status_code == 200
    prompt_body = prompt.json()
    assert prompt_body["target"] == "docx"
    assert prompt_body["document_ir"]["source_type"] == "xlsx"
    assert "WordIR" in prompt_body["prompt"]
    assert "只输出一个 JSON 对象" in prompt_body["prompt"]

    rendered = client.post(
        "/api/aicoding/render",
        json={
            "target": "docx",
            "original_filename": "finance.xlsx",
            "model_output": {
                "title": "收入分析",
                "subtitle": "AICoding 手动桥接",
                "blocks": [
                    {"type": "heading", "text": "一、总体情况", "level": 1},
                    {"type": "paragraph", "text": "收入整体增长，成本保持稳定。"},
                    {"type": "table", "table_headers": ["月份", "收入"], "table_rows": [["一月", "120"], ["二月", "150"]]},
                ],
            },
        },
    )

    assert rendered.status_code == 200
    task_id = rendered.json()["task_id"]
    history = client.get(f"/api/generate/history/{task_id}")
    assert history.status_code == 200
    assert history.json()["status"] == "completed"
    assert history.json()["compliance_report"]["score"] > 0

    download = client.get(f"/api/generate/download/{task_id}")
    assert download.status_code == 200
    output = tmp_path / "download.docx"
    output.write_bytes(download.content)
    text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert "收入分析" in text


def test_aicoding_render_rejects_invalid_json(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "outputs"))
    client = TestClient(app)

    response = client.post("/api/aicoding/render", json={"target": "docx", "model_output": "not json"})

    assert response.status_code == 400
    assert "JSON" in response.json()["detail"]
