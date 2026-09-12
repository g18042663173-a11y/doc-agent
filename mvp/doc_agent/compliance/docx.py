from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx import Document

from doc_agent.compliance.pptx import ComplianceItem, ComplianceReport


class WordComplianceChecker:
    allowed_fonts = {"arial", "microsoft yahei", "微软雅黑"}

    def check(self, docx_path: str | Path) -> ComplianceReport:
        items: list[ComplianceItem] = []
        try:
            document = Document(str(docx_path))
        except Exception as exc:
            return self._report([ComplianceItem(severity="Error", code="docx.invalid", message=f"DOCX cannot be opened: {exc}")])

        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        if not paragraphs:
            items.append(ComplianceItem(severity="Error", code="docx.empty", message="Word document contains no readable text."))

        headings = [paragraph for paragraph in paragraphs if (paragraph.style.name if paragraph.style else "").startswith("Heading")]
        if not headings:
            items.append(ComplianceItem(severity="Warning", code="heading.missing", message="Word document should contain heading styles for scanability."))
        elif self._heading_level(headings[0]) > 1:
            items.append(ComplianceItem(severity="Warning", code="heading.first_level", message="The first heading should normally be level 1."))

        empty_paragraphs = sum(1 for paragraph in document.paragraphs if not paragraph.text.strip())
        if paragraphs and empty_paragraphs > len(paragraphs) * 2:
            items.append(ComplianceItem(severity="Info", code="paragraph.empty_many", message="Document contains many empty paragraphs."))

        font_names: Counter[str] = Counter()
        font_sizes: set[int] = set()
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                if run.font.name:
                    font_names[run.font.name.strip().lower()] += 1
                if run.font.size:
                    font_sizes.add(round(run.font.size.pt))
        for font_name, count in font_names.items():
            if font_name not in self.allowed_fonts:
                items.append(ComplianceItem(severity="Info", code="font.unexpected", message=f"{count} run(s) use unexpected font {font_name}."))
        if len(font_sizes) > 5:
            items.append(ComplianceItem(severity="Warning", code="font.size.too_many", message=f"Document uses {len(font_sizes)} font sizes; keep Word styles concise."))

        for table_index, table in enumerate(document.tables, start=1):
            if not table.rows or not table.columns:
                items.append(ComplianceItem(severity="Warning", code="table.empty", message=f"Table {table_index} is empty."))

        return self._report(items)

    @staticmethod
    def _heading_level(paragraph) -> int:
        style_name = paragraph.style.name if paragraph.style else ""
        for token in reversed(style_name.split()):
            if token.isdigit():
                return int(token)
        return 1

    @staticmethod
    def _report(items: list[ComplianceItem]) -> ComplianceReport:
        errors = sum(1 for item in items if item.severity == "Error")
        warnings = sum(1 for item in items if item.severity == "Warning")
        info = sum(1 for item in items if item.severity == "Info")
        score = max(0, 100 - errors * 20 - warnings * 6 - info * 2)
        summary = f"{errors} Error, {warnings} Warning, {info} Info" if items else "No Word format issues found"
        return ComplianceReport(score=score, summary=summary, items=items)
