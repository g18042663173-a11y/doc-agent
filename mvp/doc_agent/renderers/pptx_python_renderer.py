from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from doc_agent.config import get_settings, load_style_profile
from doc_agent.ir.schemas import DeckIR, SlideIR
from doc_agent.renderers.base import BaseRenderer
from doc_agent.utils.file_utils import ensure_parent
from doc_agent.validators.deck_validator import DeckValidator


class PythonPptxRenderer(BaseRenderer):
    def __init__(self, style_profile: dict[str, Any] | None = None) -> None:
        self.settings = get_settings()
        self.style = style_profile or load_style_profile()
        self.colors = self.style.get("colors", {})
        self.fonts = self.style.get("fonts", {})

    def render(self, ir: dict[str, Any], output_path: str | Path) -> Path:
        deck = DeckIR.model_validate(ir)
        deck, errors, _ = DeckValidator(self.style).validate_and_fix(deck)
        if errors:
            raise ValueError("; ".join(errors))

        prs = Presentation()
        slide_size = self.style.get("slide_size", {})
        prs.slide_width = Inches(float(slide_size.get("width_inches", 13.333)))
        prs.slide_height = Inches(float(slide_size.get("height_inches", 7.5)))

        for index, slide_ir in enumerate(deck.slides, start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._set_background(slide)
            self._render_slide(slide, slide_ir, index, len(deck.slides), deck)

        output = ensure_parent(output_path)
        prs.save(str(output))
        return output

    def _render_slide(self, slide, slide_ir: SlideIR, index: int, total: int, deck: DeckIR) -> None:
        layout = slide_ir.layout
        if layout == "cover":
            self._render_cover(slide, slide_ir)
        elif layout == "agenda":
            self._render_agenda(slide, slide_ir)
        elif layout == "section":
            self._render_section(slide, slide_ir, index)
        elif layout == "two_column":
            self._render_two_column(slide, slide_ir)
        elif layout == "table":
            self._render_table(slide, slide_ir)
        elif layout == "cards":
            self._render_cards(slide, slide_ir)
        elif layout == "chart":
            self._render_chart(slide, slide_ir)
        elif layout == "image":
            self._render_image(slide, slide_ir)
        elif layout == "conclusion":
            self._render_conclusion(slide, slide_ir)
        else:
            self._render_title_bullets(slide, slide_ir)
        self._footer(slide, deck, slide_ir, index, total)

    def _render_cover(self, slide, slide_ir: SlideIR) -> None:
        self._accent_bar(slide)
        self._text_box(slide, Inches(0.9), Inches(1.95), Inches(11.3), Inches(1.0), slide_ir.title, 34, "title", True)
        if slide_ir.subtitle:
            self._text_box(slide, Inches(0.95), Inches(3.05), Inches(9.8), Inches(0.5), slide_ir.subtitle, 18, "muted")
        self._text_box(slide, Inches(0.95), Inches(5.85), Inches(5.8), Inches(0.4), "结构化生成 · 本地渲染 · 可编辑输出", 13, "muted")

    def _render_agenda(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        items = slide_ir.bullets[:6] or ["背景", "方案", "实施", "结论"]
        for idx, item in enumerate(items, start=1):
            top = Inches(1.45 + (idx - 1) * 0.68)
            self._number_badge(slide, idx, Inches(1.05), top)
            self._text_box(slide, Inches(1.75), top, Inches(9.5), Inches(0.42), item, 18, "body", False)

    def _render_section(self, slide, slide_ir: SlideIR, index: int) -> None:
        self._accent_bar(slide)
        self._text_box(slide, Inches(0.95), Inches(2.35), Inches(10.9), Inches(0.8), f"{index:02d}", 30, "secondary", True)
        self._text_box(slide, Inches(0.95), Inches(3.05), Inches(10.9), Inches(0.75), slide_ir.title, 30, "title", True)

    def _render_title_bullets(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        self._bullets(slide, slide_ir.bullets, Inches(1.15), Inches(1.55), Inches(10.8), Inches(4.7), 20)

    def _render_two_column(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        self._column(slide, Inches(0.9), Inches(1.45), Inches(5.55), slide_ir.left_title or "左侧", slide_ir.left_bullets)
        self._column(slide, Inches(6.9), Inches(1.45), Inches(5.55), slide_ir.right_title or "右侧", slide_ir.right_bullets)

    def _render_table(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        headers = slide_ir.table_headers[:5] or ["项目", "说明"]
        rows = slide_ir.table_rows[:6] or [["暂无数据", "输入内容未提供表格"]]
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.9), Inches(1.45), Inches(11.55), Inches(4.8))
        table = table_shape.table
        for col, header in enumerate(headers):
            cell = table.cell(0, col)
            cell.text = header
            self._format_cell(cell, bold=True, fill="primary", color="white")
        for row_index, row in enumerate(rows, start=1):
            for col_index, value in enumerate(row[: len(headers)]):
                cell = table.cell(row_index, col_index)
                cell.text = value
                self._format_cell(cell, bold=False, fill="background", color="body")

    def _render_cards(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        cards = slide_ir.cards[:4]
        if not cards:
            cards = []
        count = max(len(cards), 1)
        card_width = 10.9 / count
        for idx, card in enumerate(cards or []):
            left = Inches(0.95 + idx * card_width)
            shape = slide.shapes.add_shape(1, left, Inches(1.55), Inches(card_width - 0.25), Inches(3.75))
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb("white")
            shape.line.color.rgb = self._rgb("secondary")
            text_frame = shape.text_frame
            text_frame.clear()
            title = text_frame.paragraphs[0]
            title.text = card.title
            title.font.name = self.fonts.get("zh", "Microsoft YaHei")
            title.font.size = Pt(16)
            title.font.bold = True
            title.font.color.rgb = self._rgb("primary")
            if card.body:
                body = text_frame.add_paragraph()
                body.text = card.body
                body.font.name = self.fonts.get("zh", "Microsoft YaHei")
                body.font.size = Pt(12)
                body.font.color.rgb = self._rgb("body")
            for item in card.bullets[:3]:
                paragraph = text_frame.add_paragraph()
                paragraph.text = f"- {item}"
                paragraph.font.name = self.fonts.get("zh", "Microsoft YaHei")
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = self._rgb("body")
        if not cards:
            self._bullets(slide, slide_ir.bullets, Inches(1.15), Inches(1.65), Inches(10.8), Inches(3.8), 18)

    def _render_chart(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        chart = slide_ir.chart
        labels = chart.labels[:6] if chart else []
        series = chart.series[0].values[: len(labels)] if chart and chart.series else []
        if not labels:
            labels = ["项1", "项2", "项3"]
            series = [1.0, 2.0, 3.0]
        max_value = max(series or [1.0])
        for idx, label in enumerate(labels):
            top = Inches(1.55 + idx * 0.55)
            self._text_box(slide, Inches(1.0), top, Inches(1.8), Inches(0.28), label, 11, "body")
            width = 0.1 if max_value <= 0 else max(0.25, 7.4 * float(series[idx] if idx < len(series) else 0) / max_value)
            bar = slide.shapes.add_shape(1, Inches(2.9), top, Inches(width), Inches(0.28))
            bar.fill.solid()
            bar.fill.fore_color.rgb = self._rgb("primary")
            bar.line.fill.background()
        if chart and chart.summary:
            self._text_box(slide, Inches(1.0), Inches(5.45), Inches(10.8), Inches(0.45), chart.summary, 13, "muted")

    def _render_image(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        visual = slide_ir.visuals[0] if slide_ir.visuals else None
        image_path = Path(visual.path) if visual and visual.path else None
        if image_path and image_path.exists():
            slide.shapes.add_picture(str(image_path), Inches(1.05), Inches(1.45), width=Inches(7.4), height=Inches(4.55))
        else:
            placeholder = slide.shapes.add_shape(1, Inches(1.05), Inches(1.45), Inches(7.4), Inches(4.55))
            placeholder.fill.solid()
            placeholder.fill.fore_color.rgb = self._rgb("white")
            placeholder.line.color.rgb = self._rgb("secondary")
            paragraph = placeholder.text_frame.paragraphs[0]
            paragraph.text = (visual.alt_text if visual and visual.alt_text else "图片占位")
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = self.fonts.get("zh", "Microsoft YaHei")
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = self._rgb("muted")
        caption = visual.caption if visual and visual.caption else slide_ir.subtitle
        if caption:
            self._text_box(slide, Inches(8.8), Inches(1.65), Inches(3.2), Inches(2.8), caption, 16, "body", True)

    def _render_conclusion(self, slide, slide_ir: SlideIR) -> None:
        self._title(slide, slide_ir.title)
        bullets = slide_ir.bullets[:3] or ["默认本地运行", "输出文件可编辑", "后续可接入 GLM"]
        for idx, item in enumerate(bullets, start=1):
            top = Inches(1.65 + (idx - 1) * 1.15)
            self._number_badge(slide, idx, Inches(1.05), top)
            self._text_box(slide, Inches(1.85), top, Inches(9.8), Inches(0.5), item, 20, "body", True)

    def _title(self, slide, title: str) -> None:
        self._accent_bar(slide, height=0.12)
        self._text_box(slide, Inches(0.85), Inches(0.55), Inches(11.6), Inches(0.55), title, 24, "title", True)

    def _column(self, slide, left, top, width, title: str, bullets: list[str]) -> None:
        self._text_box(slide, left, top, width, Inches(0.45), title, 17, "primary", True)
        self._bullets(slide, bullets[:4], left, top + Inches(0.65), width, Inches(3.7), 17)

    def _bullets(self, slide, bullets: list[str], left, top, width, height, font_size: int) -> None:
        text_frame = slide.shapes.add_textbox(left, top, width, height).text_frame
        text_frame.clear()
        for idx, bullet in enumerate(bullets or ["内容已整理"]):
            paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.space_after = Pt(8)
            paragraph.font.size = Pt(font_size)
            paragraph.font.name = self.fonts.get("zh", "Microsoft YaHei")
            paragraph.font.color.rgb = self._rgb("body")

    def _text_box(self, slide, left, top, width, height, text: str, font_size: int, color_key: str, bold: bool = False) -> None:
        box = slide.shapes.add_textbox(left, top, width, height)
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.size = Pt(font_size)
        paragraph.font.name = self.fonts.get("zh", "Microsoft YaHei")
        paragraph.font.bold = bold
        paragraph.font.color.rgb = self._rgb(color_key)

    def _number_badge(self, slide, number: int, left, top) -> None:
        shape = slide.shapes.add_shape(1, left, top, Inches(0.42), Inches(0.42))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb("primary")
        shape.line.color.rgb = self._rgb("primary")
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = str(number)
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(13)
        paragraph.font.bold = True
        paragraph.font.color.rgb = self._rgb("white")

    def _accent_bar(self, slide, height: float = 0.18) -> None:
        shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb("primary")
        shape.line.fill.background()

    def _footer(self, slide, deck: DeckIR, slide_ir: SlideIR, index: int, total: int) -> None:
        footer = slide_ir.footer or deck.footer
        confidentiality = slide_ir.confidentiality or (footer.confidentiality if footer else None) or deck.confidentiality or "HUAWEI CONFIDENTIAL"
        text = footer.text if footer and footer.text else deck.deck_title
        self._text_box(slide, Inches(0.85), Inches(6.95), Inches(6.0), Inches(0.25), f"{text} · {confidentiality}", 9, "muted")
        if footer and not footer.show_page_number:
            return
        box = slide.shapes.add_textbox(Inches(11.8), Inches(6.95), Inches(0.8), Inches(0.25))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = f"{index}/{total}"
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph.font.size = Pt(9)
        paragraph.font.color.rgb = self._rgb("muted")

    def _set_background(self, slide) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self._rgb("background")

    def _format_cell(self, cell, bold: bool, fill: str, color: str) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = self._rgb(fill)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = self.fonts.get("zh", "Microsoft YaHei")
            paragraph.font.size = Pt(11)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = self._rgb(color)

    def _rgb(self, key: str) -> RGBColor:
        value = self.colors.get(key, "111827")
        return RGBColor.from_string(value)
