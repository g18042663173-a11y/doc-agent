from __future__ import annotations

from pathlib import Path

from doc_agent.ir.schemas import DocumentBlock, DocumentIR
from doc_agent.parsers.base import BaseParser
from doc_agent.utils.text_utils import normalize_space


class PptxParser(BaseParser):
    def parse(self, path: str | Path) -> DocumentIR:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-pptx is required to parse .pptx files") from exc

        source = Path(path)
        presentation = Presentation(str(source))
        blocks: list[DocumentBlock] = []
        title: str | None = None

        for slide_index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            if slide.shapes.title is not None and getattr(slide.shapes.title, "text", ""):
                title_candidate = normalize_space(slide.shapes.title.text)
                texts.append(title_candidate)
                if title is None:
                    title = title_candidate
            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue
                text = normalize_space(shape.text)
                if text and text not in texts:
                    texts.append(text)
            blocks.append(
                DocumentBlock(
                    id=f"s{slide_index}",
                    type="slide",
                    text="\n".join(texts),
                    meta={"slide_number": slide_index},
                )
            )

        return DocumentIR(
            source_file=str(source),
            source_type="pptx",
            title=title,
            blocks=blocks,
            meta={"parser": "python-pptx", "slide_count": len(presentation.slides)},
        )
