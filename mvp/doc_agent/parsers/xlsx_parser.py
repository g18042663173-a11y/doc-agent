from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import Any

from doc_agent.ir.schemas import DocumentBlock, DocumentIR
from doc_agent.parsers.base import BaseParser
from doc_agent.utils.text_utils import normalize_space, truncate_text


class XlsxParser(BaseParser):
    max_preview_rows = 30
    max_preview_cols = 12
    max_stat_rows = 500
    max_stat_cols = 20

    def parse(self, path: str | Path) -> DocumentIR:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is required to parse .xlsx files") from exc

        source = Path(path)
        workbook = load_workbook(source, read_only=False, data_only=False)
        blocks: list[DocumentBlock] = []
        workbook_meta: dict[str, Any] = {
            "parser": "openpyxl",
            "sheet_count": len(workbook.worksheets),
            "sheets": [],
        }
        block_id = 1

        for sheet in workbook.worksheets:
            sheet_meta = {
                "sheet_name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_cell_count": len(list(sheet.merged_cells.ranges)),
            }
            preview_rows = self._preview_rows(sheet)
            stats = self._numeric_stats(sheet, preview_rows[0] if preview_rows else [])
            formulas = self._formula_count(sheet)
            sheet_meta["formula_count"] = formulas
            sheet_meta["numeric_stats"] = stats
            workbook_meta["sheets"].append(sheet_meta)

            blocks.append(
                DocumentBlock(
                    id=f"b{block_id}",
                    type="heading",
                    text=f"Sheet: {sheet.title}",
                    level=1,
                    meta=sheet_meta,
                )
            )
            block_id += 1

            if preview_rows:
                blocks.append(
                    DocumentBlock(
                        id=f"b{block_id}",
                        type="table",
                        rows=preview_rows,
                        meta={
                            "sheet_name": sheet.title,
                            "preview_rows": len(preview_rows),
                            "preview_columns": max(len(row) for row in preview_rows),
                        },
                    )
                )
                block_id += 1

            summary = self._sheet_summary(sheet.title, sheet_meta, stats)
            blocks.append(
                DocumentBlock(
                    id=f"b{block_id}",
                    type="note",
                    text=summary,
                    meta={"sheet_name": sheet.title},
                )
            )
            block_id += 1

        return DocumentIR(
            source_file=str(source),
            source_type="xlsx",
            title=source.stem,
            blocks=blocks,
            meta=workbook_meta,
        )

    def _preview_rows(self, sheet) -> list[list[str]]:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row or 1, self.max_preview_rows),
            max_col=min(sheet.max_column or 1, self.max_preview_cols),
        ):
            values = [self._cell_to_text(cell.value) for cell in row]
            if any(value for value in values):
                rows.append(values)
        return rows

    def _numeric_stats(self, sheet, header_row: list[str]) -> list[dict[str, Any]]:
        columns: dict[int, list[float]] = {}
        for row in sheet.iter_rows(
            min_row=2 if header_row else 1,
            max_row=min(sheet.max_row or 1, self.max_stat_rows),
            max_col=min(sheet.max_column or 1, self.max_stat_cols),
            values_only=True,
        ):
            for index, value in enumerate(row, start=1):
                if isinstance(value, bool):
                    continue
                if isinstance(value, int | float):
                    columns.setdefault(index, []).append(float(value))

        stats: list[dict[str, Any]] = []
        for index, values in columns.items():
            if not values:
                continue
            label = header_row[index - 1] if index - 1 < len(header_row) and header_row[index - 1] else f"Column {index}"
            stats.append(
                {
                    "column": label,
                    "count": len(values),
                    "min": round(min(values), 4),
                    "max": round(max(values), 4),
                    "avg": round(mean(values), 4),
                }
            )
        return stats[:12]

    def _formula_count(self, sheet) -> int:
        count = 0
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row or 1, self.max_stat_rows),
            max_col=min(sheet.max_column or 1, self.max_stat_cols),
            values_only=True,
        ):
            for value in row:
                if isinstance(value, str) and value.startswith("="):
                    count += 1
        return count

    @staticmethod
    def _cell_to_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime | date):
            return value.isoformat()
        return truncate_text(normalize_space(str(value)), 120)

    @staticmethod
    def _sheet_summary(sheet_name: str, sheet_meta: dict[str, Any], stats: list[dict[str, Any]]) -> str:
        parts = [
            f"{sheet_name} 包含 {sheet_meta['max_row']} 行、{sheet_meta['max_column']} 列。",
            f"合并单元格 {sheet_meta['merged_cell_count']} 个，公式单元格 {sheet_meta['formula_count']} 个。",
        ]
        if stats:
            stat_text = "；".join(
                f"{item['column']} count={item['count']} min={item['min']} max={item['max']} avg={item['avg']}"
                for item in stats[:6]
            )
            parts.append(f"数值列统计：{stat_text}。")
        return "".join(parts)
