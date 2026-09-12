from __future__ import annotations

from pathlib import Path

from doc_agent.ir.schemas import DocumentBlock, DocumentIR
from doc_agent.parsers.base import BaseParser
from doc_agent.utils.text_utils import normalize_space


class DocxParser(BaseParser):
    def parse(self, path: str | Path) -> DocumentIR:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required to parse .docx files") from exc

        source = Path(path)
        doc = Document(str(source))
        blocks: list[DocumentBlock] = []
        title: str | None = None
        block_id = 1

        for paragraph in doc.paragraphs:
            text = normalize_space(paragraph.text)
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                level = self._heading_level(style_name)
                if title is None and level == 1:
                    title = text
                blocks.append(
                    DocumentBlock(
                        id=f"b{block_id}",
                        type="heading",
                        text=text,
                        level=level,
                    )
                )
            elif "List Bullet" in style_name:
                blocks.append(DocumentBlock(id=f"b{block_id}", type="bullet_list", items=[text]))
            else:
                if title is None:
                    title = text
                blocks.append(DocumentBlock(id=f"b{block_id}", type="paragraph", text=text))
            block_id += 1

        for table in doc.tables:
            rows = [[normalize_space(cell.text) for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(DocumentBlock(id=f"b{block_id}", type="table", rows=rows))
                block_id += 1

        return DocumentIR(
            source_file=str(source),
            source_type="docx",
            title=title,
            blocks=blocks,
            meta={"parser": "python-docx"},
        )

    @staticmethod
    def _heading_level(style_name: str) -> int:
        for token in reversed(style_name.split()):
            if token.isdigit():
                return min(max(int(token), 1), 6)
        return 1
