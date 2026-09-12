from __future__ import annotations

import argparse
import base64
import shutil
import zipfile
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches as DocxInches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = ROOT / "samples" / "input" / "synthetic"
DEFAULT_LARGE_ROWS = 50_000
PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate synthetic dirty office files for parser boundary testing.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--large-rows", type=int, default=DEFAULT_LARGE_ROWS)
    return parser


def generate_dirty_samples(output_dir: Path, *, large_rows: int = DEFAULT_LARGE_ROWS) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    generated = [
        _make_xlsx_merged_ledger(output_dir),
        _make_xlsx_cross_sheet_refs(output_dir),
        _make_xlsx_formula_report(output_dir),
        _make_xlsx_large_ledger(output_dir, large_rows),
        _make_docx_revisions_comments(output_dir),
        _make_docx_nested_lists_table_sections(output_dir),
        _make_docx_embedded_image(output_dir),
        _make_pptx_nested_groups(output_dir),
        _make_pptx_table_notes(output_dir),
        _make_pptx_smartart_equivalent(output_dir),
    ]
    return generated


def _make_xlsx_merged_ledger(output_dir: Path) -> Path:
    path = output_dir / "xlsx_01_merged_ledger.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目台账"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "部门台账"
    sheet.append(["部门", "事项", "负责人", "状态"])
    sheet.merge_cells("A3:A4")
    sheet["A3"] = "无线网络部"
    sheet["B3"] = "站点巡检"
    sheet["C3"] = "张三"
    sheet["D3"] = "推进中"
    sheet["B4"] = "隐患闭环"
    sheet["C4"] = "李四"
    sheet["D4"] = "待复核"
    sheet.merge_cells("B6:C6")
    sheet["A6"] = "备注"
    sheet["B6"] = "跨列合并说明"
    sheet["D6"] = "需周五前反馈"
    workbook.save(path)
    return path


def _make_xlsx_cross_sheet_refs(output_dir: Path) -> Path:
    path = output_dir / "xlsx_02_cross_sheet_refs.xlsx"
    workbook = Workbook()
    base = workbook.active
    base.title = "基础数据"
    base.append(["月份", "预算", "实际"])
    for month, budget, actual in [("一月", 120, 118), ("二月", 135, 142), ("三月", 150, 153)]:
        base.append([month, budget, actual])

    ledger = workbook.create_sheet("月度台账")
    ledger.append(["月份", "预算", "实际", "偏差"])
    for row in range(2, 5):
        ledger.cell(row=row, column=1, value=f"=基础数据!A{row}")
        ledger.cell(row=row, column=2, value=f"=基础数据!B{row}")
        ledger.cell(row=row, column=3, value=f"=基础数据!C{row}")
        ledger.cell(row=row, column=4, value=f"=C{row}-B{row}")

    summary = workbook.create_sheet("汇总")
    summary.append(["指标", "值"])
    summary.append(["预算合计", "=SUM(月度台账!B2:B4)"])
    summary.append(["实际合计", "=SUM(月度台账!C2:C4)"])
    summary.append(["偏差合计", "=SUM(月度台账!D2:D4)"])
    workbook.save(path)
    return path


def _make_xlsx_formula_report(output_dir: Path) -> Path:
    path = output_dir / "xlsx_03_formula_report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营报表"
    sheet.append(["区域", "收入", "成本", "毛利", "毛利率"])
    for region, revenue, cost in [("华东", 320, 210), ("华南", 280, 180), ("西北", 160, 120)]:
        row = sheet.max_row + 1
        sheet.cell(row=row, column=1, value=region)
        sheet.cell(row=row, column=2, value=revenue)
        sheet.cell(row=row, column=3, value=cost)
        sheet.cell(row=row, column=4, value=f"=B{row}-C{row}")
        sheet.cell(row=row, column=5, value=f"=D{row}/B{row}")
    sheet.append(["合计", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)", "=D5/B5"])
    workbook.save(path)
    return path


def _make_xlsx_large_ledger(output_dir: Path, rows: int) -> Path:
    path = output_dir / "xlsx_04_large_ledger_50000.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("超大巡检台账")
    sheet.append(["序号", "地市", "站点", "告警数", "闭环状态", "负责人", "计划日期", "备注"])
    for index in range(1, rows + 1):
        sheet.append(
            [
                index,
                f"地市{index % 17:02d}",
                f"站点-{index:05d}",
                index % 9,
                "已闭环" if index % 3 else "处理中",
                f"同事{index % 23:02d}",
                f"2026-07-{(index % 28) + 1:02d}",
                "用于压测 read_only、预览截断与大表性能",
            ]
        )
    workbook.save(path)
    return path


def _make_docx_revisions_comments(output_dir: Path) -> Path:
    path = output_dir / "docx_01_revisions_comments.docx"
    document = Document()
    document.add_heading("周报修订样例", level=1)
    document.add_paragraph("本段用于生成含插入、删除、格式修改与批注标记的测试文件。")
    document.save(path)
    _inject_docx_revision_xml(path)
    return path


def _make_docx_nested_lists_table_sections(output_dir: Path) -> Path:
    path = output_dir / "docx_02_nested_lists_table_sections.docx"
    document = Document()
    document.add_heading("项目周报", level=1)
    document.add_paragraph("一、本周进展", style="List Number")
    document.add_paragraph("1.1 完成台账核对", style="List Number 2")
    document.add_paragraph("1.2 输出风险清单", style="List Number 2")
    document.add_paragraph("关键风险", style="List Bullet")
    document.add_paragraph("跨部门依赖未完全闭环", style="List Bullet 2")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "责任人"
    table.cell(0, 2).text = "状态"
    table.cell(1, 0).text = "需求澄清"
    table.cell(1, 1).text = "王五"
    table.cell(1, 2).text = "完成"
    table.cell(2, 0).text = "风险复盘"
    table.cell(2, 1).text = "赵六"
    table.cell(2, 2).text = "推进中"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("下周计划", level=2)
    document.add_paragraph("补齐验收证据并同步导师。")
    document.save(path)
    return path


def _make_docx_embedded_image(output_dir: Path) -> Path:
    path = output_dir / "docx_03_embedded_image.docx"
    image = output_dir / "_synthetic_pixel.png"
    image.write_bytes(PIXEL_PNG)
    try:
        document = Document()
        document.add_heading("汇报插图", level=1)
        document.add_paragraph("下方图片模拟内嵌截图,解析器只记录存在与尺寸。")
        document.add_picture(str(image), width=DocxInches(1.2))
        document.save(path)
    finally:
        image.unlink(missing_ok=True)
    return path


def _make_pptx_nested_groups(output_dir: Path) -> Path:
    path = output_dir / "pptx_01_nested_groups.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(5.5), Inches(0.5))
    title.text_frame.text = "嵌套组合形状样例"
    presentation.save(path)
    _inject_nested_group_xml(path)
    return path


def _make_pptx_table_notes(output_dir: Path) -> Path:
    path = output_dir / "pptx_02_table_notes.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "项目例会汇报"
    slide.placeholders[1].text = "本页同时包含表格与演讲备注。"
    table_shape = slide.shapes.add_table(3, 3, Inches(0.8), Inches(2.4), Inches(6.8), Inches(1.3))
    table = table_shape.table
    for col, text in enumerate(["事项", "责任人", "状态"]):
        table.cell(0, col).text = text
    table.cell(1, 0).text = "风险闭环"
    table.cell(1, 1).text = "张三"
    table.cell(1, 2).text = "推进中"
    table.cell(2, 0).text = "资源协调"
    table.cell(2, 1).text = "李四"
    table.cell(2, 2).text = "待确认"
    slide.notes_slide.notes_text_frame.text = "演讲备注:强调风险闭环和下周资源协调。"
    presentation.save(path)
    return path


def _make_pptx_smartart_equivalent(output_dir: Path) -> Path:
    path = output_dir / "pptx_03_smartart_equivalent.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(5.0), Inches(0.5)).text_frame.text = "复杂图形等价样例"
    slide.shapes.add_shape(1, Inches(1.0), Inches(1.7), Inches(1.8), Inches(0.8)).text_frame.text = "输入"
    slide.shapes.add_shape(1, Inches(3.0), Inches(1.7), Inches(1.8), Inches(0.8)).text_frame.text = "处理"
    slide.shapes.add_shape(1, Inches(5.0), Inches(1.7), Inches(1.8), Inches(0.8)).text_frame.text = "输出"
    presentation.save(path)
    _inject_smartart_marker(path)
    return path


def _inject_docx_revision_xml(path: Path) -> None:
    revision_paragraph = """
<w:p>
  <w:r><w:t>修订样例:</w:t></w:r>
  <w:ins w:id="1" w:author="Codex" w:date="2026-07-08T00:00:00Z"><w:r><w:t>新增本周风险闭环动作</w:t></w:r></w:ins>
  <w:del w:id="2" w:author="Codex" w:date="2026-07-08T00:00:00Z"><w:r><w:delText>删除旧版口径</w:delText></w:r></w:del>
  <w:r><w:rPr><w:rPrChange w:id="3" w:author="Codex" w:date="2026-07-08T00:00:00Z"><w:rPr><w:b/></w:rPr></w:rPrChange></w:rPr><w:t>格式修改痕迹</w:t></w:r>
  <w:commentRangeStart w:id="0"/><w:r><w:t>批注定位文本</w:t></w:r><w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r>
</w:p>
""".strip()
    comments_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:comment w:id="0" w:author="Codex" w:date="2026-07-08T00:00:00Z">
    <w:p><w:r><w:t>请确认风险描述是否准确。</w:t></w:r></w:p>
  </w:comment>
</w:comments>
"""
    _rewrite_zip(
        path,
        {
            "word/document.xml": lambda data: data.replace(b"</w:body>", revision_paragraph.encode("utf-8") + b"</w:body>"),
            "word/_rels/document.xml.rels": lambda data: _append_xml_before(
                data,
                b"</Relationships>",
                b'<Relationship Id="rIdDirtyComments" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>',
                guard=b"rIdDirtyComments",
            ),
            "[Content_Types].xml": lambda data: _append_xml_before(
                data,
                b"</Types>",
                b'<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>',
                guard=b"/word/comments.xml",
            ),
        },
        extra_files={"word/comments.xml": comments_xml.encode("utf-8")},
    )


def _inject_nested_group_xml(path: Path) -> None:
    group_xml = """
<p:grpSp>
  <p:nvGrpSpPr><p:cNvPr id="100" name="外层组合"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr><a:xfrm><a:off x="914400" y="1371600"/><a:ext cx="3657600" cy="1828800"/><a:chOff x="0" y="0"/><a:chExt cx="3657600" cy="1828800"/></a:xfrm></p:grpSpPr>
  <p:grpSp>
    <p:nvGrpSpPr><p:cNvPr id="101" name="内层组合"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
    <p:grpSpPr><a:xfrm><a:off x="228600" y="228600"/><a:ext cx="3200400" cy="914400"/><a:chOff x="0" y="0"/><a:chExt cx="3200400" cy="914400"/></a:xfrm></p:grpSpPr>
    <p:sp>
      <p:nvSpPr><p:cNvPr id="102" name="组合内层行动项"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
      <p:spPr><a:xfrm><a:off x="228600" y="228600"/><a:ext cx="2743200" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
      <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>组合内层行动项</a:t></a:r><a:endParaRPr lang="zh-CN"/></a:p></p:txBody>
    </p:sp>
  </p:grpSp>
</p:grpSp>
""".strip()
    _rewrite_zip(
        path,
        {
            "ppt/slides/slide1.xml": lambda data: data.replace(b"</p:spTree>", group_xml.encode("utf-8") + b"</p:spTree>"),
        },
    )


def _inject_smartart_marker(path: Path) -> None:
    marker = (
        b'<p:extLst><p:ext uri="{codex-smartart-equivalent}">'
        b'<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
        b'r:dm="rId999" r:lo="rId998" r:qs="rId997" r:cs="rId996"/>'
        b"</p:ext></p:extLst>"
    )
    _rewrite_zip(
        path,
        {
            "ppt/slides/slide1.xml": lambda data: data.replace(b"</p:sld>", marker + b"</p:sld>"),
        },
    )


def _append_xml_before(data: bytes, marker: bytes, insertion: bytes, *, guard: bytes) -> bytes:
    if guard in data:
        return data
    return data.replace(marker, insertion + marker)


def _rewrite_zip(path: Path, replacements: dict[str, Callable[[bytes], bytes]], *, extra_files: dict[str, bytes] | None = None) -> None:
    temp = path.with_suffix(path.suffix + ".tmp")
    extra_files = extra_files or {}
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        existing = set()
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in replacements:
                data = replacements[item.filename](data)
            target.writestr(item, data)
            existing.add(item.filename)
        for filename, data in extra_files.items():
            if filename not in existing:
                target.writestr(filename, data)
    shutil.move(temp, path)


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    generated = generate_dirty_samples(args.output_dir, large_rows=args.large_rows)
    for path in generated:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
