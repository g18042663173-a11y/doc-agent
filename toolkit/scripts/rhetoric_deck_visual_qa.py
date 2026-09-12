"""Text-level empty/leak checks for rhetoric-deck-workflow sample decks.

Optional PowerPoint COM export writes PNGs when Office is installed.
This script lives in the repository `scripts/` tree and is not packed into the skill ZIP.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys

from pptx import Presentation


FORBIDDEN_DEFAULT = (
    "源件高度敏感",
    "源件封面机密",
    "源件卡片标题",
    "源件表格机密",
    "源件节点框",
    "源件密级内部资料不可进入产物",
)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--deck", type=Path, required=True)
    parser.add_argument("--forbidden", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--export-png", action="store_true")
    return parser


def inspect_deck(path: Path, forbidden: list[str] | None = None) -> dict:
    presentation = Presentation(path)
    phrases = forbidden or list(FORBIDDEN_DEFAULT)
    texts = _visible_texts(presentation)
    blob = "\n".join(texts)
    empty_slots = [item for item in texts if not item.strip()]
    hits = [phrase for phrase in phrases if phrase and phrase in blob]
    footer = [item for item in _master_texts(presentation) if item.strip()]
    return {
        "slides": len(presentation.slides),
        "text_blocks": len(texts),
        "empty_text_blocks": len(empty_slots),
        "forbidden_hits": hits,
        "footer_texts": footer,
        "ok": not hits and len(presentation.slides) >= 1,
    }


def export_pngs(deck: Path, outdir: Path) -> list[Path]:
    if sys.platform != "win32":
        return []
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return []
    outdir.mkdir(parents=True, exist_ok=True)
    application = None
    try:
        application = win32com.client.Dispatch("PowerPoint.Application")
        presentation = application.Presentations.Open(str(deck.resolve()), WithWindow=False)
        try:
            presentation.SaveAs(str(outdir.resolve()), 18)
        finally:
            presentation.Close()
    except Exception:
        return []
    finally:
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
    return sorted(outdir.glob("*.PNG")) + sorted(outdir.glob("*.png"))


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    forbidden = None
    if args.forbidden and args.forbidden.is_file():
        forbidden = [line.strip() for line in args.forbidden.read_text(encoding="utf-8").splitlines() if line.strip()]
    report = inspect_deck(args.deck, forbidden)
    if args.export_png and args.out:
        report["pngs"] = [str(path) for path in export_pngs(args.deck, args.out)]
    elif args.out:
        args.out.mkdir(parents=True, exist_ok=True)
        (args.out / "visual_qa_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


def _visible_texts(presentation) -> list[str]:
    texts: list[str] = []
    for slide in presentation.slides:
        texts.extend(_shape_texts(slide.shapes))
    texts.extend(_master_texts(presentation))
    return texts


def _master_texts(presentation) -> list[str]:
    texts: list[str] = []
    for master in presentation.slide_masters:
        texts.extend(_shape_texts(master.shapes))
        for layout in master.slide_layouts:
            texts.extend(_shape_texts(layout.shapes))
    return texts


def _shape_texts(shapes) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            texts.append(shape.text or "")
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    texts.append(cell.text or "")
        if getattr(shape, "shape_type", None) == 6:
            texts.extend(_shape_texts(shape.shapes))
    return texts


if __name__ == "__main__":
    raise SystemExit(main())
