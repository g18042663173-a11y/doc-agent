from __future__ import annotations

import importlib.util
import os
from pathlib import Path
import subprocess
import sys

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]


def _module():
    path = ROOT / "scripts" / "template_inspect.py"
    spec = importlib.util.spec_from_file_location("template_inspect", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_template_inspect_writes_profile_and_human_readable_structure(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)

    paths = _module().inspect_template(template, tmp_path / "audit")

    assert set(paths) == {"profile", "structure", "structure_markdown"}
    assert all(path.is_file() for path in paths.values())
    assert paths["structure_markdown"].read_text(encoding="utf-8").startswith("# PPT 模板结构：")


def test_template_inspect_cli_sets_its_own_backend_import_path(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    output_dir = tmp_path / "audit"
    environment = os.environ.copy()
    environment.pop("PYTHONPATH", None)

    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "template_inspect.py"),
            str(template),
            "--output-dir",
            str(output_dir),
        ],
        cwd=ROOT,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )

    assert completed.returncode == 0, completed.stderr
    assert (output_dir / "template_profile.json").is_file()
    assert (output_dir / "template_structure.md").is_file()
