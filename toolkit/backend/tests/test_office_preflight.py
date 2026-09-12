from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.parsers.errors import ParseFailure
from app.parsers.pptx_parser import parse_pptx
from app.parsers.docx_parser import parse_docx
from app.security.office_package import OfficePackageError, preflight_office_package


def test_office_preflight_accepts_real_type_and_rejects_extension_spoofing(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)

    report = preflight_office_package(source, purpose="source")
    assert report.office_kind == "pptx"
    assert report.safe

    disguised = tmp_path / "disguised.docx"
    disguised.write_bytes(source.read_bytes())
    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(disguised, purpose="source")
    assert captured.value.reason == "extension_mismatch"
    assert captured.value.loc == "input_file"


@pytest.mark.parametrize("unsafe_name", ["../escape.xml", "/absolute.xml"])
def test_office_preflight_rejects_unsafe_member_paths(tmp_path: Path, unsafe_name: str) -> None:
    source = _docx_bytes(tmp_path)
    package = _rewrite_package(source, extra={unsafe_name: b"unsafe"})
    path = tmp_path / "unsafe.docx"
    path.write_bytes(package)

    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(path, purpose="source")
    assert captured.value.reason == "unsafe_package_path"


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
def test_office_preflight_rejects_utf16_xml_entities(tmp_path: Path, encoding: str) -> None:
    """A UTF-16 encoded DTD/ENTITY must be caught, not bypassed by the ASCII-only scan."""
    source = _docx_bytes(tmp_path)

    def to_utf16(data: bytes) -> bytes:
        text = "<!DOCTYPE x [<!ENTITY y SYSTEM 'file:///x'>]>" + data.decode("utf-8")
        return text.encode(encoding)

    package = _rewrite_package(source, transforms={"word/document.xml": to_utf16})
    path = tmp_path / f"entity-{encoding}.docx"
    path.write_bytes(package)

    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(path, purpose="source")
    assert captured.value.reason == "unsafe_xml_declaration"


def test_office_preflight_rejects_duplicate_casefolded_parts_and_xml_entities(tmp_path: Path) -> None:
    source = _docx_bytes(tmp_path)
    duplicate = _rewrite_package(source, extra={"WORD/DOCUMENT.XML": b"duplicate"})
    duplicate_path = tmp_path / "duplicate.docx"
    duplicate_path.write_bytes(duplicate)
    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(duplicate_path, purpose="source")
    assert captured.value.reason == "duplicate_package_part"

    entity = _rewrite_package(
        source,
        transforms={"word/document.xml": lambda data: b"<!DOCTYPE x [<!ENTITY y SYSTEM 'file:///x'>]>" + data},
    )
    entity_path = tmp_path / "entity.docx"
    entity_path.write_bytes(entity)
    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(entity_path, purpose="source")
    assert captured.value.reason == "unsafe_xml_declaration"


def test_office_preflight_source_warns_but_template_rejects_active_content(tmp_path: Path) -> None:
    source = _pptx_bytes(tmp_path)
    package = _rewrite_package(
        source,
        extra={"ppt/embeddings/oleObject1.bin": b"unsafe"},
        transforms={
            "ppt/slides/_rels/slide1.xml.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/oleObject" '
                b'Target="../embeddings/oleObject1.bin"/></Relationships>',
            )
        },
    )
    path = tmp_path / "active.pptx"
    path.write_bytes(package)

    report = preflight_office_package(path, purpose="source")
    assert report.safe
    assert report.active_parts == ["ppt/embeddings/oleObject1.bin"]
    assert any("active content ignored" in warning for warning in report.warnings)

    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(path, purpose="template")
    assert captured.value.reason == "unsafe_ole"
    assert captured.value.context == "第 1 页"


def test_office_preflight_rejects_suspicious_compression_ratio(tmp_path: Path) -> None:
    source = _docx_bytes(tmp_path)
    package = _rewrite_package(source, extra={"word/large.xml": b"A" * (2 * 1024 * 1024)})
    path = tmp_path / "ratio.docx"
    path.write_bytes(package)

    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(path, purpose="source")
    assert captured.value.reason == "suspicious_compression_ratio"


def test_office_preflight_allows_only_chart_owned_embedded_xlsx(tmp_path: Path) -> None:
    directory_only = tmp_path / "embeddings-directory.pptx"
    directory_only.write_bytes(
        _rewrite_package(_pptx_bytes(tmp_path), extra={"ppt/embeddings/": b""})
    )
    directory_report = preflight_office_package(directory_only, purpose="output")
    assert directory_report.safe
    assert directory_report.active_parts == []

    chart_deck = tmp_path / "chart.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Series", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(3),
        chart_data,
    )
    presentation.save(chart_deck)

    report = preflight_office_package(chart_deck, purpose="output")
    assert report.safe
    assert report.active_parts == []

    unowned = tmp_path / "unowned-workbook.pptx"
    unowned.write_bytes(
        _rewrite_package(chart_deck.read_bytes(), extra={"ppt/embeddings/unowned.xlsx": _xlsx_bytes(tmp_path)})
    )
    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(unowned, purpose="output")
    assert captured.value.reason == "unsafe_ole"
    assert "unowned.xlsx" in captured.value.message


def test_office_preflight_rejects_chart_workbook_with_external_relationship(tmp_path: Path) -> None:
    chart_deck = tmp_path / "chart-external.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Series", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.LINE,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(3),
        chart_data,
    )
    presentation.save(chart_deck)

    with ZipFile(chart_deck) as package:
        workbook_name = next(name for name in package.namelist() if name.startswith("ppt/embeddings/"))
        workbook = package.read(workbook_name)
    unsafe_workbook = _rewrite_package(
        workbook,
        transforms={
            "xl/_rels/workbook.xml.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdExternal" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/externalLink" Target="https://example.invalid/data.xlsx" '
                b'TargetMode="External"/></Relationships>',
            )
        },
    )
    unsafe_deck = tmp_path / "unsafe-chart-workbook.pptx"
    unsafe_deck.write_bytes(
        _rewrite_package(chart_deck.read_bytes(), transforms={workbook_name: lambda _data: unsafe_workbook})
    )

    with pytest.raises(OfficePackageError) as captured:
        preflight_office_package(unsafe_deck, purpose="output")
    assert captured.value.reason == "external_relationship"
    assert workbook_name in captured.value.message


def test_office_parsers_apply_preflight_before_opening_spoofed_packages(tmp_path: Path) -> None:
    pptx = tmp_path / "base.pptx"
    pptx.write_bytes(_pptx_bytes(tmp_path))
    disguised = tmp_path / "disguised.docx"
    disguised.write_bytes(pptx.read_bytes())

    with pytest.raises(ParseFailure) as captured:
        parse_docx(disguised)
    assert captured.value.code == "E003"
    assert "实际 PPTX" in captured.value.message


def test_pptx_source_parser_records_ignored_active_content_without_copying_it(tmp_path: Path) -> None:
    package = _rewrite_package(
        _pptx_bytes(tmp_path),
        extra={"ppt/embeddings/oleObject1.bin": b"unsafe"},
    )
    path = tmp_path / "source-with-ole.pptx"
    path.write_bytes(package)

    document = parse_pptx(path)
    assert any("active content ignored" in warning for warning in document.warnings)


def _docx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("content")
    document.save(path)
    return path.read_bytes()


def _pptx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "base.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return path.read_bytes()


def _xlsx_bytes(tmp_path: Path) -> bytes:
    from openpyxl import Workbook

    path = tmp_path / "base.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "safe"
    workbook.save(path)
    return path.read_bytes()


def _rewrite_package(
    source: bytes,
    *,
    extra: dict[str, bytes] | None = None,
    transforms: dict[str, object] | None = None,
) -> bytes:
    output = BytesIO()
    transforms = transforms or {}
    with ZipFile(BytesIO(source)) as current, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in current.infolist():
            data = current.read(info.filename)
            transform = transforms.get(info.filename)
            if callable(transform):
                data = transform(data)
            target.writestr(info.filename, data)
        for name, data in (extra or {}).items():
            target.writestr(name, data)
    return output.getvalue()
