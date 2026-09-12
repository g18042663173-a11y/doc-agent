from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]


def test_render_cli_writes_word_docx(tmp_path: Path) -> None:
    output = tmp_path / "sample.docx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "word",
            "samples/ir/word_valid_03_table.json",
            "--output",
            str(output),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    assert output.exists()
    assert Document(str(output)).tables[0].cell(0, 0).text == "风险"


def test_render_cli_reports_validation_code_for_invalid_word_ir(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "word",
            "samples/ir/word_invalid_e004_ragged_table.json",
            "--output",
            str(tmp_path / "bad.docx"),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 1
    assert "E004" in result.stderr


def test_render_cli_rejects_template_for_word_with_a_structured_argument_error(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "word",
            "samples/ir/word_valid_01_plain.json",
            "--template",
            str(tmp_path / "ignored.pptx"),
            "--output",
            str(tmp_path / "word.docx"),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 2
    assert '"code": "E001"' in result.stderr
    assert "Traceback" not in result.stderr


def test_render_cli_writes_deck_pptx_and_runs_lint(tmp_path: Path) -> None:
    ir_path = tmp_path / "deck.json"
    output = tmp_path / "sample.pptx"
    ir_path.write_text(
        """
{
  "ir_type": "deck",
  "ir_version": "1.7",
  "meta": {"title": "CLI Deck", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
  "slides": [
    {"layout": "cover", "title": "CLI Deck"},
    {"layout": "agenda", "items": ["背景", "进展"]},
    {"layout": "title_bullets", "title": "进展", "bullets": [{"text": "render 命令直接产出 PPTX", "level": 1}]}
  ]
}
""".strip()
        + "\n",
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "deck",
            str(ir_path),
            "--output",
            str(output),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    assert output.exists()
    assert len(Presentation(str(output)).slides) == 3


def test_render_cli_reports_validation_code_for_invalid_deck_ir(tmp_path: Path) -> None:
    ir_path = tmp_path / "bad_deck.json"
    ir_path.write_text(
        """
{
  "ir_type": "deck",
  "ir_version": "1.7",
  "meta": {"title": "坏 Deck"},
  "slides": [{"layout": "unknown", "title": "无法渲染"}]
}
""".strip()
        + "\n",
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "deck",
            str(ir_path),
            "--output",
            str(tmp_path / "bad.pptx"),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 1
    assert "D003" in result.stderr
