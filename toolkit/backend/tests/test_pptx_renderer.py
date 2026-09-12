from __future__ import annotations

import ast
import json
import shutil
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_render_deck_ir_p0_layouts_are_editable(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "Q3 业务汇报", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {"layout": "cover", "title": "Q3 业务汇报", "subtitle": "命令行文档工具链", "presenter": "张三"},
                {"layout": "agenda", "items": ["业务回顾", "关键进展", "风险与对策"]},
                {"layout": "section", "index": 1, "title": "关键进展"},
                {"layout": "title_bullets", "title": "主链路进展", "bullets": [{"text": "stub 模式可离线运行", "level": 1}]},
                {
                    "layout": "table",
                    "title": "风险清单",
                    "table": {
                        "header": ["风险", "等级"],
                        "rows": [["模型 JSON 不稳定", "中"]],
                    },
                },
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "deck.pptx")
    prs = Presentation(str(output))
    all_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))

    assert len(prs.slides) == 5
    assert "Q3 业务汇报" in all_text
    assert "业务回顾" in all_text
    assert "关键进展" in all_text
    assert "stub 模式可离线运行" in all_text
    assert "HUAWEI CONFIDENTIAL" in all_text
    assert any(shape.has_table for shape in prs.slides[4].shapes)


def test_render_deck_ir_p0_layouts_lint_with_zero_errors(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "P0 五版式", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {"layout": "cover", "title": "P0 五版式"},
                {"layout": "agenda", "items": ["章节一", "章节二"]},
                {"layout": "section", "index": 1, "title": "章节一"},
                {"layout": "title_bullets", "title": "要点", "bullets": [{"text": "P0 版式可编辑", "level": 1}]},
                {"layout": "table", "title": "表格", "table": {"header": ["项", "值"], "rows": [["状态", "通过"]]}},
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "p0-lint.pptx")
    report = check_pptx(output, classification=deck.meta.classification)

    assert report.summary["errors"] == 0


def test_render_deck_ir_decision_matrix_table_matches_source_spec(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "方案对比", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "table",
                    "title": "方案对比表",
                    "table": {
                        "header": ["方案", "可靠性", "成本", "交付", "结论"],
                        "column_groups": [{"label": "评估维度", "start_col": 1, "span": 3}],
                        "row_groups": [{"label": "主推方案", "start_row": 0, "span": 2}],
                        "cell_spans": [{"area": "body", "row": 0, "col": 2, "rowspan": 1, "colspan": 2}],
                        "conclusion_col": 4,
                        "rows": [
                            [
                                {"text": "方案A"},
                                "高",
                                {"items": ["周期短", "依赖少"]},
                                "",
                                {"text": "推荐", "emphasis": "yellow"},
                            ],
                            ["方案B", "中", "成本可控", "交付较慢", {"text": "备选", "emphasis": "cyan"}],
                        ],
                    },
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "decision-matrix.pptx")
    prs = Presentation(str(output))
    table_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_table", False))
    table = table_shape.table

    assert len(table.rows) == 5
    assert len(table.columns) == 5
    assert round(table.rows[1].height / 914400, 2) == 0.50
    assert round(table.rows[3].height / 914400, 2) == 0.40
    assert _cell_fill_hex(table.cell(1, 0)) == "C7000B"
    header_run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert header_run.font.name == "Arial"
    assert header_run.font.size == Pt(11)
    assert header_run.font.bold is True
    assert str(header_run.font.color.rgb) == "FFFFFF"
    assert _cell_fill_hex(table.cell(3, 0)) == "FFFFFF"
    assert _cell_fill_hex(table.cell(4, 0)) == "F5F5F5"
    assert _cell_fill_hex(table.cell(3, 4)) == "FCC800"
    assert _cell_fill_hex(table.cell(4, 4)) == "30B5C5"
    assert "• 周期短" in table.cell(3, 2).text
    assert "• 依赖少" in table.cell(3, 2).text
    assert "主推方案" in table.cell(2, 0).text
    assert _cell_fill_hex(table.cell(0, 0)) != "C7000B"
    assert _cell_fill_hex(table.cell(0, 4)) != "C7000B"
    assert 'w="12700"' in table.cell(1, 0)._tc.xml
    assert "DDDDDD" in table.cell(1, 0)._tc.xml


def test_render_deck_ir_wide_table_widths_fill_content_area(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    content_width_in = theme["layouts"]["table"]["width_in"]
    cases = (("weighted", [0.18, 0.34, 0.24, 0.24]), ("auto", None), ("extreme", [0.01, 0.33, 0.33, 0.33]))
    for case_name, col_widths in cases:
        table_data = {
            "header": ["能力域", "目标与方法", "目标值·测试条件", "现阶段证据与判断"],
            "rows": [
                [
                    "电磁环境感知",
                    "联合时频图结合深度模型识别复杂干扰与小样本干扰",
                    "需核验目标值、样本规模及测试信噪比",
                    "已开展训练与仿真分析，尚不能据此判定达成",
                ]
            ],
        }
        if col_widths is not None:
            table_data["col_widths"] = col_widths
        deck = DeckIR.model_validate(
            {
                "ir_type": "deck",
                "ir_version": "1.7",
                "meta": {"title": "宽表回归", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
                "slides": [{"layout": "table", "title": "研究目标必须由可量化指标约束", "table": table_data}],
            }
        )

        output = render_deck_ir(deck, tmp_path / f"{case_name}-wide-table.pptx")
        prs = Presentation(str(output))
        table_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_table", False))
        widths_in = [column.width / 914400 for column in table_shape.table.columns]
        dynamic_min_width_in = content_width_in / (2 * len(widths_in))

        assert abs(sum(widths_in) - content_width_in) < 0.01
        assert all(width >= dynamic_min_width_in for width in widths_in)
        assert all(_cell_fill_hex(table_shape.table.cell(0, col)) == "C7000B" for col in range(4))


def test_render_deck_ir_legal_extremes_stay_in_page_and_shrink_text(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    long_title = "超长结论标题用于验证换行缩放与版心边界" * 8
    long_text = "这是仍然合法但明显超长的内容，用于验证固定版心内的换行、缩小和防重叠。" * 10
    nodes = [
        {"id": f"n{index}", "text": f"节点{index} {long_text}", "type": "secondary"}
        for index in range(12)
    ]
    slides = [
        {"layout": "cover", "title": long_title, "subtitle": long_text, "presenter": "测试人员"},
        {"layout": "agenda", "items": [f"议题{index} {long_text}" for index in range(8)]},
        {"layout": "section", "index": 99, "title": long_title, "subtitle": long_text},
        {
            "layout": "title_bullets",
            "title": long_title,
            "bullets": [{"text": long_text, "level": 1} for _ in range(7)],
        },
        {
            "layout": "table",
            "title": long_title,
            "table": {
                "header": [f"第{index}列超长表头" for index in range(8)],
                "rows": [[f"第{row}行第{col}列 {long_text}" for col in range(8)] for row in range(12)],
                "col_widths": [0.01, 0.02, 0.03, 0.04, 0.1, 0.2, 0.25, 0.35],
                "row_groups": [
                    {"label": f"分组{index}", "start_row": index, "span": 1}
                    for index in range(8)
                ],
            },
        },
        {"layout": "architecture_diagram", "title": long_title, "nodes": nodes, "edges": [], "groups": []},
        {
            "layout": "cards",
            "title": long_title,
            "cards": [
                {"title": f"卡片{index} {long_title}", "desc": long_text * 2, "tag": long_text}
                for index in range(4)
            ],
        },
        {
            "layout": "two_column",
            "title": long_title,
            "left": {
                "heading": long_title,
                "text": long_text * 2,
                "bullets": [{"text": long_text, "level": 1} for _ in range(7)],
            },
            "right": {"heading": "短标题", "text": "短内容"},
        },
        {"layout": "conclusion", "title": long_title, "bullets": [long_text] * 5, "cta": long_text},
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "极端版式", "classification": "公开", "theme": "hw_v1"},
            "slides": slides,
        }
    )

    output = render_deck_ir(deck, tmp_path / "legal-extremes.pptx")
    prs = Presentation(str(output))
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= slide_width
            assert shape.top + shape.height <= slide_height
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                assert shape.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            assert cell.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    table_shape = next(shape for shape in prs.slides[4].shapes if getattr(shape, "has_table", False))
    table_bottom_in = (table_shape.top + table_shape.height) / 914400
    expected_bottom_in = theme["slide"]["footer_top_in"] - theme["grid"]["min_gap_in"]
    assert table_bottom_in <= expected_bottom_in + 0.01
    assert len(table_shape.table.rows) == 21


def test_render_architecture_dense_edge_labels_do_not_overlap(tmp_path: Path) -> None:
    from itertools import combinations

    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    nodes = [{"id": f"n{index}", "text": f"节点 {index}", "type": "secondary"} for index in range(12)]
    edges = [
        {
            "from": f"n{index}",
            "to": f"n{(index + offset) % 12}",
            "label": f"连接 {index}-{offset} 的较长说明文字",
            "direction": "forward",
        }
        for offset in (1, 2)
        for index in range(12)
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "密集架构图", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "密集边标签必须避让节点和其它标签",
                    "nodes": nodes,
                    "edges": edges,
                    "groups": [],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "dense-edge-labels.pptx")
    slide = Presentation(str(output)).slides[0]
    labels = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_EDGE_LABEL:")]
    node_shapes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    edge_segments = _architecture_edge_segments(slide)

    assert len(labels) == 24
    assert not any(_overlap(first, second) for first, second in combinations(labels, 2))
    assert not any(_overlap(label, node) for label in labels for node in node_shapes)
    assert all(_is_orthogonal_connector(segment) for segment in edge_segments)
    assert not any(
        _connector_crosses_unrelated_node(segment, node)
        for segment in edge_segments
        for node in node_shapes
    )
    density_warnings = [
        item
        for item in check_pptx(output, classification="公开").items
        if item.code == "HW-W03" and item.message.startswith("架构图过密·")
    ]
    assert [item.message for item in density_warnings] == [
        "架构图过密·边数 24 超过阈值 12,建议人工调整或拆分。"
    ]


def test_render_deck_ir_uses_16_by_9_page_size(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "尺寸"},
            "slides": [{"layout": "cover", "title": "尺寸"}],
        }
    )

    output = render_deck_ir(deck, tmp_path / "size.pptx")
    prs = Presentation(str(output))

    assert round(prs.slide_width / 914400, 2) == 13.34
    assert round(prs.slide_height / 914400, 1) == 7.5


def test_render_process_flow_as_equal_editable_shapes_in_both_orientations(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    steps = [
        {"id": "input", "title": "材料输入", "description": "收集原始报告"},
        {"id": "parse", "title": "结构解析", "description": "形成 DocumentIR"},
        {"id": "review", "title": "合规复核", "description": "输出报告"},
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "流程", "classification": "公开"},
            "slides": [
                {"layout": "process_flow", "title": "横向流程", "orientation": "horizontal", "steps": steps},
                {"layout": "process_flow", "title": "纵向流程", "orientation": "vertical", "steps": steps},
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "process-flow.pptx")
    prs = Presentation(str(output))
    for slide in prs.slides:
        step_shapes = [shape for shape in slide.shapes if shape.name.startswith("HW_PROCESS_STEP:")]
        connectors = [shape for shape in slide.shapes if shape.name.startswith("HW_PROCESS_CONNECTOR:")]
        assert len(step_shapes) == 3
        assert len(connectors) == 2
        assert len({(shape.width, shape.height) for shape in step_shapes}) == 1
        assert all(getattr(shape, "has_text_frame", False) for shape in step_shapes)
        assert all(shape.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE for shape in step_shapes)
        assert all(shape.top + shape.height < Inches(6.9) for shape in step_shapes)
        assert not any("PICTURE" in str(shape.shape_type) for shape in slide.shapes)


def test_render_timeline_as_equal_editable_milestones_in_both_orientations(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    milestones = [
        {"label": "阶段一", "title": "基线冻结", "description": "Schema 与样例", "status": "completed"},
        {"label": "阶段二", "title": "能力验证", "description": "渲染与 lint", "status": "current"},
        {"label": "阶段三", "title": "内网终审", "description": "真实字体与 CI", "status": "planned"},
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "时间线", "classification": "公开"},
            "slides": [
                {"layout": "timeline", "title": "横向路线", "orientation": "horizontal", "milestones": milestones},
                {"layout": "timeline", "title": "纵向路线", "orientation": "vertical", "milestones": milestones},
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "timeline.pptx")
    prs = Presentation(str(output))
    for slide in prs.slides:
        cards = [shape for shape in slide.shapes if shape.name.startswith("HW_TIMELINE_MILESTONE:")]
        markers = [shape for shape in slide.shapes if shape.name.startswith("HW_TIMELINE_MARKER:")]
        axes = [shape for shape in slide.shapes if shape.name == "HW_TIMELINE_AXIS"]
        assert len(cards) == 3
        assert len(markers) == 3
        assert len(axes) == 1
        assert len({(shape.width, shape.height) for shape in cards}) == 1
        assert all(card.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE for card in cards)
        assert all(card.top + card.height < Inches(6.9) for card in cards)
        assert not any("PICTURE" in str(shape.shape_type) for shape in slide.shapes)


def test_render_horizontal_data_bar_kpi_and_image_slot_as_editable_objects(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    payload = json.loads((ROOT / "samples/ir/deck_valid_09_data_bar_kpi.json").read_text(encoding="utf-8"))
    deck = DeckIR.model_validate(payload)
    output = render_deck_ir(deck, tmp_path / "data-bar-kpi.pptx")
    prs = Presentation(str(output))

    kpi_shapes = [shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_RENDERED_TEXT:KPI_CARD:")]
    assert len(kpi_shapes) == 3
    assert len({(shape.width, shape.height) for shape in kpi_shapes}) == 1
    assert [shape.text_frame.paragraphs[0].text for shape in kpi_shapes] == ["96.8%", "42ms", "98.2%"]
    assert all(shape.text_frame.paragraphs[0].runs[0].font.name == "Arial" for shape in kpi_shapes)

    chart_shape = next(shape for shape in prs.slides[1].shapes if getattr(shape, "has_chart", False))
    threshold_line = next(shape for shape in prs.slides[1].shapes if shape.name.startswith("HW_THRESHOLD_LINE"))
    threshold_label = next(shape for shape in prs.slides[1].shapes if shape.name.startswith("HW_THRESHOLD_LABEL"))
    assert chart_shape.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
    assert chart_shape.chart.category_axis.reverse_order is True
    assert chart_shape.chart.value_axis.minimum_scale == 0
    assert chart_shape.chart.value_axis.maximum_scale == pytest.approx(100.0)
    assert threshold_line.height > threshold_line.width
    assert "通过率目标阈值 75%" in threshold_label.text
    assert threshold_label.text_frame.margin_left == 0
    assert threshold_label.text_frame.margin_right == 0
    assert threshold_label.text_frame.margin_top == 0
    assert threshold_label.text_frame.margin_bottom == 0
    assert threshold_label.top + threshold_label.height < chart_shape.top

    theme = json.loads(
        (ROOT / "backend" / "app" / "rendering" / "themes" / "hw_theme.json").read_text(encoding="utf-8")
    )
    plot = theme["layouts"]["chart"]["plot"]
    threshold = theme["layouts"]["chart"]["threshold"]
    plot_left = plot["left_in"] + threshold["plot_left_offset_in"]
    plot_width = plot["width_in"] - threshold["plot_left_offset_in"] - threshold["plot_right_offset_in"]
    expected_threshold_x = plot_left + plot_width * 75 / chart_shape.chart.value_axis.maximum_scale
    assert threshold_line.left / 914400 == pytest.approx(expected_threshold_x, abs=0.02)

    image_slot = next(shape for shape in prs.slides[2].shapes if shape.name == "HW_IMAGE_PLACEHOLDER")
    assert abs((image_slot.width / image_slot.height) - (16 / 9)) < 0.01
    assert "等比放入" in image_slot.text and "禁止随意裁切" in image_slot.text
    assert not any("PICTURE" in str(shape.shape_type) for shape in prs.slides[2].shapes)

    report = check_pptx(output, classification=deck.meta.classification)
    assert report.summary["errors"] == 0
    assert not {"HW-W01", "HW-W02", "HW-W06", "HW-W07", "HW-W09"} & {item.code for item in report.items}


def test_render_kpi_value_keeps_number_and_unit_on_one_line(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.2",
            "meta": {"title": "KPI 单位", "classification": "公开"},
            "slides": [
                {
                    "layout": "cards",
                    "title": "KPI 数值与单位保持同一视觉单元",
                    "variant": "kpi",
                        "cards": [
                            {"title": "下行峰值吞吐", "desc": "1.18 Gbps", "tag": "目标 ≥ 1.0 Gbps"},
                            {"title": "协同感知时延", "desc": "16.4 ms", "tag": "目标 ≤ 20 ms"},
                            {"title": "高速切换中断", "desc": "38 ms", "tag": "目标 ≤ 50 ms"},
                            {"title": "典型业务功耗", "desc": "10.7 W", "tag": "目标 ≤ 12 W"},
                        ],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "kpi-unit.pptx")
    slide = Presentation(str(output)).slides[0]
    value_paragraphs = [
        shape.text_frame.paragraphs[0]
        for shape in slide.shapes
        if shape.name.startswith("HW_RENDERED_TEXT:KPI_CARD:")
    ]
    values = [paragraph.text for paragraph in value_paragraphs]
    sizes = [paragraph.runs[0].font.size.pt for paragraph in value_paragraphs]

    assert values == ["1.18\u00a0Gbps", "16.4\u00a0ms", "38\u00a0ms", "10.7\u00a0W"]
    assert all("\n" not in value for value in values)
    assert 20 <= sizes[0] < sizes[1] <= 28
    report = check_pptx(output, classification="公开")
    assert report.summary["warnings"] == 0


def test_render_deck_ir_uses_theme_layout_coordinates(tmp_path: Path, monkeypatch) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    import app.rendering.theme as theme_module
    from app.rendering.pptx_renderer import render_deck_ir

    source_theme = json.loads((ROOT / "backend" / "app" / "rendering" / "themes" / "hw_theme.json").read_text(encoding="utf-8"))
    title_layout = source_theme.setdefault("layouts", {}).setdefault("title", {})
    title_layout["left_in"] = 1.2
    title_layout["top_in"] = 0.9
    title_layout["width_in"] = 9.5
    title_layout["height_in"] = 0.7
    theme_dir = tmp_path / "themes"
    theme_dir.mkdir()
    (theme_dir / "custom.json").write_text(json.dumps(source_theme), encoding="utf-8")
    monkeypatch.setattr(theme_module, "THEMES_DIR", theme_dir)
    monkeypatch.setitem(theme_module.THEME_REGISTRY, "custom", "custom")

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "主题坐标", "theme": "custom"},
            "slides": [{"layout": "title_bullets", "title": "可校准标题", "bullets": [{"text": "坐标来自 theme", "level": 1}]}],
        }
    )

    output = render_deck_ir(deck, tmp_path / "theme-layout.pptx")
    prs = Presentation(str(output))
    title_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text == "可校准标题")

    assert round(title_shape.left / 914400, 2) == 1.2
    assert round(title_shape.top / 914400, 2) == 0.9
    assert round(title_shape.width / 914400, 2) == 9.5
    assert round(title_shape.height / 914400, 2) == 0.7
    report = check_pptx(output, classification=deck.meta.classification, theme_name="custom")
    assert "HW-W08" not in {item.code for item in report.items}


def test_render_deck_ir_applies_source_file_theme_font_footer_and_card_tokens(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.theme import load_theme
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "主题映射", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {"layout": "section", "index": 1, "title": "阶段一", "subtitle": "技术评审"},
                {
                    "layout": "cards",
                    "title": "关键内容",
                    "cards": [
                        {"title": "方案", "desc": "红色强调由主题控制"},
                        {"title": "风险", "desc": "边框和底色由主题控制"},
                    ],
                },
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "official-theme.pptx")
    prs = Presentation(str(output))
    theme = load_theme("hw_v1")

    section_index = next(
        shape for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text == "01"
    )
    section_index_run = section_index.text_frame.paragraphs[0].runs[0]
    assert section_index_run.font.name == "Arial"

    section_title = next(
        shape for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text == "阶段一"
    )
    section_title_run = section_title.text_frame.paragraphs[0].runs[0]
    assert section_title_run.font.size == Pt(28)
    assert section_title.text_frame.paragraphs[0].line_spacing == 1.3
    assert "typeface=\"Arial\"" in section_title_run._r.xml
    assert "typeface=\"微软雅黑\"" in section_title_run._r.xml

    card_shapes = [
        shape
        for shape in prs.slides[1].shapes
        if round(shape.height / 914400, 1) == 2.1
        and round(shape.width / 914400, 1) > 5.0
        and f"#{shape.fill.fore_color.rgb}" == "#DDDDDD"
    ]
    assert card_shapes
    card = card_shapes[0]
    assert f"#{card.fill.fore_color.rgb}" == "#DDDDDD"
    assert f"#{card.line.color.rgb}" == "#DDDDDD"
    assert card.line.width == Pt(0.5)

    all_text = "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    assert "Security Level: HUAWEI CONFIDENTIAL" in all_text
    assert "1/2" in all_text
    assert "2/2" in all_text
    assert theme["footer"]["copyright"] in all_text


def test_render_ppt_typography_uses_large_semantic_hierarchy_and_compact_tables(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "字号层级", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {"layout": "cover", "title": "暑期实践团打印签字材料说明", "subtitle": "活动前完成阅读与签署", "date": "2026年7月"},
                {
                    "layout": "title_bullets",
                    "title": "本次签署材料覆盖保密、出行与现场安全三项要求",
                    "bullets": [
                        {"text": "保密承诺书明确活动纪律", "level": 1},
                        {"text": "交通安全告知书明确出行要求", "level": 1},
                        {"text": "安全生产告知书明确现场要求", "level": 1},
                    ],
                },
                {
                    "layout": "two_column",
                    "title": "保密承诺要求落实到个人责任",
                    "left": {"heading": "需作出的承诺", "bullets": [{"text": "遵守保密规定", "level": 1}]},
                    "right": {"heading": "活动期间注意事项", "bullets": [{"text": "按要求管理电子设备", "level": 1}]},
                },
                {"layout": "table", "title": "紧凑表格保持实测字号", "table": {"header": ["材料", "状态"], "rows": [["承诺书", "待签署"]]}},
                {
                    "layout": "title_bullets",
                    "title": "长" * 35,
                    "bullets": [{"text": "长标题应降档但正文仍保持可读", "level": 1}],
                },
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "typography.pptx")
    prs = Presentation(str(output))

    def font_size(shape) -> float:
        return shape.text_frame.paragraphs[0].runs[0].font.size.pt

    def named(slide, name: str):
        return next(shape for shape in slide.shapes if shape.name == name)

    cover = prs.slides[0]
    assert font_size(named(cover, "HW_RENDERED_TEXT:COVER_TITLE")) == 40
    assert font_size(named(cover, "HW_RENDERED_TEXT:COVER_SUBTITLE")) == 20
    assert font_size(named(cover, "HW_RENDERED_TEXT:NOTE")) == 12

    bullets = prs.slides[1]
    assert font_size(named(bullets, "HW_RENDERED_TEXT:TITLE")) == 28
    assert {font_size(shape) for shape in bullets.shapes if shape.name == "HW_RENDERED_TEXT:BODY"} == {16}

    columns = prs.slides[2]
    assert {font_size(shape) for shape in columns.shapes if shape.name == "HW_RENDERED_TEXT:SUBHEADING"} == {18}
    assert {font_size(shape) for shape in columns.shapes if shape.name == "HW_RENDERED_TEXT:BODY"} == {16}

    table = next(shape.table for shape in prs.slides[3].shapes if getattr(shape, "has_table", False))
    assert table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size == Pt(11)
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size == Pt(10)

    long_title_slide = prs.slides[4]
    long_title = named(long_title_slide, "HW_RENDERED_TEXT:TITLE")
    title_bar = named(long_title_slide, "HW_DECORATION:RED_BAR")
    assert font_size(long_title) == 24
    assert long_title.top + long_title.height < title_bar.top


def test_title_decoration_tracks_wrapped_title_and_preserves_single_line_position(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    short_title = "单行结论标题"
    wrapped_title = (
        "CNN-LSTM在1000km轨道高度达到NMSE -10dB；"
        "edRVFL仍需量化验证，当前输入仍缺少跨场景对比结果"
    )
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "标题装饰线回归", "classification": "PUBLIC"},
            "slides": [
                {"layout": "title_bullets", "title": short_title, "bullets": [{"text": "短标题基线", "level": 1}]},
                {"layout": "title_bullets", "title": wrapped_title, "bullets": [{"text": "长标题换行", "level": 1}]},
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "title-decoration.pptx")
    prs = Presentation(str(output))

    def named(slide, name: str):
        return next(shape for shape in slide.shapes if shape.name == name)

    short_title_shape = named(prs.slides[0], "HW_RENDERED_TEXT:TITLE")
    short_bar = named(prs.slides[0], "HW_DECORATION:RED_BAR")
    wrapped_title_shape = named(prs.slides[1], "HW_RENDERED_TEXT:TITLE")
    wrapped_bar = named(prs.slides[1], "HW_DECORATION:RED_BAR")
    base_title = theme["layouts"]["title"]
    base_bar = theme["layouts"]["title_bar"]
    expected_gap_in = base_bar["gap_after_title_in"]

    assert abs(short_bar.top / 914400 - base_bar["top_in"]) < 0.001
    assert abs(short_title_shape.height / 914400 - base_title["height_in"]) < 0.001
    assert wrapped_title_shape.height > short_title_shape.height
    assert wrapped_bar.top > short_bar.top
    assert abs((wrapped_bar.top - wrapped_title_shape.top - wrapped_title_shape.height) / 914400 - expected_gap_in) < 0.001
    assert wrapped_title_shape.top + wrapped_title_shape.height < wrapped_bar.top


def test_pptx_renderer_keeps_layout_dimensions_in_theme() -> None:
    source_path = ROOT / "backend" / "app" / "rendering" / "pptx_renderer.py"
    source = source_path.read_text(encoding="utf-8")
    tree = ast.parse(source)
    raw_literal_calls = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in {"Inches", "Pt"}:
            for arg in node.args:
                if isinstance(arg, ast.Constant) and isinstance(arg.value, (int, float)):
                    raw_literal_calls.append((node.func.id, arg.value, node.lineno))

    assert raw_literal_calls == []
    assert source.count(".add_picture(") == 2
    assert "HW_ASSET_IMAGE:" in source
    assert "shadow" not in source.lower()
    assert "gradient" not in source.lower()


def test_render_deck_ir_p1_layouts_and_downgrades(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "P1 版式"},
            "slides": [
                {
                    "layout": "two_column",
                    "title": "双栏对比",
                    "left": {"heading": "优势", "bullets": [{"text": "离线可测", "level": 1}]},
                    "right": {"heading": "风险", "text": "视觉终审仍需人工"},
                },
                {
                    "layout": "cards",
                    "title": "三大能力",
                    "cards": [
                        {"title": "解析", "desc": "多格式输入"},
                        {"title": "渲染", "desc": "可编辑产物"},
                    ],
                },
                {"layout": "conclusion", "title": "结论", "bullets": ["主链路继续推进"], "cta": "进入内网校准"},
                {
                    "layout": "chart",
                    "title": "趋势",
                    "chart": {"kind": "bar", "categories": ["W1"], "series": [{"name": "完成数", "values": [3]}]},
                },
                {"layout": "image", "title": "架构图", "placeholder": "黄区替换真图", "caption": "工具链架构"},
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "p1.pptx")
    prs = Presentation(str(output))
    all_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))

    assert "双栏对比" in all_text
    assert "离线可测" in all_text
    assert "三大能力" in all_text
    assert "多格式输入" in all_text
    assert "主链路继续推进" in all_text
    assert any(getattr(shape, "has_chart", False) for shape in prs.slides[3].shapes)
    assert "图表待内网 Skill / 后续版本生成" not in all_text
    assert "图片占位" in all_text


def test_render_deck_ir_performance_chart_matches_source_spec(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "性能图表", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "性能趋势",
                    "chart": {
                        "kind": "bar",
                        "unit": "ms",
                        "categories": ["1月", "2月", "3月"],
                        "series": [
                            {"name": "方案A", "values": [62, 58, 55], "emphasis": True},
                            {"name": "方案B", "values": [75, 69, 64]},
                        ],
                        "show_data_labels": True,
                        "legend_position": "right",
                        "thresholds": [{"value": 60, "label": "目标阈值"}],
                        "side_conclusion": "方案A 2月起低于60ms目标阈值。",
                        "side_table": {"header": ["指标", "结论"], "rows": [["时延", "2月起达标"], ["稳定性", "需关注"]]},
                    },
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "performance-chart.pptx")
    prs = Presentation(str(output))
    slide = prs.slides[0]
    chart_shape = next(shape for shape in slide.shapes if getattr(shape, "has_chart", False))
    all_text = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))

    assert round(chart_shape.left / 914400, 2) == 1.21
    assert round(chart_shape.top / 914400, 2) == 1.80
    assert round(chart_shape.width / 914400, 2) == 8.15
    assert round(chart_shape.height / 914400, 2) == 4.55
    assert chart_shape.chart.has_legend is True
    assert chart_shape.chart.plots[0].data_labels.font.size == Pt(9)
    assert str(chart_shape.chart.series[0].format.fill.fore_color.rgb) == "C7000B"
    assert "方案A目标阈值 60ms" in all_text
    assert "方案A 2月起低于60ms目标阈值。" in all_text
    side_table = next(shape.table for shape in slide.shapes if getattr(shape, "has_table", False))
    assert side_table.cell(0, 0).text == "指标"
    assert side_table.cell(1, 1).text == "2月起达标"
    threshold_line = next(shape for shape in slide.shapes if shape.name.startswith("HW_THRESHOLD_LINE"))
    threshold_label = next(shape for shape in slide.shapes if shape.name.startswith("HW_THRESHOLD_LABEL"))
    shape_names = [shape.name for shape in slide.shapes]
    assert shape_names.index(threshold_line.name) < shape_names.index(chart_shape.name)
    assert threshold_label.top + threshold_label.height < chart_shape.top
    assert str(threshold_line.line.color.rgb) == "C7000B"
    assert threshold_line.line.width == Pt(0.75)


def test_render_table_cell_tcpr_children_follow_ooxml_order(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir
    from pptx.oxml.ns import qn

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "表格顺序", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "table",
                    "title": "比较表",
                    "table": {
                        "header": ["方案", "可靠性"],
                        "rows": [["方案A", "高"]],
                    },
                }
            ],
        }
    )
    output = render_deck_ir(deck, tmp_path / "tcpr-order.pptx")
    prs = Presentation(str(output))
    table = next(shape.table for shape in prs.slides[0].shapes if getattr(shape, "has_table", False))

    cell = table.cell(0, 0)
    tc_pr = cell._tc.get_or_add_tcPr()
    tags = [child.tag for child in tc_pr]
    border_indexes = [tags.index(qn(f"a:{edge}")) for edge in ("lnL", "lnR", "lnT", "lnB")]
    fill_index = tags.index(qn("a:solidFill"))
    assert border_indexes == sorted(border_indexes)
    assert max(border_indexes) < fill_index, f"ln* must precede fill: {tags}"


def test_render_combo_chart_draws_threshold_lines_on_primary_axis(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "组合图阈值", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "收入增长且利润率改善",
                    "chart": {
                        "kind": "combo",
                        "categories": ["一月", "二月", "三月"],
                        "series": [
                            {"name": "收入", "values": [13, 14, 15], "chart_type": "bar", "axis": "primary", "unit": "万元", "emphasis": True},
                            {"name": "利润率", "values": [0.2, 0.22, 0.25], "chart_type": "line", "axis": "secondary", "unit": "%"},
                        ],
                        "thresholds": [{"value": 12, "label": "收入目标 12 万元"}],
                        "side_conclusion": "收入连续三个月高于12。",
                        "show_data_labels": False,
                    },
                }
            ],
        }
    )
    output = render_deck_ir(deck, tmp_path / "combo-threshold.pptx")
    prs = Presentation(str(output))
    slide = prs.slides[0]

    threshold_lines = [shape for shape in slide.shapes if shape.name.startswith("HW_THRESHOLD_LINE")]
    threshold_labels = [shape for shape in slide.shapes if shape.name.startswith("HW_THRESHOLD_LABEL")]
    assert len(threshold_lines) == 1
    assert len(threshold_labels) == 1
    assert "收入目标" in threshold_labels[0].text_frame.paragraphs[0].text


def test_render_deck_ir_vertical_chart_threshold_labels_do_not_overlap(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "多阈值图", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "性能趋势",
                    "chart": {
                        "kind": "bar",
                        "unit": "ms",
                        "categories": ["1月", "2月", "3月"],
                        "series": [{"name": "方案A", "values": [62, 58, 55], "emphasis": True}],
                        "thresholds": [
                            {"value": 60, "label": "下限 60ms"},
                            {"value": 80, "label": "上限 80ms"},
                        ],
                    },
                }
            ],
        }
    )
    output = render_deck_ir(deck, tmp_path / "multi-threshold.pptx")
    prs = Presentation(str(output))

    labels = [shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_THRESHOLD_LABEL")]
    assert len(labels) == 2
    tops = sorted(label.top / 914400 for label in labels)
    assert len(set(tops)) == 2, f"两个阈值标签共用同一 y 坐标: {tops}"
    assert (tops[1] - tops[0]) == pytest.approx(0.222222 + 0.08, abs=0.02)


def test_render_deck_ir_threshold_label_expands_short_label_with_series_and_unit(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "阈值标签", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "单方案趋势",
                    "chart": {
                        "kind": "line",
                        "unit": "ms",
                        "categories": ["1月", "2月"],
                        "series": [{"name": "方案A", "values": [58, 57]}],
                        "thresholds": [{"value": 57, "label": "目标"}],
                    },
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "threshold-label.pptx")
    prs = Presentation(str(output))
    all_text = "\n".join(shape.text for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False))

    assert "方案A目标阈值 57ms" in all_text


def test_render_deck_ir_threshold_line_stays_inside_plot_area(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "阈值线位置", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "简单趋势图",
                    "chart": {
                        "kind": "line",
                        "unit": "ms",
                        "categories": ["1月", "2月", "3月"],
                        "series": [{"name": "方案A", "values": [58, 56, 54]}],
                        "thresholds": [{"value": 57, "label": "目标"}],
                    },
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "threshold-inside-plot.pptx")
    prs = Presentation(str(output))
    chart_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_chart", False))
    threshold_line = next(shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_THRESHOLD_LINE"))

    assert threshold_line.top > chart_shape.top + int(1.0 * 914400)
    assert threshold_line.top + threshold_line.height < chart_shape.top + chart_shape.height - int(0.45 * 914400)


def test_render_architecture_diagram_as_editable_shapes(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "架构骨架", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "可编辑架构骨架",
                    "nodes": [
                        {"id": "portal", "text": "门户", "type": "primary", "group": "access"},
                        {"id": "api", "text": "接口服务", "type": "secondary", "group": "service"},
                        {"id": "engine", "text": "计算引擎", "type": "emphasis", "group": "service"},
                        {"id": "store", "text": "数据存储", "type": "data", "group": "data"},
                    ],
                    "edges": [
                        {"from": "portal", "to": "api", "label": "请求", "style": "solid", "direction": "forward"},
                        {"from": "api", "to": "engine", "style": "dashed", "direction": "both"},
                        {"from": "engine", "to": "store", "label": "读写", "style": "solid", "direction": "forward"},
                    ],
                    "groups": [
                        {"id": "access", "label": "接入层", "node_ids": ["portal"]},
                        {"id": "service", "label": "服务层", "node_ids": ["api", "engine"]},
                        {"id": "data", "label": "数据层", "node_ids": ["store"]},
                    ],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "architecture.pptx")
    prs = Presentation(str(output))
    slide = prs.slides[0]
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    edges = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_EDGE:")]
    edge_segments = _architecture_edge_segments(slide)
    groups = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_GROUP:")]
    labels = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_EDGE_LABEL:")]

    assert len(nodes) == 4
    assert len(edges) == 3
    assert len(groups) == 3
    assert {shape.text for shape in nodes} == {"门户", "接口服务", "计算引擎", "数据存储"}
    assert {shape.text for shape in groups} == {"接入层", "服务层", "数据层"}
    assert all(shape.text_frame.vertical_anchor == MSO_ANCHOR.TOP for shape in groups)
    assert {str(shape.fill.fore_color.rgb) for shape in nodes} == {"C7000B", "FFFFFF"}
    assert {str(shape.line.color.rgb) for shape in nodes} == {"30B5C5", "FCC800", "C7000B", "61B230"}
    assert all("<p:cxnSp" in shape.element.xml for shape in edges)
    assert "<a:tailEnd" in edges[0].element.xml
    assert any("<a:headEnd" in shape.element.xml for shape in _architecture_edge_segments(slide, edge_index=2))
    assert any("<a:tailEnd" in shape.element.xml for shape in _architecture_edge_segments(slide, edge_index=2))
    assert all("<a:prstDash" not in shape.element.xml for shape in groups)
    assert all(str(shape.fill.fore_color.rgb) == "F5F5F5" for shape in groups)
    assert all(not _overlap(first, second) for index, first in enumerate(nodes) for second in nodes[index + 1 :])
    assert all(_is_orthogonal_connector(segment) for segment in edge_segments)
    assert not any(
        _connector_crosses_unrelated_node(segment, node)
        for segment in edge_segments
        for node in nodes
    )
    assert all(_label_is_near_own_edge(label, slide) for label in labels)

    same_layer_segments = _architecture_edge_segments(slide, edge_index=2)
    assert len(same_layer_segments) == 1
    assert same_layer_segments[0].height == 0


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_render_graphviz_architecture_long_text_stays_inside_editable_node(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir
    from app.rendering.typography import estimate_text_height_in

    long_text = (
        "质量闸门在生成与渲染之间执行结构校验和有限修复，"
        "非法内容不得进入Office输出，所有告警必须保留证据定位"
    )
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "长文字节点", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "Graphviz 按文字内容计算节点尺寸",
                    "nodes": [
                        {"id": "gate", "text": long_text, "type": "emphasis"},
                        {"id": "output", "text": "可编辑Office产物", "type": "data"},
                    ],
                    "edges": [
                        {"from": "gate", "to": "output", "label": "校验通过", "direction": "forward"}
                    ],
                    "groups": [],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "graphviz-long-node.pptx")
    slide = Presentation(str(output)).slides[0]
    gate = next(shape for shape in slide.shapes if shape.name == "HW_ARCH_NODE:gate")
    output_node = next(shape for shape in slide.shapes if shape.name == "HW_ARCH_NODE:output")
    label = next(shape for shape in slide.shapes if shape.name == "HW_ARCH_EDGE_LABEL:1")
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    layout = theme["layouts"]["architecture_diagram"]
    font_size_pt = gate.text_frame.paragraphs[0].runs[0].font.size.pt
    estimated_height_in = estimate_text_height_in(
        gate.text,
        font_size_pt,
        width_in=gate.width / 914400,
        line_spacing=theme["typography"]["line_spacing"],
        baseline_pt=theme["grid"]["baseline_pt"],
        horizontal_margin_in=layout["node_text_margin_in"],
        vertical_margin_in=layout["node_text_margin_in"],
    )

    assert gate.text.replace("\n", "") == long_text
    assert estimated_height_in <= gate.height / 914400
    assert not _overlap(gate, output_node)
    assert not _overlap(label, gate)
    assert not _overlap(label, output_node)
    assert "<p:sp" in gate.element.xml
    assert all("<p:cxnSp" in shape.element.xml for shape in _architecture_edge_segments(slide))


def test_render_architecture_falls_back_with_warning_when_graphviz_is_unavailable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.architecture_graphviz import GraphvizLayoutUnavailable
    import app.rendering.pptx_renderer as renderer

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "fallback", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "Graphviz 不可用仍可渲染",
                    "nodes": [
                        {"id": "portal", "text": "门户", "type": "primary", "group": "access"},
                        {"id": "api", "text": "接口服务", "type": "secondary", "group": "service"},
                        {"id": "engine", "text": "计算引擎", "type": "emphasis", "group": "service"},
                        {"id": "store", "text": "数据存储", "type": "data", "group": "data"},
                    ],
                    "edges": [
                        {"from": "portal", "to": "api", "label": "请求", "direction": "forward"},
                        {
                            "from": "api",
                            "to": "engine",
                            "label": "双向调用",
                            "style": "dashed",
                            "direction": "both",
                        },
                        {"from": "engine", "to": "store", "label": "读写", "direction": "forward"},
                    ],
                    "groups": [
                        {"id": "access", "label": "接入层", "node_ids": ["portal"]},
                        {"id": "service", "label": "服务层", "node_ids": ["api", "engine"]},
                        {"id": "data", "label": "数据层", "node_ids": ["store"]},
                    ],
                }
            ],
        }
    )

    def unavailable(*_args, **_kwargs):
        raise GraphvizLayoutUnavailable("dot executable is missing")

    monkeypatch.setattr(renderer, "layout_architecture_with_graphviz", unavailable)
    with pytest.warns(RuntimeWarning, match="using deterministic fallback: dot executable is missing"):
        output = renderer.render_deck_ir(deck, tmp_path / "architecture-fallback.pptx")

    slide = Presentation(str(output)).slides[0]
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    labels = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_EDGE_LABEL:")]
    edge_segments = _architecture_edge_segments(slide)
    assert {shape.text for shape in nodes} == {"门户", "接口服务", "计算引擎", "数据存储"}
    assert {shape.text for shape in labels} == {"请求", "双向调用", "读写"}
    assert all("<p:cxnSp" in shape.element.xml for shape in edge_segments)
    assert all(_is_orthogonal_connector(shape) for shape in edge_segments)
    assert not any(_connector_crosses_unrelated_node(segment, node) for segment in edge_segments for node in nodes)
    assert all(_label_is_near_own_edge(label, slide) for label in labels)


def test_render_architecture_fallback_expands_ungrouped_technical_nodes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.architecture_graphviz import GraphvizLayoutUnavailable
    from app.rendering.typography import estimate_text_height_in
    import app.rendering.pptx_renderer as renderer

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.2",
            "meta": {"title": "技术架构", "classification": "公开", "theme": "hw-report"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "四层平台贯通业务接口、协议处理、数字基带与射频收发",
                    "nodes": [
                        {"id": "business", "text": "业务接口层\n以太网｜PCIe｜CAN FD", "type": "primary"},
                        {"id": "protocol", "text": "协议软件层\nRRC｜PDCP｜RLC｜MAC", "type": "module"},
                        {
                            "id": "baseband",
                            "text": "数字基带层\n调制解调｜信道编码｜MIMO 检测",
                            "type": "emphasis",
                        },
                        {"id": "rf", "text": "射频收发层\n多频段收发器｜天线切换单元", "type": "data"},
                    ],
                    "edges": [
                        {"from": "business", "to": "protocol", "label": "数据汇聚 / 分发", "direction": "both"},
                        {"from": "protocol", "to": "baseband", "label": "协议处理", "direction": "both"},
                        {"from": "baseband", "to": "rf", "label": "物理层处理", "direction": "both"},
                    ],
                    "groups": [],
                }
            ],
        }
    )

    def unavailable(*_args, **_kwargs):
        raise GraphvizLayoutUnavailable("dot executable is missing")

    monkeypatch.setattr(renderer, "layout_architecture_with_graphviz", unavailable)
    with pytest.warns(RuntimeWarning, match="using deterministic fallback"):
        output = renderer.render_deck_ir(deck, tmp_path / "architecture-technical-fallback.pptx")

    slide = Presentation(str(output)).slides[0]
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw-report.json").read_text(encoding="utf-8"))
    layout = theme["layouts"]["architecture_diagram"]

    assert len(nodes) == 4
    assert all(shape.width / 914400 >= layout["node_width_in"] * 1.34 for shape in nodes)
    assert all(shape.height / 914400 >= layout["node_height_in"] * 1.39 for shape in nodes)
    assert not any(_overlap(first, second) for index, first in enumerate(nodes) for second in nodes[index + 1 :])
    for shape in nodes:
        font_size_pt = shape.text_frame.paragraphs[0].runs[0].font.size.pt
        estimated_height = estimate_text_height_in(
            shape.text,
            font_size_pt,
            width_in=shape.width / 914400,
            line_spacing=theme["typography"]["line_spacing"],
            baseline_pt=theme["grid"]["baseline_pt"],
            horizontal_margin_in=layout["node_text_margin_in"],
            vertical_margin_in=layout["node_text_margin_in"],
        )
        assert estimated_height <= shape.height / 914400


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_render_architecture_node_colors_follow_theme_type_mapping_and_default(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "语义配色", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "type 决定节点颜色",
                    "nodes": [
                        {"id": "job", "text": "Job", "type": "job"},
                        {"id": "module", "text": "Module", "type": "module"},
                        {"id": "data", "text": "Data", "type": "data"},
                        {"id": "unknown", "text": "Unknown", "type": "unregistered_type"},
                    ],
                    "edges": [
                        {"from": "job", "to": "module"},
                        {"from": "module", "to": "data"},
                        {"from": "unknown", "to": "module"},
                    ],
                    "groups": [],
                }
            ],
        }
    )
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    color_map = theme["layouts"]["architecture_diagram"]["node_type_colors"]

    slide = Presentation(str(render_deck_ir(deck, tmp_path / "semantic-colors.pptx"))).slides[0]
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    # 新视觉契约：普通节点白底（background），类型语义色体现在边框上
    borders = {shape.name.removeprefix("HW_ARCH_NODE:"): str(shape.line.color.rgb) for shape in nodes}
    fills = {shape.name.removeprefix("HW_ARCH_NODE:"): str(shape.fill.fore_color.rgb) for shape in nodes}

    assert borders == {
        "job": theme["colors"][color_map["job"]].removeprefix("#"),
        "module": theme["colors"][color_map["module"]].removeprefix("#"),
        "data": theme["colors"][color_map["data"]].removeprefix("#"),
        "unknown": theme["colors"][color_map["default"]].removeprefix("#"),
    }
    assert set(fills.values()) == {theme["colors"]["background"].removeprefix("#")}


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_render_composite_table_and_architecture_as_editable_region_bound_shapes(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "组合页", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "方案结论与数据流可以在一页联读",
                    "regions": [
                        {
                            "slot": "left",
                            "component": {
                                "layout": "table",
                                "title": "方案对比",
                                "table": {
                                    "header": ["方案", "时延", "结论"],
                                    "rows": [["方案A", "57ms", "推荐"], ["方案B", "72ms", "备选"]],
                                    "conclusion_col": 2,
                                },
                            },
                        },
                        {
                            "slot": "right",
                            "component": {
                                "layout": "architecture_diagram",
                                "title": "数据流骨架",
                                "nodes": [
                                    {"id": "input", "text": "输入材料", "type": "primary"},
                                    {"id": "gate", "text": "质量闸门", "type": "emphasis"},
                                    {"id": "output", "text": "可编辑输出", "type": "data"},
                                ],
                                "edges": [
                                    {"from": "input", "to": "gate", "label": "校验"},
                                    {"from": "gate", "to": "output", "label": "通过"},
                                ],
                                "groups": [],
                            },
                        },
                    ],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "composite-table-architecture.pptx")
    slide = Presentation(str(output)).slides[0]
    table_shape = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    edges = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_EDGE:")]
    region_titles = [
        shape for shape in slide.shapes if shape.name.startswith("HW_RENDERED_TEXT:COMPOSITE_REGION_TITLE:")
    ]
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    composite = theme["layouts"]["composite"]
    column_width = (
        theme["slide"]["width_in"] - theme["slide"]["margin_left_in"] - theme["slide"]["margin_right_in"]
    ) / theme["grid"]["columns"]
    left_box = (
        composite["content"]["left_in"],
        composite["content"]["left_in"] + composite["left_span_columns"] * column_width,
    )
    right_left = left_box[1] + composite["gap_columns"] * column_width
    right_box = (right_left, right_left + composite["right_span_columns"] * column_width)

    assert len(region_titles) == 2
    assert table_shape.left / 914400 >= left_box[0] - 0.01
    assert (table_shape.left + table_shape.width) / 914400 <= left_box[1] + 0.01
    assert len(nodes) == 3
    assert edges
    architecture_shapes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_")]
    assert all(shape.left / 914400 >= right_box[0] - 0.01 for shape in architecture_shapes)
    assert all((shape.left + shape.width) / 914400 <= right_box[1] + 0.01 for shape in architecture_shapes)
    assert (table_shape.top + table_shape.height) / 914400 <= composite["content"]["top_in"] + composite["content"]["height_in"]
    assert any(shape.has_table for shape in slide.shapes)
    assert all("<p:sp" in shape.element.xml for shape in nodes)
    assert all("<p:cxnSp" in shape.element.xml for shape in edges)


def test_render_composite_title_bullets_and_cards_reuses_component_shapes_without_grid_warnings(
    tmp_path: Path,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "文字与卡片组合", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "结论与依据可以分区联读",
                    "regions": [
                        {
                            "slot": "left",
                            "component": {
                                "layout": "title_bullets",
                                "title": "核心结论",
                                "bullets": [
                                    {"text": "组合页只负责区域编排", "level": 1},
                                    {"text": "组件仍保持原有语义", "level": 1},
                                ],
                            },
                        },
                        {
                            "slot": "right",
                            "component": {
                                "layout": "cards",
                                "title": "验证依据",
                                "cards": [
                                    {"title": "契约", "desc": "递归校验"},
                                    {"title": "产物", "desc": "原生可编辑"},
                                ],
                            },
                        },
                    ],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "composite-bullets-cards.pptx")
    slide = Presentation(str(output)).slides[0]
    body_text = [shape.text for shape in slide.shapes if shape.name == "HW_RENDERED_TEXT:BODY"]
    cards = [shape for shape in slide.shapes if shape.name == "HW_RENDERED_TEXT:CARD"]
    report = check_pptx(output, classification="公开")

    assert body_text == ["• 组合页只负责区域编排", "• 组件仍保持原有语义"]
    assert len(cards) == 2
    assert {shape.text.replace("\x0b", "\n") for shape in cards} == {
        "契约\n递归校验",
        "产物\n原生可编辑",
    }
    assert not {"HW-W06", "HW-W07"} & {item.code for item in report.items}


def test_v17_single_component_composite_keeps_v18_single_component_geometry(tmp_path: Path) -> None:
    import copy

    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    legacy_payload = {
        "ir_type": "deck",
        "ir_version": "1.7",
        "meta": {"title": "单块兼容", "classification": "公开", "theme": "hw_v1"},
        "slides": [
            {
                "layout": "composite",
                "title": "单组件区域保持既有几何",
                "regions": [
                    {
                        "slot": "left",
                        "component": {
                            "layout": "title_bullets",
                            "title": "左栏",
                            "bullets": [{"text": "原有单块区域", "level": 1}],
                        },
                    },
                    {
                        "slot": "right",
                        "component": {
                            "layout": "cards",
                            "title": "右栏",
                            "cards": [{"title": "契约", "desc": "兼容"}, {"title": "渲染", "desc": "不变"}],
                        },
                    },
                ],
            }
        ],
    }
    current_payload = copy.deepcopy(legacy_payload)
    current_payload["ir_version"] = "1.8"
    for region in current_payload["slides"][0]["regions"]:
        region["components"] = [region.pop("component")]

    legacy = DeckIR.model_validate(legacy_payload)
    current = DeckIR.model_validate(current_payload)
    legacy_slide = Presentation(str(render_deck_ir(legacy, tmp_path / "legacy-single.pptx"))).slides[0]
    current_slide = Presentation(str(render_deck_ir(current, tmp_path / "current-single.pptx"))).slides[0]

    def geometry(slide) -> list[tuple[str, str, int, int, int, int]]:
        return [
            (
                shape.name,
                getattr(shape, "text", ""),
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
            )
            for shape in slide.shapes
        ]

    assert geometry(legacy_slide) == geometry(current_slide)


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_render_composite_stacks_table_architecture_and_bullets_inside_left_region(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "堆叠组合页", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "左栏可按证据链堆叠多个组件",
                    "regions": [
                        {
                            "slot": "left",
                            "components": [
                                {
                                    "layout": "table",
                                    "title": "阶段状态",
                                    "table": {"header": ["阶段", "状态"], "rows": [["校验", "通过"]]},
                                },
                                {
                                    "layout": "architecture_diagram",
                                    "title": "质量链路",
                                    "nodes": [
                                        {"id": "input", "text": "输入", "type": "primary"},
                                        {"id": "output", "text": "产物", "type": "data"},
                                    ],
                                    "edges": [{"from": "input", "to": "output"}],
                                    "groups": [],
                                },
                                {
                                    "layout": "title_bullets",
                                    "title": "结论",
                                    "bullets": [{"text": "按顺序保留证据链", "level": 1}],
                                },
                            ],
                        },
                        {
                            "slot": "right",
                            "components": [
                                {
                                    "layout": "cards",
                                    "title": "交付依据",
                                    "cards": [{"title": "契约", "desc": "Schema"}, {"title": "产物", "desc": "可编辑"}],
                                }
                            ],
                        },
                    ],
                }
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "composite-stacked.pptx")
    slide = Presentation(str(output)).slides[0]
    blocks = sorted(
        [shape for shape in slide.shapes if shape.name.startswith("HW_COMPOSITE_BLOCK:left:")],
        key=lambda shape: shape.top,
    )
    table = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
    nodes = [shape for shape in slide.shapes if shape.name.startswith("HW_ARCH_NODE:")]
    bullet = next(shape for shape in slide.shapes if shape.name == "HW_RENDERED_TEXT:BODY")
    report = check_pptx(output, classification="公开")

    assert len(blocks) == 3
    assert all(blocks[index].top + blocks[index].height <= blocks[index + 1].top for index in range(2))
    tolerance_emu = 2
    assert blocks[0].top <= table.top and table.top + table.height <= blocks[0].top + blocks[0].height + tolerance_emu
    assert nodes and all(blocks[1].top <= node.top and node.top + node.height <= blocks[1].top + blocks[1].height + tolerance_emu for node in nodes)
    assert blocks[2].top <= bullet.top and bullet.top + bullet.height <= blocks[2].top + blocks[2].height + tolerance_emu
    assert not {"HW-W06", "HW-W07"} & {item.code for item in report.items}


def _composite_stack_deck(left_components: list[dict]) -> dict:
    return {
        "ir_type": "deck",
        "ir_version": "1.9",
        "meta": {"title": "栏内留白", "classification": "公开", "theme": "hw_v1"},
        "slides": [
            {
                "layout": "composite",
                "title": "栏内仅回收多余留白",
                "regions": [
                    {"slot": "left", "components": left_components},
                    {
                        "slot": "right",
                        "components": [
                            {
                                "layout": "title_bullets",
                                "title": "对照",
                                "bullets": [{"text": "右栏保持正常间距。", "level": 1}],
                            }
                        ],
                    },
                ],
            }
        ],
    }


def _left_composite_blocks(slide) -> list:
    return sorted(
        [shape for shape in slide.shapes if shape.name.startswith("HW_COMPOSITE_BLOCK:left:")],
        key=lambda shape: shape.top,
    )


def test_composite_stack_preserves_default_gap_when_column_fits(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        _composite_stack_deck(
            [
                {"layout": "table", "title": "参数", "table": {"header": ["项", "值"], "rows": [["范围", "正常"]]}},
                {"layout": "title_bullets", "title": "结论", "bullets": [{"text": "正常内容不触发间距回收。", "level": 1}]},
            ]
        )
    )
    output = render_deck_ir(deck, tmp_path / "composite-default-gap.pptx")
    slide = Presentation(str(output)).slides[0]
    blocks = _left_composite_blocks(slide)
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    layout = theme["layouts"]["composite"]
    baseline_step = theme["grid"]["baseline_pt"] / 72
    first_bottom = (blocks[0].top + blocks[0].height) / 914400
    expected_second_top = round((first_bottom + layout["stack_gap_in"]) / baseline_step) * baseline_step
    report = check_pptx(output, classification="公开")

    assert len(blocks) == 2
    assert blocks[1].top / 914400 == pytest.approx(expected_second_top)
    assert not any(item.code == "HW-W03" and "组合页该栏内容过多" in item.message for item in report.items)


def test_composite_stack_reclaims_gaps_for_slight_overflow_without_changing_content_or_fonts(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    bullet_text = "完整说明不截断，也不缩小字号。"
    deck = DeckIR.model_validate(
        _composite_stack_deck(
            [
                {
                    "layout": "table",
                    "title": "七行参数",
                    "table": {"header": ["项", "值"], "rows": [[f"参数{index}", "已保留"] for index in range(1, 8)]},
                },
                {"layout": "title_bullets", "title": "结论", "bullets": [{"text": bullet_text, "level": 1}]},
            ]
        )
    )
    output = render_deck_ir(deck, tmp_path / "composite-slight-overflow.pptx")
    presentation = Presentation(str(output))
    slide = presentation.slides[0]
    blocks = _left_composite_blocks(slide)
    table_shape = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
    body = next(shape for shape in slide.shapes if getattr(shape, "text", "") == f"• {bullet_text}")
    block_title = next(shape for shape in slide.shapes if shape.name == "HW_RENDERED_TEXT:COMPOSITE_BLOCK_TITLE:left:2")
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    layout = theme["layouts"]["composite"]
    actual_gap = (blocks[1].top - blocks[0].top - blocks[0].height) / 914400
    content_bottom = layout["content"]["top_in"] + layout["content"]["height_in"]
    report = check_pptx(output, classification="公开")

    assert len(presentation.slides) == 1
    assert layout["stack_min_gap_in"] <= actual_gap < layout["stack_gap_in"]
    assert (blocks[-1].top + blocks[-1].height) / 914400 <= content_bottom + 0.01
    assert "参数7" in table_shape.table.cell(7, 0).text
    assert body.text == f"• {bullet_text}"
    assert table_shape.table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size == Pt(11)
    assert table_shape.table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size == Pt(10)
    assert block_title.text_frame.paragraphs[0].runs[0].font.size == Pt(18)
    assert body.text_frame.paragraphs[0].runs[0].font.size == Pt(16)
    assert not {"HW-W03", "HW-W06", "HW-W07"} & {item.code for item in report.items}

    body.left += Inches(0.2)
    presentation.save(output)
    assert any(item.code == "HW-W06" for item in check_pptx(output, classification="公开").items)


def test_composite_stack_keeps_minimum_gap_and_warns_when_content_is_far_too_tall(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    bullet_text = "保留完整事实" * 40
    deck = DeckIR.model_validate(
        _composite_stack_deck(
            [
                {
                    "layout": "table",
                    "title": "七行参数",
                    "table": {"header": ["项", "值"], "rows": [[f"参数{index}", "已保留"] for index in range(1, 8)]},
                },
                {"layout": "title_bullets", "title": "完整说明", "bullets": [{"text": bullet_text, "level": 1}]},
            ]
        )
    )
    output = render_deck_ir(deck, tmp_path / "composite-large-overflow.pptx")
    presentation = Presentation(str(output))
    slide = presentation.slides[0]
    blocks = _left_composite_blocks(slide)
    body = next(shape for shape in slide.shapes if getattr(shape, "text", "") == f"• {bullet_text}")
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    layout = theme["layouts"]["composite"]
    actual_gap = (blocks[1].top - blocks[0].top - blocks[0].height) / 914400
    content_bottom = layout["content"]["top_in"] + layout["content"]["height_in"]
    report = check_pptx(output, classification="公开")

    assert len(presentation.slides) == 1
    assert actual_gap >= layout["stack_min_gap_in"]
    assert actual_gap < layout["stack_gap_in"]
    assert (blocks[-1].top + blocks[-1].height) / 914400 > content_bottom
    assert body.text == f"• {bullet_text}"
    assert body.text_frame.paragraphs[0].runs[0].font.size == Pt(16)
    assert any(item.code == "HW-W03" and "left 栏" in item.message for item in report.items)


def _overlap(first, second) -> bool:
    return (
        first.left < second.left + second.width
        and second.left < first.left + first.width
        and first.top < second.top + second.height
        and second.top < first.top + first.height
    )


def _architecture_edge_segments(slide, *, edge_index: int | None = None) -> list:
    prefixes = (
        "HW_ARCH_EDGE:",
        "HW_ARCH_EDGE_SEGMENT:",
    )
    segments = [shape for shape in slide.shapes if shape.name.startswith(prefixes)]
    if edge_index is None:
        return segments
    return [shape for shape in segments if _architecture_edge_index(shape.name) == edge_index]


def _architecture_edge_index(name: str) -> int:
    parts = name.split(":")
    return int(parts[1])


def _architecture_edge_node_ids(name: str) -> set[str]:
    route = name.rsplit(":", 1)[-1]
    return set(route.split("->"))


def _is_orthogonal_connector(shape) -> bool:
    return shape.width == 0 or shape.height == 0


def _connector_crosses_unrelated_node(connector, node) -> bool:
    node_id = node.name.split(":", 1)[1]
    if node_id in _architecture_edge_node_ids(connector.name):
        return False
    if connector.height == 0:
        y = connector.top
        return (
            node.top < y < node.top + node.height
            and connector.left < node.left + node.width
            and node.left < connector.left + connector.width
        )
    if connector.width == 0:
        x = connector.left
        return (
            node.left < x < node.left + node.width
            and connector.top < node.top + node.height
            and node.top < connector.top + connector.height
        )
    return True


def _label_is_near_own_edge(label, slide) -> bool:
    edge_index = int(label.name.rsplit(":", 1)[1])
    center_x = label.left + label.width / 2
    center_y = label.top + label.height / 2
    theme = json.loads((ROOT / "backend/app/rendering/themes/hw_theme.json").read_text(encoding="utf-8"))
    max_distance = Inches(theme["layouts"]["architecture_diagram"]["edge_label_max_distance_in"])
    return any(
        _point_to_connector_distance(center_x, center_y, segment) <= max_distance
        for segment in _architecture_edge_segments(slide, edge_index=edge_index)
    )


def _point_to_connector_distance(x: float, y: float, connector) -> float:
    if connector.height == 0:
        nearest_x = min(max(x, connector.left), connector.left + connector.width)
        return ((x - nearest_x) ** 2 + (y - connector.top) ** 2) ** 0.5
    if connector.width == 0:
        nearest_y = min(max(y, connector.top), connector.top + connector.height)
        return ((x - connector.left) ** 2 + (y - nearest_y) ** 2) ** 0.5
    return float("inf")


def _cell_fill_hex(cell) -> str:
    rgb = cell.fill.fore_color.rgb
    assert isinstance(rgb, RGBColor)
    return str(rgb)


def test_render_scatter_chart_with_default_data_labels_does_not_crash(tmp_path: Path) -> None:
    import zipfile

    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR(
        ir_type="deck",
        ir_version="2.1",
        meta={"title": "T", "classification": "HUAWEI CONFIDENTIAL"},
        slides=[
            {
                "layout": "chart",
                "title": "scatter",
                "chart": {
                    "kind": "scatter",
                    "categories": [],
                    "series": [{"name": "s", "values": [1.0, 2.0], "x_values": [0.0, 1.0]}],
                },
            }
        ],
    )
    output = render_deck_ir(deck, tmp_path / "scatter.pptx")

    with zipfile.ZipFile(output) as package:
        chart_xml = next(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if name.startswith("ppt/charts/") and name.endswith(".xml")
        )
    assert "dLbls" in chart_xml


def test_deck_ir_rejects_body_span_crossing_row_group_label(tmp_path: Path) -> None:
    from app.ir.validation import validate_deck_ir

    payload = {
        "ir_type": "deck",
        "ir_version": "2.1",
        "meta": {"title": "T", "classification": "HUAWEI CONFIDENTIAL"},
        "slides": [
            {
                "layout": "table",
                "title": "t",
                "table": {
                    "header": ["H1", "H2"],
                    "rows": [["a", "b"], ["", "d"], ["", "f"]],
                    "row_groups": [{"start_row": 1, "span": 1, "label": "G"}],
                    "cell_spans": [{"area": "body", "row": 0, "col": 0, "rowspan": 3, "colspan": 1}],
                },
            }
        ],
    }
    result = validate_deck_ir(payload)
    assert not result.ok
    assert any(item.code == "D005" for item in result.errors)


def test_deck_ir_accepts_span_on_row_group_own_row(tmp_path: Path) -> None:
    from app.rendering.pptx_renderer import render_deck_ir
    from app.ir.validation import validate_deck_ir

    payload = {
        "ir_type": "deck",
        "ir_version": "2.1",
        "meta": {"title": "T", "classification": "HUAWEI CONFIDENTIAL"},
        "slides": [
            {
                "layout": "table",
                "title": "t",
                "table": {
                    "header": ["H1", "H2"],
                    "rows": [["", "b"], ["c", "d"]],
                    "row_groups": [{"start_row": 0, "span": 1, "label": "G"}],
                    "cell_spans": [{"area": "body", "row": 0, "col": 0, "rowspan": 1, "colspan": 1}],
                },
            }
        ],
    }
    result = validate_deck_ir(payload)
    assert result.ok
    if result.value is not None:
        render_deck_ir(result.value, tmp_path / "table.pptx")
