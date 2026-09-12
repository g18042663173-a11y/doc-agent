"""Extract a presentation's reusable layout skeleton without copying its style."""

from __future__ import annotations

from dataclasses import dataclass, replace
import json
from pathlib import Path
from typing import Literal

from pptx import Presentation

from app.ir.document_ir import SlideSummary
from app.parsers.pptx_parser import parse_pptx


EMU_PER_INCH = 914400
BlockKind = Literal["title", "subtitle", "text", "placeholder", "table", "chart", "image"]
SlideKind = Literal[
    "title_slide",
    "blank",
    "table",
    "chart",
    "image",
    "composite",
    "two_column",
    "multi_column",
    "title_and_content",
    "title_and_bullets",
    "content",
]


class PptTemplateStructureError(RuntimeError):
    """Raised when a PPT cannot produce a consistent parser/shape structure."""


@dataclass(frozen=True)
class TemplateBlock:
    kind: BlockKind
    left_ratio: float
    top_ratio: float
    width_ratio: float
    height_ratio: float
    text_preview: str | None = None
    rows: int | None = None
    columns: int | None = None
    placeholder_type: str | None = None

    def to_dict(self) -> dict:
        data = {
            "kind": self.kind,
            "left_ratio": self.left_ratio,
            "top_ratio": self.top_ratio,
            "width_ratio": self.width_ratio,
            "height_ratio": self.height_ratio,
        }
        if self.text_preview:
            data["text_preview"] = self.text_preview
        if self.rows is not None:
            data["rows"] = self.rows
        if self.columns is not None:
            data["columns"] = self.columns
        if self.placeholder_type:
            data["placeholder_type"] = self.placeholder_type
        return data


@dataclass(frozen=True)
class TemplateSlide:
    index: int
    master_layout: str | None
    detected_layout: SlideKind
    columns: int
    title: str | None
    blocks: tuple[TemplateBlock, ...]
    parser_body_count: int
    parser_table_count: int
    notes_present: bool
    parser_warnings: tuple[str, ...]

    def to_dict(self) -> dict:
        return {
            "index": self.index,
            "master_layout": self.master_layout,
            "detected_layout": self.detected_layout,
            "columns": self.columns,
            "title": self.title,
            "blocks": [block.to_dict() for block in self.blocks],
            "parser_summary": {
                "body_count": self.parser_body_count,
                "table_count": self.parser_table_count,
                "notes_present": self.notes_present,
                "warnings": list(self.parser_warnings),
            },
        }


@dataclass(frozen=True)
class PptTemplateStructure:
    filename: str
    width_in: float
    height_in: float
    slides: tuple[TemplateSlide, ...]
    parser_warnings: tuple[str, ...]

    def to_dict(self) -> dict:
        return {
            "filename": self.filename,
            "slide_size_in": {"width": self.width_in, "height": self.height_in},
            "slides": [slide.to_dict() for slide in self.slides],
            "parser_warnings": list(self.parser_warnings),
        }

    def to_markdown(self) -> str:
        lines = [
            f"# PPT 模板结构：{self.filename}",
            "",
            f"页面尺寸：{self.width_in:.2f} x {self.height_in:.2f} in",
            "",
        ]
        for slide in self.slides:
            title = f"，标题：{slide.title}" if slide.title else ""
            master = slide.master_layout or "未命名母版"
            lines.append(
                f"## 第 {slide.index} 页：{slide.detected_layout}（母版：{master}，{slide.columns} 栏）{title}"
            )
            for block in slide.blocks:
                details = _block_description(block)
                lines.append(f"- {block.kind}：{details}")
            if slide.notes_present:
                lines.append("- notes：存在演讲备注")
            if slide.parser_warnings:
                lines.append(f"- parser warnings：{'；'.join(slide.parser_warnings)}")
            lines.append("")
        if self.parser_warnings:
            lines.extend(["## 全局解析提示", ""])
            lines.extend(f"- {warning}" for warning in self.parser_warnings)
            lines.append("")
        return "\n".join(lines)


def extract_ppt_template_structure(path: Path) -> PptTemplateStructure:
    """Return a structural description by reusing the existing PPT parser plus shape geometry."""

    document = parse_pptx(path)
    presentation = Presentation(str(path))
    summaries = document.content.slides
    if len(summaries) != len(presentation.slides):
        raise PptTemplateStructureError("PPT parser summary count does not match presentation slide count")
    width = int(presentation.slide_width)
    height = int(presentation.slide_height)
    slides = tuple(
        _extract_slide(index, slide, summary, width, height)
        for index, (slide, summary) in enumerate(zip(presentation.slides, summaries), start=1)
    )
    return PptTemplateStructure(
        filename=path.name,
        width_in=round(width / EMU_PER_INCH, 2),
        height_in=round(height / EMU_PER_INCH, 2),
        slides=slides,
        parser_warnings=tuple(document.warnings),
    )


def write_ppt_template_structure(structure: PptTemplateStructure, output_path: Path) -> Path:
    """Write a JSON or Markdown structure report based on the output suffix."""

    suffix = output_path.suffix.lower()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".json":
        output_path.write_text(
            json.dumps(structure.to_dict(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    elif suffix == ".md":
        output_path.write_text(structure.to_markdown(), encoding="utf-8")
    else:
        raise PptTemplateStructureError("output path must end with .json or .md")
    return output_path


def _extract_slide(index: int, slide, summary: SlideSummary, width: int, height: int) -> TemplateSlide:
    title_shape = getattr(slide.shapes, "title", None)
    title_shape_id = getattr(title_shape, "shape_id", None)
    raw_blocks = tuple(
        block
        for shape in _iter_shapes(slide.shapes)
        if (block := _shape_block(shape, title_shape_id, width, height)) is not None
    )
    blocks = tuple(block for block in raw_blocks if not _is_decorative_or_footer(block))
    blocks = _promote_top_title(blocks, summary.title)
    columns = _column_count(blocks)
    return TemplateSlide(
        index=index,
        master_layout=summary.layout_name,
        detected_layout=_detect_layout(index, summary, blocks, columns),
        columns=columns,
        title=summary.title,
        blocks=blocks,
        parser_body_count=len(summary.bodies),
        parser_table_count=len(summary.tables),
        notes_present=summary.notes is not None,
        parser_warnings=tuple(summary.shape_warnings),
    )


def _iter_shapes(shapes):
    for shape in shapes:
        children = getattr(shape, "shapes", None)
        if children is not None:
            yield from _iter_shapes(children)
        else:
            yield shape


def _shape_block(shape, title_shape_id: int | None, width: int, height: int) -> TemplateBlock | None:
    left_ratio = round(int(shape.left) / width, 4)
    top_ratio = round(int(shape.top) / height, 4)
    width_ratio = round(int(shape.width) / width, 4)
    height_ratio = round(int(shape.height) / height, 4)
    if getattr(shape, "has_table", False):
        return TemplateBlock(
            kind="table",
            left_ratio=left_ratio,
            top_ratio=top_ratio,
            width_ratio=width_ratio,
            height_ratio=height_ratio,
            rows=len(shape.table.rows),
            columns=len(shape.table.columns),
        )
    if getattr(shape, "has_chart", False):
        return TemplateBlock(
            kind="chart",
            left_ratio=left_ratio,
            top_ratio=top_ratio,
            width_ratio=width_ratio,
            height_ratio=height_ratio,
        )
    if "PICTURE" in str(getattr(shape, "shape_type", "")):
        return TemplateBlock(
            kind="image",
            left_ratio=left_ratio,
            top_ratio=top_ratio,
            width_ratio=width_ratio,
            height_ratio=height_ratio,
        )
    if not getattr(shape, "has_text_frame", False) and not getattr(shape, "is_placeholder", False):
        return None
    text = getattr(shape, "text", "").strip()
    placeholder_type = _placeholder_type(shape)
    if getattr(shape, "shape_id", None) == title_shape_id:
        kind: BlockKind = "title"
    elif placeholder_type and "SUBTITLE" in placeholder_type:
        kind = "subtitle"
    elif text:
        kind = "text"
    else:
        kind = "placeholder"
    return TemplateBlock(
        kind=kind,
        left_ratio=left_ratio,
        top_ratio=top_ratio,
        width_ratio=width_ratio,
        height_ratio=height_ratio,
        text_preview=_preview(text) if text else None,
        placeholder_type=placeholder_type,
    )


def _placeholder_type(shape) -> str | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return str(shape.placeholder_format.type).split(" ", 1)[0]
    except (AttributeError, ValueError):
        return "placeholder"


def _promote_top_title(blocks: tuple[TemplateBlock, ...], parser_title: str | None) -> tuple[TemplateBlock, ...]:
    if any(block.kind == "title" for block in blocks):
        return blocks
    candidates = [
        (index, block)
        for index, block in enumerate(blocks)
        if block.kind == "text"
        and block.top_ratio <= 0.32
        and block.width_ratio >= 0.4
        and (parser_title is None or block.text_preview == _preview(parser_title))
    ]
    if not candidates:
        return blocks
    index, block = min(candidates, key=lambda candidate: candidate[1].top_ratio)
    promoted = list(blocks)
    promoted[index] = replace(block, kind="title")
    return tuple(promoted)


def _column_count(blocks: tuple[TemplateBlock, ...]) -> int:
    content = [block for block in blocks if block.kind not in {"title", "subtitle"}]
    if len(content) < 2:
        return 1
    centers = sorted(block.left_ratio + block.width_ratio / 2 for block in content)
    column_breaks = sum(
        current - previous > 0.15
        for previous, current in zip(centers, centers[1:])
    )
    return min(column_breaks + 1, 4)


def _detect_layout(
    index: int,
    summary: SlideSummary,
    blocks: tuple[TemplateBlock, ...],
    columns: int,
) -> SlideKind:
    if not blocks:
        return "blank"
    layout_name = (summary.layout_name or "").casefold()
    kinds = {block.kind for block in blocks}
    content = [block for block in blocks if block.kind not in {"title", "subtitle"}]
    has_title = "title" in kinds
    if any(token in layout_name for token in ("title slide", "标题幻灯片", "封面")):
        return "title_slide"
    if has_title and len(content) <= 1 and "subtitle" in kinds:
        return "title_slide"
    if index == 1 and has_title and not {"table", "chart"} & kinds:
        return "title_slide"
    if "table" in kinds and len(content) <= 2:
        return "table"
    if "table" in kinds and len(content) > 2:
        return "composite"
    if "chart" in kinds:
        return "chart"
    if "image" in kinds and len(content) == 1:
        return "image"
    if columns >= 3:
        return "multi_column"
    if columns == 2:
        return "two_column"
    text_count = sum(block.kind in {"text", "placeholder"} for block in content)
    if has_title and text_count >= 2:
        return "title_and_bullets"
    if has_title and content:
        return "title_and_content"
    return "content"


def _is_decorative_or_footer(block: TemplateBlock) -> bool:
    if block.top_ratio >= 0.9 and block.height_ratio <= 0.08:
        return True
    return block.kind == "placeholder" and not block.text_preview and block.height_ratio <= 0.025


def _preview(text: str, limit: int = 80) -> str:
    normalized = " ".join(text.split())
    return normalized if len(normalized) <= limit else normalized[: limit - 1] + "…"


def _block_description(block: TemplateBlock) -> str:
    geometry = (
        f"x={block.left_ratio:.2f}, y={block.top_ratio:.2f}, "
        f"w={block.width_ratio:.2f}, h={block.height_ratio:.2f}"
    )
    details = [geometry]
    if block.rows is not None and block.columns is not None:
        details.append(f"{block.rows}x{block.columns}")
    if block.text_preview:
        details.append(block.text_preview)
    if block.placeholder_type:
        details.append(f"placeholder={block.placeholder_type}")
    return "；".join(details)
