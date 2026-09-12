from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from pptx import Presentation

from engine.extract.draft import write_skeleton_draft
from engine.extract.recover import diagram_frames, diagram_texts, recover_source, write_recoveries, xfrm_box
from engine.security.office_package import OfficePackageError, preflight_office_package
from engine.shared.common import RdwError, SKILL_ROOT, sha256, write_json, write_text
from engine.native.package import inspect_package

BLOCKED_SKIP_REASONS = frozenset()


def skip_report(*, pages: list[dict]) -> dict:
    return {"format": "rdw_skip_report", "version": "1.0", "pages": pages}


def _write_slot_candidates(pack: Path, pages: list[tuple[str, list[dict]]]) -> None:
    payload = {"format": "rdw_slot_candidates", "version": "1.0", "pages": []}
    for page_id, elements in pages:
        candidates = []
        for item in elements:
            text = (item.get("text") or "").strip()
            if not text and item.get("kind") != "chart":
                continue
            kind = item.get("kind") or "shape"
            if kind not in {"table_cell", "smartart_node", "chart"}:
                kind = "shape"
            candidates.append(
                {
                    "shape_ref": item["shape_ref"],
                    "kind": kind,
                    "chars": len(text),
                    "bbox": item.get("bbox") or {},
                }
            )
        payload["pages"].append({"page_id": page_id, "candidates": candidates})
    write_json(pack / "slot_candidates.json", payload)


def extract_source(source: Path, workdir: Path) -> dict:
    source = source.resolve()
    workdir = workdir.resolve()
    if workdir.exists() and any(workdir.iterdir()):
        raise RdwError("RD-E002", "out", "工作目录必须为空。", "使用新的空目录，避免覆盖或混入旧源文。")
    try:
        report = preflight_office_package(source, purpose="source")
    except OfficePackageError as exc:
        raise RdwError("RD-E002", exc.loc, exc.message, "换用安全、未加密且无异常活动内容的 PPTX。") from exc
    if report.office_kind != "pptx":
        raise RdwError("RD-E002", "source", "extract 只接受 PPTX。", "提供可编辑的 .pptx 源件。")
    try:
        presentation = Presentation(source)
    except Exception as exc:
        raise RdwError("RD-E002", "source", f"python-pptx 无法读取源件: {exc}", "确认文件完整且不是截图封装或损坏包。") from exc
    if not 1 <= len(presentation.slides) <= 200:
        raise RdwError("RD-E002", "source", "源件页数必须在 1-200 页。", "拆分后重试。")
    workdir.mkdir(parents=True, exist_ok=True)
    temporary_source = workdir / "source.pptx"
    shutil.copy2(source, temporary_source)
    reference = workdir / "source_reference.pptx"
    shutil.copy2(source, reference)
    reference.chmod(0o444)
    recoveries = recover_source(temporary_source)
    presentation = Presentation(temporary_source)
    pack = workdir / "extract_pack"
    pages_dir = pack / "pages"
    pages_dir.mkdir(parents=True)
    inventory = inspect_package(temporary_source)
    write_json(pack / "native_inventory.json", inventory)
    skipped_pages: list[dict] = []
    page_elements: list[tuple[str, list[dict]]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        page_id = f"p{index:02d}"
        elements = _slide_elements(slide, presentation.slide_width, presentation.slide_height)
        native = inventory["pages"][index - 1]["objects"]
        elements = [item for item in elements if item.get("kind") != "chart"]
        elements.extend({**item, "paragraphs": []} for item in native if item["kind"] in {"chart", "smartart_node"})
        if index == 1:
            elements.extend({**item, "kind": "shared_text", "paragraphs": []} for item in inventory["shared_objects"] if item.get("text", "").strip() and item.get("action") == "replace")
        page = {
            "page_id": page_id,
            "slide_index": index,
            "width_emu": presentation.slide_width,
            "height_emu": presentation.slide_height,
            "elements": elements,
        }
        write_json(pages_dir / f"{page_id}.json", page)
        page_elements.append((page_id, elements))
        reason = _slide_skip_reason(slide, elements)
        if reason:
            skipped_pages.append({"page_id": page_id, "reason": reason})
    schema_dir = pack / "schema"
    schema_dir.mkdir(parents=True)
    shutil.copy2(SKILL_ROOT / "schemas/deck_skeleton.schema.json", schema_dir / "deck_skeleton.schema.json")
    shutil.copy2(SKILL_ROOT / "prompts/extract.md", pack / "INSTRUCTIONS.md")
    write_text(
        pack / "TASK.md",
        "阅读 INSTRUCTIONS.md、skeleton_draft.json、pages/、slot_candidates.json、skip_hints.json、recoveries.json 与 schema。"
        "以 skeleton_draft.json 为起点改角色，写出工作目录下的 skeleton.json。"
        "保持全部源页与页序；每页全部可编辑候选均须绑定。截图像素包含的旧字作为保留项记录。\n",
    )
    _write_slot_candidates(pack, page_elements)
    write_skeleton_draft(pack, page_elements, skipped_pages)
    hints = skip_report(pages=skipped_pages)
    write_json(pack / "skip_hints.json", hints)
    write_recoveries(pack / "recoveries.json", recoveries)
    write_json(
        workdir / "extract_state.json",
        {
            "format": "rdw_extract_state", "version": "1.0", "source_sha256": sha256(source),
            "source_name": source.name, "page_count": len(presentation.slides),
            "preflight": report.model_dump(), "temporary_source": "source.pptx",
            "recovered_pages": [item["page_id"] for item in recoveries],
        },
    )
    return {
        "workdir": str(workdir), "page_count": len(presentation.slides), "warnings": report.warnings,
        "extract_pack": str(pack), "skipped_pages": skipped_pages, "recovered_pages": recoveries,
    }


def _slide_skip_reason(slide, elements: list[dict]) -> str | None:
    return None


def _slide_elements(slide, slide_width: int, slide_height: int) -> list[dict]:
    elements: list[dict] = []
    for shape in slide.shapes:
        elements.extend(_shape_elements(shape, slide_width, slide_height))
    return elements


def _shape_elements(shape, slide_width: int, slide_height: int) -> list[dict]:
    if getattr(shape, "shape_type", None) != 6 and diagram_frames(shape._element):
        return []
    kind = "group" if getattr(shape, "shape_type", None) == 6 else "shape"
    if getattr(shape, "has_table", False):
        kind = "table"
    elif getattr(shape, "has_chart", False):
        kind = "chart"
    elif getattr(shape, "shape_type", None) == 13:
        kind = "picture"
    base = {
        "shape_ref": f"sp_{shape.shape_id}", "name": shape.name, "kind": kind,
        "bbox": {"left": shape.left, "top": shape.top, "width": shape.width, "height": shape.height},
        "ratio": {
            "left": round(shape.left / slide_width, 5), "top": round(shape.top / slide_height, 5),
            "width": round(shape.width / slide_width, 5), "height": round(shape.height / slide_height, 5),
        },
        "z_order": _z_order(shape),
    }
    results: list[dict] = []
    if getattr(shape, "has_text_frame", False):
        item = deepcopy(base)
        item["text"] = shape.text
        item["paragraphs"] = _paragraphs(shape.text_frame)
        results.append(item)
    if getattr(shape, "has_table", False):
        for row_index, row in enumerate(shape.table.rows):
            for column_index, cell in enumerate(row.cells):
                item = deepcopy(base)
                item["kind"] = "table_cell"
                item["shape_ref"] = f"sp_{shape.shape_id}.cell_{row_index}_{column_index}"
                item["text"] = cell.text
                item["paragraphs"] = _paragraphs(cell.text_frame)
                results.append(item)
    if kind == "group":
        for child in shape.shapes:
            results.extend(_shape_elements(child, slide_width, slide_height))
    if not results:
        base["text"] = ""
        base["paragraphs"] = []
        results.append(base)
    return results


def _diagram_elements(slide, slide_width: int, slide_height: int) -> list[dict]:
    results: list[dict] = []
    for frame in diagram_frames(slide.element):
        props = frame.xpath('.//*[local-name()="cNvPr"]')
        try:
            frame_id = int(props[0].get("id") or 0) if props else 0
        except (TypeError, ValueError):
            frame_id = 0
        left, top, width, height = xfrm_box(frame)
        texts = diagram_texts(slide, frame)
        if not texts:
            texts = [""]
        for index, text in enumerate(texts):
            results.append(
                {
                    "shape_ref": f"sp_{frame_id}.dgm_{index}",
                    "name": props[0].get("name") if props else "SmartArt",
                    "kind": "smartart_node",
                    "bbox": {"left": left, "top": top, "width": width, "height": height},
                    "ratio": {
                        "left": round(left / slide_width, 5),
                        "top": round(top / slide_height, 5),
                        "width": round(width / slide_width, 5),
                        "height": round(height / slide_height, 5),
                    },
                    "z_order": 0,
                    "text": text,
                    "paragraphs": [],
                }
            )
    return results


def _paragraphs(frame) -> list[dict]:
    result = []
    for paragraph in frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            color = None
            try:
                color = str(run.font.color.rgb) if run.font.color.rgb is not None else None
            except (AttributeError, TypeError, ValueError):
                pass
            runs.append(
                {
                    "text": run.text, "font_name": run.font.name,
                    "font_size_pt": run.font.size.pt if run.font.size is not None else None,
                    "bold": run.font.bold, "italic": run.font.italic, "color_hex": color,
                }
            )
        result.append({"level": paragraph.level, "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None, "runs": runs})
    return result


def _z_order(shape) -> int:
    parent = shape._element.getparent()
    return list(parent).index(shape._element) if parent is not None else 0
