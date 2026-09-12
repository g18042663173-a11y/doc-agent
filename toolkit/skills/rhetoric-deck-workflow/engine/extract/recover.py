"""Keep source visuals; only strip animation. Text is sanitized later during seal."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation

from engine.shared.common import write_json


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def recover_source(path: Path) -> list[dict]:
    presentation = Presentation(path)
    recoveries: list[dict] = []
    stripped = False
    for index, slide in enumerate(presentation.slides, start=1):
        actions: list[dict] = []
        if _strip_motion(slide):
            stripped = True
            actions.append({"action": "animation_stripped"})
        frames = diagram_frames(slide.element)
        if frames:
            node_count = sum(len(diagram_text_nodes(slide, frame)) for frame in frames)
            actions.append({"action": "smartart_kept", "nodes": node_count, "graphics": len(frames)})
        pictures = [shape for shape in slide.shapes if getattr(shape, "shape_type", None) == 13]
        if pictures:
            actions.append({"action": "screenshot_kept", "pictures": len(pictures)})
        if actions:
            recoveries.append({"page_id": f"p{index:02d}", "actions": actions})
    if stripped:
        presentation.save(path)
    return recoveries


def write_recoveries(path: Path, recoveries: list[dict]) -> None:
    write_json(path, {"format": "rdw_recoveries", "version": "1.0", "pages": recoveries})


def recovered_page_ids(recoveries: list[dict]) -> set[str]:
    return {item["page_id"] for item in recoveries}


def slide_has_diagram(slide) -> bool:
    return bool(diagram_frames(slide.element))


def diagram_frames(root) -> list:
    frames = []
    seen: set[int] = set()
    nodes = []
    if getattr(root, "tag", None) is not None and etree.QName(root).localname == "graphicFrame":
        nodes.append(root)
    nodes.extend(root.xpath('.//*[local-name()="graphicFrame"]'))
    for frame in nodes:
        ident = id(frame)
        if ident in seen:
            continue
        seen.add(ident)
        for data in frame.xpath('.//*[local-name()="graphicData"]'):
            if (data.get("uri") or "") == DIAGRAM_URI:
                frames.append(frame)
                break
    return frames


def diagram_text_nodes(slide, frame) -> list:
    root = _diagram_root(slide, frame)
    if root is None:
        return []
    return root.xpath(f'.//*[namespace-uri()="{A_NS}" and local-name()="t"]')


def diagram_texts(slide, frame) -> list[str]:
    return [(node.text or "") for node in diagram_text_nodes(slide, frame)]


def write_diagram_node(slide, frame_id: int, node_index: int, text: str) -> bool:
    for frame in diagram_frames(slide.element):
        if _frame_id(frame) != frame_id:
            continue
        part_root = _diagram_part(slide, frame)
        if part_root is None:
            return False
        part, root = part_root
        nodes = root.xpath(f'.//*[namespace-uri()="{A_NS}" and local-name()="t"]')
        if node_index >= len(nodes):
            return False
        nodes[node_index].text = text
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return True
    return False


def xfrm_box(frame) -> tuple[int, int, int, int]:
    off = frame.xpath('.//*[local-name()="off"]')
    ext = frame.xpath('.//*[local-name()="ext"]')
    left = int(off[0].get("x") or 0) if off else 0
    top = int(off[0].get("y") or 0) if off else 0
    width = max(int(ext[0].get("cx") or 1) if ext else 1, 1)
    height = max(int(ext[0].get("cy") or 1) if ext else 1, 1)
    return left, top, width, height


def _strip_motion(slide) -> bool:
    changed = False
    for tag in ("timing", "transition"):
        element = slide.element.find(f"{{{P_NS}}}{tag}")
        if element is not None:
            slide.element.remove(element)
            changed = True
    return changed


def _frame_id(frame) -> int:
    props = frame.xpath('.//*[local-name()="cNvPr"]')
    if not props:
        return 0
    try:
        return int(props[0].get("id") or 0)
    except (TypeError, ValueError):
        return 0


def _diagram_part(slide, frame):
    rel_ids = frame.xpath('.//*[local-name()="relIds"]')
    if not rel_ids:
        return None
    dm = rel_ids[0].get(f"{{{R_NS}}}dm")
    if not dm:
        return None
    try:
        part = slide.part.rels[dm].target_part
        root = etree.fromstring(part.blob)
    except Exception:
        return None
    return part, root


def _diagram_root(slide, frame):
    loaded = _diagram_part(slide, frame)
    return None if loaded is None else loaded[1]
