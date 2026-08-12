from __future__ import annotations

import base64
import re
import zipfile
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_parse_docx_extracts_headings_paragraph_lists_and_table(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "需求说明.docx"
    doc = Document()
    doc.add_heading("需求说明", level=1)
    doc.add_paragraph("本文件描述输入解析能力。")
    doc.add_heading("功能列表", level=2)
    doc.add_paragraph("解析 DOCX 标题", style="List Bullet")
    doc.add_paragraph("提取表格预览", style="List Bullet 2")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "title"
    table.cell(1, 1).text = "文档标题"
    doc.save(path)

    ir = parse_docx(path)

    assert ir.source.format == "docx"
    assert [item.text for item in ir.content.outline] == ["需求说明", "功能列表"]
    assert [block.type for block in ir.content.blocks] == ["heading", "paragraph", "heading", "bullet_list", "table"]
    assert ir.content.blocks[-1].header == ["字段", "说明"]
    assert ir.content.blocks[-1].rows == [["title", "文档标题"]]


def test_parse_docx_preserves_monospace_code_paragraph_indentation(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "code.docx"
    doc = Document()
    code = doc.add_paragraph(style="Normal")
    run = code.add_run("if (ready) {\n    send();\n}")
    run.font.name = "Consolas"
    doc.save(path)

    ir = parse_docx(path)

    assert ir.ir_version == "1.2"
    assert ir.content.blocks[0].type == "code_block"
    assert ir.content.blocks[0].code == "if (ready) {\n    send();\n}"


def test_parse_docx_truncates_table_preview_to_20_rows(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "large.docx"
    doc = Document()
    doc.add_heading("大表", level=1)
    table = doc.add_table(rows=26, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    for row in range(1, 26):
        table.cell(row, 0).text = f"A{row}"
        table.cell(row, 1).text = f"B{row}"
    doc.save(path)

    ir = parse_docx(path)

    table_block = ir.content.blocks[-1]
    assert table_block.type == "table"
    assert len(table_block.rows) == 20
    assert any("truncated to 20 rows" in warning for warning in ir.warnings)


def test_parse_docx_warns_when_nested_table_content_is_skipped(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "nested-table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "外层表头"
    outer_cell = table.cell(1, 0)
    outer_cell.text = "外层说明"
    nested = outer_cell.add_table(rows=2, cols=1)
    nested.cell(0, 0).text = "内层表头"
    nested.cell(1, 0).text = "内层敏感内容"
    doc.save(path)

    ir = parse_docx(path)

    assert any("nested tables unsupported: 1" in warning for warning in ir.warnings)


def test_parse_docx_detects_heading_from_outline_level_without_heading_style(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "outline-level.docx"
    doc = Document()
    paragraph = doc.add_paragraph("XML 大纲标题")
    ppr = paragraph._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    ppr.append(outline)
    doc.save(path)

    ir = parse_docx(path)

    assert ir.content.blocks[0].type == "heading"
    assert ir.content.blocks[0].level == 2
    assert ir.content.outline[0].text == "XML 大纲标题"


def test_parse_docx_clamps_negative_outline_level_to_heading_1(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "negative-outline.docx"
    doc = Document()
    paragraph = doc.add_paragraph("负值大纲")
    ppr = paragraph._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "-1")
    ppr.append(outline)
    doc.save(path)

    ir = parse_docx(path)

    assert ir.content.blocks[0].type == "heading"
    assert ir.content.blocks[0].level == 1


def test_parse_docx_records_image_presence_and_dimensions(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))
    path = tmp_path / "image.docx"
    doc = Document()
    doc.add_paragraph("带图片")
    doc.add_picture(str(image))
    doc.save(path)

    ir = parse_docx(path)

    assert ir.stats.images == 1
    assert any("docx image present" in warning for warning in ir.warnings)


def test_parse_docx_tolerates_inline_shape_without_extent(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))
    path = tmp_path / "no-extent.docx"
    doc = Document()
    doc.add_paragraph("带损坏图片")
    doc.add_picture(str(image))
    doc.save(path)

    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    document_xml = parts["word/document.xml"].decode("utf-8")
    assert "<wp:extent" in document_xml
    parts["word/document.xml"] = re.sub(r"<wp:extent[^>]*/>", "", document_xml).encode("utf-8")
    with zipfile.ZipFile(path, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)

    ir = parse_docx(path)

    assert ir.stats.images == 1
    assert any("W103" in warning and "unreadable geometry" in warning for warning in ir.warnings)
    assert ir.content.blocks[0].text == "带损坏图片"


def test_parse_docx_records_unsupported_word_features_from_xml(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "unsupported.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.save(path)
    with zipfile.ZipFile(path, "a") as package:
        package.writestr(
            "word/unsupported.xml",
            "<w:commentRangeStart/><w:ins/><w:txbxContent/><dgm:relIds/>",
        )

    ir = parse_docx(path)

    assert any("unsupported comments" in warning for warning in ir.warnings)
    assert any("unsupported revisions" in warning for warning in ir.warnings)
    assert any("unsupported text boxes" in warning for warning in ir.warnings)
    assert any("unsupported SmartArt" in warning for warning in ir.warnings)
    assert all("at part word/unsupported.xml" in warning for warning in ir.warnings if warning.startswith("unsupported"))


def test_parse_docx_truncates_long_paragraph_and_cell_with_locations(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "long-text.docx"
    doc = Document()
    doc.add_paragraph("段" * 2100)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "正常"
    table.cell(1, 1).text = "格" * 2101
    doc.save(path)

    ir = parse_docx(path)

    assert len(ir.content.blocks[0].text) == 2000
    assert len(ir.content.blocks[1].rows[0][1]) == 2000
    assert any("docx body paragraph 1" in warning for warning in ir.warnings)
    assert any("docx table 1 row 1 column 2" in warning for warning in ir.warnings)
    assert not any("row 1 column 1" in warning for warning in ir.warnings)


def test_parse_docx_records_embedded_ole_part_and_size(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "embedded.docx"
    doc = Document()
    doc.add_paragraph("带嵌入对象")
    doc.save(path)
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("word/embeddings/oleObject1.bin", b"fake-ole")

    ir = parse_docx(path)

    assert any(
        "embedded/OLE objects" in warning
        and "word/embeddings/oleObject1.bin" in warning
        and "total_bytes=8" in warning
        for warning in ir.warnings
    )


def test_parse_docx_truncates_wide_table_to_contract_limit(tmp_path: Path) -> None:
    from app.parsers.docx_parser import parse_docx

    path = tmp_path / "wide.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=13)
    for column in range(13):
        table.cell(0, column).text = f"列{column + 1}"
        table.cell(1, column).text = f"值{column + 1}"
    doc.save(path)

    ir = parse_docx(path)

    assert len(ir.content.blocks[0].header) == 12
    assert len(ir.content.blocks[0].rows[0]) == 12
    assert any("table 1 preview truncated to 12 columns" in warning for warning in ir.warnings)
