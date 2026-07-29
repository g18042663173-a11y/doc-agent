from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation

from app.assets.pipeline import load_asset_manifest, normalize_assets
from app.ir.deck_ir import DeckIR
from app.lint.pptx_lint import check_pptx
from app.rendering.pptx_renderer import render_deck_ir


def _make_assets(tmp_path: Path):
    files = []
    for index, size in enumerate(((1600, 900), (900, 1600), (1200, 1200)), start=1):
        path = tmp_path / f"image-{index}.png"
        Image.new("RGB", size, (220, 20 * index, 50 * index)).save(path)
        files.append(path)
    normalize_assets(files, tmp_path / "assets")
    return load_asset_manifest(tmp_path / "assets" / "asset_manifest.json")


def test_render_v2_images_infographics_scatter_and_combo_as_native_objects(tmp_path: Path) -> None:
    registry = _make_assets(tmp_path)
    refs = [asset.asset_id for asset in registry.manifest.assets]
    image = {"image_ref": refs[0], "fit": "cover", "focal_x": 0.2, "alt": "现场", "credit": "内部项目组"}
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.0",
            "meta": {"title": "视觉能力", "classification": "公开"},
            "slides": [
                {"layout": "image_text", "title": "图片与结论并列", "image": image, "text": "现场证据支撑结论"},
                {
                    "layout": "image_grid",
                    "title": "多图保持独立编辑",
                    "images": [image, {**image, "image_ref": refs[1], "fit": "contain"}],
                },
                {
                    "layout": "infographic",
                    "title": "漏斗阶段逐步收敛",
                    "infographic": {"kind": "funnel", "stages": [{"label": "识别"}, {"label": "筛选"}, {"label": "闭环"}]},
                },
                {
                    "layout": "infographic",
                    "title": "方案按影响和难度定位",
                    "infographic": {"kind": "quadrant", "x_axis": "影响", "y_axis": "难度", "items": [{"label": "方案A", "x": 0.8, "y": 0.2}]},
                },
                {
                    "layout": "infographic",
                    "title": "流程形成持续闭环",
                    "infographic": {"kind": "cycle", "stages": [{"label": "计划", "icon": "target"}, {"label": "执行", "icon": "process"}, {"label": "复盘", "icon": "quality"}]},
                },
                {
                    "layout": "infographic",
                    "title": "任务按优先级分类",
                    "infographic": {"kind": "matrix", "row_labels": ["高", "低"], "column_labels": ["急", "缓"], "cells": [["立即", "计划"], ["授权", "观察"]]},
                },
                {
                    "layout": "chart",
                    "title": "投入与产出存在正相关",
                    "chart": {"kind": "scatter", "categories": [], "series": [{"name": "样本", "x_values": [1, 2, 3], "values": [2, 4, 5]}], "show_data_labels": False},
                },
                {
                    "layout": "chart",
                    "title": "收入增长且利润率改善",
                    "chart": {
                        "kind": "combo",
                        "categories": ["一月", "二月", "三月"],
                        "series": [
                            {"name": "收入", "values": [10, 12, 15], "chart_type": "bar", "axis": "primary", "unit": "万元"},
                            {"name": "利润率", "values": [0.2, 0.22, 0.25], "chart_type": "line", "axis": "secondary", "unit": "%"},
                        ],
                        "number_format": "0.0",
                        "secondary_number_format": "0.0%",
                        "source": "经营月报",
                        "methodology": "利润率=利润/收入",
                        "show_data_labels": False,
                    },
                },
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "visual-v2.pptx", asset_registry=registry)
    presentation = Presentation(output)

    assert sum(shape.shape_type == 13 for slide in presentation.slides for shape in slide.shapes) == 3
    assert sum(getattr(shape, "has_chart", False) for slide in presentation.slides for shape in slide.shapes) == 3
    assert any(shape.name == "HW_INFOGRAPHIC:matrix:table" and shape.has_table for shape in presentation.slides[5].shapes)
    assert any(shape.name.startswith("HW_INFOGRAPHIC:funnel:") for shape in presentation.slides[2].shapes)
    assert any(shape.name == "HW_SEMANTIC_ICON:target" for shape in presentation.slides[4].shapes)
    assert {shape.name for shape in presentation.slides[7].shapes if getattr(shape, "has_chart", False)} == {
        "HW_COMBO_CHART:primary",
        "HW_COMBO_CHART:secondary",
    }
    lint_codes = {item.code for item in check_pptx(output, classification="公开").items}
    assert "HW-W13" not in lint_codes
    assert "HW-W16" not in lint_codes


def test_visual_lint_reports_low_resolution_collision_and_broken_combo(tmp_path: Path) -> None:
    path = tmp_path / "tiny.png"
    Image.new("RGB", (100, 100), (200, 20, 20)).save(path)
    normalize_assets([path], tmp_path / "assets")
    registry = load_asset_manifest(tmp_path / "assets" / "asset_manifest.json")
    ref = registry.manifest.assets[0].asset_id
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.0",
            "meta": {"title": "视觉 lint", "classification": "公开"},
            "slides": [
                {"layout": "image", "title": "低分辨率", "image_ref": ref, "alt": "低分辨率测试"},
                {
                    "layout": "infographic",
                    "title": "碰撞",
                    "infographic": {
                        "kind": "quadrant",
                        "x_axis": "影响",
                        "y_axis": "难度",
                        "items": [{"label": "A", "x": 0.5, "y": 0.5}, {"label": "B", "x": 0.5, "y": 0.5}],
                    },
                },
                {
                    "layout": "chart",
                    "title": "组合图",
                    "chart": {
                        "kind": "combo",
                        "categories": ["一", "二"],
                        "series": [
                            {"name": "金额", "values": [1, 2], "chart_type": "bar", "axis": "primary", "unit": "元"},
                            {"name": "比例", "values": [0.1, 0.2], "chart_type": "line", "axis": "secondary", "unit": "%"},
                        ],
                        "number_format": "0",
                        "secondary_number_format": "0%",
                        "show_data_labels": False,
                    },
                },
            ],
        }
    )
    output = render_deck_ir(deck, tmp_path / "lint.pptx", asset_registry=registry)
    presentation = Presentation(output)
    next(shape for shape in presentation.slides[2].shapes if shape.name == "HW_COMBO_CHART:secondary").name = "BROKEN_COMBO"
    presentation.save(output)

    codes = {item.code for item in check_pptx(output, classification="公开").items}
    assert {"HW-W13", "HW-W15", "HW-W16"} <= codes
