from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

from docx import Document
from PIL import Image
from pptx import Presentation
import pytest

from scripts import package_huawei_doc_skill
from scripts.portable_skill_runtime import run_skill


ROOT = Path(__file__).resolve().parents[2]
SKILL = ROOT / "skills" / "huawei-doc-workflow"
WORKFLOW = SKILL / "scripts" / "workflow.py"
ENGINE = SKILL / "scripts" / "engine"


def _run(
    workflow: Path,
    *arguments: object,
    cwd: Path | None = None,
    env_overrides: dict[str, str] | None = None,
) -> subprocess.CompletedProcess[str]:
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    if env_overrides:
        environment.update(env_overrides)
    return subprocess.run(
        [sys.executable, str(workflow), *(str(argument) for argument in arguments)],
        cwd=cwd or ROOT,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def _assert_ok(result: subprocess.CompletedProcess[str]) -> None:
    assert result.returncode == 0, result.stdout + result.stderr


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def test_skill_metadata_snapshot_and_policy_are_standalone() -> None:
    manifest = package_huawei_doc_skill.verify_skill_source()
    skill_md = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    openai_yaml = (SKILL / "agents" / "openai.yaml").read_text(encoding="utf-8")
    deck_schema = json.loads((ENGINE / "schemas" / "deck_ir.schema.json").read_text(encoding="utf-8"))
    schema_text = json.dumps(deck_schema, ensure_ascii=False)

    assert skill_md.startswith("---\nname: huawei-doc-workflow\n")
    assert "$rhetoric-deck-workflow" in skill_md
    assert "master_redraw" in skill_md
    assert "allow_implicit_invocation: false" in openai_yaml
    assert manifest["source_product_version"] == "2.2.0"
    assert not (ENGINE / "app" / "generators").exists()
    assert not (ENGINE / "app" / "web_api.py").exists()
    assert not (ENGINE / "app" / "static").exists()
    assert not (SKILL / "wheelhouse").exists()
    for layout in (
        "cover",
        "agenda",
        "section",
        "title_bullets",
        "two_column",
        "table",
        "cards",
        "chart",
        "architecture_diagram",
        "composite",
        "process_flow",
        "timeline",
        "image",
        "image_text",
        "image_grid",
        "infographic",
        "conclusion",
    ):
        assert f'"{layout}"' in schema_text


def test_doctor_passes_with_required_host_dependencies() -> None:
    result = _run(WORKFLOW, "doctor", "--json")
    _assert_ok(result)
    report = json.loads(result.stdout)
    assert report["pass"] is True
    assert report["skill_version"] == "1.1.0"
    graphviz = next(item for item in report["checks"] if item["name"] == "graphviz")
    if not graphviz["ok"]:
        assert "必须进行视觉复核" in graphviz["message"]
    assert {item["name"] for item in report["checks"] if item["required"]} == {
        "python",
        "pydantic",
        "python-docx",
        "openpyxl",
        "python-pptx",
        "Pillow",
        "engine_snapshot",
    }


@pytest.mark.parametrize(
    ("relative_path", "source_format"),
    [
        ("samples/input/quarterly_report.md", "md"),
        ("samples/input/需求说明.docx", "docx"),
        ("samples/input/销售台账.xlsx", "xlsx"),
        ("samples/input/项目汇报.pptx", "pptx"),
    ],
)
def test_prepare_parses_all_supported_inputs(
    tmp_path: Path,
    relative_path: str,
    source_format: str,
) -> None:
    brief = tmp_path / "brief.md"
    brief.write_text("基于输入资料形成一份事实可追溯的项目汇报。\n", encoding="utf-8")
    run_dir = tmp_path / "run"
    result = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--input",
        ROOT / relative_path,
        "--theme",
        "hw-report",
        "--pages",
        8,
        "--depth",
        "标准",
        "--output-dir",
        run_dir,
    )
    _assert_ok(result)
    packet = json.loads((run_dir / "generation_packet.json").read_text(encoding="utf-8"))
    document = json.loads((run_dir / "document_ir.json").read_text(encoding="utf-8"))

    assert packet["source"]["summary"]["source"]["format"] == source_format
    assert packet["contract"]["ir_version"] == "2.2"
    assert packet["visual_plan"] == "visual_plan.json"
    assert document["source"]["format"] == source_format
    assert (run_dir / "visual_plan.json").is_file()


def test_prepare_supports_topic_only_without_fake_document_ir(tmp_path: Path) -> None:
    brief = tmp_path / "brief.md"
    brief.write_text("生成一份八页的无线项目阶段总结，未知数据必须明确标注。\n", encoding="utf-8")
    run_dir = tmp_path / "run"
    result = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--pages",
        8,
        "--output-dir",
        run_dir,
    )
    _assert_ok(result)
    packet = json.loads((run_dir / "generation_packet.json").read_text(encoding="utf-8"))
    assert packet["source"] is None
    assert packet["visual_plan"] is None
    assert not (run_dir / "document_ir.json").exists()


def test_prepare_does_not_treat_unlabeled_images_as_semantic_evidence(tmp_path: Path) -> None:
    brief = tmp_path / "brief.md"
    brief.write_text("根据输入报告生成汇报；没有图片说明时不要猜测图片内容。\n", encoding="utf-8")
    image = tmp_path / "opaque.png"
    Image.new("RGB", (640, 360), "#1F4E79").save(image)
    run_dir = tmp_path / "run"

    prepared = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--input",
        ROOT / "samples" / "input" / "quarterly_report.md",
        "--asset",
        image,
        "--output-dir",
        run_dir,
    )
    _assert_ok(prepared)

    packet = json.loads((run_dir / "generation_packet.json").read_text(encoding="utf-8"))
    plan = json.loads((run_dir / "visual_plan.json").read_text(encoding="utf-8"))
    semantics = packet["asset_semantics"]
    assert semantics["mode"] == "text_grounded_only"
    assert semantics["engine_image_understanding"] is False
    assert semantics["grounded_assets"] == []
    assert len(semantics["ungrounded_assets"]) == 1
    assert all(not item["asset_ids"] for item in plan["opportunities"])


def test_validation_blocks_invalid_ir_after_two_repairs(tmp_path: Path) -> None:
    draft = ROOT / "samples" / "ir" / "deck_invalid_d002_missing_title.json"
    run_dir = tmp_path / "run"
    for expected_remaining in (2, 1, 0):
        result = _run(
            WORKFLOW,
            "validate",
            "--target",
            "deck",
            "--draft",
            draft,
            "--output-dir",
            run_dir,
        )
        assert result.returncode == 2
        report = json.loads((run_dir / "validation_report.json").read_text(encoding="utf-8"))
        assert report["summary"]["pass"] is False
        assert report["remaining_repairs"] == expected_remaining
        assert report["items"][0]["code"] == "D002"

    blocked = _run(
        WORKFLOW,
        "validate",
        "--target",
        "deck",
        "--draft",
        draft,
        "--output-dir",
        run_dir,
    )
    assert blocked.returncode == 2
    assert "SKILL-E020" in blocked.stderr
    assert not (run_dir / "validated_ir.json").exists()
    assert not (run_dir / "deck.pptx").exists()


def test_word_and_full_deck_finalize_through_validated_ir(tmp_path: Path) -> None:
    word_dir = tmp_path / "word"
    word_validate = _run(
        WORKFLOW,
        "validate",
        "--target",
        "word",
        "--draft",
        ENGINE / "examples" / "word_valid_03_table.json",
        "--output-dir",
        word_dir,
    )
    _assert_ok(word_validate)
    word_finalize = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "word",
        "--ir",
        word_dir / "validated_ir.json",
        "--output-dir",
        word_dir,
    )
    _assert_ok(word_finalize)
    word_document = Document(word_dir / "document.docx")
    assert len(word_document.tables) == 1
    assert json.loads((word_dir / "workflow_manifest.json").read_text(encoding="utf-8"))["lint"]["pass"] is True

    deck_dir = tmp_path / "deck"
    deck_validate = _run(
        WORKFLOW,
        "validate",
        "--target",
        "deck",
        "--draft",
        ENGINE / "examples" / "deck_valid_full.json",
        "--output-dir",
        deck_dir,
    )
    _assert_ok(deck_validate)
    deck_finalize = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "deck",
        "--ir",
        deck_dir / "validated_ir.json",
        "--output-dir",
        deck_dir,
        env_overrides={"PATH": ""},
    )
    _assert_ok(deck_finalize)
    assert "deterministic fallback" in deck_finalize.stderr
    presentation = Presentation(deck_dir / "deck.pptx")
    assert len(presentation.slides) == 14
    manifest = json.loads((deck_dir / "workflow_manifest.json").read_text(encoding="utf-8"))
    assert manifest["ir"]["ir_version"] == "2.2"
    assert manifest["lint"]["errors"] == 0
    assert manifest["visual_review"] == {
        "engine_pixel_inspection_performed": False,
        "note": "确定性引擎不理解或评审渲染像素；预览图仅供人工终审。",
        "status": "human_required",
    }
    assert any("deterministic fallback" in item["message"] for item in manifest["runtime_warnings"])

    audit_dir = tmp_path / "deck-audit"
    deck_audit = _run(
        WORKFLOW,
        "audit",
        deck_dir / "deck.pptx",
        "--output-dir",
        audit_dir,
        "--theme",
        "hw_v1",
    )
    _assert_ok(deck_audit)
    audit_report = json.loads((audit_dir / "report.json").read_text(encoding="utf-8"))
    assert audit_report["summary"]["pass"] is True
    assert (audit_dir / "report.md").is_file()

    original_hash = _sha256(deck_dir / "deck.pptx")
    overwrite_refused = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "deck",
        "--ir",
        deck_dir / "validated_ir.json",
        "--output-dir",
        deck_dir,
    )
    assert overwrite_refused.returncode == 3
    assert "SKILL-E003" in overwrite_refused.stderr
    assert _sha256(deck_dir / "deck.pptx") == original_hash


def test_real_images_and_template_modes_emit_audits(tmp_path: Path) -> None:
    brief = tmp_path / "brief.md"
    brief.write_text("用本地证据图生成可编辑演示文稿。\n", encoding="utf-8")
    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    Image.new("RGB", (1200, 800), "#C7000B").save(first)
    Image.new("RGB", (900, 1200), "#1F4E79").save(second)

    image_dir = tmp_path / "image-run"
    prepared = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--asset",
        first,
        "--asset",
        second,
        "--output-dir",
        image_dir,
    )
    _assert_ok(prepared)
    asset_manifest = json.loads(
        (image_dir / "assets" / "asset_manifest.json").read_text(encoding="utf-8")
    )
    first_id, second_id = [item["asset_id"] for item in asset_manifest["assets"]]
    deck = {
        "ir_type": "deck",
        "ir_version": "2.2",
        "meta": {
            "title": "图片证据链",
            "classification": "HUAWEI CONFIDENTIAL",
            "theme": "hw_v1",
        },
        "slides": [
            {"layout": "cover", "title": "图片证据链"},
            {
                "layout": "image_text",
                "title": "单图证据",
                "image": {"image_ref": first_id, "fit": "contain", "alt": "红色测试图"},
                "heading": "已规范化",
                "bullets": ["图片来自本地测试资产", "哈希和使用位置已审计"],
            },
            {
                "layout": "image_grid",
                "title": "多图对照",
                "images": [
                    {"image_ref": first_id, "alt": "横向测试图"},
                    {"image_ref": second_id, "alt": "纵向测试图"},
                ],
            },
            {
                "layout": "infographic",
                "title": "资产处理形成闭环",
                "infographic": {
                    "kind": "funnel",
                    "stages": [
                        {"label": "安全解码"},
                        {"label": "规范化"},
                        {"label": "使用审计"},
                    ],
                },
            },
            {"layout": "conclusion", "title": "结论", "bullets": ["资产链路可追溯"]},
        ],
    }
    draft = image_dir / "draft_ir.json"
    draft.write_text(json.dumps(deck, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _assert_ok(
        _run(
            WORKFLOW,
            "validate",
            "--target",
            "deck",
            "--draft",
            draft,
            "--output-dir",
            image_dir,
        )
    )
    finalized = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "deck",
        "--ir",
        image_dir / "validated_ir.json",
        "--asset-manifest",
        image_dir / "assets" / "asset_manifest.json",
        "--output-dir",
        image_dir,
    )
    _assert_ok(finalized)
    assert (image_dir / "deck.pptx").is_file()
    assert (image_dir / "asset_usage_audit.json").is_file()
    assert len(json.loads((image_dir / "asset_usage_audit.json").read_text(encoding="utf-8"))["usages"]) == 3

    template_dir = tmp_path / "template-run"
    template = ROOT / "samples" / "input" / "项目汇报.pptx"
    template_prepared = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--template",
        template,
        "--output-dir",
        template_dir,
    )
    _assert_ok(template_prepared)
    _assert_ok(
        _run(
            WORKFLOW,
            "validate",
            "--target",
            "deck",
            "--draft",
            ENGINE / "examples" / "deck_valid_01_minimal.json",
            "--output-dir",
            template_dir,
        )
    )
    template_finalized = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "deck",
        "--ir",
        template_dir / "validated_ir.json",
        "--template",
        template,
        "--output-dir",
        template_dir,
    )
    _assert_ok(template_finalized)
    for name in (
        "template_profile.json",
        "template_plan.json",
        "template_structure.json",
        "template_replacement_audit.json",
        "pptx_package_report.json",
    ):
        assert (template_dir / name).is_file()


def test_dangerous_template_and_missing_asset_never_publish_artifact(tmp_path: Path) -> None:
    brief = tmp_path / "brief.md"
    brief.write_text("安全失败测试。\n", encoding="utf-8")
    dangerous = tmp_path / "dangerous.pptx"
    with zipfile.ZipFile(dangerous, "w") as archive:
        archive.writestr("../escape.xml", "<unsafe/>")
    template_dir = tmp_path / "template-run"
    rejected = _run(
        WORKFLOW,
        "prepare",
        "--target",
        "deck",
        "--brief-file",
        brief,
        "--template",
        dangerous,
        "--output-dir",
        template_dir,
    )
    assert rejected.returncode == 3
    assert '"code": "E003"' in rejected.stderr
    assert not (template_dir / "generation_packet.json").exists()
    assert not (template_dir / "deck.pptx").exists()

    missing_asset_dir = tmp_path / "missing-asset"
    missing_asset_dir.mkdir()
    deck = {
        "ir_type": "deck",
        "ir_version": "2.2",
        "meta": {"title": "缺图", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
        "slides": [{"layout": "image", "title": "缺失引用", "image_ref": "asset-missing"}],
    }
    validated = missing_asset_dir / "validated_ir.json"
    validated.write_text(json.dumps(deck, ensure_ascii=False) + "\n", encoding="utf-8")
    manifest = {
        "manifest_version": "1.0",
        "assets": [],
    }
    asset_manifest = missing_asset_dir / "asset_manifest.json"
    asset_manifest.write_text(json.dumps(manifest) + "\n", encoding="utf-8")
    missing = _run(
        WORKFLOW,
        "finalize",
        "--target",
        "deck",
        "--ir",
        validated,
        "--asset-manifest",
        asset_manifest,
        "--output-dir",
        missing_asset_dir,
    )
    assert missing.returncode == 3
    assert '"code": "A005"' in missing.stderr
    assert not (missing_asset_dir / "deck.pptx").exists()


def test_packaged_zip_runs_without_repository_imports(tmp_path: Path) -> None:
    archive, digest_path = package_huawei_doc_skill.build_archive(tmp_path, overwrite=True)
    assert digest_path.read_text(encoding="utf-8").split()[0] == _sha256(archive)

    with zipfile.ZipFile(archive) as bundle:
        names = bundle.namelist()
        assert "huawei-doc-workflow/SKILL_PACKAGE_MANIFEST.json" in names
        assert "huawei-doc-workflow/runtime/python/python.exe" in names
        assert "huawei-doc-workflow/runtime/graphviz/bin/dot.exe" in names
        assert "huawei-doc-workflow/runtime/requirements.lock" in names
        assert "huawei-doc-workflow/run.cmd" in names
        package_manifest = json.loads(bundle.read("huawei-doc-workflow/SKILL_PACKAGE_MANIFEST.json"))
        assert package_manifest["schema_version"] == 2
        assert package_manifest["dependency_bundle_included"] is True
        assert package_manifest["graphviz_included"] is True
        assert package_manifest["python_version"] == "3.12.10"
        manifest_paths = {item["path"] for item in package_manifest["files"]}
        assert "runtime/python/python.exe" in manifest_paths
        assert "runtime/graphviz/bin/dot.exe" in manifest_paths
        assert not any("/generators/" in name or "/desktop/" in name or "/wheelhouse/" in name for name in names)
        install_root = tmp_path / "installed"
        bundle.extractall(install_root)

    installed_skill = install_root / "huawei-doc-workflow"
    for item in package_manifest["files"]:
        file = installed_skill / item["path"]
        assert file.stat().st_size == item["bytes"]
        assert _sha256(file) == item["sha256"]
    outside_cwd = tmp_path / "outside"
    outside_cwd.mkdir()
    runtime = run_skill(installed_skill, "--verify-runtime", cwd=outside_cwd)
    _assert_ok(runtime)
    doctor = run_skill(installed_skill, "doctor", "--json", cwd=outside_cwd)
    _assert_ok(doctor)
    assert json.loads(doctor.stdout)["pass"] is True

    run_dir = outside_cwd / "run"
    validated = run_skill(
        installed_skill,
        "validate",
        "--target",
        "word",
        "--draft",
        str(installed_skill / "scripts" / "engine" / "examples" / "word_valid_01_plain.json"),
        "--output-dir",
        str(run_dir),
        cwd=outside_cwd,
    )
    _assert_ok(validated)
    finalized = run_skill(
        installed_skill,
        "finalize",
        "--target",
        "word",
        "--ir",
        str(run_dir / "validated_ir.json"),
        "--output-dir",
        str(run_dir),
        cwd=outside_cwd,
    )
    _assert_ok(finalized)
    assert (run_dir / "document.docx").is_file()
