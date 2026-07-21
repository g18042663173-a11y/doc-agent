from __future__ import annotations

import sys
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_check_docx_reports_unreadable_file(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = tmp_path / "bad.docx"
    path.write_text("not a docx", encoding="utf-8")

    report = check_docx(path)

    assert report.summary["pass"] is False
    assert report.items[0].code == "E001"


def test_check_docx_reports_empty_document(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = tmp_path / "empty.docx"
    Document().save(path)

    report = check_docx(path)

    assert report.summary["pass"] is False
    assert report.items[0].code == "E006"


def test_check_docx_reports_missing_classification_footer(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = tmp_path / "missing-classification.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(path)

    report = check_docx(path, classification="内部公开")

    assert report.summary["pass"] is False
    assert "E002" in [item.code for item in report.items]


def test_check_docx_accepts_theme_compliant_renderer_output(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.lint.docx_lint import check_docx
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "合规", "classification": "HUAWEI CONFIDENTIAL"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "合规标题"},
                {"type": "paragraph", "text": "正文"},
                {"type": "paragraph", "style": "quote", "text": "引用"},
                {"type": "paragraph", "style": "note", "text": "提示"},
                {"type": "code_block", "language": "c", "code": "if (ready) {\n    send();\n}"},
                {"type": "table", "header": ["项", "值"], "rows": [["主题", "通过"], ["表格", "通过"]]},
                {"type": "image_placeholder", "ref": "arch.png", "caption": "图1: 架构图"},
            ],
        }
    )
    path = render_word_ir(ir, tmp_path / "ok.docx")

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is True
    assert report.items == []


def test_check_docx_rejects_legacy_theme_and_plain_image_placeholder(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = tmp_path / "legacy.docx"
    document = Document()
    document.add_heading("旧蓝标题", level=1)
    document.add_paragraph("正文")
    document.add_paragraph("[图片占位] 图1: 架构图")
    document.sections[0].footer.paragraphs[0].text = "HUAWEI CONFIDENTIAL"
    document.save(path)

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")
    codes = [item.code for item in report.items]

    assert report.summary["pass"] is False
    assert "E004" in codes


def test_check_docx_rejects_non_theme_font_or_size(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = tmp_path / "bad-font.docx"
    document = Document()
    run = document.add_paragraph().add_run("不合规字体")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    document.save(path)

    report = check_docx(path)

    assert report.summary["pass"] is False
    assert "HW-E02" in [item.code for item in report.items]


def test_check_docx_rejects_code_block_font_or_preserved_spacing_drift(tmp_path: Path) -> None:
    path = _render_compliant_docx(
        tmp_path,
        blocks=[{"type": "code_block", "code": "if (ready) {\n    send();\n}"}],
    )
    _rewrite_docx_xml(path, lambda text: text.replace("Consolas", "Arial").replace('xml:space="preserve"', ""))

    from app.lint.docx_lint import check_docx

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is False
    assert "E004" in [item.code for item in report.items]


def test_check_docx_rejects_table_header_or_border_drift(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = _render_compliant_docx(
        tmp_path,
        blocks=[
            {"type": "heading", "level": 1, "text": "表格"},
            {"type": "table", "header": ["项", "值"], "rows": [["颜色", "错误"]]},
        ],
    )
    _rewrite_docx_xml(path, lambda text: text.replace('w:fill="C7000B"', 'w:fill="FFFFFF"'))

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is False
    assert "E004" in [item.code for item in report.items]


def test_check_docx_rejects_quote_without_left_border(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = _render_compliant_docx(
        tmp_path,
        blocks=[{"type": "paragraph", "style": "quote", "text": "引用"}],
    )
    _rewrite_docx_xml(path, lambda text: text.replace("<w:left", "<w:right", 1))

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is False
    assert "E004" in [item.code for item in report.items]


def test_check_docx_rejects_note_without_theme_fill(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = _render_compliant_docx(
        tmp_path,
        blocks=[{"type": "paragraph", "style": "note", "text": "提示"}],
    )
    _rewrite_docx_xml(path, lambda text: text.replace('w:fill="F5F5F5"', 'w:fill="FFFFFF"', 1))

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is False
    assert "E004" in [item.code for item in report.items]


def test_check_docx_rejects_footer_without_page_field(tmp_path: Path) -> None:
    from app.lint.docx_lint import check_docx

    path = _render_compliant_docx(tmp_path, blocks=[{"type": "paragraph", "text": "正文"}])
    _rewrite_docx_xml(path, lambda text: text.replace("PAGE", ""))

    report = check_docx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["pass"] is False
    assert "HW-E01" in [item.code for item in report.items]


def _render_compliant_docx(tmp_path: Path, *, blocks: list[dict]) -> Path:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "合规", "classification": "HUAWEI CONFIDENTIAL"},
            "blocks": blocks,
        }
    )
    return render_word_ir(ir, tmp_path / "compliant.docx")


def _rewrite_docx_xml(path: Path, replace_text) -> None:
    temp_path = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                data = replace_text(data.decode("utf-8")).encode("utf-8")
            target.writestr(info, data)
    temp_path.replace(path)
