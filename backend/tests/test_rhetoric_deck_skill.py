from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

from lxml import etree
import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[2]
SKILL = ROOT / "skills/rhetoric-deck-workflow"
RDW = SKILL / "bin/rdw.py"


def _rdw_python() -> str:
    candidates = [sys.executable, shutil.which("python")]
    for candidate in dict.fromkeys(item for item in candidates if item):
        probe = subprocess.run(
            [candidate, "-c", "import pptx,jsonschema,lxml"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if probe.returncode == 0:
            return candidate
    pytest.skip("rhetoric-deck-workflow runtime dependencies are not installed")


def _run(*args: str | Path, expected: int = 0) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [_rdw_python(), "-X", "utf8", str(RDW), *(str(arg) for arg in args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode == expected, result.stderr or result.stdout
    return result


def _material(path: Path) -> Path:
    path.write_text(
        """# 用户技术方案

背景：现网流程存在重复录入和定位慢的痛点，目标是缩短处理链路。
- 方案 A：保留现有接口，优点是改造小，缺点是验证链长。
- 方案 B：增加统一校验层，优点是问题提前暴露，风险是接口适配。
- 推荐方案 B，因为校验职责集中且回归范围清晰。

实现流程：解析输入、校验契约、生成产物、执行审计。
组件包括解析器、校验器和交付模块，约束是 30 秒内完成 20 个样例。
测试覆盖正常输入、缺失字段、异常包和超长内容，重点验证非法输入不产生结果。
问题是字段不一致；现象是下游失败；分析为缺少统一契约；措施是 Schema 前置。
阶段一投入接口梳理，产出冻结契约；阶段二投入自动化，产出回归证据。
当前目标完成，进展为全链路通过，风险是目标机字体差异，下一步进行人工视觉终审。
密级：HUAWEI CONFIDENTIAL
""",
        encoding="utf-8",
    )
    return path


def _content_from_fill_pack(workdir: Path) -> dict:
    skeleton = json.loads((workdir / "fill_pack/skeleton.json").read_text(encoding="utf-8"))
    material = json.loads((workdir / "fill_pack/material.json").read_text(encoding="utf-8"))
    evidence = material["evidence"]
    facts = [
        (value, item["id"])
        for value in ("用户技术方案", "统一校验", "解析输入", "校验契约", "生成产物", "执行审计", "冻结契约", "回归证据")
        for item in evidence if value in item["text"]
    ]
    if not facts:
        facts = [(item["text"].splitlines()[0].strip("# ")[:8], item["id"]) for item in evidence if item["text"].strip()]
    assert facts
    pages = []
    counter = 0
    for page in skeleton["pages"]:
        slots = []
        for slot in page["slots"]:
            value, evidence_ref = facts[0 if slot["type"] == "title" else counter % len(facts)]
            if (slot.get("shape_ref") or "").startswith("part_"):
                value = material["classification"]
                evidence_ref = next(item["id"] for item in evidence if value in item["text"])
            minimum = int(slot["cardinality"]["min"])
            if minimum > 1:
                value = [value] * minimum
            slots.append({"slot_id": slot["slot_id"], "value": value, "evidence_refs": [evidence_ref]})
            counter += 1
        pages.append({"page_id": page["page_id"], "slots": slots})
    return {"format": "fill_content", "version": "2.0", "pages": pages}


def _complete_skeleton(workdir: Path, overrides: dict) -> dict:
    """Keep explicit fixture roles while retaining every extracted page, including empty pages."""
    draft = json.loads((workdir / "extract_pack/skeleton_draft.json").read_text(encoding="utf-8"))
    replacements = {page["page_id"]: page for page in overrides["pages"]}
    draft["deck_pattern"] = overrides["deck_pattern"]
    merged = []
    for page in draft["pages"]:
        replacement = deepcopy(replacements.get(page["page_id"], page))
        bound = {slot["shape_ref"] for slot in replacement["slots"]}
        for slot in page["slots"]:
            if slot["shape_ref"] not in bound:
                added = deepcopy(slot)
                added["slot_id"] = f"s{len(replacement['slots']) + 1}"
                replacement["slots"].append(added)
        merged.append(replacement)
    draft["pages"] = merged
    assert draft["page_policy"] == "preserve"
    assert draft["page_count"] == len(draft["pages"])
    return draft


def _source_deck(path: Path) -> tuple[Path, dict]:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[0]) if len(presentation.slides) else None
    refs = []
    source_texts = [
        ("源件高度敏感背景说明不可进入产物", "源件高度敏感痛点数据不可进入产物"),
        ("源件高度敏感状态说明不可进入产物", "源件高度敏感风险数据不可进入产物"),
    ]
    for title, body in source_texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_shape = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        title_shape.text = title
        body_shape = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
        body_shape.text = body
        refs.append((title_shape.shape_id, body_shape.shape_id))
    presentation.save(path)
    pages = []
    patterns = [
        ("context_pain", "context_pain_goal", "context.background"),
        ("status_progress", "goal_progress_risk_next", "status.progress"),
    ]
    for index, ((title_ref, body_ref), (pattern, flow, semantic)) in enumerate(zip(refs, patterns), start=1):
        pages.append(
            {
                "page_id": f"p{index:02d}", "page_pattern": pattern, "argument_flow": flow, "confidence": 0.9,
                "structure": {"items": {"count": 2, "expandable": True, "max": 6}},
                "slots": [
                    {"slot_id": "s1", "semantic": "page.title", "type": "title", "cardinality": {"min": 1, "max": 1}, "capacity": {"chars_cjk": 40, "lines": 2}, "required": True, "shape_ref": f"sp_{title_ref}"},
                    {"slot_id": "s2", "semantic": semantic, "type": "bullet_list", "cardinality": {"min": 1, "max": 6}, "capacity": {"chars_cjk": 120, "lines": 7}, "required": True, "shape_ref": f"sp_{body_ref}"},
                ],
                "emphasis": [],
                "diagram": {"type": "none", "groups": 0, "nodes_per_group": 0, "flow": "none", "annotations": 0, "labels": None},
                "structural_labels": {"_display": []},
            }
        )
    skeleton = {"format": "deck_skeleton", "version": "1.0", "source_kind": "extracted", "deck_pattern": "report_progress", "page_count": 2, "pages": pages}
    return path, skeleton


def test_doctor_and_builtin_deck_ir_end_to_end(tmp_path: Path) -> None:
    doctor = json.loads(_run("doctor", "--json").stdout)
    assert doctor["ok"] is True
    assert doctor["external_dependencies"] == ["python-pptx", "jsonschema", "lxml", "pydantic", "pypdf", "fonttools"]

    workdir = tmp_path / "deck-ir"
    workdir.mkdir()
    material = _material(tmp_path / "material.md")
    _run("plan", "--workdir", workdir, "--pattern", "review_solution", "--material", material, "--render-mode", "deck-ir")
    content = _content_from_fill_pack(workdir)
    content_path = tmp_path / "content.json"
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    invalid = {"format": "fill_content", "version": "1.0", "pages": [{"page_id": content["pages"][0]["page_id"], "slots": [content["pages"][0]["slots"][0]]}]}
    invalid_path = tmp_path / "invalid-content.json"
    invalid_path.write_text(json.dumps(invalid, ensure_ascii=False), encoding="utf-8")
    invalid_out = tmp_path / "invalid-out"
    invalid_result = _run("finalize", "--workdir", workdir, "--content", invalid_path, "--out", invalid_out, expected=50)
    assert '"code": "RD-E030"' in invalid_result.stderr
    assert not invalid_out.exists()
    for label, replacement in (
        ("missing", {"slot_id": content["pages"][0]["slots"][0]["slot_id"], "status": "missing", "reason": "用户材料没有这个事实"}),
        ("invalid-reference", {**content["pages"][0]["slots"][0], "evidence_refs": ["d99:b9999"]}),
        ("no-reference", {key: value for key, value in content["pages"][0]["slots"][0].items() if key != "evidence_refs"}),
    ):
        incomplete = deepcopy(content)
        incomplete["pages"][0]["slots"][0] = replacement
        incomplete_path = tmp_path / f"{label}.json"
        incomplete_path.write_text(json.dumps(incomplete, ensure_ascii=False), encoding="utf-8")
        incomplete_out = tmp_path / label
        blocked = _run("finalize", "--workdir", workdir, "--content", incomplete_path, "--out", incomplete_out, expected=50)
        assert '"code": "RD-E030"' in blocked.stderr
        assert not incomplete_out.exists()
    outdir = tmp_path / "out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    deck = json.loads((outdir / "deck_ir.json").read_text(encoding="utf-8"))
    assert deck["ir_version"] == "2.2"
    assert not (outdir / "deck.pptx").exists()
    assert json.loads((outdir / "manifest.json").read_text(encoding="utf-8"))["artifact"] == "deck_ir.json"


def test_extract_seal_source_shell_and_leak_gate(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "source.pptx")
    workdir = tmp_path / "source-shell"
    _run("extract", "--source", source, "--out", workdir)
    candidates = json.loads((workdir / "extract_pack/slot_candidates.json").read_text(encoding="utf-8"))
    assert candidates["pages"][0]["candidates"]
    assert "源件高度敏感" in (workdir / "extract_pack/pages/p01.json").read_text(encoding="utf-8")
    skeleton_path = tmp_path / "skeleton.json"
    skeleton_path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    sealed = json.loads(_run("seal", "--workdir", workdir, "--skeleton", skeleton_path).stdout)
    assert sealed["deleted"] == []
    assert set(sealed["internal_only"]) == {"source.pptx", "extract_pack/pages", "sealed/shell.pptx"}
    assert (workdir / "source.pptx").is_file()
    assert (workdir / "source_reference.pptx").read_bytes() == source.read_bytes()
    assert (workdir / "extract_pack/pages").is_dir()
    assert "源件高度敏感" in (workdir / "extract_pack/pages/p01.json").read_text(encoding="utf-8")
    sealed_text = (workdir / "sealed/skeleton.sealed.json").read_text(encoding="utf-8")
    fingerprint_text = (workdir / "sealed/source_ngrams.json").read_text(encoding="utf-8")
    assert "源件高度敏感" not in sealed_text
    assert "源件高度敏感" not in fingerprint_text
    shell = Presentation(workdir / "sealed/shell.pptx")
    assert all(not shape.text.strip() for slide in shell.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))

    material = _material(tmp_path / "source-material.md")
    _run("plan", "--workdir", workdir, "--skeleton", workdir / "sealed/skeleton.sealed.json", "--material", material, "--render-mode", "source-shell")
    content = _content_from_fill_pack(workdir)
    content_path = tmp_path / "source-content.json"
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "source-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    output_text = "\n".join(shape.text for slide in Presentation(outdir / "deck.pptx").slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert "用户技术方案" in output_text
    assert "源件高度敏感" not in output_text

    blocked_workdir = tmp_path / "blocked"
    source2, skeleton2 = _source_deck(tmp_path / "source2.pptx")
    _run("extract", "--source", source2, "--out", blocked_workdir)
    skeleton2_path = tmp_path / "skeleton2.json"
    skeleton2_path.write_text(json.dumps(skeleton2, ensure_ascii=False), encoding="utf-8")
    _run("seal", "--workdir", blocked_workdir, "--skeleton", skeleton2_path)
    _run("plan", "--workdir", blocked_workdir, "--skeleton", blocked_workdir / "sealed/skeleton.sealed.json", "--material", material, "--render-mode", "source-shell")
    leaked = _content_from_fill_pack(blocked_workdir)
    leaked["pages"][0]["slots"][0]["value"] = "源件高度敏感背景说明不可进入产物"
    leaked_path = tmp_path / "leaked.json"
    leaked_path.write_text(json.dumps(leaked, ensure_ascii=False), encoding="utf-8")
    blocked_out = tmp_path / "blocked-out"
    result = _run("finalize", "--workdir", blocked_workdir, "--content", leaked_path, "--out", blocked_out, expected=60)
    assert '"code": "RD-E040"' in result.stderr
    leak = json.loads((blocked_out / "leak_report.json").read_text(encoding="utf-8"))
    assert leak["hits"]
    assert leak["hits"][0]["loc"]
    assert not (blocked_out / "deck.pptx").exists()


def test_seal_rejects_cjk_outside_display_without_destroying_extract_pack(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "source.pptx")
    workdir = tmp_path / "bad"
    _run("extract", "--source", source, "--out", workdir)
    skeleton["pages"][0]["slots"][0]["semantic"] = "源文残留"
    path = tmp_path / "bad-skeleton.json"
    path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    result = _run("seal", "--workdir", workdir, "--skeleton", path, expected=30)
    assert '"code": "RD-E010"' in result.stderr or '"code": "RD-E011"' in result.stderr
    assert (workdir / "source.pptx").is_file()
    assert (workdir / "extract_pack/pages").is_dir()
    assert not (workdir / "sealed/shell.pptx").exists()


@pytest.mark.parametrize(
    "office_path",
    [
        ROOT / "samples/input/需求说明.docx",
        ROOT / "samples/input/parser_samples/xlsx_sample_01.xlsx",
        ROOT / "samples/input/项目汇报.pptx",
    ],
)
def test_plan_parses_supported_office_material_formats(tmp_path: Path, office_path: Path) -> None:
    rich_material = _material(tmp_path / "support.md")
    workdir = tmp_path / office_path.suffix.lstrip(".")
    workdir.mkdir()
    _run(
        "plan", "--workdir", workdir, "--pattern", "review_solution",
        "--material", office_path, "--material", rich_material, "--render-mode", "deck-ir",
    )
    material = json.loads((workdir / "fill_pack/material.json").read_text(encoding="utf-8"))
    kinds = {document["kind"] for document in material["documents"]}
    assert office_path.suffix.lstrip(".") in kinds


_TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
)


def _zip_parts(path: Path, prefix: str) -> list[str]:
    with zipfile.ZipFile(path) as package:
        return [name for name in package.namelist() if name.startswith(prefix) and not name.endswith("/")]


def _media_payloads(path: Path) -> list[bytes]:
    with zipfile.ZipFile(path) as package:
        return [package.read(name) for name in _zip_parts(path, "ppt/media/")]


def _xml_blob(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as package:
        return "\n".join(package.read(name).decode("utf-8", errors="replace") for name in _zip_parts(path, prefix))


def _blank_page_fields(page_id: str, pattern: str, flow: str, slots: list[dict]) -> dict:
    return {
        "page_id": page_id,
        "page_pattern": pattern,
        "argument_flow": flow,
        "confidence": 0.9,
        "structure": {"items": {"count": 1, "expandable": False, "max": 1}},
        "slots": slots,
        "emphasis": [],
        "diagram": {"type": "none", "groups": 0, "nodes_per_group": 0, "flow": "none", "annotations": 0, "labels": None},
        "structural_labels": {"_display": []},
    }


def _cover_and_screenshot_deck(path: Path) -> tuple[Path, int]:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[0]) if len(presentation.slides) else None
    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_shape = cover.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
    title_shape.text = "源件封面口号不可进入产物"
    shot = presentation.slides.add_slide(presentation.slide_layouts[6])
    shot_title = shot.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.6))
    shot_title.text = "源件截图页标题不可进入产物"
    png = path.with_name("shot.png")
    png.write_bytes(_TINY_PNG)
    shot.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(4), Inches(3))
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return path, title_shape.shape_id


def test_extract_recovers_screenshot_and_preserves_empty_page_without_fake_slot(tmp_path: Path) -> None:
    source, title_ref = _cover_and_screenshot_deck(tmp_path / "mixed.pptx")
    workdir = tmp_path / "skip"
    payload = json.loads(_run("extract", "--source", source, "--out", workdir).stdout)
    hints = json.loads((workdir / "extract_pack/skip_hints.json").read_text(encoding="utf-8"))
    skipped = {item["page_id"]: item["reason"] for item in hints["pages"]}
    recoveries = json.loads((workdir / "extract_pack/recoveries.json").read_text(encoding="utf-8"))
    recovered_actions = {
        page["page_id"]: {action["action"] for action in page["actions"]}
        for page in recoveries["pages"]
    }
    assert skipped == {}
    assert "screenshot_kept" in recovered_actions["p02"]
    assert payload["skipped_pages"] == hints["pages"]
    draft = json.loads((workdir / "extract_pack/skeleton_draft.json").read_text(encoding="utf-8"))
    assert draft["page_policy"] == "preserve"
    assert [page["page_id"] for page in draft["pages"]] == ["p01", "p02", "p03"]
    assert draft["pages"][2]["slots"] == []

    illegal = {
        "format": "deck_skeleton",
        "version": "1.0",
        "source_kind": "extracted",
        "deck_pattern": "report_progress",
        "page_count": 1,
        "pages": [
            _blank_page_fields(
                "p03",
                "context_pain",
                "context_pain_goal",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": f"sp_{title_ref}",
                    }
                ],
            )
        ],
    }
    illegal_path = tmp_path / "illegal.json"
    illegal_path.write_text(json.dumps(illegal), encoding="utf-8")
    result = _run("seal", "--workdir", workdir, "--skeleton", illegal_path, expected=30)
    assert '"code": "RD-E010"' in result.stderr
    assert (workdir / "extract_pack/pages").is_dir()


def test_struct_cover_source_shell_fills_from_material(tmp_path: Path) -> None:
    source, title_ref = _cover_and_screenshot_deck(tmp_path / "cover.pptx")
    workdir = tmp_path / "cover-run"
    _run("extract", "--source", source, "--out", workdir)
    skeleton = {
        "format": "deck_skeleton",
        "version": "1.0",
        "source_kind": "extracted",
        "deck_pattern": "pitch_proposal",
        "page_count": 1,
        "pages": [
            _blank_page_fields(
                "p01",
                "struct_cover",
                "cover_title_subtitle",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": f"sp_{title_ref}",
                    }
                ],
            )
        ],
    }
    skeleton_path = tmp_path / "cover-skeleton.json"
    skeleton_path.write_text(json.dumps(skeleton), encoding="utf-8")
    omitted = _run("seal", "--workdir", workdir, "--skeleton", skeleton_path, expected=30)
    assert '"code": "RD-E010"' in omitted.stderr
    assert not (workdir / "sealed/shell.pptx").exists()
    skeleton = _complete_skeleton(workdir, skeleton)
    skeleton_path.write_text(json.dumps(skeleton), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    skip_ids = [item["page_id"] for item in json.loads((workdir / "sealed/skip_report.json").read_text(encoding="utf-8"))["pages"]]
    assert skip_ids == []
    material = _material(tmp_path / "cover-material.md")
    _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        material,
        "--render-mode",
        "source-shell",
    )
    content_path = tmp_path / "cover-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "cover-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    deck = Presentation(outdir / "deck.pptx")
    assert len(deck.slides) == 3
    output_text = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "用户技术方案" in output_text
    assert "源件封面口号" not in output_text
    manifest = json.loads((outdir / "manifest.json").read_text(encoding="utf-8"))
    skip_report = json.loads((outdir / "skip_report.json").read_text(encoding="utf-8"))
    assert skip_report["pages"] == []
    assert manifest["skipped_pages"] == []
    assert manifest["page_ids"] == ["p01", "p02", "p03"]
    assert len(deck.slides[2].shapes) == 0
    assert _TINY_PNG in _media_payloads(outdir / "deck.pptx")


def test_extract_strips_animation_without_skipping_text_pages(tmp_path: Path) -> None:
    source, _skeleton = _source_deck(tmp_path / "animated.pptx")
    presentation = Presentation(source)
    etree.SubElement(
        presentation.slides[0].element,
        "{http://schemas.openxmlformats.org/presentationml/2006/main}timing",
    )
    presentation.save(source)
    workdir = tmp_path / "anim"
    payload = json.loads(_run("extract", "--source", source, "--out", workdir).stdout)
    skipped = {item["page_id"] for item in payload["skipped_pages"]}
    assert "p01" not in skipped
    recoveries = json.loads((workdir / "extract_pack/recoveries.json").read_text(encoding="utf-8"))
    actions = {
        action["action"]
        for page in recoveries["pages"]
        if page["page_id"] == "p01"
        for action in page["actions"]
    }
    assert "animation_stripped" in actions
    assert "<p:timing" not in Presentation(workdir / "source.pptx").slides[0].element.xml


def test_recovered_screenshot_page_fills_from_material(tmp_path: Path) -> None:
    source, _title_ref = _cover_and_screenshot_deck(tmp_path / "shot-fill.pptx")
    workdir = tmp_path / "shot-run"
    _run("extract", "--source", source, "--out", workdir)
    page = json.loads((workdir / "extract_pack/pages/p02.json").read_text(encoding="utf-8"))
    title_ref = next(item["shape_ref"] for item in page["elements"] if "源件截图页标题" in (item.get("text") or ""))
    assert any(item.get("kind") == "picture" for item in page["elements"])
    skeleton = {
        "format": "deck_skeleton",
        "version": "1.0",
        "source_kind": "extracted",
        "deck_pattern": "pitch_proposal",
        "page_count": 1,
        "pages": [
            _blank_page_fields(
                "p02",
                "context_pain",
                "context_pain_goal",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": title_ref,
                    }
                ],
            )
        ],
    }
    skeleton_path = tmp_path / "shot-skeleton.json"
    skeleton = _complete_skeleton(workdir, skeleton)
    skeleton_path.write_text(json.dumps(skeleton), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    assert _TINY_PNG in _media_payloads(workdir / "sealed/shell.pptx")
    material = _material(tmp_path / "shot-material.md")
    _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        material,
        "--render-mode",
        "source-shell",
    )
    content_path = tmp_path / "shot-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "shot-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    deck = Presentation(outdir / "deck.pptx")
    assert len(deck.slides) == 3
    output_text = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "用户技术方案" in output_text
    assert "源件截图页标题" not in output_text
    assert _TINY_PNG in _media_payloads(outdir / "deck.pptx")
    assert json.loads((outdir / "recoveries.json").read_text(encoding="utf-8"))["pages"]


def test_recovered_smartart_page_fills_from_material(tmp_path: Path) -> None:
    source = _smartart_deck(tmp_path / "smartart.pptx")
    workdir = tmp_path / "smart-run"
    _run("extract", "--source", source, "--out", workdir)
    page = json.loads((workdir / "extract_pack/pages/p01.json").read_text(encoding="utf-8"))
    texts = [item.get("text", "") for item in page["elements"]]
    assert any("源件节点甲不可进入产物" in text for text in texts)
    recoveries = json.loads((workdir / "extract_pack/recoveries.json").read_text(encoding="utf-8"))
    assert any(
        action["action"] == "smartart_kept"
        for item in recoveries["pages"]
        for action in item["actions"]
    )
    title_ref = next(item["shape_ref"] for item in page["elements"] if "源件SmartArt标题" in (item.get("text") or ""))
    node_items = [item for item in page["elements"] if item.get("kind") == "smartart_node" and (item.get("text") or "").strip()]
    assert len(node_items) >= 2
    skeleton = {
        "format": "deck_skeleton",
        "version": "1.0",
        "source_kind": "extracted",
        "deck_pattern": "review_solution",
        "page_count": 1,
        "pages": [
            _blank_page_fields(
                "p01",
                "method_walkthrough",
                "steps_compare_benefit",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": title_ref,
                    },
                    *[
                        {
                            "slot_id": f"s{index + 2}",
                            "semantic": f"method.node_{index + 1}",
                            "type": "node_label",
                            "cardinality": {"min": 1, "max": 1},
                            "capacity": {"chars_cjk": 40, "lines": 2},
                            "required": True,
                            "shape_ref": item["shape_ref"],
                        }
                        for index, item in enumerate(node_items)
                    ],
                ],
            )
        ],
    }
    skeleton_path = tmp_path / "smart-skeleton.json"
    skeleton_path.write_text(json.dumps(skeleton), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    assert _zip_parts(workdir / "sealed/shell.pptx", "ppt/diagrams/")
    material = _material(tmp_path / "smart-material.md")
    _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        material,
        "--render-mode",
        "source-shell",
    )
    content_path = tmp_path / "smart-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "smart-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    output_text = "\n".join(
        shape.text
        for slide in Presentation(outdir / "deck.pptx").slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "用户技术方案" in output_text
    assert "源件节点甲" not in output_text
    assert "源件SmartArt标题" not in output_text
    diagram_xml = _xml_blob(outdir / "deck.pptx", "ppt/diagrams/")
    assert "graphicFrame" in Presentation(outdir / "deck.pptx").slides[0].element.xml
    assert "统一校验" in diagram_xml and "解析输入" in diagram_xml
    assert "源件节点甲" not in diagram_xml
    assert "源件节点乙" not in diagram_xml


def test_source_shell_keeps_smartart_and_screenshot_pixels(tmp_path: Path) -> None:
    source = tmp_path / "visual.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    if len(presentation.slides._sldIdLst):
        presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[0])
    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover_title = cover.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
    cover_title.text = "源件封面口号不可进入产物"
    shot = presentation.slides.add_slide(presentation.slide_layouts[6])
    shot_title = shot.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.6))
    shot_title.text = "源件截图页标题不可进入产物"
    png = tmp_path / "visual.png"
    png.write_bytes(_TINY_PNG)
    shot.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(4), Inches(3))
    smart = presentation.slides.add_slide(presentation.slide_layouts[6])
    smart_title = smart.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
    smart_title.text = "源件SmartArt标题不可进入产物"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    _inject_smartart(source, ["源件节点甲不可进入产物", "源件节点乙不可进入产物"], slide_index=2)

    workdir = tmp_path / "visual-run"
    payload = json.loads(_run("extract", "--source", source, "--out", workdir).stdout)
    skipped = {item["page_id"]: item["reason"] for item in payload["skipped_pages"]}
    recoveries = json.loads((workdir / "extract_pack/recoveries.json").read_text(encoding="utf-8"))
    actions = {
        page["page_id"]: {action["action"] for action in page["actions"]}
        for page in recoveries["pages"]
    }
    assert skipped == {}
    assert "screenshot_kept" in actions["p02"]
    assert "smartart_kept" in actions["p03"]

    p01 = json.loads((workdir / "extract_pack/pages/p01.json").read_text(encoding="utf-8"))
    p02 = json.loads((workdir / "extract_pack/pages/p02.json").read_text(encoding="utf-8"))
    p03 = json.loads((workdir / "extract_pack/pages/p03.json").read_text(encoding="utf-8"))
    skeleton = {
        "format": "deck_skeleton",
        "version": "1.0",
        "source_kind": "extracted",
        "deck_pattern": "review_solution",
        "page_count": 3,
        "pages": [
            _blank_page_fields(
                "p01",
                "struct_cover",
                "cover_title_subtitle",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": next(item["shape_ref"] for item in p01["elements"] if "源件封面口号" in (item.get("text") or "")),
                    }
                ],
            ),
            _blank_page_fields(
                "p02",
                "context_pain",
                "context_pain_goal",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": next(item["shape_ref"] for item in p02["elements"] if "源件截图页标题" in (item.get("text") or "")),
                    }
                ],
            ),
            _blank_page_fields(
                "p03",
                "method_walkthrough",
                "steps_compare_benefit",
                [
                    {
                        "slot_id": "s1",
                        "semantic": "page.title",
                        "type": "title",
                        "cardinality": {"min": 1, "max": 1},
                        "capacity": {"chars_cjk": 40, "lines": 2},
                        "required": True,
                        "shape_ref": next(item["shape_ref"] for item in p03["elements"] if "源件SmartArt标题" in (item.get("text") or "")),
                    },
                    *[
                        {
                            "slot_id": f"s{index + 2}",
                            "semantic": f"method.node_{index + 1}",
                            "type": "node_label",
                            "cardinality": {"min": 1, "max": 1},
                            "capacity": {"chars_cjk": 40, "lines": 2},
                            "required": True,
                            "shape_ref": item["shape_ref"],
                        }
                        for index, item in enumerate(
                            [element for element in p03["elements"] if element.get("kind") == "smartart_node" and (element.get("text") or "").strip()]
                        )
                    ],
                ],
            ),
        ],
    }
    skeleton_path = tmp_path / "visual-skeleton.json"
    skeleton = _complete_skeleton(workdir, skeleton)
    skeleton_path.write_text(json.dumps(skeleton), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    shell = workdir / "sealed/shell.pptx"
    assert _TINY_PNG in _media_payloads(shell)
    assert _zip_parts(shell, "ppt/diagrams/")
    assert "源件节点甲" not in _xml_blob(shell, "ppt/")
    material = _material(tmp_path / "visual-material.md")
    _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        material,
        "--render-mode",
        "source-shell",
    )
    content_path = tmp_path / "visual-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "visual-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    deck = outdir / "deck.pptx"
    assert len(Presentation(deck).slides) == 4
    assert len(Presentation(deck).slides[3].shapes) == 0
    output_text = "\n".join(
        shape.text
        for slide in Presentation(deck).slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "用户技术方案" in output_text
    assert "源件封面口号" not in output_text
    assert "源件截图页标题" not in output_text
    assert "源件SmartArt标题" not in output_text
    assert _TINY_PNG in _media_payloads(deck)
    assert "graphicFrame" in Presentation(deck).slides[2].element.xml
    diagram_xml = _xml_blob(deck, "ppt/diagrams/")
    assert any(text in diagram_xml for text in ("统一校验", "解析输入", "校验契约", "生成产物", "执行审计"))
    assert "源件节点甲" not in diagram_xml
    leaked = json.loads(content_path.read_text(encoding="utf-8"))
    leaked["pages"][2]["slots"][1]["value"] = "源件节点甲不可进入产物"
    leaked_path = tmp_path / "visual-leaked.json"
    leaked_path.write_text(json.dumps(leaked, ensure_ascii=False), encoding="utf-8")
    blocked = tmp_path / "visual-blocked"
    result = _run("finalize", "--workdir", workdir, "--content", leaked_path, "--out", blocked, expected=60)
    assert '"code": "RD-E040"' in result.stderr
    assert not (blocked / "deck.pptx").exists()


def _smartart_deck(path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[0]) if len(presentation.slides) else None
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_shape = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
    title_shape.text = "源件SmartArt标题不可进入产物"
    presentation.save(path)
    _inject_smartart(path, ["源件节点甲不可进入产物", "源件节点乙不可进入产物"])
    return path


def _inject_smartart(path: Path, texts: list[str], *, slide_index: int = 0) -> None:
    pkg_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    dgm_ns = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    with zipfile.ZipFile(path) as incoming:
        contents = {name: incoming.read(name) for name in incoming.namelist()}
    slide_name = sorted(name for name in contents if name.startswith("ppt/slides/slide") and name.endswith(".xml"))[slide_index]
    rels_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
    rels = etree.fromstring(contents[rels_name])
    rid = "rIdSA1"
    etree.SubElement(
        rels,
        f"{{{pkg_ns}}}Relationship",
        Id=rid,
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData",
        Target="../diagrams/data1.xml",
    )
    slide = etree.fromstring(contents[slide_name])
    tree = slide.find(f".//{{{p_ns}}}spTree")
    assert tree is not None
    frame = etree.SubElement(tree, f"{{{p_ns}}}graphicFrame")
    nv = etree.SubElement(frame, f"{{{p_ns}}}nvGraphicFramePr")
    etree.SubElement(nv, f"{{{p_ns}}}cNvPr", id="99", name="Diagram")
    etree.SubElement(nv, f"{{{p_ns}}}cNvGraphicFramePr")
    etree.SubElement(nv, f"{{{p_ns}}}nvPr")
    xfrm = etree.SubElement(frame, f"{{{p_ns}}}xfrm")
    etree.SubElement(xfrm, f"{{{a_ns}}}off", x="914400", y="1371600")
    etree.SubElement(xfrm, f"{{{a_ns}}}ext", cx="8229600", cy="4114800")
    graphic = etree.SubElement(frame, f"{{{a_ns}}}graphic")
    data = etree.SubElement(
        graphic,
        f"{{{a_ns}}}graphicData",
        uri="http://schemas.openxmlformats.org/drawingml/2006/diagram",
    )
    rel_ids = etree.SubElement(data, f"{{{dgm_ns}}}relIds")
    rel_ids.set(f"{{{r_ns}}}dm", rid)
    rel_ids.set(f"{{{r_ns}}}lo", rid)
    rel_ids.set(f"{{{r_ns}}}qs", rid)
    rel_ids.set(f"{{{r_ns}}}cs", rid)
    points = "".join(
        f'<dgm:pt modelId="{index}"><dgm:t><a:p><a:r><a:t>{text}</a:t></a:r></a:p></dgm:t></dgm:pt>'
        for index, text in enumerate(texts, start=1)
    )
    data_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<dgm:ptLst>{points}</dgm:ptLst></dgm:dataModel>"
    )
    types = etree.fromstring(contents["[Content_Types].xml"])
    etree.SubElement(
        types,
        f"{{{ct_ns}}}Override",
        PartName="/ppt/diagrams/data1.xml",
        ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
    )
    contents[slide_name] = etree.tostring(slide, xml_declaration=True, encoding="UTF-8", standalone=True)
    contents[rels_name] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    contents["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    contents["ppt/diagrams/data1.xml"] = data_xml.encode("utf-8")
    with zipfile.ZipFile(path, "w") as outgoing:
        for name, data_bytes in contents.items():
            outgoing.writestr(name, data_bytes)


def test_rhetoric_deck_ir_validates_in_huawei_doc_workflow(tmp_path: Path) -> None:
    workdir = tmp_path / "handoff"
    workdir.mkdir()
    material = _material(tmp_path / "handoff-material.md")
    _run("plan", "--workdir", workdir, "--pattern", "review_solution", "--material", material, "--render-mode", "deck-ir")
    content_path = tmp_path / "handoff-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    rhetoric_out = tmp_path / "rhetoric-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", rhetoric_out)
    hw_run = tmp_path / "hw-run"
    hw_run.mkdir()
    workflow = ROOT / "skills/huawei-doc-workflow/scripts/workflow.py"
    validate = subprocess.run(
        [sys.executable, str(workflow), "validate", "--target", "deck", "--draft", str(rhetoric_out / "deck_ir.json"), "--output-dir", str(hw_run)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert validate.returncode == 0, validate.stderr or validate.stdout
    assert (hw_run / "validated_ir.json").is_file()


def test_skill_docs_stay_split_and_install_copies_rhetoric(tmp_path: Path) -> None:
    rhetoric = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    huawei = (ROOT / "skills/huawei-doc-workflow/SKILL.md").read_text(encoding="utf-8")
    assert "Use when explicitly invoked as $rhetoric-deck-workflow" in rhetoric
    assert "$huawei-doc-workflow" in rhetoric
    assert "保留全部源页及顺序" in rhetoric
    assert "evidence_refs" in rhetoric
    assert "slot_candidates.json" in rhetoric
    assert "skeleton_draft.json" in rhetoric
    assert "native_inventory.json" in rhetoric
    assert "status: missing" in rhetoric
    assert "--preview" in rhetoric
    assert "$rhetoric-deck-workflow" in huawei
    assert "master_redraw" in huawei
    assert "install.ps1" in huawei
    assert "run.cmd" in huawei
    assert "runtime/python/python.exe" in huawei
    assert "install_agent_kit.py" not in huawei
    assert "Docker" in huawei
    from scripts.package_rhetoric_deck_skill import install_skill

    target = tmp_path / "codex-skills"
    installed = install_skill(target)
    assert installed == target / "rhetoric-deck-workflow"
    assert (installed / "SKILL.md").is_file()
    assert (installed / "bin/rdw.py").is_file()


def _add_master_footer(path: Path, text: str) -> None:
    presentation = Presentation(path)
    donor = presentation.slides[0]
    shape = donor.shapes.add_textbox(Inches(0.4), Inches(7.15), Inches(6.4), Inches(0.28))
    shape.text = text
    element = deepcopy(shape._element)
    donor.shapes._spTree.remove(shape._element)
    nvpr = element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    if nvpr is not None:
        nvpr.set("id", "4096")
        nvpr.set("name", "FooterClassification")
    tree = presentation.slide_masters[0]._element.find(qn("p:cSld")).find(qn("p:spTree"))
    tree.append(element)
    presentation.save(path)


def _master_text(path: Path) -> str:
    presentation = Presentation(path)
    parts = []
    for master in presentation.slide_masters:
        for shape in master.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text or "")
    return "\n".join(parts)


def _deck_text(path: Path) -> str:
    presentation = Presentation(path)
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text or "")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text or "")
    parts.append(_master_text(path))
    return "\n".join(parts)


def test_seal_keeps_master_footer_and_finalize_rewrites_classification(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "footer-source.pptx")
    _add_master_footer(source, "源件密级内部资料不可进入产物")
    workdir = tmp_path / "footer-run"
    _run("extract", "--source", source, "--out", workdir)
    skeleton_path = tmp_path / "footer-skeleton.json"
    skeleton = _complete_skeleton(workdir, skeleton)
    skeleton_path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    shell_text = _master_text(workdir / "sealed/shell.pptx")
    with zipfile.ZipFile(workdir / "sealed/shell.pptx") as package:
        master = etree.fromstring(package.read("ppt/slideMasters/slideMaster1.xml"))
    footer = master.xpath('.//*[local-name()="sp"][.//*[local-name()="cNvPr"][@id="4096"]]')
    assert len(footer) == 1
    assert all(not (node.text or "").strip() for node in footer[0].xpath('.//*[local-name()="t"]'))
    assert "源件密级内部资料不可进入产物" not in shell_text
    material = tmp_path / "footer-material.md"
    material.write_text(_material(tmp_path / "footer-base.md").read_text(encoding="utf-8").replace("密级：HUAWEI CONFIDENTIAL", "密级：内部公开"), encoding="utf-8")
    _run("plan", "--workdir", workdir, "--skeleton", workdir / "sealed/skeleton.sealed.json", "--material", material, "--render-mode", "source-shell")
    content_path = tmp_path / "footer-content.json"
    content_path.write_text(json.dumps(_content_from_fill_pack(workdir), ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "footer-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    final_text = _deck_text(outdir / "deck.pptx")
    assert "内部公开" in final_text
    assert "源件密级内部资料不可进入产物" not in final_text
    assert "源件高度敏感" not in final_text


def test_seal_rejects_unbound_text_candidates_with_rd_e010(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "gap-source.pptx")
    workdir = tmp_path / "gap-run"
    _run("extract", "--source", source, "--out", workdir)
    skeleton["pages"] = [skeleton["pages"][0]]
    skeleton["pages"][0]["slots"] = skeleton["pages"][0]["slots"][:1]
    skeleton["page_count"] = 1
    path = tmp_path / "gap-skeleton.json"
    path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    result = _run("seal", "--workdir", workdir, "--skeleton", path, expected=30)
    assert '"code": "RD-E010"' in result.stderr
    assert "未绑定" in result.stderr
    assert (workdir / "extract_pack/pages").is_dir()
    assert not (workdir / "sealed/shell.pptx").exists()


def test_extract_writes_coverage_complete_skeleton_draft(tmp_path: Path) -> None:
    source, _skeleton = _source_deck(tmp_path / "draft-source.pptx")
    workdir = tmp_path / "draft-run"
    _run("extract", "--source", source, "--out", workdir)
    draft_path = workdir / "extract_pack/skeleton_draft.json"
    assert draft_path.is_file()
    draft = json.loads(draft_path.read_text(encoding="utf-8"))
    assert draft["page_count"] == len(draft["pages"])
    assert draft["pages"]
    candidates = json.loads((workdir / "extract_pack/slot_candidates.json").read_text(encoding="utf-8"))
    included = {page["page_id"]: page for page in draft["pages"]}
    assert set(included) == {page["page_id"] for page in candidates["pages"]}
    for page in candidates["pages"]:
        spec = included[page["page_id"]]
        bound = {slot.get("shape_ref") for slot in spec["slots"] if slot.get("shape_ref")}
        for item in page["candidates"]:
            assert item["shape_ref"] in bound
    _run("seal", "--workdir", workdir, "--skeleton", draft_path)
    assert (workdir / "sealed/skeleton.sealed.json").is_file()


def test_plan_role_hints_never_drop_pages_and_missing_content_blocks_finalize(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "mixed-source.pptx")
    skeleton["pages"][1]["page_pattern"] = "test_matrix"
    skeleton["pages"][1]["argument_flow"] = "intent_matrix_emphasis"
    workdir = tmp_path / "mixed-run"
    _run("extract", "--source", source, "--out", workdir)
    skeleton_path = tmp_path / "mixed-skeleton.json"
    skeleton_path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    material = tmp_path / "pain-only.md"
    material.write_text("# 用户技术方案\n背景：现网存在重复录入的痛点，目标是缩短处理链路。\n", encoding="utf-8")
    _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        material,
        "--render-mode",
        "source-shell",
    )
    fit = {page["page_id"]: page for page in json.loads((workdir / "fill_pack/fit_report.json").read_text(encoding="utf-8"))["pages"]}
    assert fit["p01"]["verdict"] == fit["p02"]["verdict"] == "needs_content"
    assert fit["p01"]["role_hint_hits"] > 0
    assert fit["p02"]["role_hint_hits"] == 0
    assert fit["p01"]["score"] is None and fit["p02"]["score"] is None
    content = _content_from_fill_pack(workdir)
    content["pages"][1]["slots"] = [
        {"slot_id": slot["slot_id"], "status": "missing", "reason": "素材没有测试矩阵证据"}
        for slot in skeleton["pages"][1]["slots"]
    ]
    content_path = tmp_path / "mixed-content.json"
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "mixed-out"
    blocked = _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir, expected=50)
    assert '"code": "RD-E030"' in blocked.stderr
    assert not outdir.exists()
    preview = tmp_path / "mixed-preview"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", preview, "--preview")
    assert len(Presentation(preview / "draft.pptx").slides) == 2
    assert not (preview / "deck.pptx").exists()
    manifest = json.loads((preview / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "draft"
    assert manifest["evidence_pass"] is False
    assert manifest["page_ids"] == ["p01", "p02"]
    assert manifest["skipped_pages"] == []


def test_unrelated_material_does_not_prove_readiness_or_authorize_page_removal(tmp_path: Path) -> None:
    source, skeleton = _source_deck(tmp_path / "reject-source.pptx")
    workdir = tmp_path / "reject-run"
    _run("extract", "--source", source, "--out", workdir)
    skeleton_path = tmp_path / "reject-skeleton.json"
    skeleton_path.write_text(json.dumps(skeleton, ensure_ascii=False), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", skeleton_path)
    thin = tmp_path / "unrelated.md"
    thin.write_text("alpha beta gamma delta epsilon", encoding="utf-8")
    result = _run(
        "plan",
        "--workdir",
        workdir,
        "--skeleton",
        workdir / "sealed/skeleton.sealed.json",
        "--material",
        thin,
        "--render-mode",
        "source-shell",
    )
    planned = json.loads(result.stdout)
    assert planned["accepted_pages"] == ["p01", "p02"]
    assert planned["rejected_pages"] == []
    fit = json.loads((workdir / "fill_pack/fit_report.json").read_text(encoding="utf-8"))
    assert fit["ready_for_final"] is False
    assert all(p["verdict"] == "needs_content" and p["role_hint_hits"] == 0 for p in fit["pages"])
    content = _content_from_fill_pack(workdir)
    content["pages"].pop()
    path = tmp_path / "omitted-page.json"
    path.write_text(json.dumps(content), encoding="utf-8")
    blocked = _run("finalize", "--workdir", workdir, "--content", path, "--out", tmp_path / "omitted-out", expected=50)
    assert '"code": "RD-E030"' in blocked.stderr
    assert not (tmp_path / "omitted-out").exists()


def test_rich_sample_fills_table_nodes_cards_and_blocks_cell_leak(tmp_path: Path) -> None:
    source = SKILL / "samples/source_sample.pptx"
    material = SKILL / "samples/material_sample.docx"
    skeleton = SKILL / "samples/expected/skeleton.json"
    assert source.is_file() and material.is_file() and skeleton.is_file()
    workdir = tmp_path / "sample-run"
    _run("extract", "--source", source, "--out", workdir)
    candidates = json.loads((workdir / "extract_pack/slot_candidates.json").read_text(encoding="utf-8"))
    draft = json.loads((workdir / "extract_pack/skeleton_draft.json").read_text(encoding="utf-8"))
    assert draft["page_count"] == len(draft["pages"])
    draft_pages = {page["page_id"]: page for page in draft["pages"]}
    included = {page["page_id"] for page in json.loads(skeleton.read_text(encoding="utf-8"))["pages"]}
    for page in candidates["pages"]:
        spec = draft_pages.get(page["page_id"])
        if spec is not None:
            bound = {slot.get("shape_ref") for slot in spec["slots"] if slot.get("shape_ref")}
            for item in page["candidates"]:
                assert item["shape_ref"] in bound
        if page["page_id"] in included:
            assert page["candidates"]
            kinds = {item["kind"] for item in page["candidates"]}
            if page["page_id"] == "p03":
                assert "table_cell" in kinds
            if page["page_id"] == "p04":
                assert "shape" in kinds
    complete = _complete_skeleton(workdir, json.loads(skeleton.read_text(encoding="utf-8")))
    complete_path = tmp_path / "sample-complete-skeleton.json"
    complete_path.write_text(json.dumps(complete, ensure_ascii=False), encoding="utf-8")
    _run("seal", "--workdir", workdir, "--skeleton", complete_path)
    assert (workdir / "sealed/slot_candidates.json").is_file()
    _run("plan", "--workdir", workdir, "--skeleton", workdir / "sealed/skeleton.sealed.json", "--material", material, "--render-mode", "source-shell")
    content = _content_from_fill_pack(workdir)
    content_path = tmp_path / "sample-content.json"
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    outdir = tmp_path / "sample-out"
    _run("finalize", "--workdir", workdir, "--content", content_path, "--out", outdir)
    deck = outdir / "deck.pptx"
    presentation = Presentation(deck)
    assert len(presentation.slides) == 4
    blob = _deck_text(deck)
    assert "统一校验" in blob
    assert "源件封面机密标题不可进入产物" not in blob
    assert "源件表格机密单元格甲不可进入产物" not in blob
    assert "源件节点框甲不可进入产物" not in blob
    assert "源件密级内部资料不可进入产物" not in blob
    assert "HUAWEI CONFIDENTIAL" in _master_text(deck)
    table_texts = []
    for shape in presentation.slides[2].shapes:
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    table_texts.append((cell.text or "").strip())
    assert table_texts
    assert all(table_texts)
    assert all("源件表格机密" not in item for item in table_texts)
    node_texts = [
        (shape.text or "").strip()
        for shape in presentation.slides[3].shapes
        if getattr(shape, "has_text_frame", False)
    ]
    assert any(any(fact in item for fact in ("统一校验", "解析输入", "校验契约", "生成产物", "执行审计", "回归证据")) for item in node_texts)
    assert all("源件节点框" not in item for item in node_texts)
    card_titles = [
        (shape.text or "").strip()
        for shape in presentation.slides[1].shapes
        if getattr(shape, "has_text_frame", False) and (shape.text or "").strip()
    ]
    assert len(card_titles) >= 9
    assert all("源件卡片" not in item for item in card_titles)
    assert any("统一校验" in item for item in card_titles)

    leaked = json.loads(content_path.read_text(encoding="utf-8"))
    table_slot = next(
        item
        for page in leaked["pages"] if page["page_id"] == "p03"
        for item in page["slots"]
        if item["slot_id"] != "s1"
    )
    table_slot["value"] = "源件表格机密单元格甲不可进入产物"
    leaked_path = tmp_path / "sample-leaked.json"
    leaked_path.write_text(json.dumps(leaked, ensure_ascii=False), encoding="utf-8")
    blocked = tmp_path / "sample-blocked"
    result = _run("finalize", "--workdir", workdir, "--content", leaked_path, "--out", blocked, expected=60)
    assert '"code": "RD-E040"' in result.stderr
    assert not (blocked / "deck.pptx").exists()

    spec = ROOT / "scripts/rhetoric_deck_visual_qa.py"
    qa = subprocess.run(
        [_rdw_python(), "-X", "utf8", str(spec), "--deck", str(deck), "--forbidden", str(SKILL / "samples/expected/forbidden_source_phrases.txt")],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert qa.returncode == 0, qa.stderr or qa.stdout
    report = json.loads(qa.stdout)
    assert report["ok"] is True
    assert report["slides"] == 4
