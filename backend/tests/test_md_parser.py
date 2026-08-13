from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_parse_markdown_extracts_blocks_outline_and_stats() -> None:
    from app.parsers.md_parser import parse_markdown

    ir = parse_markdown(ROOT / "samples" / "input" / "quarterly_report.md")

    assert ir.ir_type == "document"
    assert ir.source.format == "md"
    assert ir.stats.headings == 4
    assert ir.stats.paragraphs == 1
    assert ir.stats.tables == 1
    assert [item.text for item in ir.content.outline] == ["Q3 业务汇报", "关键进展", "风险清单", "下步计划"]
    assert [block.type for block in ir.content.blocks] == [
        "heading",
        "paragraph",
        "heading",
        "bullet_list",
        "heading",
        "table",
        "heading",
        "numbered_list",
    ]


def test_parse_markdown_round_trips_to_docx(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.parsers.md_parser import parse_markdown
    from app.rendering.docx_renderer import render_word_ir

    document_ir = parse_markdown(ROOT / "samples" / "input" / "quarterly_report.md")
    word_ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "Q3 业务汇报"},
            "blocks": [block.model_dump() for block in document_ir.content.blocks],
        }
    )

    output = render_word_ir(word_ir, tmp_path / "roundtrip.docx")
    doc = Document(str(output))

    assert "Q3 业务汇报" in [paragraph.text for paragraph in doc.paragraphs]
    assert doc.tables[0].cell(0, 0).text == "风险"


def test_parse_markdown_falls_back_to_gbk_with_warning(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "gbk.md"
    path.write_bytes("# 周报\n\n本周完成台账核对。".encode("gbk"))

    ir = parse_markdown(path)

    assert ir.content.outline[0].text == "周报"
    assert any("gb18030 fallback" in warning for warning in ir.warnings)


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
def test_parse_markdown_decodes_utf16_without_bom(tmp_path: Path, encoding: str) -> None:
    """A BOM-less UTF-16 file must decode correctly, not silently corrupt."""
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "utf16.md"
    path.write_bytes("# 标题\n\n正文内容 test".encode(encoding))

    ir = parse_markdown(path)

    assert ir.content.outline[0].text == "标题"
    assert ir.content.blocks[-1].text == "正文内容 test"
    assert any("UTF-16 (no BOM)" in warning for warning in ir.warnings)


def test_parse_markdown_utf16_no_bom_ascii_only_is_not_nul_corrupted(tmp_path: Path) -> None:
    """Pure-ASCII UTF-16LE no-BOM bytes are valid UTF-8 too; must still decode as UTF-16."""
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "ascii-utf16.md"
    path.write_bytes("# Heading\n\nBody text".encode("utf-16-le"))

    ir = parse_markdown(path)

    assert ir.content.outline[0].text == "Heading"
    assert "\x00" not in ir.content.outline[0].text


def test_parse_markdown_warns_for_malformed_table_rows(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "malformed-table.md"
    path.write_text("| 项目 | 状态 |\n| --- | --- |\n| 仅一列 |\n| 正常 | 完成 |\n", encoding="utf-8")

    ir = parse_markdown(path)

    table = next(block for block in ir.content.blocks if block.type == "table")
    assert table.rows == [["正常", "完成"]]
    assert any("malformed table rows" in warning for warning in ir.warnings)


def test_parse_markdown_fills_empty_header_cells(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "empty-header.md"
    path.write_text("|  | 状态 |\n| --- | --- |\n| 甲 | 完成 |\n", encoding="utf-8")

    ir = parse_markdown(path)
    table = next(block for block in ir.content.blocks if block.type == "table")
    assert table.header == ["Column 1", "状态"]


def test_parse_markdown_handles_escaped_pipe_in_cell(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "escaped-pipe.md"
    path.write_text("| 项目 | 状态 |\n| --- | --- |\n| A\\|B | 完成 |\n", encoding="utf-8")

    ir = parse_markdown(path)
    table = next(block for block in ir.content.blocks if block.type == "table")
    assert table.rows == [["A|B", "完成"]]


def test_parse_markdown_stops_table_at_second_separator_row(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "second-separator.md"
    path.write_text("| 项目 | 状态 |\n| --- | --- |\n| 甲 | 完成 |\n| --- | --- |\n\n后续正文。\n", encoding="utf-8")

    ir = parse_markdown(path)
    table = next(block for block in ir.content.blocks if block.type == "table")
    assert table.rows == [["甲", "完成"]]
    assert any(block.type == "paragraph" and "后续正文" in block.text for block in ir.content.blocks)


def test_parse_markdown_caps_wide_table_at_12_columns_with_warning(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    header = " | ".join(f"c{index}" for index in range(1, 14))
    row = " | ".join(str(index) for index in range(1, 14))
    path = tmp_path / "wide-table.md"
    path.write_text(f"| {header} |\n| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n| {row} |\n", encoding="utf-8")

    ir = parse_markdown(path)

    table = next(block for block in ir.content.blocks if block.type == "table")
    assert table.header == [f"c{index}" for index in range(1, 13)]
    assert table.rows[0] == [str(index) for index in range(1, 13)]
    assert any("truncated to 12 columns" in warning for warning in ir.warnings)


def test_parse_markdown_preserves_fenced_code_block_indentation() -> None:
    from app.parsers.md_parser import parse_markdown

    ir = parse_markdown(ROOT / "samples" / "input" / "synthetic" / "md_03_code_fence.md")
    code = next(block for block in ir.content.blocks if block.type == "code_block")

    assert ir.ir_version == "1.2"
    assert code.language == "c"
    assert code.code == "typedef struct {\n    uint16_t frame_id;\n    uint8_t payload[32];\n} FrameHeader;"


def test_parse_markdown_enforces_file_resource_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.parsers import md_parser
    from app.parsers.errors import ParseFailure

    monkeypatch.setattr(md_parser, "MAX_MD_FILE_BYTES", 8)
    path = tmp_path / "large.md"
    path.write_bytes(b"123456789")

    with pytest.raises(ParseFailure) as captured:
        md_parser.parse_markdown(path)
    assert captured.value.code == "E001"
    assert captured.value.loc == "source.resource"


def test_parse_markdown_limits_long_text_with_locations(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    long_text = "文" * 2100
    path = tmp_path / "long.md"
    path.write_text(f"# 标题\n\n{long_text}\n\n- {long_text}\n", encoding="utf-8")

    result = parse_markdown(path)
    paragraph = next(block for block in result.content.blocks if block.type == "paragraph")
    bullet = next(block for block in result.content.blocks if block.type == "bullet_list")
    assert len(paragraph.text) == 2000
    assert len(bullet.items[0].text) == 2000
    assert any("markdown paragraph" in warning for warning in result.warnings)
    assert any("markdown list item" in warning for warning in result.warnings)


def test_parse_markdown_skips_whitespace_only_headings(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "empty-heading.md"
    path.write_text("#   \n##  \n#  \u00a0\n\nreal content\n", encoding="utf-8")

    result = parse_markdown(path)
    assert [block.type for block in result.content.blocks] == ["paragraph"]
    assert result.content.blocks[0].text == "real content"
    assert result.content.outline == []


def test_parse_markdown_skips_whitespace_only_list_items(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "empty-items.md"
    path.write_text("- a\n-  \n- \u00a0\n- b\n\n1. x\n2.  \n3. y\n", encoding="utf-8")

    result = parse_markdown(path)
    bullet = next(block for block in result.content.blocks if block.type == "bullet_list")
    numbered = next(block for block in result.content.blocks if block.type == "numbered_list")
    assert [item.text for item in bullet.items] == ["a", "b"]
    assert [item.text for item in numbered.items] == ["x", "y"]


def test_parse_markdown_skips_all_empty_list_without_emitting_block(tmp_path: Path) -> None:
    from app.parsers.md_parser import parse_markdown

    path = tmp_path / "all-empty.md"
    path.write_text("-  \n- \u00a0\n\nbody\n", encoding="utf-8")

    result = parse_markdown(path)
    assert all(block.type != "bullet_list" for block in result.content.blocks)
