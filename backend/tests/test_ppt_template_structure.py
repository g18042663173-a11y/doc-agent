from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.template.ppt_template_structure import (
    PptTemplateStructureError,
    extract_ppt_template_structure,
    write_ppt_template_structure,
)


ROOT = Path(__file__).resolve().parents[2]


def test_extract_ppt_template_structure_classifies_common_template_skeletons(tmp_path: Path) -> None:
    source = tmp_path / "template.pptx"
    presentation = Presentation()

    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "项目模板"
    cover.placeholders[1].text = "副标题"

    two_columns = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(two_columns, "方案对比", 0.6, 0.25, 9.8, 0.55)
    _textbox(two_columns, "现状", 0.7, 1.4, 4.4, 3.2)
    _textbox(two_columns, "目标", 8.2, 1.4, 4.4, 3.2)

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(table_slide, "参数表", 0.6, 0.25, 9.8, 0.55)
    table = table_slide.shapes.add_table(3, 2, Inches(0.8), Inches(1.4), Inches(8.6), Inches(3.2)).table
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "name"
    table.cell(1, 1).text = "名称"

    blank = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(two_columns, "Footer", 4.0, 7.0, 1.0, 0.2)
    presentation.save(source)

    structure = extract_ppt_template_structure(source)

    assert structure.filename == "template.pptx"
    assert structure.width_in > 0 and structure.height_in > 0
    assert [slide.detected_layout for slide in structure.slides] == [
        "title_slide",
        "two_column",
        "table",
        "blank",
    ]
    assert structure.slides[1].columns == 2
    assert [block.kind for block in structure.slides[2].blocks] == ["title", "table"]
    assert structure.slides[2].blocks[1].rows == 3
    assert structure.slides[2].blocks[1].columns == 2
    assert structure.slides[0].title == "项目模板"
    assert structure.slides[2].parser_table_count == 1
    assert all(block.text_preview != "Footer" for block in structure.slides[1].blocks)


def test_template_structure_reuses_parser_summary_and_writes_json_and_markdown(tmp_path: Path) -> None:
    source = tmp_path / "summary.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "实现状态"
    slide.placeholders[1].text = "正文要点"
    slide.notes_slide.notes_text_frame.text = "演讲备注"
    presentation.save(source)

    structure = extract_ppt_template_structure(source)
    json_path = write_ppt_template_structure(structure, tmp_path / "structure.json")
    markdown_path = write_ppt_template_structure(structure, tmp_path / "structure.md")

    assert structure.slides[0].parser_body_count >= 1
    assert structure.slides[0].notes_present is True
    assert '"detected_layout"' in json_path.read_text(encoding="utf-8")
    markdown = markdown_path.read_text(encoding="utf-8")
    assert "# PPT 模板结构：summary.pptx" in markdown
    assert "第 1 页" in markdown
    assert "存在演讲备注" in markdown


def test_template_structure_rejects_unknown_report_extension(tmp_path: Path) -> None:
    source = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)

    structure = extract_ppt_template_structure(source)
    with pytest.raises(PptTemplateStructureError, match=".json or .md"):
        write_ppt_template_structure(structure, tmp_path / "structure.txt")


@pytest.mark.parametrize(
    ("relative_path", "expected_first_layout", "expected_block"),
    [
        ("samples/input/项目汇报.pptx", "title_slide", "table"),
        ("samples/input/parser_samples/pptx_sample_01.pptx", "table", "table"),
    ],
)
def test_extracts_structure_from_existing_ppt_samples(
    relative_path: str,
    expected_first_layout: str,
    expected_block: str,
) -> None:
    structure = extract_ppt_template_structure(ROOT / relative_path)

    assert structure.slides
    assert structure.slides[0].detected_layout == expected_first_layout
    assert any(
        block.kind == expected_block
        for slide in structure.slides
        for block in slide.blocks
    )


def _textbox(slide, text: str, left: float, top: float, width: float, height: float) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
