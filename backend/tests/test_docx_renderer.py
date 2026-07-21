from __future__ import annotations

import sys
import zipfile
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_render_word_ir_creates_editable_docx_with_core_blocks(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "项目周报", "classification": "内部公开"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "本周进展"},
                {"type": "paragraph", "text": "完成 DOCX 渲染核心。"},
                {"type": "bullet_list", "items": [{"text": "标题与正文", "level": 1}, {"text": "两级列表", "level": 2}]},
                {"type": "numbered_list", "items": [{"text": "第一步", "level": 1}]},
                {"type": "paragraph", "style": "quote", "text": "引用段落"},
                {"type": "paragraph", "style": "note", "text": "提示段落"},
                {"type": "page_break"},
                {"type": "heading", "level": 2, "text": "下周计划"},
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "weekly.docx")

    assert output.exists()
    doc = Document(str(output))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    styles = [paragraph.style.name for paragraph in doc.paragraphs]

    assert "本周进展" in texts
    assert "完成 DOCX 渲染核心。" in texts
    assert "标题与正文" in texts
    assert "两级列表" in texts
    assert "下周计划" in texts
    assert "Heading 1" in styles
    assert "Heading 2" in styles
    assert "List Bullet" in styles
    assert "List Bullet 2" in styles
    assert "List Number" in styles
    assert "IR Quote" in styles
    assert "IR Note" in styles


def test_render_word_ir_writes_header_footer_and_page_field(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "带页眉页脚", "classification": "内部公开"},
            "blocks": [{"type": "paragraph", "text": "正文"}],
        }
    )

    output = render_word_ir(ir, tmp_path / "header-footer.docx")
    doc = Document(str(output))
    section = doc.sections[0]

    assert "带页眉页脚" in section.header.paragraphs[0].text
    assert "内部公开" in section.footer.paragraphs[0].text
    with zipfile.ZipFile(output) as package:
        footer_xml = "\n".join(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
    assert "PAGE" in footer_xml


def test_render_word_ir_writes_document_control_tables_before_body(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.2",
            "meta": {
                "title": "链路控制模块详细设计",
                "classification": "内部公开",
                "document_control": {
                    "product_name": "星载基带传输平台",
                    "document_name": "链路控制模块详细设计",
                    "version": "V1.0",
                    "prepared": {"name": "张三", "date": "2026-07-21"},
                    "reviewed": {"name": "李四", "date": "2026-07-22"},
                    "approved": {"name": "王五", "date": "2026-07-23"},
                },
            },
            "blocks": [{"type": "heading", "level": 1, "text": "功能设计"}],
        }
    )

    output = render_word_ir(ir, tmp_path / "document-control.docx")
    doc = Document(str(output))
    xml = _word_package_xml(output)

    assert doc.core_properties.category == "HW_DOCUMENT_CONTROL"
    assert len(doc.tables) == 2
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [
        ["产品名称", "星载基带传输平台", "文档名称", "链路控制模块详细设计"],
        ["密级", "内部公开", "版本号", "V1.0"],
    ]
    assert [[cell.text for cell in row.cells] for row in doc.tables[1].rows] == [
        ["角色", "姓名", "日期"],
        ["拟制", "张三", "2026-07-21"],
        ["审核", "李四", "2026-07-22"],
        ["批准", "王五", "2026-07-23"],
    ]
    assert "HW_DOCUMENT_CONTROL_INFO" in xml
    assert "HW_DOCUMENT_CONTROL_APPROVAL" in xml
    assert 'w:fill="F5F5F5"' in xml
    assert "DDDDDD" in xml


def test_render_word_ir_keeps_blank_document_control_signoff_cells(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.2",
            "meta": {
                "title": "部分签核详设",
                "author": "张三",
                "document_control": {
                    "product_name": "产品",
                    "document_name": "部分签核详设",
                    "version": "V1.0",
                    "prepared": {"date": "2026-07-21"},
                },
            },
            "blocks": [{"type": "paragraph", "text": "正文"}],
        }
    )

    doc = Document(str(render_word_ir(ir, tmp_path / "partial-control.docx")))

    approval = doc.tables[1]
    assert len(approval.rows) == 4
    assert approval.cell(1, 1).text == "张三"
    assert not approval.cell(2, 1).text.strip()
    assert not approval.cell(2, 2).text.strip()
    assert [approval.cell(row, 0).text for row in range(1, 4)] == ["拟制", "审核", "批准"]


def test_render_word_ir_includes_page_break_xml(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "分页"},
            "blocks": [
                {"type": "paragraph", "text": "第一页"},
                {"type": "page_break"},
                {"type": "paragraph", "text": "第二页"},
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "page-break.docx")

    with zipfile.ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert 'w:type="page"' in document_xml or "w:type='page'" in document_xml


def test_render_word_ir_formats_table_header_and_widths(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "表格"},
            "blocks": [
                {
                    "type": "table",
                    "caption": "风险清单",
                    "header": ["风险", "等级"],
                    "rows": [["模型 JSON 不稳定", "中"]],
                    "col_widths": [3, 1],
                }
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "table.docx")
    doc = Document(str(output))

    assert doc.tables[0].cell(0, 0).text == "风险"
    assert doc.tables[0].cell(1, 0).text == "模型 JSON 不稳定"
    assert doc.tables[0].cell(0, 0).paragraphs[0].runs[0].bold is True
    with zipfile.ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "w:tblHeader" in document_xml
    assert 'w:fill="C7000B"' in document_xml
    assert 'w:color w:val="FFFFFF"' in document_xml
    assert "DDDDDD" in document_xml
    widths = [int(value) for value in re.findall(r'<w:gridCol w:w="(\d+)"', document_xml)]
    assert widths[:2][0] > widths[:2][1]


def test_render_word_ir_preserves_code_block_indentation_and_theme_tokens(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    code = "typedef struct {\n    uint16_t frame_id;\n    uint8_t payload[32];\n} FrameHeader;"
    function = "static bool frame_header_valid(const FrameHeader *header) {\n    return header != NULL && header->frame_id != 0U;\n}"
    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.1",
            "meta": {"title": "代码详设"},
            "blocks": [
                {"type": "code_block", "language": "c", "code": code},
                {"type": "code_block", "language": "c", "code": function},
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "code.docx")
    doc = Document(str(output))
    paragraphs = [item for item in doc.paragraphs if item.style.name == "IR Code"]
    xml = _word_package_xml(output)

    assert [paragraph.text for paragraph in paragraphs] == [code, function]
    assert all(paragraph.runs[0].font.name == "Consolas" for paragraph in paragraphs)
    assert all(paragraph.runs[0].font.size == Pt(9) for paragraph in paragraphs)
    assert 'xml:space="preserve"' in xml
    assert "w:br" in xml
    assert 'w:fill="F5F5F5"' in xml
    assert "DDDDDD" in xml


def test_render_word_ir_uses_hw_theme_tokens_not_word_blue(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "主题", "classification": "HUAWEI CONFIDENTIAL"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "主题标题"},
                {"type": "heading", "level": 2, "text": "二级标题"},
                {"type": "paragraph", "text": "正文"},
                {"type": "table", "header": ["项", "值"], "rows": [["颜色", "主题红"]]},
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "theme.docx")
    doc = Document(str(output))
    xml = _word_package_xml(output)

    assert "C7000B" in xml
    assert "4F81BD" not in xml
    assert str(doc.styles["Heading 1"].font.color.rgb) == "C7000B"
    assert doc.styles["Heading 1"].font.size == Pt(14)
    assert doc.styles["Heading 2"].font.size == Pt(12)
    assert str(doc.styles["Normal"].font.color.rgb) == "1D1D1A"
    assert doc.styles["Normal"].font.size == Pt(10)


def test_render_word_ir_draws_bordered_image_placeholder_and_caption(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "图片", "classification": "HUAWEI CONFIDENTIAL"},
            "blocks": [{"type": "image_placeholder", "ref": "arch.png", "caption": "图1: 架构图"}],
        }
    )

    output = render_word_ir(ir, tmp_path / "image-placeholder.docx")
    doc = Document(str(output))
    xml = _word_package_xml(output)

    assert len(doc.tables) == 1
    assert "图片占位" in doc.tables[0].cell(0, 0).text
    assert any(paragraph.text == "图1: 架构图" for paragraph in doc.paragraphs)
    assert 'w:val="single"' in xml
    assert "DDDDDD" in xml


def test_render_word_ir_quote_has_left_border(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "引用", "classification": "HUAWEI CONFIDENTIAL"},
            "blocks": [{"type": "paragraph", "style": "quote", "text": "引用段落"}],
        }
    )

    output = render_word_ir(ir, tmp_path / "quote.docx")
    with zipfile.ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    assert "<w:left" in document_xml
    assert 'w:color="666666"' in document_xml


def test_render_word_ir_handles_100_by_12_table(tmp_path: Path) -> None:
    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    ir = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "极限表"},
            "blocks": [
                {
                    "type": "table",
                    "header": [f"H{i}" for i in range(12)],
                    "rows": [[f"R{row}C{col}" for col in range(12)] for row in range(100)],
                }
            ],
        }
    )

    output = render_word_ir(ir, tmp_path / "large-table.docx")
    doc = Document(str(output))

    assert len(doc.tables[0].rows) == 101
    assert len(doc.tables[0].columns) == 12


def _word_package_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        return "\n".join(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
