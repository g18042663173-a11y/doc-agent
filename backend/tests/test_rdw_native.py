"""Native-package preservation tests, including an Office-authored SmartArt case."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import base64
from pathlib import Path
import subprocess
import sys

import pytest
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SKILL = ROOT / "skills/rhetoric-deck-workflow"


@pytest.fixture(scope="module")
def native():
    sys.path.insert(0, str(SKILL))
    from engine.native import package
    yield package
    sys.path.remove(str(SKILL))


def fixture_deck(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(.4), Inches(4), Inches(.4))
    text.text = "Native editable source"
    text.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    image = BytesIO()
    screenshot = Image.new("RGB", (256, 128), (45, 89, 130))
    ImageDraw.Draw(screenshot).text((8, 12), "SOURCE SCREENSHOT 123", fill="white")
    screenshot.save(image, "PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(3), Inches(2))
    picture.crop_left = .1
    data = CategoryChartData()
    data.categories = ["Old A", "Old B"]
    data.add_series("Old series", [80, 100])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(3), Inches(3), data).chart
    chart.value_axis.major_unit = 20
    deck.core_properties.title = "Preserve metadata and types"
    deck.save(path)
    return path


def test_nonstandard_chart_part_and_pixels_are_preserved(tmp_path, native):
    source = fixture_deck(tmp_path / "source.pptx")
    parts = native.read_parts(source)
    # Legal part identity is independent of conventional directory naming.
    old = "ppt/charts/chart1.xml"
    new = "ppt/slides/charts/chart1.xml"
    parts[new] = parts.pop(old)
    old_rels = native.rels_name(old)
    new_rels = native.rels_name(new)
    rels = native.xml(parts.pop(old_rels))
    for rel in rels:
        if rel.get("Target", "").startswith("../embeddings/"):
            rel.set("Target", "../../embeddings/" + rel.get("Target").rsplit("/", 1)[-1])
    parts[new_rels] = native.serialized(rels)
    slide_rels = native.xml(parts["ppt/slides/_rels/slide1.xml.rels"])
    for rel in slide_rels:
        if rel.get("Type", "").endswith("/chart"):
            rel.set("Target", "charts/chart1.xml")
    parts["ppt/slides/_rels/slide1.xml.rels"] = native.serialized(slide_rels)
    ct = native.xml(parts["[Content_Types].xml"])
    for entry in ct:
        if entry.get("PartName") == "/" + old:
            entry.set("PartName", "/" + new)
    parts["[Content_Types].xml"] = native.serialized(ct)
    native.write_parts(source, parts)
    before = native.inspect_package(source)
    chart = next(obj for obj in before["pages"][0]["objects"] if obj["kind"] == "chart")
    assert chart["chart_part"] == new
    shell = tmp_path / "shell.pptx"
    native.sanitize_package(source, shell)
    after = native.inspect_package(shell)
    assert before["media"] == after["media"]
    assert native.read_parts(shell)["docProps/core.xml"] == parts["docProps/core.xml"]
    assert next(obj for obj in after["pages"][0]["objects"] if obj["kind"] == "chart")["chart_data"] == chart["chart_data"]
    assert all(not obj.get("text") for obj in after["pages"][0]["objects"] if obj["kind"] not in {"chart", "picture"})
    assert native.audit_output(source, shell)["pass"]
    source_pic = next(obj for obj in before["pages"][0]["objects"] if obj["kind"] == "picture")
    result_pic = next(obj for obj in after["pages"][0]["objects"] if obj["kind"] == "picture")
    assert source_pic["media"] == result_pic["media"] and source_pic["bbox"] == result_pic["bbox"] and source_pic["crop"] == result_pic["crop"]
    assert source_pic["crop"]


def test_chart_replacement_updates_workbook_and_readback(tmp_path, native):
    from engine.render.source_shell import fill_source_shell
    source = fixture_deck(tmp_path / "source.pptx")
    shell = tmp_path / "shell.pptx"
    native.sanitize_package(source, shell)
    inventory = native.inspect_package(source)
    objects = [obj for obj in inventory["pages"][0]["objects"] if obj["kind"] in {"shape", "chart"} and (obj.get("text") or obj["kind"] == "chart")]
    slots = [{"slot_id": f"s{i}", "shape_ref": obj["shape_ref"], "type": "chart" if obj["kind"] == "chart" else "caption"} for i, obj in enumerate(objects)]
    skeleton = {"pages": [{"page_id": "p01", "slots": slots}]}
    data = {"categories": ["A", "B", "C"], "series": [{"name": "New evidence", "values": [.43, .51, .69]}, {"name": "Baseline", "values": [.4, .48, .6]}]}
    content = {"pages": [{"page_id": "p01", "slots": [{"slot_id": slot["slot_id"], "value": data if slot["type"] == "chart" else "New caption"} for slot in slots]}]}
    output = tmp_path / "result.pptx"
    fill_source_shell(shell, skeleton, content, output)
    assert native.readback_audit(output, skeleton, content)["pass"]
    assert native.audit_output(source, output)["pass"]
    chart = next(s.chart for s in Presentation(output).slides[0].shapes if s.has_chart)
    assert chart.value_axis.maximum_scale == 1
    assert chart.value_axis.major_unit is None
    assert chart.has_legend
    from openpyxl import load_workbook
    parts = native.read_parts(output)
    workbook_data = next(data for name, data in parts.items() if name.endswith(".xlsx"))
    rows = list(load_workbook(BytesIO(workbook_data), data_only=True).active.values)
    assert rows[1] == ("A", .43, .4)
    assert rows[3] == ("C", .69, .6)
    damaged = deepcopy(content)
    damaged["pages"][0]["slots"][0]["value"] = "Different text"
    assert not native.readback_audit(output, skeleton, damaged)["pass"]


def test_imitation_keeps_nine_point_caption(tmp_path, native):
    from engine.fit.text_fit import replace_text_preserving_style
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(.25))
    shape.text = "Before"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    assert replace_text_preserving_style(shape, "After", role="title", preserve_source_size=True)
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 9
    assert not replace_text_preserving_style(shape, "A paragraph much too long for the source box. " * 60, role="body", preserve_source_size=True)


def test_doughnut_small_points_hide_labels_without_changing_data(tmp_path, native):
    from engine.render.source_shell import _format_replaced_chart
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    payload = {"categories": ["Main 97.40%", "Small 2.50%", "Tiny 0.10%"], "series": [{"name": "Share", "values": [.974, .025, .001]}]}
    data = CategoryChartData()
    data.categories = payload["categories"]
    data.add_series("Share", payload["series"][0]["values"])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(1), Inches(6), Inches(4), data).chart
    _format_replaced_chart(chart, payload)
    path = tmp_path / "donut.pptx"
    deck.save(path)
    actual = next(obj for obj in native.inspect_package(path)["pages"][0]["objects"] if obj["kind"] == "chart")
    assert actual["chart_data"] == payload
    deletes = chart._chartSpace.xpath('.//*[local-name()="dLbl"]/*[local-name()="delete"]')
    assert len(deletes) == 2 and all(node.get("val") == "1" for node in deletes)
    assert all([etree.QName(child).localname for child in node.getparent()] == ["idx", "delete"] for node in deletes)
    assert chart.has_legend
    assert native.chart_workbook_audit(path)["pass"]
    _assert_office_opens(path)


def test_merged_cells_mixed_runs_and_shared_master_replace(tmp_path, native):
    from engine.render.source_shell import fill_source_shell, _CellProxy
    from engine.shared.common import write_json
    from pptx.dml.color import RGBColor
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    mixed = slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(5), Inches(.7))
    paragraph = mixed.text_frame.paragraphs[0]
    first = paragraph.add_run(); first.text = "AA"; first.font.bold = True; first.font.size = Pt(12)
    second = paragraph.add_run(); second.text = "BBBBBB"; second.font.italic = True; second.font.size = Pt(12); second.font.color.rgb = RGBColor(199, 0, 10)
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    table.table.cell(0, 0).merge(table.table.cell(1, 1))
    table.table.cell(0, 0).text = "Merged source text"
    proxy = _CellProxy(table, 0, 0)
    assert proxy.width == table.width and proxy.height == table.height
    shared = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(5), Inches(.5))
    shared.text = "Original visible master claim"
    element = deepcopy(shared._element)
    element.xpath('.//*[local-name()="cNvPr"]')[0].set("id", "9000")
    slide.slide_layout.slide_master.shapes._spTree.append(element)
    slide.shapes._spTree.remove(shared._element)
    source = tmp_path / "source.pptx"; deck.save(source)
    original = native.inspect_package(source)
    shared_objects = [obj for obj in original["shared_objects"] if obj.get("text") == "Original visible master claim"]
    assert len(shared_objects) == 1 and shared_objects[0]["action"] == "replace"
    objects = [obj for obj in original["pages"][0]["objects"] if obj.get("text")] + shared_objects
    slots = [{"slot_id": f"s{i}", "shape_ref": obj["shape_ref"], "type": "caption"} for i, obj in enumerate(objects)]
    skeleton = {"pages": [{"page_id": "p01", "slots": slots}]}
    content = {"pages": [{"page_id": "p01", "slots": [{"slot_id": slot["slot_id"], "value": "ABCDEFGH" if slot["shape_ref"] == f"sp_{mixed.shape_id}" else "User replacement"} for slot in slots]}]}
    shell = tmp_path / "shell.pptx"; native.sanitize_package(source, shell)
    write_json(tmp_path / "native_inventory.json", original)
    output = tmp_path / "result.pptx"; fill_source_shell(shell, skeleton, content, output)
    assert native.readback_audit(output, skeleton, content)["pass"]
    actual_mixed = next(shape for shape in Presentation(output).slides[0].shapes if shape.shape_id == mixed.shape_id)
    runs = actual_mixed.text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["AB", "CDEFGH"]
    assert runs[0].font.bold and runs[1].font.italic and str(runs[1].font.color.rgb) == "C7000A"
    assert native.audit_output(source, output)["pass"]


def _assert_office_opens(path):
    if sys.platform != "win32":
        return
    quoted = str(path.resolve()).replace("'", "''")
    script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject PowerPoint.Application
$presentation = $null
try {{
  $presentation = $app.Presentations.Open('{quoted}',-1,0,0)
  if ($presentation.Slides.Count -lt 1) {{ throw 'No slides' }}
}} finally {{
  if ($null -ne $presentation) {{ $presentation.Close() }}
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app)
}}
"""
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
        capture_output=True, encoding="utf-8", errors="replace", timeout=90,
    )
    assert result.returncode == 0, result.stderr[-700:]


@pytest.fixture(scope="module")
def office_smartart(tmp_path_factory):
    if sys.platform != "win32":
        pytest.skip("Office-authored fixture requires Windows PowerPoint")
    root = tmp_path_factory.mktemp("rdw-office-smartart")
    path = root / "smartart.pptx"
    quoted = str(path).replace("'", "''")
    script = f"""
$ErrorActionPreference='Stop'
$app = $null
$presentation = $null
try {{
  $app = New-Object -ComObject PowerPoint.Application
  # PowerPoint loads the built-in SmartArt catalog only after a window exists.
  $presentation = $app.Presentations.Add(-1)
  $slide = $presentation.Slides.Add(1,12)
  $layout = $app.SmartArtLayouts.Item(1)
  $shape = $slide.Shapes.AddSmartArt($layout,72,72,500,300)
  for ($index=1; $index -le $shape.SmartArt.AllNodes.Count; $index++) {{
    $shape.SmartArt.AllNodes.Item($index).TextFrame2.TextRange.Text = 'Original node ' + $index
  }}
  $presentation.SaveAs('{quoted}',24)
}} finally {{
  if ($null -ne $presentation) {{ $presentation.Close() }}
  if ($null -ne $app) {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }}
}}
"""
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
        capture_output=True, encoding="utf-8", errors="replace", timeout=90,
    )
    if result.returncode:
        pytest.skip("PowerPoint COM is unavailable: " + result.stderr[-300:])
    return path


def test_office_smartart_data_and_drawing_cache_stay_synchronized(tmp_path, native, office_smartart):
    source = office_smartart
    inventory = native.inspect_package(source)
    nodes = [obj for obj in inventory["pages"][0]["objects"] if obj["kind"] == "smartart_node" and obj.get("text")]
    assert len(nodes) >= 2
    output = tmp_path / "smartart-result.pptx"
    native.sanitize_package(source, output)
    replacements = [{"diagram_part": node["diagram_part"], "node_id": node["node_id"], "text": f"Replacement node {i}"} for i, node in enumerate(nodes)]
    results = native.update_smartart_file(output, replacements)
    assert all(item["cache_shapes_updated"] >= 1 for item in results)
    parts = native.read_parts(output)
    types = native.part_types(parts)
    cache_text = "\n".join("\n".join(native.xml(parts[name]).xpath('.//a:t/text()', namespaces=native.NS)) for name, kind in types.items() if kind.endswith("diagramDrawing+xml"))
    for item in replacements:
        assert item["text"] in cache_text
    assert "Original node" not in cache_text
    assert native.audit_output(source, output)["pass"]
    assert native.smartart_cache_audit(output)["pass"]
    _assert_office_opens(output)
    # A stale cache is a detectable failure even when the logical data is right.
    drawing_part = next(name for name, kind in types.items() if kind.endswith("diagramDrawing+xml"))
    drawing = native.xml(parts[drawing_part])
    drawing.find(".//a:t", native.NS).text = "Stale cached source text"
    parts[drawing_part] = native.serialized(drawing)
    native.write_parts(output, parts)
    assert not native.smartart_cache_audit(output)["pass"]


def test_authored_business_label_is_not_rewritten_as_template_chrome(tmp_path, native):
    from engine.render.source_shell import fill_source_shell
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shape.text = "Original business label"
    source = tmp_path / "source.pptx"
    deck.save(source)
    shell = tmp_path / "shell.pptx"
    native.sanitize_package(source, shell)
    skeleton = {"pages": [{"page_id": "p01", "slots": [{"slot_id": "s1", "shape_ref": f"sp_{shape.shape_id}", "type": "caption"}]}]}
    content = {"pages": [{"page_id": "p01", "slots": [{"slot_id": "s1", "value": "Classification: internal decision process"}]}]}
    result = tmp_path / "result.pptx"
    fill_source_shell(shell, skeleton, content, result, classification="HUAWEI CONFIDENTIAL")
    assert native.readback_audit(result, skeleton, content)["pass"]
