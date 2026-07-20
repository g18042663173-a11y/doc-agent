from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_parser_fixture_manifest_has_five_real_categories_per_format() -> None:
    from app.cli.parse import parse_file
    from app.parsers.errors import ParseFailure

    manifest = _json("samples/input/parser_matrix/manifest.json")
    expected_categories = {"normal", "empty", "corrupt", "large", "degraded"}

    assert set(manifest["formats"]) == {"md", "docx", "xlsx", "pptx"}
    for source_format, cases in manifest["formats"].items():
        assert {case["category"] for case in cases} == expected_categories
        assert len(cases) == 5
        for case in cases:
            path = ROOT / case["path"]
            assert path.exists() and path.is_file()
            if case["outcome"] == "error":
                with pytest.raises(ParseFailure) as exc_info:
                    parse_file(path)
                assert exc_info.value.code == case["expected_error"]
                continue
            document_ir = parse_file(path)
            assert document_ir.source.format == source_format
            marker = case.get("expected_warning_contains")
            if marker:
                warnings = [*document_ir.warnings]
                warnings.extend(value for slide in document_ir.content.slides for value in slide.shape_warnings)
                assert any(marker in value for value in warnings), (source_format, case["category"], warnings)


def test_appendix_b_fixed_office_samples_match_checked_expected() -> None:
    from app.cli.parse import parse_file

    expected = _json("samples/expected/appendix_b_document_ir.expected.json")
    assert set(expected) == {"需求说明.docx", "销售台账.xlsx", "项目汇报.pptx"}
    for filename, facts in expected.items():
        assert _document_ir_facts(parse_file(ROOT / "samples" / "input" / filename)) == facts


def test_eight_slide_pptx_expected_is_page_by_page_and_feature_complete() -> None:
    expected = _json("samples/expected/项目汇报.pptx.expected.json")
    slides = expected["slides"]

    assert len(slides) == 8
    assert [slide["index"] for slide in slides] == list(range(1, 9))
    assert any(slide["table_count"] > 0 for slide in slides)
    assert sum(slide["notes"] is not None for slide in slides) >= 2
    assert slides[-1] == {"index": 8, "title": None, "body_count": 0, "table_count": 0, "notes": None}

    presentation = Presentation(ROOT / "samples" / "input" / "项目汇报.pptx")
    assert len(presentation.slides) == 8
    assert any(getattr(shape, "has_chart", False) for shape in presentation.slides[5].shapes)


def test_deck_invalid_manifest_covers_d001_through_d006() -> None:
    from app.ir.shell import validate_deck_ir_text

    manifest = _json("samples/ir/deck_invalid_manifest.json")
    assert {case["code"] for case in manifest["cases"]} == {f"D00{index}" for index in range(1, 7)}
    assert len({case["path"] for case in manifest["cases"]}) == len(manifest["cases"])
    for case in manifest["cases"]:
        path = ROOT / case["path"]
        result = validate_deck_ir_text(path.read_text(encoding="utf-8"))
        assert result.value is None
        assert result.errors[0].code == case["code"]


def test_expected_hash_manifest_matches_checked_in_json() -> None:
    manifest = _json("samples/expected/manifest.json")
    listed = {record["path"] for record in manifest["assets"]}
    expected = {
        path.relative_to(ROOT).as_posix()
        for path in (ROOT / "samples" / "expected").glob("*.json")
        if path.name != "manifest.json"
    }
    assert listed == expected
    for record in manifest["assets"]:
        path = ROOT / record["path"]
        assert path.stat().st_size == record["bytes"]
        assert _sha256(path) == record["sha256"]


def test_official_word_outputs_match_structural_expected() -> None:
    expected = _json("samples/expected/word_official_outputs.expected.json")
    assert set(expected) == {"word_valid_01_plain", "word_valid_02_list", "word_valid_03_table"}
    for stem, facts in expected.items():
        path = ROOT / "samples" / "output" / "word" / f"{stem}.docx"
        assert _docx_facts(path) == facts
        assert facts["has_theme_red"] is True
        assert facts["has_legacy_word_blue"] is False


def test_full_deck_output_matches_structural_expected() -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx

    expected = _json("samples/expected/deck_valid_full.expected.json")
    deck = DeckIR.model_validate_json((ROOT / "samples" / "ir" / "deck_valid_full.json").read_text(encoding="utf-8"))
    path = ROOT / "samples" / "output" / "deck" / "deck_valid_full.pptx"
    presentation = Presentation(path)
    report = check_pptx(path, classification=deck.meta.classification)

    actual = {
        "slide_count": len(presentation.slides),
        "layouts": [slide.layout for slide in deck.slides],
        "titles": [getattr(slide, "title", None) for slide in deck.slides],
        "editable_table_slides": [
            index
            for index, slide in enumerate(presentation.slides, start=1)
            if any(getattr(shape, "has_table", False) for shape in slide.shapes)
        ],
        "editable_chart_slides": [
            index
            for index, slide in enumerate(presentation.slides, start=1)
            if any(getattr(shape, "has_chart", False) for shape in slide.shapes)
        ],
        "lint": {
            "errors": report.summary["errors"],
            "warnings": report.summary["warnings"],
            "infos": report.summary["infos"],
            "pass": report.summary["pass"],
        },
    }
    assert actual == expected


def test_lint_violation_asset_hits_declared_codes() -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx

    declaration = _json("samples/ir/deck_lint_violation.inject.json")
    expected = _json("samples/expected/deck_lint_violation.expected.json")
    deck = DeckIR.model_validate_json((ROOT / declaration["source"]).read_text(encoding="utf-8"))
    path = ROOT / "samples" / "output" / "deck" / "deck_lint_violation.pptx"
    report = check_pptx(path, classification=deck.meta.classification)
    codes = [item.code for item in report.items]

    assert codes == declaration["expected_codes"] == expected["codes"]
    assert report.summary == {key: expected[key] for key in ("errors", "warnings", "infos", "pass")}


def test_windows_wheel_manifest_and_hash_lock_are_auditable() -> None:
    manifest = _json("docs/wheelhouse-win312-manifest.json")
    lock = (ROOT / "requirements-win312.lock").read_text(encoding="utf-8")
    direct_requirements = {
        line.split("==", 1)[0].lower()
        for line in (ROOT / "requirements.txt").read_text(encoding="utf-8").splitlines()
        if line and not line.startswith("#")
    }
    manifest_names = {record["name"].lower() for record in manifest["wheels"]}

    assert manifest["target"] == {"implementation": "CPython", "platform": "win_amd64", "python_version": "3.12"}
    assert manifest["wheel_count"] == len(manifest["wheels"])
    assert direct_requirements <= manifest_names
    assert manifest["source_requirements_sha256"] == _sha256(ROOT / "requirements.txt")
    for record in manifest["wheels"]:
        assert re.fullmatch(r"[0-9a-f]{64}", record["sha256"])
        assert f"{record['name']}=={record['version']} --hash=sha256:{record['sha256']}" in lock
        local_wheel = ROOT / "wheelhouse" / record["filename"]
        if local_wheel.exists():
            assert local_wheel.stat().st_size == record["bytes"]
            assert _sha256(local_wheel) == record["sha256"]
    assert "Windows true-machine validation" in manifest["verification_boundary"]


def test_demo_logs_screenshots_and_manifest_are_real_assets() -> None:
    manifest = _json("samples/demo/manifest.json")
    assert manifest["cases"]["success_word"]["returncode"] == 0
    assert manifest["cases"]["success_deck"]["returncode"] == 0
    assert manifest["cases"]["failure_deck"]["returncode"] != 0
    assert "D003" in (ROOT / "samples" / "demo" / "logs" / "failure_deck.txt").read_text(encoding="utf-8")
    for record in manifest["assets"]:
        path = ROOT / record["path"]
        assert path.stat().st_size == record["bytes"]
        assert _sha256(path) == record["sha256"]
    for name in ("success.png", "failure.png"):
        with Image.open(ROOT / "samples" / "demo" / "screenshots" / name) as image:
            assert image.size == (1600, 900)
            assert len(image.getcolors(maxcolors=1_000_000) or []) > 3


def test_delivery_docs_cover_reviews_assets_and_all_18_cards() -> None:
    asset_doc = (ROOT / "docs" / "交付资产说明.md").read_text(encoding="utf-8")
    failure_review = (ROOT / "docs" / "reviews" / "失败文案评审记录.md").read_text(encoding="utf-8")
    theme_review = (ROOT / "docs" / "reviews" / "主题校准评审记录.md").read_text(encoding="utf-8")
    index = (ROOT / "docs" / "任务卡交付索引.md").read_text(encoding="utf-8")

    assert "build_delivery_assets.py" in asset_doc
    assert "Windows true-machine" not in asset_doc or "不等于" in asset_doc
    for code in [*(f"E00{value}" for value in range(1, 7)), *(f"D00{value}" for value in range(1, 7)), "W101", "W104", "I201"]:
        assert code in failure_review
    assert "#C7000B" in theme_review and "#C00000" in theme_review
    assert sum(line.startswith("| S") for line in index.splitlines()) == 18


def test_delivery_docs_are_current_and_demo_wrappers_are_thin() -> None:
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    usage = (ROOT / "docs" / "使用说明.md").read_text(encoding="utf-8")
    intranet = (ROOT / "docs" / "内网接入.md").read_text(encoding="utf-8")
    acceptance = (ROOT / "docs" / "验收手册.md").read_text(encoding="utf-8")
    style = (ROOT / "docs" / "风格规范.md").read_text(encoding="utf-8")
    taskbook = (ROOT / "docs" / "taskbook.md").read_text(encoding="utf-8")
    shell_wrapper = (ROOT / "scripts" / "record_demo.sh").read_text(encoding="utf-8")
    powershell_wrapper = (ROOT / "scripts" / "record_demo.ps1").read_text(encoding="utf-8")

    assert "项目汇报.pptx" in readme
    assert "samples/expected/项目汇报.pptx.expected.json" in usage
    assert "requirements-win312.lock" in intranet and "--require-hashes" in intranet
    assert "test_delivery_assets.py" in acceptance and "record_demo.py" in acceptance
    assert "SOURCE_FILE_SPEC.md" in style and "#C7000B" in style and "#C00000" in style
    assert "A.2  决议与剩余内网确认" in taskbook
    assert "chart 版式使用 python-pptx 原生可编辑图表" in taskbook
    assert shell_wrapper.count("\n") <= 5 and "record_demo.py" in shell_wrapper
    assert powershell_wrapper.count("\n") <= 5 and "record_demo.py" in powershell_wrapper


def _json(relative: str) -> dict:
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _document_ir_facts(document_ir) -> dict:
    return {
        "format": document_ir.source.format,
        "stats": document_ir.stats.model_dump(mode="json"),
        "warning_fragments": sorted(_warning_fragment(value) for value in document_ir.warnings),
        "block_types": [block.type for block in document_ir.content.blocks],
        "outline": [{"level": item.level, "text": item.text} for item in document_ir.content.outline],
        "sheets": [
            {
                "name": sheet.name,
                "nrows": sheet.nrows,
                "ncols": sheet.ncols,
                "header_guess": sheet.header_guess,
                "formula_count": sheet.formula_count,
                "merged_count": sheet.merged_count,
            }
            for sheet in document_ir.content.sheets
        ],
        "slides": [
            {
                "index": slide.index,
                "title": slide.title,
                "body_count": len(slide.bodies),
                "table_count": len(slide.tables),
                "notes": slide.notes,
            }
            for slide in document_ir.content.slides
        ],
    }


def _warning_fragment(value: str) -> str:
    for fragment in ("W103", "unsupported text boxes", "chart unsupported", "SmartArt"):
        if fragment in value:
            return fragment
    return value


def _docx_facts(path: Path) -> dict:
    document = Document(path)
    with zipfile.ZipFile(path) as package:
        xml = package.read("word/styles.xml").decode("utf-8") + package.read("word/document.xml").decode("utf-8")
    return {
        "paragraphs": [paragraph.text for paragraph in document.paragraphs if paragraph.text],
        "paragraph_styles": [paragraph.style.name for paragraph in document.paragraphs if paragraph.text],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
        "header": document.sections[0].header.paragraphs[0].text,
        "footer": document.sections[0].footer.paragraphs[0].text,
        "has_theme_red": "C7000B" in xml,
        "has_legacy_word_blue": "4F81BD" in xml.upper(),
    }
