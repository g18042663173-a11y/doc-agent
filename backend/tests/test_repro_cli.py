from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
PY_ENV = {**os.environ, "PYTHONPATH": str(ROOT / "backend")}


def test_parse_cli_writes_document_ir_json(tmp_path: Path) -> None:
    output = tmp_path / "doc.json"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.parse",
            "samples/input/quarterly_report.md",
            "--output",
            str(output),
        ],
        cwd=ROOT,
        env=PY_ENV,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["ir_type"] == "document"
    assert payload["source"]["format"] == "md"


def test_prompt_cli_writes_prompt_from_context(tmp_path: Path) -> None:
    context = tmp_path / "doc.json"
    prompt = tmp_path / "prompt.txt"
    subprocess.run(
        [sys.executable, "-m", "app.cli.parse", "samples/input/quarterly_report.md", "--output", str(context)],
        cwd=ROOT,
        env=PY_ENV,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=True,
    )

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.prompt",
            "--kind",
            "word",
            "--context",
            str(context),
            "--max-output-chars",
            "4000",
            "--output",
            str(prompt),
        ],
        cwd=ROOT,
        env=PY_ENV,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    prompt_text = prompt.read_text(encoding="utf-8")
    assert "WordIR v1.2" in prompt_text
    assert "最多 13 个 blocks" in prompt_text


def test_parse_then_prompt_is_byte_deterministic_across_independent_runs(tmp_path: Path) -> None:
    document_paths: list[Path] = []
    prompt_paths: list[Path] = []
    for run in range(2):
        context = tmp_path / f"document-{run}.json"
        prompt = tmp_path / f"prompt-{run}.txt"
        parse_result = subprocess.run(
            [sys.executable, "-m", "app.cli.parse", "samples/input/quarterly_report.md", "--output", str(context)],
            cwd=ROOT,
            env=PY_ENV,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
        )
        prompt_result = subprocess.run(
            [
                sys.executable,
                "-m",
                "app.cli.prompt",
                "--kind",
                "deck",
                "--context",
                str(context),
                "--output",
                str(prompt),
            ],
            cwd=ROOT,
            env=PY_ENV,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
        )
        assert parse_result.returncode == 0, parse_result.stderr
        assert prompt_result.returncode == 0, prompt_result.stderr
        document_paths.append(context)
        prompt_paths.append(prompt)

    assert document_paths[0].read_bytes() == document_paths[1].read_bytes()
    assert prompt_paths[0].read_bytes() == prompt_paths[1].read_bytes()


def test_semantic_fact_check_rejects_marker_only_target() -> None:
    import scripts.verify as verify

    document = {
        "ir_type": "document",
        "ir_version": "1.1",
        "source": {"filename": "周报.md", "format": "md"},
        "content": {
            "outline": [{"level": 1, "text": "项目周报"}],
            "blocks": [
                {"type": "heading", "level": 1, "text": "项目周报"},
                {"type": "paragraph", "text": "关键进展已经完成"},
                {"type": "table", "header": ["事项", "状态"], "rows": [["接口联调", "完成"]]},
            ],
        },
    }
    marker_only_target = {"meta": {"title": "项目周报"}}

    facts = verify._semantic_facts(document)
    missing = verify._missing_semantic_facts(facts, json.dumps(marker_only_target, ensure_ascii=False))

    assert {fact.value for fact in missing} >= {"关键进展已经完成", "事项", "接口联调", "完成"}


def test_demo_e2e_stub_word_generates_docx(tmp_path: Path) -> None:
    output_dir = tmp_path / "out"
    result = subprocess.run(
        [
            sys.executable,
            "scripts/demo_e2e.py",
            "samples/input/quarterly_report.md",
            "--target",
            "word",
            "--generator",
            "stub",
            "--output-dir",
            str(output_dir),
        ],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    assert (output_dir / "document_ir.json").exists()
    assert (output_dir / "prompt.txt").exists()
    assert (output_dir / "word.docx").exists()


def test_demo_e2e_stub_deck_generates_linted_pptx(tmp_path: Path) -> None:
    output_dir = tmp_path / "deck-out"
    result = subprocess.run(
        [
            sys.executable,
            "scripts/demo_e2e.py",
            "samples/input/quarterly_report.md",
            "--target",
            "deck",
            "--generator",
            "stub",
            "--lint",
            "--output-dir",
            str(output_dir),
        ],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    pptx = output_dir / "deck.pptx"
    report = json.loads((output_dir / "report.json").read_text(encoding="utf-8"))
    deck_ir = json.loads((output_dir / "deck_ir.json").read_text(encoding="utf-8"))
    assert pptx.exists()
    assert 5 <= len(Presentation(str(pptx)).slides) <= 12
    layouts = {slide["layout"] for slide in deck_ir["slides"]}
    assert "process_flow" in layouts
    assert "timeline" not in layouts
    assert "完成输入解析" in json.dumps(deck_ir, ensure_ascii=False)
    shape_names = {
        shape.name
        for slide in Presentation(str(pptx)).slides
        for shape in slide.shapes
    }
    assert any(name.startswith("HW_PROCESS_STEP:") for name in shape_names)
    assert not any(name.startswith("HW_TIMELINE_MILESTONE:") for name in shape_names)
    assert report["summary"]["pass"] is True


@pytest.mark.parametrize(
    ("suffix", "sample_name"),
    [
        ("md", "md_sample_01.md"),
        ("docx", "docx_sample_01.docx"),
        ("xlsx", "xlsx_sample_01.xlsx"),
        ("pptx", "pptx_sample_01.pptx"),
    ],
)
def test_demo_e2e_four_formats_parse_prompt_render_check_for_word_and_deck(tmp_path: Path, suffix: str, sample_name: str) -> None:
    sample = ROOT / "samples" / "input" / "parser_samples" / sample_name
    for target in ("word", "deck"):
        output_dir = tmp_path / f"{suffix}-{target}"
        command = [
            sys.executable,
            "scripts/demo_e2e.py",
            str(sample),
            "--target",
            target,
            "--generator",
            "stub",
            "--output-dir",
            str(output_dir),
        ]
        if target == "deck":
            command.append("--lint")
        result = subprocess.run(
            command,
            cwd=ROOT,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
        )

        assert result.returncode == 0, result.stdout + result.stderr
        assert (output_dir / "document_ir.json").exists()
        assert (output_dir / "prompt.txt").exists()
        generated_ir = output_dir / f"{target}_ir.json"
        assert generated_ir.exists()
        artifact = output_dir / ("word.docx" if target == "word" else "deck.pptx")
        assert artifact.exists()
        document = json.loads((output_dir / "document_ir.json").read_text(encoding="utf-8"))
        import scripts.verify as verify

        facts = verify._semantic_facts(document)
        assert len(facts) >= 3
        assert verify._missing_semantic_facts(facts, generated_ir.read_text(encoding="utf-8")) == []
        assert verify._missing_semantic_facts(facts, verify._artifact_text(artifact, target)) == []
        report = json.loads((output_dir / "report.json").read_text(encoding="utf-8"))
        assert report["summary"]["pass"] is True
