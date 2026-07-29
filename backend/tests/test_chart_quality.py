from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from app.ir.deck_ir import ChartSpec, DeckIR
from app.lint.pptx_lint import check_pptx
from app.rendering.pptx_renderer import (
    _chart_axis_scale,
    _chart_number_format,
    render_deck_ir,
)


@pytest.mark.parametrize(
    ("kind", "values", "expected_zero_side"),
    [
        ("bar", [10, 91], "minimum"),
        ("bar", [-91, -10], "maximum"),
        ("bar", [-30, 70], "inside"),
    ],
)
def test_bar_axis_uses_nice_scale_and_always_includes_zero(kind: str, values: list[float], expected_zero_side: str) -> None:
    spec = ChartSpec(kind=kind, categories=["A", "B"], series=[{"name": "值", "values": values}])

    minimum, maximum, major = _chart_axis_scale(spec)

    assert minimum < maximum
    assert major in {1, 2, 2.5, 5, 10, 20, 25, 50}
    assert minimum / major == pytest.approx(round(minimum / major))
    assert maximum / major == pytest.approx(round(maximum / major))
    if expected_zero_side == "minimum":
        assert minimum == 0
    elif expected_zero_side == "maximum":
        assert maximum == 0
    else:
        assert minimum < 0 < maximum


def test_line_narrow_range_uses_readable_nonzero_window() -> None:
    spec = ChartSpec(kind="line", categories=["1月", "2月", "3月"], series=[{"name": "可用率", "values": [98.1, 98.4, 98.7]}])

    minimum, maximum, major = _chart_axis_scale(spec)

    assert 90 < minimum < 98.1
    assert maximum > 98.7
    assert major > 0


def test_number_format_uses_literal_percent_suffix_for_whole_percent_values() -> None:
    percent = ChartSpec(kind="bar", categories=["A", "B"], series=[{"name": "通过率", "values": [75, 92]}], unit="%")
    decimal = ChartSpec(kind="line", categories=["A", "B"], series=[{"name": "时延", "values": [1.2, 2.35]}], unit="ms")

    assert _chart_number_format(percent) == '#,##0"%"'
    assert _chart_number_format(decimal) == '#,##0.00"ms"'


def test_rendered_chart_applies_formats_long_label_skip_and_pie_percent_labels(tmp_path: Path) -> None:
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "图表质量"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "长标签柱图",
                    "chart": {
                        "kind": "bar",
                        "categories": [f"第{index}个非常长的中文类别标签" for index in range(1, 14)],
                        "series": [{"name": "通过率", "values": list(range(70, 83))}],
                        "unit": "%",
                        "legend_position": "none",
                    },
                },
                {
                    "layout": "chart",
                    "title": "构成占比",
                    "chart": {
                        "kind": "pie",
                        "categories": ["研发", "测试", "交付"],
                        "series": [{"name": "投入", "values": [50, 30, 20]}],
                        "legend_position": "right",
                    },
                },
            ],
        }
    )

    output = render_deck_ir(deck, tmp_path / "charts.pptx")
    presentation = Presentation(output)
    bar = next(shape.chart for shape in presentation.slides[0].shapes if getattr(shape, "has_chart", False))
    pie = next(shape.chart for shape in presentation.slides[1].shapes if getattr(shape, "has_chart", False))

    assert bar.value_axis.tick_labels.number_format == '#,##0"%"'
    assert bar.value_axis.major_unit is not None
    assert bar.category_axis._element.find(qn("c:tickLblSkip")).get("val") == "2"
    assert pie.plots[0].data_labels.show_category_name is True
    assert pie.plots[0].data_labels.show_percentage is True
    assert pie.plots[0].data_labels.show_value is False


def test_chart_lint_reports_density_and_pie_suitability(tmp_path: Path) -> None:
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "图表风险"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "类别过多",
                    "chart": {
                        "kind": "pie",
                        "categories": [f"类别{index}" for index in range(7)],
                        "series": [{"name": "数量", "values": [5, 4, 3, 2, 1, 0, -1]}],
                    },
                }
            ],
        }
    )
    output = render_deck_ir(deck, tmp_path / "pie-risk.pptx")

    report = check_pptx(output, classification=deck.meta.classification)

    assert any(item.code == "HW-W12" for item in report.items)
