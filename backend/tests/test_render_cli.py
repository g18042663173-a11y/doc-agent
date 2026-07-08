from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]


def test_render_cli_writes_word_docx(tmp_path: Path) -> None:
    output = tmp_path / "sample.docx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "word",
            "samples/ir/word_valid_03_table.json",
            "--output",
            str(output),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    assert output.exists()
    assert Document(str(output)).tables[0].cell(0, 0).text == "风险"


def test_render_cli_reports_validation_code_for_invalid_word_ir(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.render",
            "--type",
            "word",
            "samples/ir/word_invalid_e004_ragged_table.json",
            "--output",
            str(tmp_path / "bad.docx"),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT / "backend")},
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 1
    assert "E004" in result.stderr
