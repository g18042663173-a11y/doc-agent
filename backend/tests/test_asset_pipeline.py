from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from app.assets.contracts import AssetManifest, AssetRecord
from app.assets.errors import AssetError
from app.assets.pipeline import AssetLimits, load_asset_manifest, normalize_assets
from app.assets.extract import extract_office_image_files
from app.assets.provider import DisabledImageProvider, ImageRequest
from app.ir.deck_ir import DeckIR
from app.rendering.pptx_renderer import render_deck_ir
from app.cli.assets import main as assets_cli_main
from app.cli.render import main as render_cli_main


def _image(path: Path, *, size: tuple[int, int] = (80, 40), mode: str = "RGB", fmt: str = "PNG") -> Path:
    Image.new(mode, size, (230, 20, 60, 180) if mode == "RGBA" else (230, 20, 60)).save(path, format=fmt)
    return path


def test_normalize_assets_decodes_strips_metadata_and_loads_registry(tmp_path: Path) -> None:
    source = _image(tmp_path / "evidence.png", mode="RGBA")

    manifest = normalize_assets([source], tmp_path / "assets")

    assert manifest.manifest_version == "1.0"
    assert len(manifest.assets) == 1
    record = manifest.assets[0]
    assert record.asset_id.startswith("asset-")
    assert record.media_type == "image/png"
    assert record.width == 80 and record.height == 40
    assert record.has_alpha is True
    assert record.metadata_removed is True
    manifest_path = tmp_path / "assets" / "asset_manifest.json"
    registry = load_asset_manifest(manifest_path)
    assert registry.resolve(record.asset_id).is_file()
    assert registry.record(record.asset_id).normalized_sha256 == sha256(
        registry.resolve(record.asset_id).read_bytes()
    ).hexdigest()


def test_normalize_assets_rejects_extension_spoofing_and_pixel_limit(tmp_path: Path) -> None:
    spoofed = tmp_path / "fake.png"
    spoofed.write_bytes(b"not an image")
    with pytest.raises(AssetError, match="A002"):
        normalize_assets([spoofed], tmp_path / "spoofed")

    large = _image(tmp_path / "large.png", size=(20, 20))
    with pytest.raises(AssetError, match="A003"):
        normalize_assets([large], tmp_path / "large-assets", limits=AssetLimits(max_pixels_per_image=100))


def test_asset_manifest_rejects_duplicate_ids_and_unsafe_paths(tmp_path: Path) -> None:
    digest = "0" * 64
    base = {
        "asset_id": "asset-000000000000",
        "source_type": "upload",
        "source_filename": "a.png",
        "original_media_type": "image/png",
        "media_type": "image/png",
        "original_sha256": digest,
        "normalized_sha256": digest,
        "width": 1,
        "height": 1,
        "pixel_count": 1,
        "bytes": 1,
        "relative_path": "../escape.png",
        "has_alpha": False,
        "exif_orientation_applied": False,
        "metadata_removed": True,
        "animated": False,
    }
    with pytest.raises(ValidationError):
        AssetRecord.model_validate(base)

    base["relative_path"] = "normalized/a.png"
    with pytest.raises(ValidationError, match="unique"):
        AssetManifest.model_validate({"manifest_version": "1.0", "assets": [base, base]})


def test_load_asset_manifest_rejects_missing_or_changed_file(tmp_path: Path) -> None:
    source = _image(tmp_path / "source.png")
    manifest = normalize_assets([source], tmp_path / "assets")
    record = manifest.assets[0]
    normalized = tmp_path / "assets" / record.relative_path
    normalized.write_bytes(b"changed")

    with pytest.raises(AssetError, match="A006"):
        load_asset_manifest(tmp_path / "assets" / "asset_manifest.json")


def test_asset_manifest_schema_is_strict_and_json_serializable(tmp_path: Path) -> None:
    manifest = normalize_assets([_image(tmp_path / "source.webp", fmt="WEBP")], tmp_path / "assets")
    payload = json.loads(manifest.model_dump_json())
    assert payload["manifest_version"] == "1.0"
    assert AssetManifest.model_json_schema()["additionalProperties"] is False


def test_deck_ir_19_image_ref_renders_real_picture_and_usage_audit(tmp_path: Path) -> None:
    source = _image(tmp_path / "wide.png", size=(400, 100))
    normalize_assets([source], tmp_path / "assets")
    registry = load_asset_manifest(tmp_path / "assets" / "asset_manifest.json")
    asset_id = registry.manifest.assets[0].asset_id
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "图片证据", "classification": "公开"},
            "slides": [{"layout": "image", "title": "证据图片保持比例", "image_ref": asset_id}],
        }
    )

    output = render_deck_ir(deck, tmp_path / "deck.pptx", asset_registry=registry)

    slide = Presentation(output).slides[0]
    pictures = [shape for shape in slide.shapes if shape.name.startswith("HW_ASSET_IMAGE:")]
    assert len(pictures) == 1
    assert pictures[0].width / pictures[0].height == pytest.approx(4.0, rel=0.01)
    audit = json.loads((tmp_path / "asset_usage_audit.json").read_text(encoding="utf-8"))
    assert audit["usages"][0]["asset_id"] == asset_id
    assert audit["usages"][0]["fit"] == "contain"


def test_deck_ir_image_ref_without_registry_fails_instead_of_using_placeholder(tmp_path: Path) -> None:
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "缺失资产"},
            "slides": [
                {"layout": "image", "title": "必须解析引用", "image_ref": "asset-missing", "placeholder": "不能降级"}
            ],
        }
    )

    with pytest.raises(AssetError, match="A005"):
        render_deck_ir(deck, tmp_path / "deck.pptx")


def test_asset_cli_writes_manifest(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = _image(tmp_path / "source.png")
    output_dir = tmp_path / "cli-assets"
    assert assets_cli_main([str(source), "--output-dir", str(output_dir)]) == 0
    assert (output_dir / "asset_manifest.json").is_file()
    assert str(output_dir / "asset_manifest.json") in capsys.readouterr().out


def test_extract_office_images_uses_safe_media_members_and_appends_manifest(tmp_path: Path) -> None:
    embedded = tmp_path / "embedded.png"
    _image(embedded)
    package = tmp_path / "source.docx"
    with ZipFile(package, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/media/image1.png", embedded.read_bytes())
        archive.writestr("word/media/vector.emf", b"unsupported")
        archive.writestr("../escape.png", embedded.read_bytes())

    extracted = extract_office_image_files(package, tmp_path / "raw")
    manifest = normalize_assets(extracted, tmp_path / "assets", source_type="document")
    uploaded = _image(tmp_path / "uploaded.jpg", fmt="JPEG")
    merged = normalize_assets([uploaded], tmp_path / "assets", source_type="upload", append=True)

    assert len(extracted) == 1
    assert {asset.source_type for asset in merged.assets} == {"document", "upload"}
    assert len(manifest.assets) == 1


def test_extract_office_images_streams_members_without_zipfile_read(tmp_path: Path, monkeypatch) -> None:
    embedded = _image(tmp_path / "embedded.png")
    package = tmp_path / "source.docx"
    with ZipFile(package, "w") as archive:
        archive.writestr("word/media/image1.png", embedded.read_bytes())

    def fail_read(*_args, **_kwargs):
        raise AssertionError("media extraction must stream through ZipFile.open")

    monkeypatch.setattr(ZipFile, "read", fail_read)
    extracted = extract_office_image_files(package, tmp_path / "raw")

    assert len(extracted) == 1
    assert extracted[0].read_bytes() == embedded.read_bytes()


def test_normalize_assets_enforces_limit_against_normalized_png_size(tmp_path: Path) -> None:
    source = _image(tmp_path / "small.webp", size=(100, 100), fmt="WEBP")
    source_bytes = source.stat().st_size
    probe = normalize_assets([source], tmp_path / "probe")
    normalized_bytes = probe.assets[0].bytes
    assert normalized_bytes > source_bytes

    with pytest.raises(AssetError, match="规范化图片总大小"):
        normalize_assets(
            [source],
            tmp_path / "limited",
            limits=AssetLimits(max_total_bytes=(source_bytes + normalized_bytes) // 2),
        )


def test_render_cli_accepts_asset_manifest_for_deck(tmp_path: Path) -> None:
    source = _image(tmp_path / "source.png")
    manifest = normalize_assets([source], tmp_path / "assets")
    ir_path = tmp_path / "deck.json"
    ir_path.write_text(
        json.dumps(
            {
                "ir_type": "deck",
                "ir_version": "2.0",
                "meta": {"title": "CLI 图片", "classification": "公开"},
                "slides": [{"layout": "image", "title": "CLI 图片", "image_ref": manifest.assets[0].asset_id}],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    output = tmp_path / "deck.pptx"

    assert render_cli_main(
        ["--type", "deck", str(ir_path), "--output", str(output), "--asset-manifest", str(tmp_path / "assets" / "asset_manifest.json")]
    ) == 0
    assert output.is_file()


def test_render_cli_preserves_asset_error_code_for_missing_reference(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    ir_path = tmp_path / "deck.json"
    ir_path.write_text(
        json.dumps(
            {
                "ir_type": "deck",
                "ir_version": "2.0",
                "meta": {"title": "缺失图片", "classification": "公开"},
                "slides": [{"layout": "image", "title": "缺失图片", "image_ref": "asset-missing"}],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    result = render_cli_main(["--type", "deck", str(ir_path), "--output", str(tmp_path / "deck.pptx")])

    assert result == 1
    error = capsys.readouterr().err
    assert '"code": "A005"' in error
    assert "Traceback" not in error


def test_disabled_image_provider_preserves_offline_boundary(tmp_path: Path) -> None:
    with pytest.raises(AssetError, match="A006"):
        DisabledImageProvider().generate(ImageRequest(prompt="生成插图", width=1024, height=768, output_dir=tmp_path))
