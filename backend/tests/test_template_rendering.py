from __future__ import annotations

import json
from pathlib import Path
from typing import Callable
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
import pytest

from app.assets.pipeline import load_asset_manifest, normalize_assets
from app.ir.deck_ir import DeckIR
from app.lint.pptx_lint import check_pptx
from app.template.package import (
    TemplateInputError,
    sanitize_template_hyperlinks,
    validate_pptx_package,
    validate_template_package,
)
from app.template.planner import build_template_plan
from app.template.profile import extract_template_profile
from app.template.renderer import render_deck_ir_with_template


def test_profile_and_plan_use_regular_text_shapes_without_placeholders(tmp_path: Path) -> None:
    template = _template(tmp_path)
    profile = extract_template_profile(template)
    deck = _deck()

    plan = build_template_plan(deck, profile)

    assert profile.profile_version == "1.0"
    assert profile.source.sha256 == _sha256_from_profile_file(template, profile.source.sha256)
    assert profile.slides[0].role == "cover"
    assert profile.slides[1].role == "title_bullets"
    assert any(shape.role == "title" for shape in profile.slides[1].shapes)
    assert any(shape.role == "body" for shape in profile.slides[1].shapes)
    assert plan.slides[0].strategy == "prototype_replace"
    assert plan.slides[1].strategy == "prototype_replace"
    assert plan.slides[2].strategy == "master_redraw"
    assert any(warning.code == "W201" for warning in plan.warnings)


def test_template_renderer_preserves_image_relationships_and_redraws_native_chart(tmp_path: Path) -> None:
    template = _template(tmp_path)
    output = tmp_path / "output" / "deck.pptx"

    result = render_deck_ir_with_template(_deck(), template, output)

    assert result.artifact_path == output.resolve()
    assert result.package_report["pass"] is True
    assert json.loads(result.profile_path.read_text(encoding="utf-8"))["profile_version"] == "1.0"
    assert json.loads(result.plan_path.read_text(encoding="utf-8"))["plan_version"] == "1.1"
    assert json.loads(result.structure_json_path.read_text(encoding="utf-8"))["filename"] == template.name
    assert result.structure_markdown_path.read_text(encoding="utf-8").startswith("# PPT 模板结构：")
    replacement_audit = json.loads(result.replacement_audit_path.read_text(encoding="utf-8"))
    assert replacement_audit["audit_version"] == "1.1"
    assert all(slide["planned_shapes_exist"] for slide in replacement_audit["slides"])
    assert "交付方案" not in result.replacement_audit_path.read_text(encoding="utf-8")
    rendered = Presentation(output)
    assert len(rendered.slides) == 3
    all_text = "\n".join(shape.text for slide in rendered.slides for shape in slide.shapes if hasattr(shape, "text"))
    assert "TEMPLATE FILLER" not in all_text
    assert "交付方案" in all_text
    assert all(
        "HUAWEI CONFIDENTIAL" in "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        for slide in rendered.slides
    )
    assert any(
        getattr(shape, "shape_type", None) is not None and "PICTURE" in str(shape.shape_type)
        for shape in rendered.slides[0].shapes
    )
    assert any(getattr(shape, "has_chart", False) for shape in rendered.slides[2].shapes)
    assert validate_pptx_package(output)["pass"] is True


def test_template_replacement_audit_records_image_hash_fit_crop_and_focus(tmp_path: Path) -> None:
    template = _template(tmp_path)
    source = tmp_path / "evidence.png"
    Image.new("RGB", (800, 400), (200, 20, 30)).save(source)
    manifest = normalize_assets([source], tmp_path / "assets")
    registry = load_asset_manifest(tmp_path / "assets" / "asset_manifest.json")
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "图片审计", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "image",
                    "title": "图片审计",
                    "image_ref": manifest.assets[0].asset_id,
                    "fit": "cover",
                    "focal_x": 0.2,
                    "focal_y": 0.7,
                    "alt": "证据图片",
                }
            ],
        }
    )

    result = render_deck_ir_with_template(
        deck, template, tmp_path / "image-template.pptx", asset_registry=registry
    )
    audit = json.loads(result.replacement_audit_path.read_text(encoding="utf-8"))
    image_record = next(shape for shape in audit["slides"][0]["shapes"] if shape["action"] == "image_replaced")
    assert image_record["asset_sha256"] == manifest.assets[0].normalized_sha256
    assert image_record["image_fit"] == "cover"
    assert image_record["focal_x"] == 0.2
    assert image_record["crop_left"] + image_record["crop_right"] >= 0


def test_template_aware_lint_ignores_profiled_decorations_but_not_generated_content(tmp_path: Path) -> None:
    template = tmp_path / "decoration-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(slide, "TEMPLATE TITLE", 1.0, 1.0, 8.0, 0.8, 28)
    _textbox(slide, "TEMPLATE SUBTITLE", 1.0, 2.0, 8.0, 0.5, 18)
    decoration = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        presentation.slide_width,
        Inches(0.08),
    )
    decoration.fill.solid()
    presentation.save(template)

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "装饰忽略", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "装饰忽略", "subtitle": "正文仍在安全边距内"}],
        }
    )
    result = render_deck_ir_with_template(deck, template, tmp_path / "decorated.pptx")
    report = check_pptx(
        result.artifact_path,
        classification="HUAWEI CONFIDENTIAL",
        template_profile=result.profile,
    )

    names = [shape.name for shape in Presentation(result.artifact_path).slides[0].shapes]
    assert any(name.startswith("HW_TEMPLATE_DECORATION:") for name in names)
    assert "HW-W07" not in [item.code for item in report.items]


def test_package_validator_reports_a_dangling_relationship(tmp_path: Path) -> None:
    template = _template(tmp_path)
    broken = tmp_path / "broken.pptx"
    with zipfile.ZipFile(template) as source, zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename == "ppt/media/image1.png":
                continue
            target.writestr(info, source.read(info.filename))

    report = validate_pptx_package(broken)

    assert report["pass"] is False
    assert any("关系目标不存在" in error for error in report["errors"])


def test_package_validator_reports_xml_relationship_references_without_matching_rels(tmp_path: Path) -> None:
    template = _template(tmp_path)
    broken = tmp_path / "missing-xml-rel.pptx"
    _rewrite_package(
        template,
        broken,
        replacements={
            "ppt/slides/slide1.xml": lambda data: data.replace(
                b"<p:nvPr/>",
                b'<p:nvPr><p:custDataLst><p:tags r:id="rIdMissing"/></p:custDataLst></p:nvPr>',
                1,
            )
        },
    )

    report = validate_pptx_package(broken)

    assert report["pass"] is False
    assert any("XML 引用了不存在的关系" in error for error in report["errors"])


def test_template_renderer_rejects_plan_deck_length_mismatch_and_bad_prototype_index(tmp_path: Path) -> None:
    from app.template.package import TemplateInputError

    template = _template(tmp_path)
    deck = _deck()
    profile = extract_template_profile(template)
    plan = build_template_plan(deck, profile)

    short_plan = plan.model_copy(deep=True)
    short_plan.slides = short_plan.slides[:1]
    with pytest.raises(TemplateInputError, match="页与 Deck .* 页不一致"):
        render_deck_ir_with_template(
            deck,
            template,
            tmp_path / "mismatch.pptx",
            profile=profile,
            plan=short_plan,
        )

    out_of_range = plan.model_copy(deep=True)
    for slide_plan in out_of_range.slides:
        slide_plan.prototype_index = 999
    with pytest.raises(TemplateInputError, match="原型页"):
        render_deck_ir_with_template(
            deck,
            template,
            tmp_path / "bad-index.pptx",
            profile=profile,
            plan=out_of_range,
        )


def test_template_renderer_records_w201_and_redraws_when_replacement_overflows(tmp_path: Path) -> None:
    template = _template(tmp_path)
    presentation = Presentation(template)
    presentation.slides[1].shapes[1].height = Inches(0.25)
    presentation.save(template)
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "容量回退", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "title_bullets",
                    "title": "原型容量不足时安全重绘",
                    "bullets": [
                        {"text": f"第{index}条" + "容量边界内容" * 9, "level": 1}
                        for index in range(1, 8)
                    ],
                }
            ],
        }
    )
    profile = extract_template_profile(template)
    plan = build_template_plan(deck, profile)
    assert plan.slides[0].strategy == "prototype_replace"

    result = render_deck_ir_with_template(
        deck,
        template,
        tmp_path / "overflow.pptx",
        profile=profile,
        plan=plan,
    )

    persisted_plan = json.loads(result.plan_path.read_text(encoding="utf-8"))
    assert result.plan.slides[0].strategy == "master_redraw"
    assert persisted_plan["slides"][0]["strategy"] == "master_redraw"
    assert any(warning.code == "W201" and "最小字号" in warning.message for warning in result.plan.warnings)


def test_missing_template_shape_font_is_replaced_and_recorded_as_w202(tmp_path: Path) -> None:
    template = _template(tmp_path)
    presentation = Presentation(template)
    for run in presentation.slides[1].shapes[1].text_frame.paragraphs[0].runs:
        run.font.name = "Missing Corporate Font"
    presentation.save(template)
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "字体替代", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "title_bullets",
                    "title": "缺失字体有显式审计记录",
                    "bullets": [{"text": "正文使用当前系统可用的兼容字体", "level": 1}],
                }
            ],
        }
    )
    profile = extract_template_profile(template)
    plan = build_template_plan(deck, profile)

    result = render_deck_ir_with_template(
        deck,
        template,
        tmp_path / "font-substitution.pptx",
        profile=profile,
        plan=plan,
    )

    assert any(
        warning.code == "W202" and "Missing Corporate Font" in warning.message
        for warning in result.plan.warnings
    )
    rendered = Presentation(result.artifact_path)
    rendered_fonts = {
        run.font.name
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    }
    assert "Missing Corporate Font" not in rendered_fonts


def test_template_package_rejects_external_relationship_path_traversal_and_embedded_objects(
    tmp_path: Path,
) -> None:
    template = _template(tmp_path)

    external = tmp_path / "external.pptx"
    _rewrite_package(
        template,
        external,
        replacements={
            "_rels/.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdExternal" Type="http://example.invalid/external" '
                b'Target="https://example.invalid/asset" TargetMode="External"/></Relationships>',
            )
        },
    )
    with pytest.raises(TemplateInputError, match="E003 at template_file:.*外部关系"):
        validate_template_package(external)

    traversal = tmp_path / "traversal.pptx"
    _rewrite_package(template, traversal, extra={"../escape.xml": b"<escape/>"})
    with pytest.raises(TemplateInputError, match="E003 at template_file:.*不安全路径"):
        validate_template_package(traversal)

    embedded = tmp_path / "embedded.pptx"
    _rewrite_package(template, embedded, extra={"ppt/embeddings/oleObject1.bin": b"unsafe"})
    with pytest.raises(TemplateInputError, match="E003 at template_file:.*不安全嵌入部件"):
        validate_template_package(embedded)

    referenced_ole = tmp_path / "referenced-ole.pptx"
    _rewrite_package(
        template,
        referenced_ole,
        replacements={
            "ppt/slides/_rels/slide1.xml.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/oleObject" '
                b'Target="../embeddings/oleObject1.bin"/></Relationships>',
            )
        },
        extra={"ppt/embeddings/oleObject1.bin": b"unsafe"},
    )
    with pytest.raises(
        TemplateInputError,
        match=r"E003 at template_file:.*第 1 页.*ppt/embeddings/oleObject1\.bin",
    ):
        validate_template_package(referenced_ole)


def test_template_hyperlink_sanitizer_creates_valid_copy_without_mutating_original(tmp_path: Path) -> None:
    template = _template(tmp_path)
    linked = tmp_path / "linked-template.pptx"
    safe_copy = tmp_path / "linked-template-safe-copy.pptx"
    original_bytes = template.read_bytes()

    def add_hyperlink_reference(data: bytes) -> bytes:
        root_marker = b"<p:sldLayout "
        assert root_marker in data
        data = data.replace(
            root_marker,
            b'<p:sldLayout xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            b'xmlns:test="urn:codex:test" mc:Ignorable="test" ',
            1,
        )
        marker = b'<p:cNvPr id="1" name=""/>'
        assert marker in data
        return data.replace(
            marker,
            b'<p:cNvPr id="1" name=""><a:hlinkClick r:id="rIdHyperlink"/></p:cNvPr>',
            1,
        )

    _rewrite_package(
        template,
        linked,
        replacements={
            "ppt/slideLayouts/slideLayout1.xml": add_hyperlink_reference,
            "ppt/slideLayouts/_rels/slideLayout1.xml.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdHyperlink" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/hyperlink" '
                b'Target="https://example.invalid/template" TargetMode="External"/></Relationships>',
            ),
        },
    )

    with pytest.raises(TemplateInputError, match="外部关系"):
        validate_template_package(linked)

    removed = sanitize_template_hyperlinks(linked, safe_copy)

    assert removed == 1
    assert template.read_bytes() == original_bytes
    assert validate_template_package(safe_copy)["pass"] is True
    with zipfile.ZipFile(safe_copy) as package:
        relationships = package.read("ppt/slideLayouts/_rels/slideLayout1.xml.rels")
        layout = package.read("ppt/slideLayouts/slideLayout1.xml")
    assert b"TargetMode=\"External\"" not in relationships
    assert b"rIdHyperlink" not in layout
    assert b'xmlns:test="urn:codex:test"' in layout
    assert b'Ignorable="test"' in layout


def test_template_hyperlink_sanitizer_refuses_non_hyperlink_external_relationship(tmp_path: Path) -> None:
    template = _template(tmp_path)
    external = tmp_path / "external-link.pptx"
    safe_copy = tmp_path / "external-link-safe-copy.pptx"
    _rewrite_package(
        template,
        external,
        replacements={
            "ppt/slides/_rels/slide1.xml.rels": lambda data: data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdExternal" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/externalLink" '
                b'Target="https://example.invalid/data" TargetMode="External"/></Relationships>',
            )
        },
    )

    with pytest.raises(TemplateInputError, match="仅支持移除外部超链接"):
        sanitize_template_hyperlinks(external, safe_copy)

    assert not safe_copy.exists()


def test_template_package_enforces_compressed_uncompressed_slide_and_shape_limits(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app.template.package as package_module
    import app.template.profile as profile_module

    template = _template(tmp_path)
    monkeypatch.setattr(package_module, "MAX_TEMPLATE_BYTES", 1)
    with pytest.raises(TemplateInputError, match="压缩包超过 50 MB"):
        validate_template_package(template)

    monkeypatch.setattr(package_module, "MAX_TEMPLATE_BYTES", 50 * 1024 * 1024)
    monkeypatch.setattr(package_module, "MAX_TEMPLATE_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(TemplateInputError, match="解包体积超过 500 MB"):
        validate_template_package(template)

    monkeypatch.setattr(package_module, "MAX_TEMPLATE_UNCOMPRESSED_BYTES", 500 * 1024 * 1024)
    monkeypatch.setattr(package_module, "MAX_TEMPLATE_SLIDES", 2)
    with pytest.raises(TemplateInputError, match="页数必须在 1-200 页"):
        validate_template_package(template)

    monkeypatch.setattr(package_module, "MAX_TEMPLATE_SLIDES", 200)
    monkeypatch.setattr(profile_module, "MAX_TEMPLATE_SHAPES", 1)
    with pytest.raises(TemplateInputError, match="形状总数超过 5000"):
        extract_template_profile(template)


def test_template_transition_and_timing_are_not_copied_to_output(tmp_path: Path) -> None:
    template = _template(tmp_path)
    animated = tmp_path / "animated.pptx"
    _rewrite_package(
        template,
        animated,
        replacements={
            "ppt/slides/slide1.xml": lambda data: data.replace(
                b"</p:sld>",
                b"<p:transition/><p:timing><p:tnLst/></p:timing></p:sld>",
            )
        },
    )
    profile = extract_template_profile(animated)
    assert profile.slides[0].has_transition is True
    assert profile.slides[0].has_timing is True

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "无动画输出", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "无动画输出", "subtitle": "模板原型只复制安全形状"}],
        }
    )
    result = render_deck_ir_with_template(deck, animated, tmp_path / "no-animation.pptx", profile=profile)

    with zipfile.ZipFile(result.artifact_path) as package:
        slide_parts = [
            package.read(name)
            for name in package.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
    assert slide_parts
    assert all(b"<p:transition" not in slide_xml for slide_xml in slide_parts)
    assert all(b"<p:timing" not in slide_xml for slide_xml in slide_parts)


def test_stripped_tag_relationship_does_not_leave_invalid_empty_tags_element(tmp_path: Path) -> None:
    template = _template(tmp_path)
    tagged = tmp_path / "tagged.pptx"

    def add_tag_reference(data: bytes) -> bytes:
        return data.replace(
            b"<p:nvPr/>",
            b'<p:nvPr><p:custDataLst><p:tags r:id="rIdTag"/></p:custDataLst></p:nvPr>',
            1,
        )

    def add_tag_relationship(data: bytes) -> bytes:
        return data.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdTag" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags" '
            b'Target="../tags/tag1.xml"/></Relationships>',
        )

    def add_tag_content_type(data: bytes) -> bytes:
        return data.replace(
            b"</Types>",
            b'<Override PartName="/ppt/tags/tag1.xml" '
            b'ContentType="application/vnd.openxmlformats-officedocument.presentationml.tags+xml"/></Types>',
        )

    _rewrite_package(
        template,
        tagged,
        replacements={
            "ppt/slides/slide1.xml": add_tag_reference,
            "ppt/slides/_rels/slide1.xml.rels": add_tag_relationship,
            "[Content_Types].xml": add_tag_content_type,
        },
        extra={
            "ppt/tags/tag1.xml": (
                b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                b'<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
            )
        },
    )
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "安全标签清理", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "安全标签清理", "subtitle": "不复制模板 tags 关系"}],
        }
    )

    result = render_deck_ir_with_template(deck, tagged, tmp_path / "tag-clean.pptx")

    with zipfile.ZipFile(result.artifact_path) as package:
        output_slide = next(
            package.read(name)
            for name in package.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
    assert b"<p:tags" not in output_slide
    assert b"<p:custDataLst" not in output_slide


def _template(tmp_path: Path) -> Path:
    image_path = tmp_path / "brand.png"
    Image.new("RGB", (80, 40), "#3494BA").save(image_path)
    path = tmp_path / "template.pptx"
    presentation = Presentation()

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(cover, "TEMPLATE FILLER COVER", 1.0, 1.7, 8.8, 1.0, 32)
    _textbox(cover, "TEMPLATE FILLER SUBTITLE", 1.0, 3.0, 8.0, 0.6, 18)
    cover.shapes.add_picture(str(image_path), Inches(10.5), Inches(0.5), Inches(1.6), Inches(0.8))

    body = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(body, "TEMPLATE FILLER TITLE", 0.7, 0.45, 11.0, 0.7, 26)
    _textbox(body, "TEMPLATE FILLER BODY WITH ENOUGH CAPACITY FOR MULTIPLE BULLETS", 0.9, 1.7, 11.2, 4.5, 18)

    generic = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(generic, "TEMPLATE FILLER GENERIC", 0.7, 0.45, 11.0, 0.7, 26)
    presentation.save(path)
    return path


def _deck() -> DeckIR:
    return DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.9",
            "meta": {"title": "交付方案", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {"layout": "cover", "title": "交付方案", "subtitle": "模板驱动生成"},
                {
                    "layout": "title_bullets",
                    "title": "关键结论",
                    "bullets": [{"text": "内容来自已校验 DeckIR", "level": 1}, {"text": "模板只影响确定性渲染", "level": 1}],
                },
                {
                    "layout": "chart",
                    "title": "交付效率持续提升",
                    "chart": {
                        "kind": "bar",
                        "categories": ["第一阶段", "第二阶段"],
                        "series": [{"name": "效率", "values": [70, 92]}],
                        "unit": "%",
                        "legend_position": "none",
                    },
                },
            ],
        }
    )


def _textbox(slide, text: str, left: float, top: float, width: float, height: float, font_size: float) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(font_size)
    shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"


def _sha256_from_profile_file(path: Path, digest: str) -> str:
    assert path.is_file()
    return digest


def _rewrite_package(
    source_path: Path,
    output_path: Path,
    *,
    replacements: dict[str, Callable[[bytes], bytes]] | None = None,
    extra: dict[str, bytes] | None = None,
) -> None:
    replacements = replacements or {}
    with zipfile.ZipFile(source_path) as source, zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            transform = replacements.get(info.filename)
            if transform is not None:
                data = transform(data)
            target.writestr(info, data)
        for name, data in (extra or {}).items():
            target.writestr(name, data)


def test_plan_falls_back_when_agenda_content_exceeds_prototype_slots(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.template.planner import build_template_plan
    from app.template.profile import extract_template_profile

    path = tmp_path / "slots.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index in range(5):
        _textbox(slide, f"AGENDA SLOT {index + 1}", 0.8, 1.0 + index * 0.6, 9.0, 0.5, 18)
    presentation.save(path)
    profile = extract_template_profile(path)

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "T", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "agenda", "title": "目录", "items": [f"item{i}" for i in range(8)]}],
        }
    )
    plan = build_template_plan(deck, profile)

    assert plan.slides[0].strategy == "master_redraw"
    assert any(warning.code == "W201" for warning in plan.slides[0].warnings)


def test_cover_optional_metadata_does_not_force_fallback_but_maps_when_slot_exists(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.template.planner import build_template_plan
    from app.template.profile import extract_template_profile

    def make_template(with_body: bool) -> Path:
        path = tmp_path / f"cover-{with_body}.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _textbox(slide, "TEMPLATE TITLE", 1.0, 0.5, 8.0, 1.0, 32)
        _textbox(slide, "TEMPLATE SUBTITLE", 1.0, 2.0, 6.0, 0.7, 18)
        if with_body:
            _textbox(slide, "TEMPLATE BODY", 1.0, 3.0, 6.0, 0.6, 14)
        presentation.save(path)
        return path

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.1",
            "meta": {"title": "T", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "T", "subtitle": "S", "presenter": "张三", "date": "2026-08-13"}],
        }
    )

    profile_without_body = extract_template_profile(make_template(with_body=False))
    plan = build_template_plan(deck, profile_without_body)
    assert plan.slides[0].strategy == "prototype_replace"

    profile_with_body = extract_template_profile(make_template(with_body=True))
    body_slots = [shape.shape_id for shape in profile_with_body.slides[0].shapes if shape.role == "body"]
    assert body_slots, (
        "template profile exposes no body slot for the presenter/date metadata; "
        f"roles={[(s.shape_id, s.role) for s in profile_with_body.slides[0].shapes]}"
    )
    plan = build_template_plan(deck, profile_with_body)
    assert any(replacement.source_path == "presenter" for replacement in plan.slides[0].replacements), (
        "presenter metadata was not mapped to a template slot; "
        f"strategy={plan.slides[0].strategy}, "
        f"replacements={[(r.source_path, r.role) for r in plan.slides[0].replacements]}, "
        f"warnings={[w.message for w in plan.slides[0].warnings]}"
    )


def test_template_renderer_replaces_cover_presenter_and_date_slots(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.template.renderer import render_deck_ir_with_template

    template = tmp_path / "cover-metadata.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(slide, "TEMPLATE TITLE", 1.0, 0.5, 8.0, 1.0, 32)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Impact"
    _textbox(slide, "TEMPLATE SUBTITLE", 1.0, 1.7, 8.0, 0.7, 18)
    _textbox(slide, "GENERIC BODY", 1.0, 2.4, 4.0, 0.6, 14)
    _textbox(slide, "汇报人 / TEMPLATE PRESENTER", 1.0, 3.0, 4.0, 0.6, 14)
    presenter_shape_id = slide.shapes[-1].shape_id
    _textbox(slide, "时间 / TEMPLATE DATE", 1.0, 4.0, 4.0, 0.6, 14)
    date_shape_id = slide.shapes[-1].shape_id
    presentation.save(template)
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "2.2",
            "meta": {"title": "T", "classification": "PUBLIC"},
            "slides": [
                {
                    "layout": "cover",
                    "title": "车载通信基带技术报告",
                    "subtitle": "模板驱动生成",
                    "presenter": "张三",
                    "date": "2026-08-25",
                }
            ],
        }
    )

    output = tmp_path / "cover-metadata-output.pptx"
    result = render_deck_ir_with_template(deck, template, output)
    replacement_by_source = {
        replacement.source_path: replacement.shape_id for replacement in result.plan.slides[0].replacements
    }
    rendered_text = "\n".join(
        shape.text
        for shape in Presentation(output).slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "张三" in rendered_text
    assert "2026-08-25" in rendered_text
    assert replacement_by_source["presenter"] == presenter_shape_id
    assert replacement_by_source["date"] == date_shape_id
    report = check_pptx(output, classification="PUBLIC", template_profile=result.profile)
    assert "HW-E02" not in {item.code for item in report.items}


def test_profile_survives_placeholder_shape_without_explicit_geometry(tmp_path: Path) -> None:
    from pptx.oxml.ns import qn
    from app.template.profile import extract_template_profile

    path = tmp_path / "no-geometry.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "TITLE"
    sp = textbox._element
    sp_pr = sp.find(qn("p:spPr"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is not None:
        sp_pr.remove(xfrm)
    presentation.save(path)

    profile = extract_template_profile(path)
    assert profile.slides[0].shapes


def test_text_fit_does_not_crash_on_shape_without_geometry(tmp_path: Path) -> None:
    from pptx.oxml.ns import qn
    from app.template.text_fit import replace_text_preserving_style

    path = tmp_path / "no-geometry.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "TITLE"
    sp = textbox._element
    sp_pr = sp.find(qn("p:spPr"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is not None:
        sp_pr.remove(xfrm)
    presentation.save(path)

    presentation = Presentation(path)
    shape = presentation.slides[0].shapes[0]
    assert replace_text_preserving_style(shape, "Longer replacement title", role="title")


def test_text_fit_does_not_double_bullet_when_template_owns_bullet_char(tmp_path: Path) -> None:
    """A template paragraph with a:buChar must not render '• • item'."""
    from lxml import etree as _etree
    from pptx.oxml.ns import qn as _qn
    from app.template.text_fit import replace_text_preserving_style

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    textbox.text = "模板正文"
    paragraph = textbox.text_frame.paragraphs[0]
    p_pr = paragraph._p.get_or_add_pPr()
    _etree.SubElement(p_pr, _qn("a:buFont")).set("typeface", "Arial")
    _etree.SubElement(p_pr, _qn("a:buChar")).set("char", "•")
    paragraph.text = "模板正文"

    assert replace_text_preserving_style(textbox, "• 第一条\n• 第二条", role="body", fallback_font_name="Arial")

    texts = [paragraph.text for paragraph in textbox.text_frame.paragraphs]
    assert texts == ["第一条", "第二条"], f"double bullet not stripped: {texts}"
