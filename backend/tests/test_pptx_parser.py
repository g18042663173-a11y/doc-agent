from __future__ import annotations

import shutil
import sys
import zipfile
import base64
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_parse_pptx_extracts_titles_bodies_tables_and_notes(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    path = tmp_path / "项目汇报.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "项目汇报"
    slide.placeholders[1].text = "主链路已经打通"
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "状态"
    table.cell(1, 0).text = "Word 输出"
    table.cell(1, 1).text = "完成"
    slide.notes_slide.notes_text_frame.text = "演讲备注"
    prs.save(path)

    ir = parse_pptx(path)

    slide_summary = ir.content.slides[0]
    assert ir.source.format == "pptx"
    assert slide_summary.index == 1
    assert slide_summary.title == "项目汇报"
    assert "主链路已经打通" in slide_summary.bodies
    assert slide_summary.tables[0]["header"] == ["事项", "状态"]
    assert slide_summary.tables[0]["rows"] == [["Word 输出", "完成"]]
    assert slide_summary.notes == "演讲备注"


def test_parse_pptx_records_transition_warning(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    path = tmp_path / "transition.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "带切换"
    prs.save(path)
    _inject_transition(path)

    ir = parse_pptx(path)

    assert any("transition" in warning for warning in ir.content.slides[0].shape_warnings)


def test_parse_pptx_records_image_presence_and_dimensions(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))
    path = tmp_path / "image.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(2), Inches(1))
    prs.save(path)

    ir = parse_pptx(path)

    assert ir.stats.images == 1
    assert any("image present width=2.00in height=1.00in" in warning for warning in ir.content.slides[0].shape_warnings)


def test_parse_pptx_records_chart_title_and_type(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    path = tmp_path / "chart.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(4), Inches(0.5)).text_frame.text = "经营汇报"
    data = CategoryChartData()
    data.categories = ["一月", "二月"]
    data.add_series("收入", (120, 150))
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(3), data)
    chart = chart_shape.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "月度收入趋势"
    prs.save(path)

    ir = parse_pptx(path)

    assert any("[chart]" in body and "月度收入趋势" in body and "COLUMN_CLUSTERED" in body for body in ir.content.slides[0].bodies)


def test_parse_pptx_truncates_long_body_table_cell_and_notes(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    path = tmp_path / "long-text.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "文" * 2100
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "正常"
    table.cell(1, 1).text = "格" * 2101
    slide.notes_slide.notes_text_frame.text = "注" * 2102
    prs.save(path)

    ir = parse_pptx(path)
    summary = ir.content.slides[0]

    assert len(summary.bodies[0]) == 2000
    assert len(summary.tables[0]["rows"][0][1]) == 2000
    assert len(summary.notes or "") == 2000
    assert any("shape 1 paragraph 1" in warning for warning in ir.warnings)
    assert any("table 2 row 1 column 2" in warning for warning in ir.warnings)
    assert any("slide 1 notes" in warning for warning in ir.warnings)


def test_parse_pptx_records_embedded_ole_part_at_slide(tmp_path: Path) -> None:
    from app.parsers.pptx_parser import parse_pptx

    path = tmp_path / "embedded.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "带嵌入对象"
    prs.save(path)
    _inject_embedding(path)

    ir = parse_pptx(path)

    assert any(
        "slide 1" in warning
        and "ppt/embeddings/oleObject1.bin" in warning
        and "bytes=8" in warning
        for warning in ir.warnings
    )
    assert any("ppt/embeddings/oleObject1.bin" in warning for warning in ir.content.slides[0].shape_warnings)


def test_slide_bodies_recurses_group_shapes() -> None:
    from app.parsers.pptx_parser import _slide_bodies

    grouped_text = _FakeShape(text="组合内正文")
    group = _FakeShape(children=[grouped_text], shape_type="GROUP")
    slide = _FakeSlide([group])

    assert _slide_bodies(slide) == ["组合内正文"]


def test_iter_shapes_warns_and_skips_broken_group_children() -> None:
    from app.parsers.pptx_parser import _iter_shapes

    warnings: list[str] = []
    shapes = list(_iter_shapes([_BrokenGroup()], warnings))

    assert len(shapes) == 1
    assert warnings == ["group shape traversal failed; child shapes skipped"]


def test_iter_shapes_warns_when_group_iterator_raises() -> None:
    from app.parsers.pptx_parser import _iter_shapes

    warnings: list[str] = []
    shapes = list(_iter_shapes(_BrokenShapeCollection(), warnings))

    assert shapes == []
    assert warnings == ["group shape traversal failed; child shapes skipped"]


class _FakeSlide:
    def __init__(self, shapes):
        self.shapes = shapes


class _FakeShape:
    def __init__(self, *, text: str = "", children=None, shape_type: str = "TEXT") -> None:
        self.text = text
        self.shapes = children or []
        self.shape_type = shape_type
        self.has_text_frame = bool(text)
        self.has_table = False
        self.has_chart = False


class _BrokenGroup:
    shape_type = "GROUP"

    @property
    def shapes(self):
        raise ValueError("broken group children")


class _BrokenShapeCollection:
    def __iter__(self):
        raise ValueError("broken shape iterator")


def _inject_transition(path: Path) -> None:
    temp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                xml = data.decode("utf-8")
                xml = xml.replace("</p:sld>", "<p:transition/></p:sld>")
                data = xml.encode("utf-8")
            target.writestr(item, data)
    shutil.move(temp, path)


def _inject_embedding(path: Path) -> None:
    temp = path.with_suffix(".tmp")
    relationship = (
        '<Relationship Id="rId999" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" '
        'Target="../embeddings/oleObject1.bin"/>'
    )
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                xml = data.decode("utf-8")
                data = xml.replace("</Relationships>", f"{relationship}</Relationships>").encode("utf-8")
            target.writestr(info, data)
        target.writestr("ppt/embeddings/oleObject1.bin", b"fake-ole")
    shutil.move(temp, path)
