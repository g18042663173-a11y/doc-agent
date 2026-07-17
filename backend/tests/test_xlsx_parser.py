from __future__ import annotations

import sys
import time
import zipfile
import base64
from datetime import date
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
import pytest

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_parse_xlsx_summarizes_sheets_formulas_and_merged_cells(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "销售台账.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["区域", "销售额", "备注"])
    sheet.append(["华东", 100, "稳定"])
    sheet.append(["华南", 120, "增长"])
    sheet["D2"] = "=B2*2"
    sheet.merge_cells("A5:B5")
    sheet["A5"] = "合并说明"
    workbook.create_sheet("空表")
    workbook.save(path)

    ir = parse_xlsx(path)

    sales = ir.content.sheets[0]
    assert ir.source.format == "xlsx"
    assert sales.name == "销售"
    assert sales.nrows == 5
    assert sales.ncols == 4
    assert sales.header_guess == ["区域", "销售额", "备注", ""]
    assert sales.preview_rows[1][:3] == ["华东", "100", "稳定"]
    assert sales.formula_count == 1
    assert sales.merged_count == 1
    assert sales.col_stats[1].type_guess == "number"
    assert sales.col_stats[1].non_empty_ratio > 0


def test_parse_xlsx_truncates_preview_to_20_by_15(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(25):
        sheet.append([f"R{row}C{col}" for col in range(18)])
    workbook.save(path)

    ir = parse_xlsx(path)
    sheet_summary = ir.content.sheets[0]

    assert len(sheet_summary.preview_rows) == 20
    assert len(sheet_summary.preview_rows[0]) == 15
    assert sheet_summary.truncated is True
    assert any("xlsx preview truncated" in warning for warning in ir.warnings)


def test_parse_xlsx_skips_hidden_rows_and_columns_with_warning(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "hidden-ledger.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["公开字段", "隐藏字段", "状态"])
    sheet.append(["站点A", "敏感标识A", "正常"])
    sheet.append(["隐藏行内容", "敏感标识B", "不应出现"])
    sheet.append(["站点C", "敏感标识C", "正常"])
    sheet.column_dimensions["B"].hidden = True
    sheet.row_dimensions[3].hidden = True
    workbook.save(path)

    ir = parse_xlsx(path)
    summary = ir.content.sheets[0]
    preview_text = "\n".join(" | ".join(row) for row in summary.preview_rows)

    assert summary.header_guess == ["公开字段", "状态"]
    assert "敏感标识" not in preview_text
    assert "隐藏行内容" not in preview_text
    assert any("hidden rows/columns" in warning and "skipped" in warning for warning in ir.warnings)


def test_parse_xlsx_keeps_visible_rows_and_columns_without_hidden_warning(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "visible-ledger.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["站点", "状态"])
    sheet.append(["站点A", "正常"])
    workbook.save(path)

    ir = parse_xlsx(path)

    assert ir.content.sheets[0].preview_rows == [["站点", "状态"], ["站点A", "正常"]]
    assert not any("hidden rows/columns" in warning for warning in ir.warnings)


def test_parse_xlsx_samples_column_stats_for_large_sheets(tmp_path: Path, monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser

    path = tmp_path / "sampled-stats.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["序号", "状态"])
    for index in range(1, 8):
        sheet.append([index, "正常"])
    workbook.save(path)
    monkeypatch.setattr(xlsx_parser, "MAX_COLUMN_STATS_ROWS", 3)
    monkeypatch.setattr(xlsx_parser, "MAX_FORMULA_SCAN_ROWS", 3)

    ir = xlsx_parser.parse_xlsx(path)

    assert any("column statistics sampled from first 3 rows" in warning for warning in ir.warnings)
    assert ir.content.sheets[0].col_stats[0].non_empty_ratio == 1.0


def test_parse_xlsx_warns_for_chart_pivot_and_vba_package_parts(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "features.xlsx"
    workbook = Workbook()
    workbook.active.append(["A", "B"])
    workbook.save(path)
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("xl/charts/chart1.xml", "<c:chartSpace/>")
        package.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTableDefinition/>")
        package.writestr("xl/vbaProject.bin", b"fake-vba")
        package.writestr("xl/media/image1.png", b"fake-image")
        package.writestr("xl/embeddings/oleObject1.bin", b"fake-ole")

    ir = parse_xlsx(path)

    assert any("chart unsupported" in warning for warning in ir.warnings)
    assert any("pivot table unsupported" in warning for warning in ir.warnings)
    assert any("VBA unsupported" in warning for warning in ir.warnings)
    assert any("embedded image" in warning for warning in ir.warnings)
    assert any("OLE object" in warning for warning in ir.warnings)
    assert any("xl/embeddings/oleObject1.bin" in warning for warning in ir.warnings)


def test_parse_xlsx_records_embedded_image_geometry_anchor_and_size(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    image_path = tmp_path / "pixel.png"
    image_path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
        )
    )
    path = tmp_path / "image.xlsx"
    workbook = Workbook()
    workbook.active.title = "周报"
    image = XlsxImage(str(image_path))
    image.width = 192
    image.height = 96
    workbook.active.add_image(image, "B3")
    workbook.save(path)

    ir = parse_xlsx(path)

    assert ir.stats.images == 1
    assert any(
        "sheet=周报" in warning
        and "anchor=B3" in warning
        and "width=2.00in" in warning
        and "height=1.00in" in warning
        and "xl/media/" in warning
        and "binary not extracted" in warning
        for warning in ir.warnings
    )


def test_parse_xlsx_warns_for_error_boolean_and_date_cells(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "special-cells.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["错误", "布尔", "日期"])
    sheet.append(["#DIV/0!", True, date(2026, 7, 10)])
    workbook.save(path)

    ir = parse_xlsx(path)

    assert any("error cells" in warning for warning in ir.warnings)
    assert any("boolean cells" in warning for warning in ir.warnings)
    assert any("date cells" in warning for warning in ir.warnings)


def test_parse_xlsx_truncates_long_cells_with_w103(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "long-cell.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["超长说明", "正常说明"])
    sheet.append(["长" * 2100, "短内容"])
    workbook.save(path)

    ir = parse_xlsx(path)

    assert len(ir.content.sheets[0].preview_rows[1][0]) == 2000
    assert ir.content.sheets[0].preview_rows[1][1] == "短内容"
    assert any("W103" in warning and "row 2 column 1" in warning for warning in ir.warnings)
    assert not any("row 2 column 2" in warning for warning in ir.warnings)


def test_parse_xlsx_rejects_suspicious_compression_ratio(tmp_path: Path) -> None:
    from app.parsers.errors import ParseFailure
    from app.parsers.xlsx_parser import parse_xlsx

    path = tmp_path / "compressed-bomb.xlsx"
    workbook = Workbook()
    workbook.active.append(["字段", "值"])
    workbook.save(path)
    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr("xl/embeddings/suspicious.bin", b"x" * (6 * 1024 * 1024))

    with pytest.raises(ParseFailure) as exc_info:
        parse_xlsx(path)

    assert exc_info.value.code == "E001"
    assert exc_info.value.loc == "source.resource"
    assert "compression ratio" in exc_info.value.message


def test_parse_xlsx_skips_optional_scans_after_time_budget(tmp_path: Path, monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser

    path = tmp_path / "budget.xlsx"
    workbook = Workbook()
    workbook.active.append(["字段", "值"])
    workbook.active.append(["事项", 1])
    workbook.save(path)
    monkeypatch.setattr(xlsx_parser, "MAX_PARSE_SECONDS", -1.0)

    ir = xlsx_parser.parse_xlsx(path)

    assert ir.content.sheets[0].preview_rows[1] == ["事项", "1"]
    assert ir.content.sheets[0].col_stats == []
    assert any("resource budget reached" in warning and "preview processed" in warning for warning in ir.warnings)


def test_parse_xlsx_large_file_hard_deadline_terminates_worker_with_processed_range(
    tmp_path: Path, monkeypatch
) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.parsers.errors import ParseFailure

    path = tmp_path / "deadline.xlsx"
    workbook = Workbook()
    workbook.active.append(["字段", "值"])
    workbook.active.append(["事项", 1])
    workbook.save(path)
    context = _FakeHardDeadlineContext(worker_result=None, poll_ready=False)
    monkeypatch.setattr(xlsx_parser.multiprocessing, "get_context", lambda _method: context)
    monkeypatch.setattr(xlsx_parser, "HARD_GUARD_FILE_BYTES", 0)
    monkeypatch.setattr(xlsx_parser, "MAX_PARSE_SECONDS", 0.001)

    started = time.perf_counter()
    with pytest.raises(ParseFailure) as exc_info:
        xlsx_parser.parse_xlsx(path)
    elapsed = time.perf_counter() - started

    assert elapsed < 1.0
    assert exc_info.value.code == "E001"
    assert exc_info.value.loc == "source.resource"
    assert "hard resource deadline" in exc_info.value.message
    assert "no worksheet content was committed" in exc_info.value.message
    assert context.process.terminated is True


def test_parse_xlsx_real_worksheet_over_10mb_under_5_seconds(tmp_path: Path) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.parsers.errors import ParseFailure

    path = tmp_path / "ten-mb.xlsx"
    _write_large_real_worksheet(path, row_count=30_000)

    started = time.perf_counter()
    try:
        ir = xlsx_parser.parse_xlsx(path)
    except ParseFailure as exc:
        assert exc.code == "E001"
        assert exc.loc == "source.resource"
        assert "hard resource deadline" in exc.message
        assert "no worksheet content was committed" in exc.message
    else:
        assert ir.content.sheets[0].nrows == 30_000
        assert ir.content.sheets[0].preview_rows[1][0].startswith("周报台账记录00002")
        assert any("column statistics sampled" in warning for warning in ir.warnings)
        assert ir.source.format == "xlsx"
    elapsed = time.perf_counter() - started

    assert path.stat().st_size >= 10 * 1024 * 1024
    assert elapsed < xlsx_parser.MAX_PARSE_SECONDS + 1.0


@pytest.mark.parametrize(
    ("worker_result", "expected_message"),
    [
        (
            (
                "parse_failure",
                {
                    "code": "E001",
                    "loc": "source.resource",
                    "message": "worker rejected workbook",
                    "suggestion": "重新导出后重试。",
                },
            ),
            "worker rejected workbook",
        ),
        (("error", "MemoryError: isolated failure"), "isolated parser failed"),
    ],
)
def test_xlsx_hard_deadline_propagates_worker_failures(
    tmp_path: Path, monkeypatch, worker_result: tuple[str, object], expected_message: str
) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.parsers.errors import ParseFailure

    path = tmp_path / "worker-failure.xlsx"
    path.write_bytes(b"placeholder")
    context = _FakeHardDeadlineContext(worker_result=worker_result)
    monkeypatch.setattr(xlsx_parser.multiprocessing, "get_context", lambda _method: context)

    with pytest.raises(ParseFailure) as exc_info:
        xlsx_parser._parse_xlsx_with_hard_deadline(path)

    assert exc_info.value.code == "E001"
    assert expected_message in exc_info.value.message
    assert context.process.terminated is True


def test_xlsx_hard_deadline_maps_worker_eof_to_structured_error(tmp_path: Path, monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.parsers.errors import ParseFailure

    path = tmp_path / "worker-eof.xlsx"
    path.write_bytes(b"placeholder")
    context = _FakeHardDeadlineContext(worker_result=None, raises_eof=True)
    monkeypatch.setattr(xlsx_parser.multiprocessing, "get_context", lambda _method: context)

    with pytest.raises(ParseFailure) as exc_info:
        xlsx_parser._parse_xlsx_with_hard_deadline(path)

    assert exc_info.value.code == "E001"
    assert exc_info.value.loc == "source.resource"
    assert "without a result" in exc_info.value.message
    assert context.process.terminated is True


def test_xlsx_isolated_worker_serializes_success_and_closes_connection(monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.ir.document_ir import DocumentIR

    result = DocumentIR.model_validate(
        {
            "ir_type": "document",
            "ir_version": "1.1",
            "source": {"filename": "worker.xlsx", "format": "xlsx", "size_kb": 1, "parsed_at": "now"},
            "stats": {},
            "content": {"sheets": []},
        }
    )
    connection = _CapturingSend()
    monkeypatch.setattr(xlsx_parser, "_parse_xlsx_impl", lambda _path: result)

    xlsx_parser._xlsx_worker("worker.xlsx", connection)

    assert connection.messages[0][0] == "ok"
    assert connection.messages[0][1]["ir_version"] == "1.1"
    assert connection.closed is True


def test_xlsx_isolated_worker_preserves_structured_parse_failure(monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser
    from app.parsers.errors import ParseFailure

    def reject(_path: Path):
        raise ParseFailure(code="E001", loc="source.resource", message="资源超限", suggestion="拆分文件。")

    connection = _CapturingSend()
    monkeypatch.setattr(xlsx_parser, "_parse_xlsx_impl", reject)

    xlsx_parser._xlsx_worker("worker.xlsx", connection)

    assert connection.messages == [
        (
            "parse_failure",
            {"code": "E001", "loc": "source.resource", "message": "资源超限", "suggestion": "拆分文件。"},
        )
    ]
    assert connection.closed is True


def test_xlsx_isolated_worker_maps_unknown_exception_without_traceback(monkeypatch) -> None:
    import app.parsers.xlsx_parser as xlsx_parser

    def crash(_path: Path):
        raise MemoryError("内存不足")

    connection = _CapturingSend()
    monkeypatch.setattr(xlsx_parser, "_parse_xlsx_impl", crash)

    xlsx_parser._xlsx_worker("worker.xlsx", connection)

    assert connection.messages == [("error", "MemoryError: 内存不足")]
    assert connection.closed is True


def test_xlsx_drawing_metadata_helpers_have_traceable_fallbacks(tmp_path: Path) -> None:
    from app.parsers.xlsx_parser import (
        _drawing_anchor_label,
        _drawing_dimensions,
        _relationship_targets,
        _rels_part,
        _resolve_part,
    )

    absolute = ElementTree.fromstring(
        '<xdr:absoluteAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
        '<xdr:ext cx="914400" cy="457200"/></xdr:absoluteAnchor>'
    )
    incomplete = ElementTree.fromstring(
        '<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
        '<xdr:from><xdr:col>1</xdr:col></xdr:from></xdr:twoCellAnchor>'
    )

    assert _drawing_anchor_label(absolute) == "absolute"
    assert _drawing_dimensions(absolute) == ("1.00in", "0.50in")
    assert _drawing_anchor_label(incomplete) == "unknown"
    assert _drawing_dimensions(incomplete) == ("unknown", "unknown")
    assert _resolve_part("xl/drawings/drawing1.xml", "/xl/media/image1.png") == "xl/media/image1.png"
    assert _rels_part("xl/drawings/drawing1.xml") == "xl/drawings/_rels/drawing1.xml.rels"

    package_path = tmp_path / "empty.zip"
    with zipfile.ZipFile(package_path, "w"):
        pass
    with zipfile.ZipFile(package_path) as package:
        assert _relationship_targets(package, "missing.rels", "xl/workbook.xml", set()) == {}


def _write_large_real_worksheet(path: Path, *, row_count: int) -> None:
    workbook = Workbook()
    workbook.active.append(["事项", "状态", "责任组", "说明"])
    workbook.save(path)

    rows = [
        '<row r="1"><c r="A1" t="inlineStr"><is><t>事项</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>状态</t></is></c>'
        '<c r="C1" t="inlineStr"><is><t>责任组</t></is></c>'
        '<c r="D1" t="inlineStr"><is><t>说明</t></is></c></row>'
    ]
    filler = "国内办公场景结构边界验证" * 8
    for row in range(2, row_count + 1):
        values = [f"周报台账记录{row:05d}-{filler}", "进行中", f"小组{row % 9}", f"可复现说明{row:05d}-{filler}"]
        cells = "".join(
            f'<c r="{column}{row}" t="inlineStr"><is><t>{value}</t></is></c>'
            for column, value in zip("ABCD", values, strict=True)
        )
        rows.append(f'<row r="{row}">{cells}</row>')
    worksheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:D{row_count}"/><sheetData>{"".join(rows)}</sheetData></worksheet>'
    ).encode("utf-8")

    original = path.with_suffix(".source.xlsx")
    path.replace(original)
    with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist():
            data = worksheet if info.filename == "xl/worksheets/sheet1.xml" else source.read(info.filename)
            target.writestr(info.filename, data)
    original.unlink()


class _FakeConnection:
    def __init__(self, result=None, *, raises_eof: bool = False, poll_ready: bool = True) -> None:
        self.result = result
        self.raises_eof = raises_eof
        self.poll_ready = poll_ready

    def poll(self, _timeout: float) -> bool:
        return self.poll_ready

    def recv(self):
        if self.raises_eof:
            raise EOFError
        return self.result

    def close(self) -> None:
        pass


class _CapturingSend:
    def __init__(self) -> None:
        self.messages: list[tuple[str, object]] = []
        self.closed = False

    def send(self, value: tuple[str, object]) -> None:
        self.messages.append(value)

    def close(self) -> None:
        self.closed = True


class _FakeProcess:
    def __init__(self) -> None:
        self.terminated = False

    def start(self) -> None:
        pass

    def join(self, timeout: float | None = None) -> None:
        pass

    def is_alive(self) -> bool:
        return not self.terminated

    def terminate(self) -> None:
        self.terminated = True


class _FakeHardDeadlineContext:
    def __init__(self, *, worker_result, raises_eof: bool = False, poll_ready: bool = True) -> None:
        self.receiver = _FakeConnection(worker_result, raises_eof=raises_eof, poll_ready=poll_ready)
        self.sender = _FakeConnection()
        self.process = _FakeProcess()

    def Pipe(self, *, duplex: bool):
        assert duplex is False
        return self.receiver, self.sender

    def Process(self, *, target, args, daemon: bool):
        assert callable(target)
        assert args
        assert daemon is True
        return self.process
