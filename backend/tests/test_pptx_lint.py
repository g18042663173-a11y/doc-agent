from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def test_check_pptx_rendered_deck_has_no_errors(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "合规", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "合规"}],
        }
    )
    path = render_deck_ir(deck, tmp_path / "ok.pptx")

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert report.summary["errors"] == 0
    assert report.summary["pass"] is True


def test_check_pptx_accepts_snapped_textbox(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "snapped.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "吸附网格", 0.6, 120 / 72, _grid_width(4), 24 / 72)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W06" not in codes


def test_check_pptx_reports_unsnapped_textbox(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "unsnapped.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "未吸附网格", 0.73, 1.63, 3.7, 0.31)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W06" in codes


def test_check_pptx_accepts_theme_key_frame_coordinates(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "关键框", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "title_bullets", "title": "关键框", "bullets": [{"text": "坐标合规"}]}],
        }
    )
    path = render_deck_ir(deck, tmp_path / "key-frame-ok.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W08" not in codes


def test_check_pptx_reports_key_frame_coordinate_drift(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "key-frame-drift.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide, "偏移标题", 1.05, 0.85, 10.2, 0.55, size=28)
    _add_required_footer(slide, top=6.55)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W08" in codes


def test_check_pptx_accepts_sufficient_text_contrast(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "contrast-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "对比度合格", 0.6, 120 / 72, _grid_width(4), 24 / 72, color="333333", fill="FFFFFF")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W09" not in codes


def test_check_pptx_reports_low_text_contrast(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "contrast-low.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "对比度过低", 0.6, 120 / 72, _grid_width(4), 24 / 72, color="333333", fill="1F1F1F")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W09" in codes


def test_check_pptx_reports_missing_footer_and_bad_font(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "bad.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Bad font"
    run.font.name = "Comic Sans MS"
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    codes = [item.code for item in report.items]

    assert "HW-E01" in codes
    assert "HW-E02" in codes


def test_check_pptx_requires_classification_in_footer_area(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "classification-not-footer.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide, "HUAWEI CONFIDENTIAL", 2.0, 2.0, 4.0, 0.4)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-E01" in codes


def test_check_pptx_does_not_treat_footer_9pt_as_body_font_violation(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "footer-size-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "正文", 0.6, 120 / 72, _grid_width(4), 24 / 72, size=16)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W01" not in codes


def test_check_pptx_reports_small_body_font_and_non_theme_color(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "font-color-warning.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "小字号非主题色", 0.6, 120 / 72, _grid_width(4), 24 / 72, size=5.5, color="00FF00")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W01" in codes
    assert "HW-W02" in codes


def test_check_pptx_accepts_three_font_size_kinds_per_slide(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "font-size-kinds-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "主标题", 0.6, 120 / 72, _grid_width(3), 24 / 72, size=28)
    _add_textbox(slide, "一级标题", 4.6, 120 / 72, _grid_width(3), 24 / 72, size=18)
    _add_textbox(slide, "正文", 8.6, 120 / 72, _grid_width(3), 24 / 72, size=16)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W01" not in codes


def test_check_pptx_uses_body_minimum_without_rejecting_compact_components(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    body_path = tmp_path / "small-body.pptx"
    body_prs = _blank_presentation()
    body_slide = body_prs.slides.add_slide(body_prs.slide_layouts[6])
    _add_required_footer(body_slide)
    _add_textbox(body_slide, "普通正文", 0.6, 120 / 72, _grid_width(4), 24 / 72, size=10)
    body_prs.save(body_path)

    compact_path = tmp_path / "compact-components.pptx"
    compact_prs = _blank_presentation()
    compact_slide = compact_prs.slides.add_slide(compact_prs.slide_layouts[6])
    _add_required_footer(compact_slide)
    compact = _add_textbox(compact_slide, "阈值标签", 0.6, 120 / 72, _grid_width(3), 24 / 72, size=8)
    compact.name = "HW_THRESHOLD_LABEL:1"
    table_shape = compact_slide.shapes.add_table(1, 1, Inches(4.6), Inches(120 / 72), Inches(3.0), Inches(0.5))
    table_shape.table.cell(0, 0).text = "表格数据"
    table_run = table_shape.table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    table_run.font.name = "Arial"
    table_run.font.size = Pt(10)
    compact_prs.save(compact_path)

    body_report = check_pptx(body_path, classification="HUAWEI CONFIDENTIAL")
    compact_report = check_pptx(compact_path, classification="HUAWEI CONFIDENTIAL")

    assert any(item.code == "HW-W01" and "10.0pt" in item.message for item in body_report.items)
    assert not any(item.code == "HW-W01" for item in compact_report.items)


def test_check_pptx_excludes_compact_annotations_from_font_size_variety(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "compact-font-variety.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "页标题", 0.6, 120 / 72, _grid_width(3), 32 / 72, size=28)
    _add_textbox(slide, "正文", 4.6, 120 / 72, _grid_width(3), 32 / 72, size=16)
    node = _add_textbox(slide, "节点", 8.6, 120 / 72, _grid_width(2), 32 / 72, size=12)
    node.name = "HW_ARCH_NODE:node"
    edge = _add_textbox(slide, "边标签", 10.6, 120 / 72, _grid_width(1), 32 / 72, size=8)
    edge.name = "HW_ARCH_EDGE_LABEL:1"
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert not any(item.code == "HW-W01" and "字号种类" in item.message for item in report.items)


def test_check_pptx_reports_more_than_three_font_size_kinds_per_slide(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "font-size-kinds-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "主标题", 0.6, 120 / 72, _grid_width(2), 24 / 72, size=14)
    _add_textbox(slide, "一级标题", 3.6, 120 / 72, _grid_width(2), 24 / 72, size=11)
    _add_textbox(slide, "正文", 6.6, 120 / 72, _grid_width(2), 24 / 72, size=10)
    _add_textbox(slide, "脚注", 9.6, 120 / 72, _grid_width(2), 24 / 72, size=9)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W01" in codes


def test_check_pptx_accepts_bullet_count_within_limit(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "bullets-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    for index in range(7):
        _add_textbox(slide, f"• 要点 {index}", 0.6, (120 + index * 36) / 72, _grid_width(4), 24 / 72)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W03" not in codes


def test_check_pptx_reports_too_many_or_too_long_bullets(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "bullets-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    for index in range(8):
        _add_textbox(slide, f"• 要点 {index}", 0.6, (120 + index * 34) / 72, _grid_width(4), 24 / 72)
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    warnings = [item for item in report.items if item.code == "HW-W03"]

    assert any(
        item.message == "要点超限·本页共 8 条要点,超过上限 7 条,建议删减或拆分。"
        for item in warnings
    )


def test_check_pptx_hw_w03_identifies_single_bullet_length_reason(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "bullet-too-long.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, f"• {'长' * 61}", 0.6, 120 / 72, _grid_width(8), 48 / 72)
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    warnings = [item for item in report.items if item.code == "HW-W03"]

    assert any(
        item.message == "要点超限·单条最长 61 字,超过上限 60 字,建议精简表述。"
        for item in warnings
    )


def test_check_pptx_accepts_non_overlapping_boxes_with_sufficient_gap(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "spacing-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "左侧", 0.6, 120 / 72, _grid_width(3), 24 / 72)
    _add_textbox(slide, "右侧", 4.6, 120 / 72, _grid_width(3), 24 / 72)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W07" not in codes


def test_check_pptx_reports_overlapping_or_too_close_boxes(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "spacing-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    _add_textbox(slide, "框一", 0.6, 120 / 72, _grid_width(3), 24 / 72)
    _add_textbox(slide, "框二", 0.8, 122 / 72, _grid_width(3), 24 / 72)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W07" in codes


def test_check_pptx_reports_transition_and_large_table_and_slide_count(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "violations.pptx"
    prs = Presentation()
    for index in range(31):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(4), Inches(0.3)).text = "HUAWEI CONFIDENTIAL"
        if index == 0:
            slide.shapes.add_table(13, 9, Inches(1), Inches(1), Inches(8), Inches(4))
    prs.save(path)
    _inject_transition(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    codes = [item.code for item in report.items]

    assert "HW-E03" in codes
    assert "HW-W04" in codes
    assert "HW-W05" in codes


def test_check_pptx_accepts_table_emphasis_accent_fills(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "table-emphasis-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    table = slide.shapes.add_table(2, 3, Inches(1.0), Inches(1.8), Inches(6.0), Inches(1.0)).table
    table.cell(0, 0).text = "方案"
    table.cell(0, 1).text = "结论"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).text = "方案A"
    table.cell(1, 1).text = "推荐"
    table.cell(1, 2).text = "备选"
    table.cell(1, 1).fill.solid()
    table.cell(1, 1).fill.fore_color.rgb = _rgb("FCC800")
    table.cell(1, 2).fill.solid()
    table.cell(1, 2).fill.fore_color.rgb = _rgb("30B5C5")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" not in codes


def test_check_pptx_reports_table_emphasis_fill_outside_accent_palette(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "table-emphasis-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    table = slide.shapes.add_table(2, 2, Inches(1.0), Inches(1.8), Inches(4.0), Inches(1.0)).table
    table.cell(0, 0).text = "方案"
    table.cell(0, 1).text = "结论"
    table.cell(1, 0).text = "方案A"
    table.cell(1, 1).text = "推荐"
    table.cell(1, 1).fill.solid()
    table.cell(1, 1).fill.fore_color.rgb = _rgb("00FF00")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" in codes


def test_check_pptx_accepts_chart_accent_series_and_threshold_line(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "性能图表", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "chart",
                    "title": "性能趋势",
                    "chart": {
                        "kind": "line",
                        "unit": "ms",
                        "categories": ["1月", "2月"],
                        "series": [{"name": "方案A", "values": [62, 58], "emphasis": True}],
                        "thresholds": [{"value": 60, "label": "目标"}],
                    },
                }
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "chart-ok.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" not in codes


def test_check_pptx_accepts_performance_chart_layout_without_grid_or_overlap_warnings(tmp_path: Path) -> None:
    import json

    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    payload = json.loads((ROOT / "samples" / "ir" / "deck_valid_03_performance_chart.json").read_text(encoding="utf-8"))
    deck = DeckIR.model_validate(payload)
    path = render_deck_ir(deck, tmp_path / "performance-chart.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W06" not in codes
    assert "HW-W07" not in codes


def test_check_pptx_accepts_image_placeholder_contrast(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "图片占位", "classification": "HUAWEI CONFIDENTIAL", "theme": "hw_v1"},
            "slides": [{"layout": "image", "title": "图片占位", "placeholder": "系统架构截图", "caption": "图1: 架构截图待补齐"}],
        }
    )
    path = render_deck_ir(deck, tmp_path / "image-placeholder.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W09" not in codes


def test_check_pptx_accepts_full_layout_sample_without_grid_or_spacing_warnings(tmp_path: Path) -> None:
    import json

    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    payload = json.loads((ROOT / "samples" / "ir" / "deck_valid_full.json").read_text(encoding="utf-8"))
    deck = DeckIR.model_validate(payload)
    path = render_deck_ir(deck, tmp_path / "full-layouts.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W06" not in codes
    assert "HW-W07" not in codes


def test_check_pptx_reports_threshold_line_color_outside_accent_palette(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "chart-threshold-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.2), Inches(2.0), Inches(8.0), Inches(2.0))
    line.name = "HW_THRESHOLD_LINE_BAD"
    line.line.color.rgb = _rgb("00FF00")
    line.line.width = Pt(1.5)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" in codes


def test_check_pptx_reports_threshold_line_width_drift(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "chart-threshold-width-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.2), Inches(2.0), Inches(8.0), Inches(2.0))
    line.name = "HW_THRESHOLD_LINE_WIDTH_BAD"
    line.line.color.rgb = _rgb("C7000B")
    line.line.width = Pt(0.5)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" in codes


def test_check_pptx_reports_chart_series_color_outside_accent_palette(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "chart-series-color-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    chart = _add_basic_chart(slide).chart
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = _rgb("00FF00")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" in codes


def test_check_pptx_reports_chart_label_font_below_minimum(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "chart-font-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    chart = _add_basic_chart(slide).chart
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = _rgb("C7000B")
    chart.category_axis.tick_labels.font.size = Pt(6)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W01" in codes


def test_check_pptx_accepts_table_and_chart_theme_fonts_and_sizes(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "font-surfaces-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    table = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1.5), Inches(3.2), Inches(1.0)).table
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "结论"
    table.cell(1, 0).text = "时延"
    table.cell(1, 1).text = "达标"
    for row in table.rows:
        for cell in row.cells:
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name = "微软雅黑"
                run.font.size = Pt(9)
                run.font.color.rgb = _rgb("1D1D1A")
    chart = _add_basic_chart(slide).chart
    chart.category_axis.tick_labels.font.name = "微软雅黑"
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.color.rgb = _rgb("666666")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-E02" not in codes
    assert "HW-W01" not in codes


def test_check_pptx_reports_table_and_chart_non_theme_fonts_and_small_table_text(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "font-surfaces-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    table = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1.5), Inches(3.2), Inches(1.0)).table
    table.cell(0, 0).text = "坏字体"
    table_run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    table_run.font.name = "Comic Sans MS"
    table_run.font.size = Pt(6)
    chart = _add_basic_chart(slide).chart
    chart.category_axis.tick_labels.font.name = "Courier New"
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert any(item.code == "HW-E02" and "表格" in item.message for item in report.items)
    assert any(item.code == "HW-E02" and "图表 X 轴" in item.message for item in report.items)
    assert any(item.code == "HW-W01" and "表格" in item.message for item in report.items)


def test_check_pptx_accepts_ordinary_shape_theme_fill_and_line(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "shape-color-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(2), Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb("FCC800")
    shape.line.color.rgb = _rgb("DDDDDD")
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W02" not in codes


def test_check_pptx_reports_ordinary_shape_non_theme_fill_or_line(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "shape-color-bad.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(2), Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb("123456")
    shape.line.color.rgb = _rgb("654321")
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert any(item.code == "HW-W02" and "普通形状" in item.message for item in report.items)


def test_check_pptx_counts_native_powerpoint_bullets(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    ok_path = tmp_path / "native-bullets-ok.pptx"
    bad_path = tmp_path / "native-bullets-bad.pptx"
    for path, count in ((ok_path, 7), (bad_path, 8)):
        prs = _blank_presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_required_footer(slide)
        shape = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5), Inches(4))
        frame = shape.text_frame
        frame.clear()
        for index in range(count):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = f"原生要点 {index + 1}"
            bullet = OxmlElement("a:buChar")
            bullet.set("char", "•")
            paragraph._p.get_or_add_pPr().append(bullet)
        prs.save(path)

    assert "HW-W03" not in _codes(check_pptx(ok_path, classification="HUAWEI CONFIDENTIAL"))
    assert "HW-W03" in _codes(check_pptx(bad_path, classification="HUAWEI CONFIDENTIAL"))


def test_check_pptx_reports_renderer_text_coordinate_regression(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "坐标回归", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "title_bullets", "title": "坐标回归", "bullets": [{"text": "正文要点"}]}],
        }
    )
    path = render_deck_ir(deck, tmp_path / "renderer-grid-drift.pptx")
    prs = Presentation(str(path))
    shape = next(shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_RENDERED_TEXT") and "正文要点" in shape.text)
    shape.left += Inches(0.2)
    actual_box = (shape.left, shape.top, shape.width, shape.height)
    shape.name = "HW_RENDERED_TEXT:" + ",".join(f"{value / 914400:.6f}" for value in actual_box)
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert any(item.code == "HW-W06" and "renderer" in item.message for item in report.items)


def test_check_pptx_reports_renderer_text_inside_forbidden_page_margin(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "页边距回归", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "title_bullets", "title": "页边距回归", "bullets": [{"text": "正文要点"}]}],
        }
    )
    path = render_deck_ir(deck, tmp_path / "renderer-edge-drift.pptx")
    prs = Presentation(str(path))
    shape = next(shape for shape in prs.slides[0].shapes if "正文要点" in getattr(shape, "text", ""))
    shape.left = Inches(0.05)
    shape.top = Inches(3.0)
    shape.width = Inches(5.0)
    shape.height = Inches(0.5)
    actual_box = (shape.left, shape.top, shape.width, shape.height)
    shape.name = "HW_RENDERED_TEXT:" + ",".join(f"{value / 914400:.6f}" for value in actual_box)
    prs.save(path)

    report = check_pptx(path, classification=deck.meta.classification)

    assert any(item.code == "HW-W07" and "页边" in item.message for item in report.items)


def test_check_pptx_accepts_table_chart_and_nontext_shape_geometry(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "all-shape-geometry-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    slide.shapes.add_table(2, 2, Inches(0.8), Inches(1.5), Inches(3), Inches(1))
    _add_basic_chart_at(slide, left=4.4, top=1.5, width=4.5, height=2.5)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), Inches(1.5), Inches(1.5), Inches(1))
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W07" not in codes


def test_check_pptx_ignores_full_bleed_background_and_valid_layout_containment(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / "layout-containment-ok.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        prs.slide_height,
    )
    background.name = "PAGE_BACKGROUND"
    container = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.0),
        Inches(1.5),
        Inches(4.0),
        Inches(2.0),
    )
    container.name = "HW_LAYOUT_CONTAINER:TEST"
    _add_textbox(slide, "容器内正文", 1.5, 2.0, 2.0, 0.5)
    _add_required_footer(slide)
    prs.save(path)

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-W07" not in codes


def test_check_pptx_reports_table_chart_and_nontext_shape_geometry_violations(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    builders = {
        "table": lambda slide: slide.shapes.add_table(2, 2, Inches(0.05), Inches(1.5), Inches(3), Inches(1)),
        "chart": lambda slide: (
            _add_basic_chart_at(slide, left=1.0, top=1.5, width=5, height=3),
            _add_textbox(slide, "压在图表上", 2.0, 2.0, 2.0, 0.5),
        ),
        "shape": lambda slide: slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.95), Inches(1.5), Inches(0.3), Inches(1)),
    }
    for name, builder in builders.items():
        path = tmp_path / f"geometry-{name}.pptx"
        prs = _blank_presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_required_footer(slide)
        builder(slide)
        prs.save(path)
        assert "HW-W07" in _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL")), name


def test_check_pptx_checks_table_and_chart_text_contrast(tmp_path: Path) -> None:
    from app.lint.pptx_lint import check_pptx

    ok_path = tmp_path / "table-chart-contrast-ok.pptx"
    bad_path = tmp_path / "table-chart-contrast-bad.pptx"
    for path, text_color in ((ok_path, "1D1D1A"), (bad_path, "F5F5F5")):
        prs = _blank_presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_required_footer(slide)
        table = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1.5), Inches(3), Inches(1)).table
        table.cell(0, 0).text = "对比度"
        table.cell(0, 0).fill.solid()
        table.cell(0, 0).fill.fore_color.rgb = _rgb("F5F5F5")
        table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(text_color)
        chart = _add_basic_chart_at(slide, left=4.4, top=1.5, width=4.5, height=2.5).chart
        chart.category_axis.tick_labels.font.color.rgb = _rgb("666666" if path == ok_path else "DDDDDD")
        prs.save(path)

    assert "HW-W09" not in _codes(check_pptx(ok_path, classification="HUAWEI CONFIDENTIAL"))
    assert "HW-W09" in _codes(check_pptx(bad_path, classification="HUAWEI CONFIDENTIAL"))


def test_check_pptx_marks_picture_background_contrast_as_manual_review(tmp_path: Path) -> None:
    import base64

    from app.lint.pptx_lint import check_pptx

    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))
    path = tmp_path / "picture-background.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    slide.shapes.add_picture(str(image), Inches(1), Inches(1.5), Inches(5), Inches(2))
    _add_textbox(slide, "图片上的文字", 1.5, 1.8, 3.0, 0.5, color="FFFFFF")
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")

    assert any(item.code == "HW-W09" and "无法自动判定" in item.message for item in report.items)


def test_write_reports_serializes_items(tmp_path: Path) -> None:
    from app.lint.pptx_lint import PptxLintItem, PptxLintReport, write_reports

    report = PptxLintReport(
        [
            PptxLintItem(
                code="HW-W02",
                level="Warning",
                slide=1,
                message="图表颜色不在色板。",
                suggestion="改用 accent 色。",
            )
        ]
    )

    json_path, md_path = write_reports(report, tmp_path)

    assert json_path.read_text(encoding="utf-8")
    assert "HW-W02" in md_path.read_text(encoding="utf-8")


def test_check_pptx_accepts_matching_agenda_and_section_counts(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "结构一致", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {"layout": "cover", "title": "结构一致"},
                {"layout": "agenda", "items": ["第一章", "第二章"]},
                {"layout": "section", "index": 1, "title": "第一章"},
                {"layout": "section", "index": 2, "title": "第二章"},
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "agenda-ok.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-I01" not in codes


def test_check_pptx_infos_when_agenda_and_section_counts_differ(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "结构不一致", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {"layout": "cover", "title": "结构不一致"},
                {"layout": "agenda", "items": ["第一章", "第二章", "第三章"]},
                {"layout": "section", "index": 1, "title": "第一章"},
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "agenda-bad.pptx")

    codes = _codes(check_pptx(path, classification="HUAWEI CONFIDENTIAL"))

    assert "HW-I01" in codes


def test_check_pptx_architecture_uses_dedicated_layout_checks(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "架构 lint", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "架构 lint",
                    "nodes": [
                        {"id": "a", "text": "节点A", "type": "primary", "group": "g1"},
                        {"id": "b", "text": "节点B", "type": "secondary", "group": "g2"},
                    ],
                    "edges": [{"from": "a", "to": "b", "style": "solid", "direction": "forward"}],
                    "groups": [
                        {"id": "g1", "label": "分组一", "node_ids": ["a"]},
                        {"id": "g2", "label": "分组二", "node_ids": ["b"]},
                    ],
                }
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "architecture-lint.pptx")

    codes = _codes(check_pptx(path, classification=deck.meta.classification))

    assert "HW-W06" not in codes
    assert "HW-W07" not in codes
    assert "HW-W01" not in codes
    assert "HW-W02" not in codes


@pytest.mark.parametrize("layout", ["process_flow", "timeline"])
def test_check_pptx_sequence_layouts_use_actual_geometry_and_theme_tokens(tmp_path: Path, layout: str) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    if layout == "process_flow":
        slide = {
            "layout": layout,
            "title": "流程元素按主题等尺寸排布",
            "steps": [
                {"id": "a", "title": "输入", "description": "材料"},
                {"id": "b", "title": "处理", "description": "规则"},
                {"id": "c", "title": "输出", "description": "结果"},
            ],
        }
        prefix = "HW_PROCESS_STEP:"
    else:
        slide = {
            "layout": layout,
            "title": "里程碑按主题等尺寸排布",
            "milestones": [
                {"label": "一", "title": "基线", "status": "completed"},
                {"label": "二", "title": "验证", "status": "current"},
                {"label": "三", "title": "终审", "status": "planned"},
            ],
        }
        prefix = "HW_TIMELINE_MILESTONE:"
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": layout, "classification": "公开"},
            "slides": [slide],
        }
    )
    path = render_deck_ir(deck, tmp_path / f"{layout}.pptx")

    report = check_pptx(path, classification="公开")
    assert not {"HW-W01", "HW-W02", "HW-W06", "HW-W07"} & set(_codes(report))

    prs = Presentation(str(path))
    target = next(shape for shape in prs.slides[0].shapes if shape.name.startswith(prefix))
    target.width += Inches(0.3)
    prs.save(path)

    tampered = check_pptx(path, classification="公开")
    assert any(item.code == "HW-W07" and "等宽等高" in item.message for item in tampered.items)


def test_check_pptx_warns_when_dense_process_flow_cannot_fit_minimum_text(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "密集流程", "classification": "公开"},
            "slides": [
                {
                    "layout": "process_flow",
                    "title": "密集流程仍不截断内容",
                    "steps": [
                        {"id": f"s{index}", "title": f"步骤{index}", "description": "超长说明" * 120}
                        for index in range(7)
                    ],
                }
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "dense-process.pptx")

    report = check_pptx(path, classification="公开")

    assert any(item.code == "HW-W03" and item.message.startswith("内容超版面·") for item in report.items)


@pytest.mark.parametrize(
    ("layout", "count", "warns"),
    [("process_flow", 5, False), ("process_flow", 6, True), ("timeline", 6, False), ("timeline", 7, True)],
)
def test_check_pptx_warns_for_dense_vertical_sequence_layouts(
    tmp_path: Path,
    layout: str,
    count: int,
    warns: bool,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    if layout == "process_flow":
        slide = {
            "layout": layout,
            "title": "纵向流程密度验证",
            "orientation": "vertical",
            "steps": [{"id": f"s{index}", "title": f"步骤{index}"} for index in range(count)],
        }
    else:
        slide = {
            "layout": layout,
            "title": "纵向时间线密度验证",
            "orientation": "vertical",
            "milestones": [
                {"label": f"M{index}", "title": f"里程碑{index}", "status": "planned"}
                for index in range(count)
            ],
        }
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "密度", "classification": "公开"},
            "slides": [slide],
        }
    )
    path = render_deck_ir(deck, tmp_path / f"{layout}-{count}.pptx")

    messages = [item.message for item in check_pptx(path, classification="公开").items if item.code == "HW-W03"]

    assert any("纵向内容过密" in message for message in messages) is warns


def test_check_pptx_detects_tampered_kpi_and_image_slot(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    payload = {
        "ir_type": "deck",
        "ir_version": "1.7",
        "meta": {"title": "增强 lint", "classification": "公开"},
        "slides": [
            {
                "layout": "cards",
                "variant": "kpi",
                "title": "KPI",
                "cards": [{"title": "及时率", "desc": "96%"}, {"title": "时延", "desc": "42ms"}],
            },
            {"layout": "image", "title": "图片", "placeholder": "截图"},
        ],
    }
    deck = DeckIR.model_validate(payload)
    path = render_deck_ir(deck, tmp_path / "enhanced-lint.pptx")
    clean = check_pptx(path, classification="公开")
    assert not {"HW-W01", "HW-W07"} & {item.code for item in clean.items}

    prs = Presentation(str(path))
    kpi = next(shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_RENDERED_TEXT:KPI_CARD:"))
    kpi.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    image_slot = next(shape for shape in prs.slides[1].shapes if shape.name == "HW_IMAGE_PLACEHOLDER")
    image_slot.width += Inches(0.5)
    prs.save(path)

    report = check_pptx(path, classification="公开")
    assert any(item.code == "HW-W01" and "KPI" in item.message for item in report.items)
    assert any(item.code == "HW-W07" and "16:9" in item.message for item in report.items)


@pytest.mark.parametrize(
    ("scenario", "expected_code", "message_fragment"),
    [
        ("kpi_size", "HW-W07", "KPI 指标块未保持等宽等高"),
        ("kpi_missing", "HW-W03", "KPI 指标缺少核心数值或指标名"),
        ("kpi_color", "HW-W02", "KPI 核心数值未使用主题红"),
        ("image_margin", "HW-W07", "图片占位槽进入页边距或页脚安全区"),
        ("image_instruction", "HW-W03", "图片占位槽缺少等比适配说明"),
    ],
)
def test_check_pptx_reports_specific_kpi_and_image_slot_tampering(
    tmp_path: Path,
    scenario: str,
    expected_code: str,
    message_fragment: str,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "增强 lint 负例", "classification": "公开"},
            "slides": [
                {
                    "layout": "cards",
                    "variant": "kpi",
                    "title": "KPI",
                    "cards": [{"title": "及时率", "desc": "96%"}, {"title": "时延", "desc": "42ms"}],
                },
                {"layout": "image", "title": "图片", "placeholder": "截图"},
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / f"enhanced-lint-{scenario}.pptx")
    prs = Presentation(str(path))
    kpi_shapes = [
        shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_RENDERED_TEXT:KPI_CARD:")
    ]
    image_slot = next(shape for shape in prs.slides[1].shapes if shape.name == "HW_IMAGE_PLACEHOLDER")

    if scenario == "kpi_size":
        kpi_shapes[1].width += Inches(0.3)
    elif scenario == "kpi_missing":
        kpi_shapes[0].text = "96%"
    elif scenario == "kpi_color":
        kpi_shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(97, 178, 48)
    elif scenario == "image_margin":
        image_slot.left = Inches(0.1)
    else:
        image_slot.text = "16:9 图片占位\n截图"
    prs.save(path)

    report = check_pptx(path, classification="公开")

    assert any(item.code == expected_code and message_fragment in item.message for item in report.items)


@pytest.mark.parametrize(("edge_count", "warns"), [(12, False), (13, True)])
def test_check_pptx_architecture_density_warning_boundary(
    tmp_path: Path,
    edge_count: int,
    warns: bool,
) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / f"architecture-{edge_count}-edges.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    for index in range(edge_count):
        edge = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(1.0),
            Inches(1.0 + index * 0.05),
            Inches(2.0),
            Inches(1.0 + index * 0.05),
        )
        edge.name = f"HW_ARCH_EDGE:e{index}"
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    density_warnings = [item for item in report.items if item.code == "HW-W03" and item.message.startswith("架构图过密·")]

    assert bool(density_warnings) is warns
    if warns:
        assert density_warnings[0].message == "架构图过密·边数 13 超过阈值 12,建议人工调整或拆分。"


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_composite_architecture_keeps_existing_density_lint(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    nodes = [{"id": f"n{index}", "text": f"节点{index}"} for index in range(7)]
    edge_pairs = [
        (0, 1), (0, 2), (0, 3), (1, 3), (1, 4), (2, 3), (2, 4),
        (2, 5), (3, 4), (3, 5), (3, 6), (4, 6), (5, 6),
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "组合页密度", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "组合页不豁免嵌入组件的合规检查",
                    "regions": [
                        {
                            "slot": "left",
                            "component": {
                                "layout": "title_bullets",
                                "title": "检查原则",
                                "bullets": [{"text": "容器只负责分区", "level": 1}],
                            },
                        },
                        {
                            "slot": "right",
                            "component": {
                                "layout": "architecture_diagram",
                                "title": "超密架构",
                                "nodes": nodes,
                                "edges": [
                                    {"from": f"n{source}", "to": f"n{target}"}
                                    for source, target in edge_pairs
                                ],
                                "groups": [],
                            },
                        },
                    ],
                }
            ],
        }
    )

    path = render_deck_ir(deck, tmp_path / "composite-dense-architecture.pptx")
    report = check_pptx(path, classification="公开")

    density_warnings = [
        item for item in report.items if item.code == "HW-W03" and item.message.startswith("架构图过密·")
    ]
    assert density_warnings
    assert "超过阈值 12" in density_warnings[0].message


@pytest.mark.parametrize(
    ("scenario", "expected_code", "message_fragment"),
    [
        ("table_color", "HW-W02", "表格单元格"),
        ("bullets_overflow", "HW-W03", "要点超限"),
        ("kpi_missing", "HW-W03", "KPI 指标缺少"),
    ],
)
def test_composite_embedded_components_keep_existing_product_lint(
    tmp_path: Path,
    scenario: str,
    expected_code: str,
    message_fragment: str,
) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    left_component = (
        {
            "layout": "table",
            "title": "方案表",
            "table": {"header": ["项", "值"], "rows": [["状态", "通过"]]},
        }
        if scenario == "table_color"
        else {
            "layout": "title_bullets",
            "title": "检查项",
            "bullets": [{"text": f"第{index}条", "level": 1} for index in range(1, 8)],
        }
    )
    right_component = {
        "layout": "cards",
        "title": "指标",
        "variant": "kpi" if scenario == "kpi_missing" else "default",
        "cards": [
            {"title": "通过率", "desc": "96%"},
            {"title": "版式数", "desc": "14"},
        ],
    }
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "递归 lint", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "容器不豁免嵌入组件规则",
                    "regions": [
                        {"slot": "left", "component": left_component},
                        {"slot": "right", "component": right_component},
                    ],
                }
            ],
        }
    )

    path = render_deck_ir(deck, tmp_path / f"composite-{scenario}.pptx")
    presentation = Presentation(str(path))
    slide = presentation.slides[0]
    if scenario == "table_color":
        table = next(shape.table for shape in slide.shapes if getattr(shape, "has_table", False))
        table.cell(1, 0).fill.solid()
        table.cell(1, 0).fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    elif scenario == "bullets_overflow":
        body = next(shape for shape in slide.shapes if shape.name == "HW_RENDERED_TEXT:BODY")
        paragraph = body.text_frame.add_paragraph()
        paragraph.text = "• 第八条"
    else:
        kpi = next(shape for shape in slide.shapes if shape.name.startswith("HW_RENDERED_TEXT:KPI_CARD:"))
        kpi.text = "96%"
    presentation.save(path)

    report = check_pptx(path, classification="公开")

    assert any(item.code == expected_code and message_fragment in item.message for item in report.items)


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_stacked_composite_architecture_keeps_existing_density_lint(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    nodes = [{"id": f"n{index}", "text": f"节点{index}"} for index in range(7)]
    edge_pairs = [
        (0, 1), (0, 2), (0, 3), (1, 3), (1, 4), (2, 3), (2, 4),
        (2, 5), (3, 4), (3, 5), (3, 6), (4, 6), (5, 6),
    ]
    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.8",
            "meta": {"title": "堆叠密度", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "堆叠容器不豁免架构图边数检查",
                    "regions": [
                        {
                            "slot": "left",
                            "components": [
                                {
                                    "layout": "title_bullets",
                                    "title": "检查原则",
                                    "bullets": [{"text": "组件规则继续执行", "level": 1}],
                                },
                                {
                                    "layout": "architecture_diagram",
                                    "title": "超密架构",
                                    "nodes": nodes,
                                    "edges": [{"from": f"n{source}", "to": f"n{target}"} for source, target in edge_pairs],
                                    "groups": [],
                                },
                            ],
                        },
                        {
                            "slot": "right",
                            "components": [
                                {
                                    "layout": "cards",
                                    "title": "依据",
                                    "cards": [{"title": "边数", "desc": "13"}, {"title": "阈值", "desc": "12"}],
                                }
                            ],
                        },
                    ],
                }
            ],
        }
    )

    report = check_pptx(render_deck_ir(deck, tmp_path / "stacked-dense.pptx"), classification="公开")

    assert any(item.code == "HW-W03" and item.message.startswith("架构图过密·边数 13") for item in report.items)


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")
def test_stacked_composite_reports_column_overflow_from_real_block_geometry(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.8",
            "meta": {"title": "栏高预警", "classification": "公开", "theme": "hw_v1"},
            "slides": [
                {
                    "layout": "composite",
                    "title": "内容过多时必须提示人工拆分",
                    "regions": [
                        {
                            "slot": "left",
                            "components": [
                                {
                                    "layout": "table",
                                    "title": "阶段表",
                                    "table": {
                                        "header": ["阶段", "状态"],
                                        "rows": [["解析", "通过"], ["生成", "通过"]],
                                    },
                                },
                                {
                                    "layout": "architecture_diagram",
                                    "title": "处理链路",
                                    "nodes": [
                                        {"id": "input", "text": "输入", "type": "primary"},
                                        {"id": "gate", "text": "校验", "type": "emphasis"},
                                        {"id": "output", "text": "输出", "type": "data"},
                                    ],
                                    "edges": [{"from": "input", "to": "gate"}, {"from": "gate", "to": "output"}],
                                    "groups": [],
                                },
                                {
                                    "layout": "title_bullets",
                                    "title": "补充结论",
                                    "bullets": [
                                        {"text": "该栏以较长的可读说明保留完整事实，不自动截断也不自动缩到不可读字号以下。", "level": 1},
                                        {"text": "当累计高度不足时，应明确交给人工拆分页面。", "level": 1},
                                        {"text": "该样例故意逼近并超过单栏可用高度。", "level": 1},
                                    ],
                                },
                            ],
                        },
                        {
                            "slot": "right",
                            "components": [
                                {
                                    "layout": "title_bullets",
                                    "title": "右栏",
                                    "bullets": [{"text": "保持正常", "level": 1}],
                                }
                            ],
                        },
                    ],
                }
            ],
        }
    )

    report = check_pptx(render_deck_ir(deck, tmp_path / "stacked-overflow.pptx"), classification="公开")
    overflow_messages = [item.message for item in report.items if item.code == "HW-W03" and item.message.startswith("组合页该栏内容过多·")]

    assert any("left 栏" in message and "建议人工拆分" in message for message in overflow_messages)


@pytest.mark.parametrize(("text", "width", "height", "warns"), [("短文本", 3.0, 0.8, False), ("超长内容" * 200, 1.0, 0.3, True)])
def test_check_pptx_warns_when_autofit_cannot_preserve_minimum_font_size(
    tmp_path: Path,
    text: str,
    width: float,
    height: float,
    warns: bool,
) -> None:
    from app.lint.pptx_lint import check_pptx

    path = tmp_path / f"autofit-{warns}.pptx"
    prs = _blank_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_required_footer(slide)
    shape = _add_textbox(slide, text, 1.2, 1.8, width, height, size=10)
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    prs.save(path)

    report = check_pptx(path, classification="HUAWEI CONFIDENTIAL")
    overflow_warnings = [item for item in report.items if item.code == "HW-W03" and item.message.startswith("内容超版面·")]

    assert bool(overflow_warnings) is warns
    if warns:
        assert overflow_warnings[0].message == (
            "内容超版面·自动缩字到主题可读字号下限 10.5pt 后仍无法容纳全部内容,建议人工拆分。"
        )
        assert "人工拆分" in overflow_warnings[0].suggestion
        assert "不会截断" in overflow_warnings[0].suggestion


def test_check_pptx_architecture_reports_node_fill_outside_accent_palette(tmp_path: Path) -> None:
    from app.ir.deck_ir import DeckIR
    from app.lint.pptx_lint import check_pptx
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "架构配色", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [
                {
                    "layout": "architecture_diagram",
                    "title": "架构配色",
                    "nodes": [{"id": "a", "text": "节点A", "type": "primary"}],
                    "edges": [],
                    "groups": [],
                }
            ],
        }
    )
    path = render_deck_ir(deck, tmp_path / "architecture-color.pptx")
    prs = Presentation(str(path))
    node = next(shape for shape in prs.slides[0].shapes if shape.name.startswith("HW_ARCH_NODE:"))
    node.fill.fore_color.rgb = _rgb("123456")
    prs.save(path)

    report = check_pptx(path, classification=deck.meta.classification)

    assert any(item.code == "HW-W02" and "架构节点" in item.message for item in report.items)


def test_check_cli_writes_json_and_markdown_reports(tmp_path: Path) -> None:
    import os
    import subprocess

    from app.ir.deck_ir import DeckIR
    from app.rendering.pptx_renderer import render_deck_ir

    deck = DeckIR.model_validate(
        {
            "ir_type": "deck",
            "ir_version": "1.7",
            "meta": {"title": "报告", "classification": "HUAWEI CONFIDENTIAL"},
            "slides": [{"layout": "cover", "title": "报告"}],
        }
    )
    pptx = render_deck_ir(deck, tmp_path / "report.pptx")
    out = tmp_path / "reports"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.check",
            str(pptx),
            "--classification",
            "HUAWEI CONFIDENTIAL",
            "--output-dir",
            str(out),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(BACKEND)},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    assert (out / "report.json").exists()
    assert (out / "report.md").exists()


def test_check_cli_accepts_docx_and_writes_reports(tmp_path: Path) -> None:
    import json
    import os
    import subprocess

    from app.ir.word_ir import WordIR
    from app.rendering.docx_renderer import render_word_ir

    word = WordIR.model_validate(
        {
            "ir_type": "word",
            "ir_version": "1.0",
            "meta": {"title": "DOCX 复检", "classification": "内部公开"},
            "blocks": [{"type": "paragraph", "text": "正文"}],
        }
    )
    docx = render_word_ir(word, tmp_path / "report.docx")
    out = tmp_path / "docx-reports"

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli.check",
            str(docx),
            "--classification",
            "内部公开",
            "--output-dir",
            str(out),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(BACKEND)},
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    payload = json.loads((out / "report.json").read_text(encoding="utf-8"))
    assert payload["summary"]["pass"] is True
    assert payload["summary"]["errors"] == 0
    assert (out / "report.md").exists()


def _inject_transition(path: Path) -> None:
    temp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                xml = data.decode("utf-8")
                xml = xml.replace("</p:sld>", "<p:transition/></p:sld>")
                data = xml.encode("utf-8")
            target.writestr(item, data)
    shutil.move(temp, path)


def _blank_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.34)
    prs.slide_height = Inches(7.5)
    return prs


def _grid_width(columns: int) -> float:
    return ((13.34 - 1.14) / 12) * columns


def _add_required_footer(slide, *, top: float = 7.0) -> None:
    _add_textbox(slide, "Security Level: HUAWEI CONFIDENTIAL", 0.57, top, 4.3, 0.22, size=9, color="666666")
    _add_textbox(slide, "Copyright (C) Huawei Technologies Co., Ltd.", 4.2, top, 5.0, 0.22, size=8, color="666666")
    _add_textbox(slide, "1/1", 12.2, top, 0.6, 0.22, size=9, color="666666")


def _add_textbox(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: float = 16,
    color: str = "333333",
    fill: str | None = None,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    return shape


def _add_basic_chart(slide):
    return _add_basic_chart_at(slide, left=1.2, top=1.8, width=6.0, height=3.0)


def _add_basic_chart_at(slide, *, left: float, top: float, width: float, height: float):
    data = CategoryChartData()
    data.categories = ["1月", "2月"]
    data.add_series("方案A", [10, 12])
    return slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
        data,
    )


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _codes(report) -> list[str]:
    return [item.code for item in report.items]
