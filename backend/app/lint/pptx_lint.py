from __future__ import annotations

from dataclasses import dataclass
import json
import math
from pathlib import Path
import re
from typing import Any

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from app.rendering.theme import load_theme
from app.template.contracts import TemplateProfile
from app.template.theme import build_template_render_theme
from app.rendering.typography import text_units


EMU_PER_INCH = 914400
PT_PER_INCH = 72
EDGE_TOLERANCE_EMU = 4


@dataclass(frozen=True)
class PptxLintItem:
    code: str
    level: str
    slide: int | None
    message: str
    suggestion: str


@dataclass(frozen=True)
class PptxLintReport:
    items: list[PptxLintItem]

    @property
    def summary(self) -> dict[str, Any]:
        errors = sum(1 for item in self.items if item.level == "Error")
        warnings = sum(1 for item in self.items if item.level == "Warning")
        infos = sum(1 for item in self.items if item.level == "Info")
        return {"errors": errors, "warnings": warnings, "infos": infos, "pass": errors == 0}

    def to_dict(self) -> dict[str, Any]:
        return {
            "summary": self.summary,
            "items": [
                {
                    "code": item.code,
                    "level": item.level,
                    "slide": item.slide,
                    "message": item.message,
                    "suggestion": item.suggestion,
                }
                for item in self.items
            ],
        }


def check_pptx(
    path: Path,
    *,
    classification: str | None = None,
    theme_name: str = "hw_v1",
    template_profile: TemplateProfile | None = None,
) -> PptxLintReport:
    theme = build_template_render_theme(theme_name, template_profile) if template_profile is not None else load_theme(theme_name)
    expected_classification = classification or theme["footer"]["default_classification"]
    prs = Presentation(str(path))
    items: list[PptxLintItem] = []

    if len(prs.slides) > theme["constraints"]["max_slides"]:
        items.append(
            PptxLintItem(
                code="HW-W05",
                level="Warning",
                slide=None,
                message=f"总页数 {len(prs.slides)} 超过 {theme['constraints']['max_slides']}。",
                suggestion="拆分演示或合并低价值页面。",
            )
        )

    for slide_index, slide in enumerate(prs.slides, start=1):
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        if not _has_footer_classification(text_shapes, theme, expected_classification):
            items.append(_item("HW-E01", "Error", slide_index, "页脚缺少密级文案。", "在页脚区写入 meta.classification。"))

        xml = slide.element.xml
        if "<p:transition" in xml or "<p:timing" in xml:
            items.append(_item("HW-E03", "Error", slide_index, "存在动画或切换效果。", "移除 transition / timing XML 节点。"))

        items.extend(_font_items(slide, slide_index, theme, expected_classification))
        items.extend(_font_size_variety_items(slide, slide_index, theme, expected_classification))
        items.extend(_shape_color_items(slide, slide_index, theme))
        items.extend(_table_items(slide, slide_index, theme))
        items.extend(_chart_items(slide, slide_index, theme))
        items.extend(_threshold_line_items(slide, slide_index, theme))
        items.extend(_architecture_items(slide, slide_index, theme))
        items.extend(_composite_items(slide, slide_index, theme))
        items.extend(_sequence_layout_items(slide, slide_index, theme))
        items.extend(_kpi_items(slide, slide_index, theme))
        items.extend(_image_placeholder_items(slide, slide_index, theme))
        items.extend(_asset_image_items(slide, slide_index, theme))
        items.extend(_infographic_items(slide, slide_index, theme))
        items.extend(_combo_chart_items(slide, slide_index))
        items.extend(_bullet_items(slide, slide_index, theme))
        items.extend(_content_overflow_items(slide, slide_index, theme, expected_classification))
        items.extend(_grid_items(slide, text_shapes, slide_index, prs, theme, expected_classification))
        items.extend(_layout_items(slide, slide_index, prs, theme, expected_classification))
        items.extend(_key_frame_items(text_shapes, slide_index, prs, theme, expected_classification))
        items.extend(_contrast_items(slide, slide_index, theme, expected_classification))

    items.extend(_structure_items(prs, theme, expected_classification))
    return PptxLintReport(items)


def write_reports(report: PptxLintReport, output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "report.json"
    md_path = output_dir / "report.md"
    payload = report.to_dict()
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_lines = [
        "# PPTX 合规检查报告",
        "",
        f"- Errors: {payload['summary']['errors']}",
        f"- Warnings: {payload['summary']['warnings']}",
        f"- Infos: {payload['summary']['infos']}",
        f"- Pass: {payload['summary']['pass']}",
        "",
    ]
    for item in report.items:
        slide = "-" if item.slide is None else str(item.slide)
        md_lines.append(f"## {item.code} ({item.level})")
        md_lines.append(f"- Slide: {slide}")
        md_lines.append(f"- Message: {item.message}")
        md_lines.append(f"- Suggestion: {item.suggestion}")
        md_lines.append("")
    md_path.write_text("\n".join(md_lines), encoding="utf-8")
    return json_path, md_path


def _font_items(slide, slide_index: int, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    whitelist = set(theme["fonts"]["whitelist"])
    palette = {value.upper() for value in theme["colors"].values()}
    for shape, run, context in _iter_slide_runs(slide):
        if _is_footer_shape(shape, theme, expected_classification):
            continue
        minimum = _minimum_font_size_for_shape(shape, theme)
        for font_name in _run_font_names(run):
            if font_name not in whitelist:
                items.append(_item("HW-E02", "Error", slide_index, f"{context}字体不在白名单: {font_name}", "改用主题字体白名单。"))
        if run.font.size and run.font.size.pt < minimum:
            items.append(_item("HW-W01", "Warning", slide_index, f"{context}字号 {run.font.size.pt:.1f}pt 小于下限。", f"提高到主题最小字号 {minimum:g}pt 以上。"))
        color = _font_color(run.font)
        if color is not None and color.upper() not in palette:
            items.append(_item("HW-W02", "Warning", slide_index, f"{context}颜色 {color} 不在主题色板。", "改用 hw_theme.json 色板。"))
    return items


def _run_font_names(run) -> list[str]:
    """Return the latin and east-asian typeface names explicitly set on the run.

    `run.font.name` only maps to `a:latin`; Chinese documents are usually set
    through `a:ea`, which used to escape the whitelist entirely.
    """
    names: list[str] = []
    rpr = run._r.find(qn("a:rPr"))
    if rpr is not None:
        for tag in ("latin", "ea"):
            node = rpr.find(qn(f"a:{tag}"))
            typeface = node.get("typeface") if node is not None else None
            if typeface:
                names.append(typeface)
    return names


def _minimum_font_size_for_shape(shape, theme: dict) -> float:
    typography = theme["ppt_typography"]
    name = getattr(shape, "name", "")
    if getattr(shape, "has_table", False) or _uses_compact_typography(shape):
        return float(typography["compact_minimum_pt"])
    if name.startswith("HW_RENDERED_TEXT:COVER_TITLE"):
        return float(typography["cover_title_candidates_pt"][-1])
    if name.startswith(("HW_RENDERED_TEXT:TITLE", "HW_RENDERED_TEXT:SECTION_TITLE")):
        return float(typography["slide_title_candidates_pt"][-1])
    return float(typography["body_minimum_pt"])


def _uses_compact_typography(shape) -> bool:
    name = getattr(shape, "name", "")
    return name.startswith(
        ("HW_THRESHOLD_", "HW_RENDERED_TEXT:COMPACT", "HW_ARCH_EDGE_LABEL:", "HW_ARCH_GROUP:")
    )


def _shape_color_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    palette = {value.upper() for value in theme["colors"].values()}
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) or getattr(shape, "has_chart", False):
            continue
        for role, color in (("填充", _shape_fill_color(shape)), ("线条", _shape_line_color(shape))):
            if color is not None and color.upper() not in palette:
                return [
                    _item(
                        "HW-W02",
                        "Warning",
                        slide_index,
                        f"普通形状{role}色 {color} 不在主题色板。",
                        "形状填充与线条使用 hw_theme.json 色板。",
                    )
                ]
    return []


def _table_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    max_rows = theme["constraints"]["max_table_rows"]
    max_cols = theme["constraints"]["max_table_cols"]
    allowed_fills = _allowed_table_fills(theme)
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            rows = len(shape.table.rows)
            cols = len(shape.table.columns)
            if rows > max_rows or cols > max_cols:
                items.append(_item("HW-W04", "Warning", slide_index, f"表格尺寸 {rows}x{cols} 超过 {max_rows}x{max_cols}。", "拆分表格或减少列。"))
            items.extend(_table_fill_items(shape.table, slide_index, allowed_fills))
    return items


def _allowed_table_fills(theme: dict) -> set[str]:
    return {value.upper() for value in theme["colors"].values()}


def _table_fill_items(table, slide_index: int, allowed_fills: set[str]) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    for row in table.rows:
        for cell in row.cells:
            fill = _cell_fill_color(cell)
            if fill is not None and fill.upper() not in allowed_fills:
                items.append(
                    _item(
                        "HW-W02",
                        "Warning",
                        slide_index,
                        f"表格单元格填充色 {fill} 不在主题基础色或 accent 色板。",
                        "重点单元格使用 accent 系颜色,普通单元格使用表头红、白底或斑马纹灰底。",
                    )
                )
                return items
            for color in _xml_rgb_colors(cell._tc.xml):
                normalized = f"#{color}".upper()
                if normalized not in allowed_fills:
                    return [
                        _item(
                            "HW-W02",
                            "Warning",
                            slide_index,
                            f"表格单元格或边框颜色 {normalized} 不在主题色板。",
                            "表格文字、填充和边框统一使用主题色板。",
                        )
                    ]
    return items


def _cell_fill_color(cell) -> str | None:
    try:
        rgb = getattr(cell.fill.fore_color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return None
    return f"#{rgb}"


def _chart_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    allowed_colors = _allowed_chart_colors(theme)
    for shape in slide.shapes:
        if not getattr(shape, "has_chart", False):
            continue
        chart = shape.chart
        items.extend(_chart_quality_items(chart, slide_index))
        color_warning_added = False
        for series in chart.series:
            for color in (_chart_fill_color(series), _chart_line_color(series)):
                if color is not None and color.upper() not in allowed_colors:
                    items.append(
                        _item(
                            "HW-W02",
                            "Warning",
                            slide_index,
                            f"图表数据系列颜色 {color} 不在 accent 色板。",
                            "图表数据使用 accent1-6;重点系列使用 hw_red。",
                        )
                    )
                    color_warning_added = True
                    break
            if color_warning_added:
                break
        items.extend(_chart_font_items(chart, slide_index, theme))
    return items


def _chart_quality_items(chart, slide_index: int) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    categories, values = _chart_cached_data(chart)
    chart_type = chart.chart_type
    pie = chart_type in {XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED}
    if pie:
        if len(categories) < 2 or len(categories) > 6 or any(value <= 0 for value in values):
            items.append(
                _item(
                    "HW-W12",
                    "Warning",
                    slide_index,
                    "饼图应包含 2-6 个正值类别，当前数据不适合可靠比较。",
                    "改用条形图，或在不改变业务数据的前提下重新选择图表类型。",
                )
            )
        return items

    try:
        minimum = chart.value_axis.minimum_scale
        maximum = chart.value_axis.maximum_scale
        major = chart.value_axis.major_unit
    except (AttributeError, ValueError):
        minimum = maximum = major = None
    bar_types = {XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED}
    misses_zero = chart_type in bar_types and (
        (minimum is not None and minimum > 0) or (maximum is not None and maximum < 0)
    )
    if minimum is None or maximum is None or major is None or misses_zero:
        items.append(
            _item(
                "HW-W10",
                "Warning",
                slide_index,
                "图表坐标轴缺少明确范围/主刻度，或柱图未包含零基线。",
                "设置可读的 minimum/maximum/major_unit；柱图必须包含零轴。",
            )
        )
    long_label = max((_chart_visual_length(category) for category in categories), default=0) > 20
    if len(chart.series) > 4 or len(categories) > 12 or long_label:
        items.append(
            _item(
                "HW-W11",
                "Warning",
                slide_index,
                "图表类别、系列或标签较密，存在图例和类目轴拥挤风险。",
                "长类目优先使用横向条形图；减少同页系列或设置跳标后人工复核。",
            )
        )
    return items


def _chart_cached_data(chart) -> tuple[list[str], list[float]]:
    categories: list[str] = []
    values: list[float] = []
    for series in chart._chartSpace.iter():
        if series.tag.rsplit("}", 1)[-1] != "ser":
            continue
        for child in series:
            local_name = child.tag.rsplit("}", 1)[-1]
            if local_name == "cat" and not categories:
                categories = [value.text or "" for value in child.iter() if value.tag.rsplit("}", 1)[-1] == "v"]
            elif local_name == "val":
                for value in child.iter():
                    if value.tag.rsplit("}", 1)[-1] != "v" or value.text is None:
                        continue
                    try:
                        values.append(float(value.text))
                    except ValueError:
                        continue
    return categories, values


def _chart_visual_length(text: str) -> int:
    return sum(2 if "\u2e80" <= character <= "\uffff" else 1 for character in text)


def _allowed_chart_colors(theme: dict) -> set[str]:
    return {
        theme["colors"][key].upper()
        for key in ("hw_red", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "secondary", "border")
        if key in theme["colors"]
    }


def _chart_fill_color(series) -> str | None:
    try:
        rgb = getattr(series.format.fill.fore_color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return None
    return f"#{rgb}"


def _chart_line_color(series) -> str | None:
    try:
        rgb = getattr(series.format.line.color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return None
    return f"#{rgb}"


def _chart_font_items(chart, slide_index: int, theme: dict) -> list[PptxLintItem]:
    minimum = theme["ppt_typography"]["compact_minimum_pt"]
    whitelist = set(theme["fonts"]["whitelist"])
    items: list[PptxLintItem] = []
    for context, font in _chart_fonts(chart):
        if font.name and font.name not in whitelist:
            return [_item("HW-E02", "Error", slide_index, f"{context}字体不在白名单: {font.name}", "改用主题字体白名单。")]
        if font.size is not None and font.size.pt < minimum:
            items.append(_item("HW-W01", "Warning", slide_index, f"{context}字号 {font.size.pt:.1f}pt 小于下限。", f"提高到主题最小字号 {minimum:g}pt 以上。"))
            return items
        color = _font_color(font)
        if color is not None:
            color = color.upper()
            if color not in {value.upper() for value in theme["colors"].values()}:
                items.append(_item("HW-W02", "Warning", slide_index, f"{context}颜色 {color} 不在主题色板。", "改用 hw_theme.json 色板。"))
                return items
    return items


def _chart_fonts(chart) -> list[tuple[str, Any]]:
    fonts: list[tuple[str, Any]] = []
    for attr, label in (("category_axis", "图表 X 轴"), ("value_axis", "图表 Y 轴")):
        try:
            axis = getattr(chart, attr)
            fonts.append((label, axis.tick_labels.font))
            if axis.has_title:
                for paragraph in axis.axis_title.text_frame.paragraphs:
                    fonts.extend((f"{label}标题", run.font) for run in paragraph.runs if run.text.strip())
        except (AttributeError, ValueError):
            continue
    try:
        if chart.has_legend:
            fonts.append(("图表图例", chart.legend.font))
    except (AttributeError, ValueError):
        pass
    try:
        for plot in chart.plots:
            if plot.has_data_labels:
                fonts.append(("图表数据标签", plot.data_labels.font))
    except (AttributeError, ValueError):
        pass
    return fonts


def _threshold_line_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    expected_color = theme["colors"]["hw_red"].upper()
    expected_width = theme["layouts"]["chart"]["threshold"]["line_width_pt"]
    horizontal_chart = any(
        getattr(shape, "has_chart", False) and shape.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
        for shape in slide.shapes
    )
    for shape in slide.shapes:
        if not getattr(shape, "name", "").startswith("HW_THRESHOLD_LINE"):
            continue
        color = _shape_line_color(shape)
        if color is not None and color.upper() != expected_color:
            return [_item("HW-W02", "Warning", slide_index, f"阈值线颜色 {color} 不是主题红。", "阈值线使用 hw_red / #C7000B。")]
        try:
            width = shape.line.width.pt
        except (AttributeError, TypeError, ValueError):
            width = expected_width
        if abs(width - expected_width) > 0.1:
            return [_item("HW-W02", "Warning", slide_index, f"阈值线线宽 {width:.1f}pt 与主题不一致。", f"阈值线线宽设为 {expected_width:g}pt。")]
        line_is_vertical = shape.height > shape.width
        if horizontal_chart != line_is_vertical:
            expected_direction = "竖向" if horizontal_chart else "横向"
            return [
                _item(
                    "HW-W07",
                    "Warning",
                    slide_index,
                    f"阈值线方向与图表数值轴不一致，当前应为{expected_direction}。",
                    "横向数据条使用竖向阈值线；纵向柱状图和折线图使用横向阈值线。",
                )
            ]
    return []


def _architecture_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    layout = theme["layouts"]["architecture_diagram"]
    allowed_node_fills = {
        theme["colors"][color_key].upper()
        for color_key in layout["node_type_colors"].values()
    }
    expected_sizes = {
        "HW_ARCH_NODE:": layout["node_font_size_pt"],
        "HW_ARCH_GROUP:": layout["group_label_font_size_pt"],
        "HW_ARCH_EDGE_LABEL:": layout["edge_label_font_size_pt"],
    }
    edge_count = sum(
        1
        for shape in slide.shapes
        if getattr(shape, "name", "").startswith("HW_ARCH_EDGE:")
    )
    edge_limit = theme["constraints"]["max_architecture_edges_before_warning"]
    if edge_count > edge_limit:
        items.append(
            _item(
                "HW-W03",
                "Warning",
                slide_index,
                f"架构图过密·边数 {edge_count} 超过阈值 {edge_limit},建议人工调整或拆分。",
                "架构图过密建议人工调整或拆分。",
            )
        )
    for shape in slide.shapes:
        name = getattr(shape, "name", "")
        if name.startswith("HW_ARCH_NODE:"):
            line = _shape_line_color(shape)
            if line is None or line.upper() not in allowed_node_fills:
                return items + [
                    _item(
                        "HW-W02",
                        "Warning",
                        slide_index,
                        f"架构节点边框色 {line or '未设置'} 不在规定 accent 色板。",
                        "节点边框色必须来自 layouts.architecture_diagram.node_type_colors 对应主题色。",
                    )
                ]
        for prefix, expected_size in expected_sizes.items():
            if not name.startswith(prefix):
                continue
            for run in _shape_runs(shape):
                if run.text.strip() and run.font.size is not None and abs(run.font.size.pt - expected_size) > 0.1:
                    return items + [
                        _item(
                            "HW-W01",
                            "Warning",
                            slide_index,
                            f"架构元素字号 {run.font.size.pt:.1f}pt 与主题规定 {expected_size:g}pt 不一致。",
                            "架构节点、分组标签和连线标签分别使用 theme 中对应字号。",
                        )
                    ]
    return items


def _composite_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    layout = theme["layouts"].get("composite", {})
    content = layout.get("content", {})
    if not content:
        return []
    content_top = float(content["top_in"])
    content_bottom = content_top + float(content["height_in"])
    overflow_by_slot: dict[str, float] = {}
    for shape in slide.shapes:
        name = getattr(shape, "name", "")
        if not name.startswith("HW_COMPOSITE_BLOCK:"):
            continue
        parts = name.split(":")
        if len(parts) != 3 or parts[1] not in {"left", "right"}:
            continue
        bottom = (shape.top + shape.height) / EMU_PER_INCH
        if bottom > content_bottom + 0.01:
            overflow_by_slot[parts[1]] = max(overflow_by_slot.get(parts[1], content_top), bottom)
    return [
        _item(
            "HW-W03",
            "Warning",
            slide_index,
            (
                f"组合页该栏内容过多·{slot} 栏累计高度 {bottom - content_top:.2f}\" "
                f"超过可用高度 {content_bottom - content_top:.2f}\",建议人工拆分。"
            ),
            "该栏内容过多,建议人工拆分；系统不会平均压缩、截断内容或自动拆页。",
        )
        for slot, bottom in sorted(overflow_by_slot.items())
    ]


def _sequence_layout_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    process_steps = [
        shape for shape in slide.shapes if getattr(shape, "name", "").startswith("HW_PROCESS_STEP:")
    ]
    timeline_cards = [
        shape for shape in slide.shapes if getattr(shape, "name", "").startswith("HW_TIMELINE_MILESTONE:")
    ]
    if not process_steps and not timeline_cards:
        return []
    shapes = process_steps or timeline_cards
    layout_name = "process_flow" if process_steps else "timeline"
    layout = theme["layouts"][layout_name]
    content = layout["content"]
    tolerance = max(float(theme["grid"]["snap_tolerance_in"]), 0.02)
    expected_width = _inches(shapes[0].width)
    expected_height = _inches(shapes[0].height)
    for shape in shapes[1:]:
        if (
            abs(_inches(shape.width) - expected_width) > tolerance
            or abs(_inches(shape.height) - expected_height) > tolerance
        ):
            return [
                _item(
                    "HW-W07",
                    "Warning",
                    slide_index,
                    f"{layout_name} 同级步骤或里程碑未保持等宽等高。",
                    "按主题内容区等分布局,让同级元素使用同一宽度和高度。",
                )
            ]
    content_left = content["left_in"]
    content_top = content["top_in"]
    content_right = content_left + content["width_in"]
    content_bottom = content_top + content["height_in"]
    for shape in shapes:
        left = _inches(shape.left)
        top = _inches(shape.top)
        right = _inches(shape.left + shape.width)
        bottom = _inches(shape.top + shape.height)
        if (
            left < content_left - tolerance
            or top < content_top - tolerance
            or right > content_right + tolerance
            or bottom > content_bottom + tolerance
        ):
            return [
                _item(
                    "HW-W07",
                    "Warning",
                    slide_index,
                    f"{layout_name} 步骤或里程碑超出主题内容区或进入页脚安全区。",
                    "按 theme.layouts 对应 content 坐标重新排布。",
                )
                ]
    lefts = [_inches(shape.left) for shape in shapes]
    tops = [_inches(shape.top) for shape in shapes]
    is_vertical = max(lefts) - min(lefts) <= tolerance and max(tops) - min(tops) > tolerance
    dense_limit = 5 if layout_name == "process_flow" else 6
    if is_vertical and len(shapes) > dense_limit:
        label = "步骤" if layout_name == "process_flow" else "里程碑"
        return [
            _item(
                "HW-W03",
                "Warning",
                slide_index,
                f"{layout_name} 纵向内容过密·{label}数 {len(shapes)} 超过可读阈值 {dense_limit},建议人工调整或拆分。",
                "保留可编辑骨架；人工拆成多页或减少单页节点，不会截断内容或静默改为横向。",
            )
        ]
    allowed_sizes = {
        float(layout["title_font_size_pt"]),
        float(layout["description_font_size_pt"]),
    }
    if layout_name == "timeline":
        allowed_sizes.add(float(layout["label_font_size_pt"]))
    for shape in shapes:
        for run in _shape_runs(shape):
            if not run.text.strip() or run.font.size is None:
                continue
            if not any(abs(run.font.size.pt - allowed) <= 0.1 for allowed in allowed_sizes):
                return [
                    _item(
                        "HW-W01",
                        "Warning",
                        slide_index,
                        f"{layout_name} 元素字号 {run.font.size.pt:.1f}pt 与主题 token 不一致。",
                        "使用 theme.layouts 对应标题、说明和标签字号。",
                    )
                ]
    if timeline_cards:
        allowed_marker_colors = {
            theme["colors"][color_key].upper()
            for color_key in layout["status_colors"].values()
        }
        for shape in slide.shapes:
            if not getattr(shape, "name", "").startswith("HW_TIMELINE_MARKER:"):
                continue
            fill = _shape_fill_color(shape)
            if fill is None or fill.upper() not in allowed_marker_colors:
                return [
                    _item(
                        "HW-W02",
                        "Warning",
                        slide_index,
                        f"timeline 状态点颜色 {fill or '未设置'} 不在主题状态色映射中。",
                        "按 completed/current/planned 对应主题色渲染状态点。",
                    )
                ]
    return []


def _kpi_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    shapes = [
        shape for shape in slide.shapes if getattr(shape, "name", "").startswith("HW_RENDERED_TEXT:KPI_CARD:")
    ]
    if not shapes:
        return []
    layout = theme["layouts"]["cards"]
    tolerance = max(float(theme["grid"]["snap_tolerance_in"]), 0.02)
    expected_width = _inches(shapes[0].width)
    expected_height = _inches(shapes[0].height)
    if any(
        abs(_inches(shape.width) - expected_width) > tolerance
        or abs(_inches(shape.height) - expected_height) > tolerance
        for shape in shapes[1:]
    ):
        return [_item("HW-W07", "Warning", slide_index, "KPI 指标块未保持等宽等高。", "按 cards 网格 token 等分 KPI 指标块。")]
    for shape in shapes:
        paragraphs = [paragraph for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
        if len(paragraphs) < 2 or not paragraphs[0].runs:
            return [_item("HW-W03", "Warning", slide_index, "KPI 指标缺少核心数值或指标名。", "每个 KPI 提供数值、指标名和可选口径。")]
        value_run = paragraphs[0].runs[0]
        expected_font = theme["fonts"]["number"][0]
        expected_size = float(layout["kpi_value_font_size_pt"])
        actual_size = value_run.font.size.pt if value_run.font.size is not None else 0
        if value_run.font.name != expected_font or abs(actual_size - expected_size) > 0.1:
            return [
                _item(
                    "HW-W01",
                    "Warning",
                    slide_index,
                    f"KPI 核心数值未使用主题 number 字体或 {expected_size:g}pt 字号。",
                    "从 hw_theme.json 的 fonts.number 和 cards.kpi_value_font_size_pt 读取样式。",
                )
            ]
        value_color = _run_color(value_run, theme["colors"]["body"])
        if value_color.upper() != theme["colors"]["hw_red"].upper():
            return [_item("HW-W02", "Warning", slide_index, "KPI 核心数值未使用主题红。", "核心数值使用 hw_red。")]
    return []


def _image_placeholder_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    placeholders = [shape for shape in slide.shapes if getattr(shape, "name", "") == "HW_IMAGE_PLACEHOLDER"]
    if not placeholders:
        return []
    expected_ratio = 16 / 9
    margin = float(theme["grid"]["min_edge_margin_in"])
    footer_top = float(theme["slide"]["footer_top_in"])
    for shape in placeholders:
        width = _inches(shape.width)
        height = _inches(shape.height)
        ratio = width / height if height else 0
        left = _inches(shape.left)
        right = left + width
        bottom = _inches(shape.top) + height
        if abs(ratio - expected_ratio) > 0.01:
            return [_item("HW-W07", "Warning", slide_index, f"图片占位槽宽高比 {ratio:.3f} 不是 16:9。", "使用 theme.layouts.image.box 的 16:9 尺寸。")]
        if left < margin or right > theme["slide"]["width_in"] - margin or bottom > footer_top - theme["grid"]["min_gap_in"]:
            return [_item("HW-W07", "Warning", slide_index, "图片占位槽进入页边距或页脚安全区。", "按主题图片槽坐标恢复安全边距。")]
        if "等比放入" not in _shape_text(shape) or "禁止随意裁切" not in _shape_text(shape):
            return [_item("HW-W03", "Warning", slide_index, "图片占位槽缺少等比适配说明。", "保留等比放入、禁止随意裁切的替换提示。")]
    return []


def _asset_image_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    for shape in slide.shapes:
        if not getattr(shape, "name", "").startswith("HW_ASSET_IMAGE:"):
            continue
        metadata = _shape_description_json(shape)
        dpi = float(metadata.get("effective_dpi", 0) or 0)
        if dpi < 150:
            items.append(
                _item(
                    "HW-W13",
                    "Warning",
                    slide_index,
                    f"图片有效分辨率仅 {dpi:g} DPI，可能在投影或打印时模糊。",
                    "替换为更高分辨率原图，或减小图片在页面中的显示尺寸。",
                )
            )
        crop = float(metadata.get("crop_fraction", 0) or 0)
        alt_present = bool(metadata.get("alt_present"))
        generated_without_credit = metadata.get("source_type") == "generated" and not metadata.get("credit_present")
        if crop > 0.55 or not alt_present or generated_without_credit:
            risks = []
            if crop > 0.55:
                risks.append("裁切比例过大")
            if not alt_present:
                risks.append("缺少替代文本")
            if generated_without_credit:
                risks.append("生成图片缺少来源")
            items.append(
                _item(
                    "HW-W14",
                    "Warning",
                    slide_index,
                    "图片存在" + "、".join(risks) + "风险。",
                    "调整 focal point/fit，并补充 alt 与必要来源信息。",
                )
            )
        margin = float(theme["grid"]["min_edge_margin_in"])
        if (
            _inches(shape.left) < margin
            or _inches(shape.top) < margin
            or _inches(shape.left + shape.width) > theme["slide"]["width_in"] - margin
            or _inches(shape.top + shape.height) > theme["slide"]["footer_top_in"]
        ):
            items.append(
                _item("HW-W14", "Warning", slide_index, "图片进入页边或页脚安全区。", "按主题图片布局恢复安全边距。")
            )
    return items


def _infographic_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    shapes = [shape for shape in slide.shapes if getattr(shape, "name", "").startswith("HW_INFOGRAPHIC:")]
    if not shapes:
        return []
    item_shapes = [shape for shape in shapes if ":item:" in shape.name or re.search(r":\d+$", shape.name)]
    long_label = any(_text_units(_shape_text(shape)) > 44 for shape in item_shapes if _has_visible_text(shape))
    overlapping = False
    for index, first in enumerate(item_shapes):
        for second in item_shapes[index + 1 :]:
            if _boxes_overlap_or_too_close(_bbox(first), _bbox(second), 0):
                overlapping = True
                break
        if overlapping:
            break
    if long_label or overlapping:
        reason = "标签过长" if long_label else "标签或节点碰撞"
        return [
            _item(
                "HW-W15",
                "Warning",
                slide_index,
                f"信息图存在{reason}，结构密度可能不适合当前版式。",
                "缩短标签、减少节点或拆分页面；不得静默截断业务内容。",
            )
        ]
    return []


def _combo_chart_items(slide, slide_index: int) -> list[PptxLintItem]:
    charts = {
        shape.name.removeprefix("HW_COMBO_CHART:"): shape
        for shape in slide.shapes
        if getattr(shape, "name", "").startswith("HW_COMBO_CHART:") and getattr(shape, "has_chart", False)
    }
    if not charts:
        return []
    if set(charts) != {"primary", "secondary"}:
        return [_item("HW-W16", "Warning", slide_index, "组合图缺少主轴或次轴原生图表对象。", "重新按组合图契约渲染主轴与次轴。")]
    primary_meta = _shape_description_json(charts["primary"])
    secondary_meta = _shape_description_json(charts["secondary"])
    axis_position = charts["secondary"].chart.value_axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}axPos")
    same_units = set(primary_meta.get("units") or []) == set(secondary_meta.get("units") or [])
    missing_format = not primary_meta.get("number_format") or not secondary_meta.get("number_format")
    if axis_position is None or axis_position.get("val") != "r" or same_units or missing_format:
        return [
            _item(
                "HW-W16",
                "Warning",
                slide_index,
                "组合图次轴位置、量纲或数字格式不完整。",
                "仅在真实量纲不同且格式显式时使用次轴，并把次轴放在右侧。",
            )
        ]
    return []


def _shape_description_json(shape) -> dict[str, Any]:
    try:
        if hasattr(shape.element, "nvPicPr"):
            raw = shape.element.nvPicPr.cNvPr.get("descr", "")
        else:
            raw = shape.element.nvGraphicFramePr.cNvPr.get("descr", "")
        parsed = json.loads(raw) if raw else {}
        return parsed if isinstance(parsed, dict) else {}
    except (AttributeError, TypeError, ValueError, json.JSONDecodeError):
        return {}


def _shape_fill_color(shape) -> str | None:
    try:
        rgb = getattr(shape.fill.fore_color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return None
    return f"#{rgb}"


def _shape_line_color(shape) -> str | None:
    try:
        rgb = getattr(shape.line.color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return None
    return f"#{rgb}"


def _font_size_variety_items(slide, slide_index: int, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    maximum = theme["constraints"].get("max_font_size_kinds_per_slide")
    if not maximum:
        return []
    sizes = set()
    for shape, run, _context in _iter_slide_runs(slide):
        if _is_footer_shape(shape, theme, expected_classification) or _uses_compact_typography(shape):
            continue
        if run.text.strip() and run.font.size is not None:
            sizes.add(round(run.font.size.pt, 1))
    for shape in slide.shapes:
        if not getattr(shape, "has_chart", False):
            continue
        for _context, font in _chart_fonts(shape.chart):
            if font.size is not None:
                sizes.add(round(font.size.pt, 1))
    if len(sizes) > maximum:
        return [
            _item(
                "HW-W01",
                "Warning",
                slide_index,
                f"单页字号种类 {len(sizes)} 种超过主题上限 {maximum} 种。",
                "合并标题、正文、辅助文字的字号层级,控制在主题上限内。",
            )
        ]
    return []


def _bullet_items(slide, slide_index: int, theme: dict) -> list[PptxLintItem]:
    bullet_lines: list[str] = []
    for _shape, paragraph, _context in _iter_slide_paragraphs(slide):
        text = paragraph.text.strip()
        if not text:
            continue
        if _paragraph_has_native_bullet(paragraph) or text.startswith(("•", "-", "–")):
            bullet_lines.append(text.lstrip("•-– "))
    count_limit = theme["constraints"]["max_bullets_per_slide"]
    char_limit = theme["constraints"]["max_bullet_chars"]
    longest = max((len(line) for line in bullet_lines), default=0)
    too_many = len(bullet_lines) > count_limit
    too_long = longest > char_limit
    if too_many or too_long:
        reasons: list[str] = []
        if too_many:
            reasons.append(f"本页共 {len(bullet_lines)} 条要点,超过上限 {count_limit} 条")
        if too_long:
            reasons.append(f"单条最长 {longest} 字,超过上限 {char_limit} 字")
        if too_many and too_long:
            action = "建议精简要点并拆分"
        elif too_many:
            action = "建议删减或拆分"
        else:
            action = "建议精简表述"
        message = f"要点超限·{'；'.join(reasons)},{action}。"
        return [_item("HW-W03", "Warning", slide_index, message, f"控制在 {count_limit} 条以内且单条不超过 {char_limit} 字。")]
    return []


def _content_overflow_items(
    slide,
    slide_index: int,
    theme: dict,
    expected_classification: str,
) -> list[PptxLintItem]:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            minimum = _minimum_font_size_for_shape(shape, theme)
            if (
                _has_visible_text(shape)
                and not _is_footer_shape(shape, theme, expected_classification)
                and _text_frame_exceeds_minimum(shape.text_frame, shape.width, shape.height, theme, minimum)
            ):
                return [_overflow_warning(slide_index, minimum)]
        if not getattr(shape, "has_table", False):
            continue
        seen_cells: set[int] = set()
        for row_index, row in enumerate(shape.table.rows):
            for column_index, cell in enumerate(row.cells):
                cell_id = id(cell._tc)
                if cell_id in seen_cells or not cell.text.strip():
                    continue
                seen_cells.add(cell_id)
                width = shape.table.columns[column_index].width
                height = shape.table.rows[row_index].height
                minimum = theme["ppt_typography"]["compact_minimum_pt"]
                if _text_frame_exceeds_minimum(cell.text_frame, width, height, theme, minimum):
                    return [_overflow_warning(slide_index, minimum)]
    return []


def _overflow_warning(slide_index: int, minimum: float) -> PptxLintItem:
    return _item(
        "HW-W03",
        "Warning",
        slide_index,
        f"内容超版面·自动缩字到主题可读字号下限 {minimum:g}pt 后仍无法容纳全部内容,建议人工拆分。",
        "该页内容过多建议人工拆分；系统不会截断内容或自动拆页。",
    )


def _text_frame_exceeds_minimum(
    text_frame,
    width_emu: int,
    height_emu: int,
    theme: dict,
    minimum: float,
) -> bool:
    if text_frame.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
        return False
    minimum = float(minimum)
    line_spacing = float(theme["typography"]["line_spacing"])
    horizontal_margins = _emu_to_points(text_frame.margin_left) + _emu_to_points(text_frame.margin_right)
    vertical_margins = _emu_to_points(text_frame.margin_top) + _emu_to_points(text_frame.margin_bottom)
    available_width = max(_emu_to_points(width_emu) - horizontal_margins, minimum)
    available_height = max(_emu_to_points(height_emu) - vertical_margins, minimum * line_spacing)
    units_per_line = max(available_width / minimum, 1)
    available_lines = max(int(available_height / (minimum * line_spacing)), 1)
    required_lines = sum(
        max(math.ceil(_text_units(line) / units_per_line), 1)
        for line in text_frame.text.splitlines()
    )
    return required_lines > available_lines


def _text_units(text: str) -> float:
    return text_units(text)


def _emu_to_points(value: int | None) -> float:
    return 0.0 if value is None else value / EMU_PER_INCH * PT_PER_INCH


def _grid_items(slide, text_shapes: list, slide_index: int, prs, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    grid_lines = _column_grid_lines(prs, theme)
    tolerance_in = theme["grid"]["snap_tolerance_in"]
    baseline_pt = theme["grid"]["baseline_pt"]
    tolerance_pt = tolerance_in * PT_PER_INCH

    for shape in text_shapes:
        if not _has_visible_text(shape) or _is_footer_shape(shape, theme, expected_classification):
            continue
        if (
            _is_architecture_shape(shape)
            or _is_threshold_shape(shape)
            or _is_sequence_shape(shape)
            or _is_image_placeholder_shape(shape)
            or _is_infographic_shape(shape)
            or _is_asset_image_shape(shape)
        ):
            continue
        allows_compacted_vertical_position = _is_intentionally_compacted_composite_text(shape, slide, theme)
        left = _inches(shape.left)
        right = _inches(shape.left + shape.width)
        top_pt = _inches(shape.top) * PT_PER_INCH
        bottom_pt = _inches(shape.top + shape.height) * PT_PER_INCH
        horizontal_ok = _near_any(left, grid_lines, tolerance_in) and _near_any(right, grid_lines, tolerance_in)
        vertical_ok = allows_compacted_vertical_position or (
            _near_step(top_pt, baseline_pt, tolerance_pt) and _near_step(bottom_pt, baseline_pt, tolerance_pt)
        )
        if not horizontal_ok or not vertical_ok:
            scope = "renderer 文本框" if _is_renderer_text_shape(shape) else "文本框"
            items.append(
                _item(
                    "HW-W06",
                    "Warning",
                    slide_index,
                    f"{scope}未吸附到 12 栏网格或 8pt 基线。",
                    "按主题 12 栏网格调整左右边界,按 8pt 基线调整上下边界。",
                )
            )
            break
    return items


def _is_intentionally_compacted_composite_text(shape, slide, theme: dict) -> bool:
    """Allow only theme-approved sub-baseline positions inside a squeezed composite block."""
    layout = theme["layouts"].get("composite", {})
    minimum_gap = float(layout.get("stack_min_gap_in", 0))
    default_gap = float(layout.get("stack_gap_in", 0))
    if minimum_gap <= 0 or minimum_gap >= default_gap:
        return False

    blocks_by_slot: dict[str, list] = {"left": [], "right": []}
    for candidate in slide.shapes:
        parts = getattr(candidate, "name", "").split(":")
        if len(parts) == 3 and parts[0] == "HW_COMPOSITE_BLOCK" and parts[1] in blocks_by_slot:
            blocks_by_slot[parts[1]].append(candidate)

    shape_box = _bbox(shape)
    for blocks in blocks_by_slot.values():
        blocks.sort(key=lambda candidate: candidate.top)
        for index, block in enumerate(blocks):
            if index == 0 or not _box_contains(_bbox(block), shape_box):
                continue
            previous = blocks[index - 1]
            actual_gap = _inches(block.top) - _inches(previous.top + previous.height)
            return minimum_gap - 0.01 <= actual_gap < default_gap - 0.01
    return False


def _is_renderer_text_shape(shape) -> bool:
    return getattr(shape, "name", "").startswith("HW_RENDERED_TEXT")


def _is_architecture_shape(shape) -> bool:
    return getattr(shape, "name", "").startswith("HW_ARCH_")


def _is_threshold_shape(shape) -> bool:
    return getattr(shape, "name", "").startswith("HW_THRESHOLD_")


def _is_sequence_shape(shape) -> bool:
    name = getattr(shape, "name", "")
    return name.startswith("HW_PROCESS_") or name.startswith("HW_TIMELINE_")


def _is_image_placeholder_shape(shape) -> bool:
    return getattr(shape, "name", "") == "HW_IMAGE_PLACEHOLDER"


def _is_infographic_shape(shape) -> bool:
    name = getattr(shape, "name", "")
    return name.startswith("HW_INFOGRAPHIC:") or name.startswith("HW_SEMANTIC_ICON:")


def _is_asset_image_shape(shape) -> bool:
    return getattr(shape, "name", "").startswith("HW_ASSET_IMAGE:")


def _layout_items(slide, slide_index: int, prs, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    width = prs.slide_width
    height = prs.slide_height
    left_margin = int(theme["slide"]["margin_left_in"] * 914400)
    right_margin = int(theme["slide"]["margin_right_in"] * 914400)
    top_margin = int(theme["slide"]["margin_top_in"] * 914400)
    bottom_margin = int(theme["slide"]["margin_bottom_in"] * 914400)
    min_gap = theme["grid"]["min_gap_in"]
    content_shapes = _layout_shapes(slide, prs, theme, expected_classification)
    for shape in content_shapes:
        if shape.left < left_margin - EDGE_TOLERANCE_EMU or shape.top < top_margin - EDGE_TOLERANCE_EMU:
            items.append(_item("HW-W07", "Warning", slide_index, "文本、表格、图表或图形过近页边。", "按主题页边距重新排布。"))
            break
        if (
            shape.left + shape.width > width - right_margin + EDGE_TOLERANCE_EMU
            or shape.top + shape.height > height - bottom_margin + EDGE_TOLERANCE_EMU
        ):
            items.append(_item("HW-W07", "Warning", slide_index, "文本、表格、图表或图形超出版心或过近页边。", "按主题页边距重新排布。"))
            break
    if items:
        return items
    for left_index, first in enumerate(content_shapes):
        first_box = _bbox(first)
        for second in content_shapes[left_index + 1 :]:
            second_box = _bbox(second)
            if _valid_containment(first, second, first_box, second_box):
                continue
            if _boxes_overlap_or_too_close(first_box, second_box, min_gap):
                items.append(_item("HW-W07", "Warning", slide_index, "文本、表格、图表或图形重叠或间距小于主题下限。", "按 theme.grid.min_gap_in 重新排布。"))
                return items
    return items


def _layout_shapes(slide, prs, theme: dict, expected_classification: str) -> list:
    shapes = []
    slide_area = int(prs.slide_width) * int(prs.slide_height)
    for shape in slide.shapes:
        name = getattr(shape, "name", "")
        if (
            _is_footer_shape(shape, theme, expected_classification)
            or name.startswith("HW_DECORATION:")
            or name.startswith("HW_TEMPLATE_DECORATION:")
            or name.startswith("HW_COMPOSITE_BLOCK:")
            or _is_architecture_shape(shape)
            or _is_threshold_shape(shape)
            or _is_sequence_shape(shape)
            or _is_infographic_shape(shape)
            or _is_asset_image_shape(shape)
            or name.startswith("HW_COMBO_CHART:")
        ):
            continue
        shape_type = str(getattr(shape, "shape_type", ""))
        if "LINE" in shape_type or "GROUP" in shape_type or "PLACEHOLDER" in shape_type:
            continue
        area = max(int(shape.width), 0) * max(int(shape.height), 0)
        if area >= slide_area * 0.9 and not _has_visible_text(shape) and not getattr(shape, "has_table", False):
            continue
        if (
            _has_visible_text(shape)
            or getattr(shape, "has_table", False)
            or getattr(shape, "has_chart", False)
            or "PICTURE" in shape_type
            or "AUTO_SHAPE" in shape_type
        ):
            shapes.append(shape)
    return shapes


def _valid_containment(first, second, first_box, second_box) -> bool:
    first_contains = _box_contains(first_box, second_box)
    second_contains = _box_contains(second_box, first_box)
    if not first_contains and not second_contains:
        return False
    container = first if first_contains else second
    child = second if first_contains else first
    if getattr(container, "name", "").startswith("HW_LAYOUT_CONTAINER:"):
        return True
    return (
        not _has_visible_text(container)
        and not getattr(container, "has_table", False)
        and not getattr(container, "has_chart", False)
        and _has_visible_text(child)
    )


def _box_contains(container: tuple[float, float, float, float], child: tuple[float, float, float, float]) -> bool:
    return (
        container[0] <= child[0]
        and container[1] <= child[1]
        and container[2] >= child[2]
        and container[3] >= child[3]
    )


def _key_frame_items(text_shapes: list, slide_index: int, prs, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    items: list[PptxLintItem] = []
    tolerance = max(theme["grid"]["snap_tolerance_in"], 0.06)
    title_layout = theme["layouts"]["title"]
    footer_layout = theme["layouts"]["footer"]
    title_left = title_layout["left_in"]
    title_top = title_layout["top_in"]
    title_width = title_layout["width_in"]
    footer_left = footer_layout["classification_left_in"]
    footer_top = theme["slide"]["footer_top_in"]

    for shape in text_shapes:
        if not _has_visible_text(shape):
            continue
        text = _shape_text(shape)
        left = _inches(shape.left)
        top = _inches(shape.top)
        width = _inches(shape.width)
        if expected_classification in text:
            if abs(left - footer_left) > tolerance or abs(top - footer_top) > tolerance:
                items.append(_item("HW-W08", "Warning", slide_index, "页脚密级框坐标偏离 theme 规定。", "将密级页脚框放回主题页脚坐标。"))
                break
        elif _looks_like_slide_title(shape, theme, expected_classification):
            if abs(left - title_left) > tolerance or abs(top - title_top) > tolerance or abs(width - title_width) > 0.12:
                items.append(_item("HW-W08", "Warning", slide_index, "标题框坐标偏离 theme 规定。", "将标题框放回主题标题区坐标。"))
                break
    return items


def _contrast_items(slide, slide_index: int, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    minimum = theme["constraints"]["contrast_ratio_min"]
    for shape in slide.shapes:
        if not _has_visible_text(shape) or _is_footer_shape(shape, theme, expected_classification):
            continue
        background = _shape_background(shape, slide, theme)
        if background is None:
            return [
                _item(
                    "HW-W09",
                    "Warning",
                    slide_index,
                    "正文位于图片、渐变或其它复杂背景上，无法自动判定对比度。",
                    "人工复核该文字区域，或改为可测量的纯色背景。",
                )
            ]
        for run in _shape_runs(shape):
            if not run.text.strip():
                continue
            foreground = _run_color(run, theme["colors"]["body"])
            ratio = _contrast_ratio(foreground, background)
            if ratio < minimum:
                return [_contrast_item(slide_index, ratio, minimum, "正文")]
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    background = _cell_fill_color(cell) or theme["colors"]["background"]
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if not run.text.strip():
                                continue
                            ratio = _contrast_ratio(_run_color(run, theme["colors"]["body"]), background)
                            if ratio < minimum:
                                return [_contrast_item(slide_index, ratio, minimum, "表格文字")]
        if getattr(shape, "has_chart", False):
            for context, font in _chart_fonts(shape.chart):
                foreground = _font_color(font) or theme["colors"]["secondary"]
                ratio = _contrast_ratio(foreground, theme["colors"]["background"])
                if ratio < minimum:
                    return [_contrast_item(slide_index, ratio, minimum, context)]
    return []


def _contrast_item(slide_index: int, ratio: float, minimum: float, context: str) -> PptxLintItem:
    return _item(
        "HW-W09",
        "Warning",
        slide_index,
        f"{context}与背景对比度 {ratio:.2f}:1 低于 {minimum}:1。",
        "提高文字与背景颜色差异,满足 WCAG AA 4.5:1。",
    )


def _column_grid_lines(prs, theme: dict) -> list[float]:
    margin = theme["slide"]["margin_in"]
    columns = theme["grid"]["columns"]
    content_width = _inches(prs.slide_width) - margin * 2
    column_width = content_width / columns
    return [margin + index * column_width for index in range(columns + 1)]


def _near_any(value: float, candidates: list[float], tolerance: float) -> bool:
    return any(abs(value - candidate) <= tolerance for candidate in candidates)


def _near_step(value: float, step: float, tolerance: float) -> bool:
    if step <= 0:
        return True
    nearest = round(value / step) * step
    return abs(value - nearest) <= tolerance


def _looks_like_slide_title(shape, theme: dict, expected_classification: str) -> bool:
    if _is_footer_shape(shape, theme, expected_classification):
        return False
    top = _inches(shape.top)
    if top > theme["slide"]["body_top_in"]:
        return False
    if getattr(shape, "name", "").startswith(("HW_RENDERED_TEXT:TITLE", "HW_RENDERED_TEXT:SECTION_TITLE")):
        return True
    minimum_title = theme["ppt_typography"]["slide_title_candidates_pt"][-1]
    return _max_font_size(shape) >= minimum_title - 1


def _max_font_size(shape) -> float:
    sizes = [run.font.size.pt for run in _shape_runs(shape) if run.font.size is not None]
    return max(sizes, default=0)


def _is_footer_shape(shape, theme: dict, expected_classification: str) -> bool:
    text = _shape_text(shape)
    if expected_classification in text:
        return _is_footer_area(shape, theme)
    top = _inches(shape.top)
    if re.fullmatch(r"\d+(?:/\d+)?", text) and abs(top - theme["slide"]["footer_top_in"]) <= 0.3:
        return True
    if text == theme["footer"].get("copyright", "") and _is_footer_area(shape, theme):
        return True
    return False


def _has_footer_classification(text_shapes: list, theme: dict, expected_classification: str) -> bool:
    for shape in text_shapes:
        if expected_classification in _shape_text(shape) and _is_footer_area(shape, theme):
            return True
    return False


def _is_footer_area(shape, theme: dict) -> bool:
    top = _inches(shape.top)
    bottom = _inches(shape.top + shape.height)
    footer_top = theme["slide"]["footer_top_in"]
    return abs(top - footer_top) <= 0.35 or (footer_top <= bottom <= theme["slide"]["height_in"] + 0.05)


def _has_visible_text(shape) -> bool:
    return bool(_shape_text(shape))


def _shape_text(shape) -> str:
    return getattr(shape, "text", "").strip()


def _shape_runs(shape) -> list:
    runs = []
    if not getattr(shape, "has_text_frame", False):
        return runs
    for paragraph in shape.text_frame.paragraphs:
        runs.extend(paragraph.runs)
    return runs


def _iter_slide_paragraphs(slide):
    for shape_index, shape in enumerate(slide.shapes, start=1):
        if getattr(shape, "has_text_frame", False):
            for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs, start=1):
                yield shape, paragraph, f"文本框 {shape_index} 段落 {paragraph_index} "
        if getattr(shape, "has_table", False):
            for row_index, row in enumerate(shape.table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    for paragraph_index, paragraph in enumerate(cell.text_frame.paragraphs, start=1):
                        yield shape, paragraph, (
                            f"表格 {shape_index} 单元格 {row_index},{column_index} 段落 {paragraph_index} "
                        )


def _iter_slide_runs(slide):
    for shape, paragraph, context in _iter_slide_paragraphs(slide):
        for run in paragraph.runs:
            if run.text.strip():
                yield shape, run, context


def _paragraph_has_native_bullet(paragraph) -> bool:
    xml = paragraph._p.xml
    return any(marker in xml for marker in ("<a:buChar", "<a:buAutoNum", "<a:buBlip")) and "<a:buNone" not in xml


def _font_color(font) -> str | None:
    try:
        rgb = getattr(font.color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    return f"#{rgb}" if rgb is not None else None


def _xml_rgb_colors(xml: str) -> set[str]:
    return {value.upper() for value in re.findall(r"<a:srgbClr\b[^>]*\bval=\"([0-9A-Fa-f]{6})\"", xml)}


def _structure_items(prs, theme: dict, expected_classification: str) -> list[PptxLintItem]:
    agenda_count: int | None = None
    section_count = 0
    section_number_min_size = theme["layouts"]["section"]["index"]["font_size_pt"] - 1
    for slide in prs.slides:
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        texts = [_shape_text(shape) for shape in text_shapes if _shape_text(shape)]
        if "目录" in texts:
            agenda_items = [
                text
                for shape, text in zip(text_shapes, [_shape_text(shape) for shape in text_shapes])
                if text
                and text != "目录"
                and not re.fullmatch(r"\d{1,2}", text)
                and not _is_footer_shape(shape, theme, expected_classification)
            ]
            agenda_count = len(agenda_items)
        if any(re.fullmatch(r"\d{2}", _shape_text(shape)) and _max_font_size(shape) >= section_number_min_size for shape in text_shapes):
            section_count += 1
    if agenda_count is not None and agenda_count != section_count:
        return [
            _item(
                "HW-I01",
                "Info",
                None,
                f"agenda 条目数 {agenda_count} 与 section 页数 {section_count} 不一致。",
                "确认目录项与章节分隔页是否一一对应;若是刻意省略可忽略。",
            )
        ]
    return []


def _bbox(shape) -> tuple[float, float, float, float]:
    left = _inches(shape.left)
    top = _inches(shape.top)
    right = _inches(shape.left + shape.width)
    bottom = _inches(shape.top + shape.height)
    return left, top, right, bottom


def _boxes_overlap_or_too_close(first: tuple[float, float, float, float], second: tuple[float, float, float, float], min_gap: float) -> bool:
    left_a, top_a, right_a, bottom_a = first
    left_b, top_b, right_b, bottom_b = second
    horizontal_overlap = left_a < right_b and left_b < right_a
    vertical_overlap = top_a < bottom_b and top_b < bottom_a
    if horizontal_overlap and vertical_overlap:
        return True
    horizontal_gap = max(left_b - right_a, left_a - right_b, 0)
    vertical_gap = max(top_b - bottom_a, top_a - bottom_b, 0)
    if vertical_overlap and horizontal_gap < min_gap:
        return True
    if horizontal_overlap and vertical_gap < min_gap:
        return True
    return False


def _shape_background(shape, slide, theme: dict) -> str | None:
    try:
        rgb = getattr(shape.fill.fore_color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is not None:
        return f"#{rgb}"
    if _is_threshold_shape(shape):
        return theme["colors"]["background"]
    underlying = _underlying_background(shape, slide)
    if underlying is not None:
        return underlying
    try:
        fill_type = shape.fill.type
    except (AttributeError, TypeError, ValueError):
        fill_type = None
    if "BACKGROUND" in str(fill_type):
        return None if _overlaps_picture(shape, slide) else theme["colors"]["background"]
    if fill_type is not None:
        return None
    return None if _overlaps_picture(shape, slide) else theme["colors"]["background"]


def _underlying_background(shape, slide) -> str | None:
    target_box = _bbox(shape)
    candidate: str | None = None
    for other in slide.shapes:
        if other is shape:
            break
        if not _box_contains(_bbox(other), target_box):
            continue
        fill = _shape_fill_color(other)
        if fill is not None:
            candidate = fill
    return candidate


def _overlaps_picture(shape, slide) -> bool:
    target = _bbox(shape)
    for other in slide.shapes:
        if other is shape or "PICTURE" not in str(getattr(other, "shape_type", "")):
            continue
        if _boxes_intersect(target, _bbox(other)):
            return True
    return False


def _boxes_intersect(first: tuple[float, float, float, float], second: tuple[float, float, float, float]) -> bool:
    return first[0] < second[2] and second[0] < first[2] and first[1] < second[3] and second[1] < first[3]


def _run_color(run, default_color: str) -> str:
    try:
        rgb = getattr(run.font.color, "rgb", None)
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return default_color
    return f"#{rgb}"


def _contrast_ratio(foreground: str, background: str) -> float:
    foreground_luminance = _relative_luminance(foreground)
    background_luminance = _relative_luminance(background)
    lighter = max(foreground_luminance, background_luminance)
    darker = min(foreground_luminance, background_luminance)
    return (lighter + 0.05) / (darker + 0.05)


def _relative_luminance(hex_color: str) -> float:
    red, green, blue = _hex_to_rgb(hex_color)
    channels = []
    for value in (red, green, blue):
        channel = value / 255
        channels.append(channel / 12.92 if channel <= 0.03928 else ((channel + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    value = hex_color.strip().lstrip("#").upper()
    if len(value) != 6:
        return 0, 0, 0
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def _inches(value) -> float:
    return int(value) / EMU_PER_INCH


def _item(code: str, level: str, slide: int | None, message: str, suggestion: str) -> PptxLintItem:
    return PptxLintItem(code=code, level=level, slide=slide, message=message, suggestion=suggestion)
