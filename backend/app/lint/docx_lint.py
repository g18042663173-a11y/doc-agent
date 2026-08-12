from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
import zipfile
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.ir.word_ir import WORD_CLASSIFICATIONS
from app.rendering.theme import load_theme


@dataclass(frozen=True)
class DocxLintItem:
    code: str
    level: str
    message: str
    suggestion: str


@dataclass(frozen=True)
class DocxLintReport:
    items: list[DocxLintItem]

    @property
    def summary(self) -> dict[str, Any]:
        errors = sum(1 for item in self.items if item.level == "Error")
        warnings = sum(1 for item in self.items if item.level == "Warning")
        infos = sum(1 for item in self.items if item.level == "Info")
        return {"errors": errors, "warnings": warnings, "infos": infos, "pass": errors == 0}

    def to_dict(self) -> dict[str, Any]:
        return {
            "summary": self.summary,
            "items": [
                {
                    "code": item.code,
                    "level": item.level,
                    "message": item.message,
                    "suggestion": item.suggestion,
                }
                for item in self.items
            ],
        }


def check_docx(path: Path, *, classification: str | None = None, theme_name: str = "hw_v1") -> DocxLintReport:
    items: list[DocxLintItem] = []
    try:
        document = Document(str(path))
        package_xml = _read_word_xml(path)
    except Exception as exc:
        return DocxLintReport(
            [
                DocxLintItem(
                    code="E001",
                    level="Error",
                    message=f"DOCX 无法打开: {exc}",
                    suggestion="确认文件是有效 .docx,并重新导出后复检。",
                )
            ]
        )

    paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_count = len(document.tables)
    if not paragraph_texts and table_count == 0:
        items.append(
            DocxLintItem(
                code="E006",
                level="Error",
                message="DOCX 没有可读正文段落或表格。",
                suggestion="确认渲染输入至少包含一个内容块。",
            )
        )
        return DocxLintReport(items)

    if classification:
        footer_text = "\n".join(
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
        )
        if classification not in footer_text:
            items.append(
                DocxLintItem(
                    code="E002",
                    level="Error",
                    message="DOCX 页脚缺少期望密级文案。",
                    suggestion="使用 WordIR meta.classification 重新渲染页脚。",
                )
            )

    theme = load_theme(theme_name)
    items.extend(_theme_items(document, package_xml, theme))
    items.extend(_font_items(document, theme))
    items.extend(_table_items(document, theme))
    items.extend(_document_control_items(document, theme, classification))
    items.extend(_placeholder_items(document))
    items.extend(_quote_note_items(document, theme))
    items.extend(_code_block_items(document, theme))
    items.extend(_footer_format_items(document, package_xml, classification))
    return DocxLintReport(items)


def _read_word_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        return "\n".join(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if name.startswith("word/")
            and name.endswith(".xml")
            and not name.startswith("word/theme/")
        )


def _theme_items(document: Document, package_xml: str, theme: dict) -> list[DocxLintItem]:
    items: list[DocxLintItem] = []
    xml_upper = package_xml.upper()
    hw_red = theme["colors"]["hw_red"].lstrip("#").upper()
    if "4F81BD" in xml_upper:
        items.append(_item("E004", "旧版 Word 蓝色 4F81BD 仍存在。", "标题与强调色应使用 hw_theme.json 的华为红。"))
    requires_accent = any(_is_heading(paragraph) for paragraph in document.paragraphs) or bool(document.tables)
    if requires_accent and hw_red not in xml_upper:
        items.append(_item("E004", f"文档未使用主题红 {theme['colors']['hw_red']}。", "从 hw_theme.json 读取颜色 token 后重新渲染。"))
    return _dedupe(items)


def _font_items(document: Document, theme: dict) -> list[DocxLintItem]:
    whitelist = set(theme["fonts"]["whitelist"]) | set(theme["fonts"].get("code", []))
    allowed_sizes = {float(value) for value in theme["font_sizes_pt"].values() if isinstance(value, (int, float))}
    for paragraph in _iter_all_paragraphs(document):
        if not paragraph.text.strip():
            continue
        visible_runs = [run for run in _iter_paragraph_runs(paragraph) if run.text.strip()]
        for run in visible_runs or [None]:
            font_names = _effective_font_names(paragraph, run)
            if not font_names:
                return [_item("HW-E02", "文本未显式设置可验证字体。", "使用主题字体白名单: 微软雅黑 / Arial。")]
            for font_name in font_names:
                if font_name not in whitelist:
                    return [_item("HW-E02", f"字体不在白名单: {font_name}", "使用主题字体白名单: 微软雅黑 / Arial。")]
            size = _effective_font_size(paragraph, run)
            if size is None or float(size) not in allowed_sizes:
                shown = "未设置" if size is None else f"{float(size):g}pt"
                return [_item("E004", f"字号 {shown} 不在主题字号体系。", "使用 hw_theme.json 的 14/12/11/10/9/8pt 体系。")]
    return []


def _table_items(document: Document, theme: dict) -> list[DocxLintItem]:
    header_fill = theme["colors"]["hw_red"].lstrip("#").upper()
    border = theme["colors"]["border"].lstrip("#").upper()
    for table in document.tables:
        if _is_image_placeholder_table(table) or _document_control_caption(table) is not None:
            continue
        if not table.rows:
            continue
        header_xml = "\n".join(cell._tc.xml for cell in table.rows[0].cells)
        table_xml = table._tbl.xml
        if f'w:fill="{header_fill}"' not in header_xml:
            return [_item("E004", "表格表头未使用主题红底。", "表头使用 #C7000B 红底和白字。")]
        if "FFFFFF" not in header_xml:
            return [_item("E004", "表格表头未使用白色文字。", "表头文字使用白色并加粗。")]
        if border not in table_xml:
            return [_item("E004", f"表格边框未使用主题边框色 {theme['colors']['border']}。", "表格边框使用主题灰色。")]
    return []


def _document_control_items(document: Document, theme: dict, classification: str | None) -> list[DocxLintItem]:
    expected_marker = document.core_properties.category == "HW_DOCUMENT_CONTROL"
    tables = {
        caption: table
        for table in document.tables
        if (caption := _document_control_caption(table)) is not None
    }
    expected = {"HW_DOCUMENT_CONTROL_INFO", "HW_DOCUMENT_CONTROL_APPROVAL"}
    if not tables and not expected_marker:
        return []
    if set(tables) != expected:
        return [
            _item(
                "E004",
                "文档控制信息表不完整。",
                "在正文前保留产品/版本表和拟制、审核、批准表，且使用完整控制信息标记。",
            )
        ]
    info = tables["HW_DOCUMENT_CONTROL_INFO"]
    approval = tables["HW_DOCUMENT_CONTROL_APPROVAL"]
    if len(info.rows) != 2 or len(info.columns) != 4:
        return [_item("E004", "产品/版本控制表尺寸错误。", "使用固定 2 行 x 4 列控制信息表。")]
    if len(approval.rows) != 4 or len(approval.columns) != 3:
        return [_item("E004", "拟制审核批准表尺寸错误。", "使用固定表头加拟制、审核、批准三行。")]
    info_labels = [[info.cell(row, col).text.strip() for col in (0, 2)] for row in range(2)]
    if info_labels != [["产品名称", "文档名称"], ["密级", "版本号"]]:
        return [_item("E004", "产品/版本控制表缺少必需标签。", "保留产品名称、文档名称、密级、版本号四项。")]
    actual_classification = info.cell(1, 1).text.strip()
    if actual_classification not in WORD_CLASSIFICATIONS:
        return [_item("E004", "文档头密级值不合法。", "使用 WordIR 约定的公开、内部公开、秘密、机密、绝密或华为密级文案。")]
    if classification is not None and actual_classification != classification:
        return [_item("E004", "文档头密级与复检期望不一致。", "使用与 meta.classification 相同的密级文案。")]
    if [approval.cell(0, col).text.strip() for col in range(3)] != ["角色", "姓名", "日期"]:
        return [_item("E004", "拟制审核批准表缺少角色、姓名、日期表头。", "使用固定三列表头。")]
    if [approval.cell(row, 0).text.strip() for row in range(1, 4)] != ["拟制", "审核", "批准"]:
        return [_item("E004", "拟制审核批准表缺少固定角色行。", "即使姓名或日期为空，也保留拟制、审核、批准三行。")]
    fill = theme["colors"][theme["word_document_control"]["label_fill_color_key"]].lstrip("#").upper()
    border = theme["colors"][theme["word_document_control"]["border_color_key"]].lstrip("#").upper()
    if fill not in info._tbl.xml or fill not in approval._tbl.xml or border not in info._tbl.xml or border not in approval._tbl.xml:
        return [_item("E004", "文档控制信息表未使用主题底纹或边框。", "使用 word_document_control 的主题 token 重新渲染。")]
    return []


def _placeholder_items(document: Document) -> list[DocxLintItem]:
    body_placeholder = any("[图片占位]" in paragraph.text for paragraph in document.paragraphs)
    table_placeholder = any("图片占位" in cell.text for table in document.tables for row in table.rows for cell in row.cells)
    if body_placeholder or (_contains_text(document, "图片占位") and not table_placeholder):
        return [_item("E004", "图片占位未渲染为带边框占位框。", "使用可编辑表格/文本框绘制边框占位区,题注独立成段。")]
    return []


def _quote_note_items(document: Document, theme: dict) -> list[DocxLintItem]:
    items: list[DocxLintItem] = []
    note_fill = theme["colors"]["table_stripe"].lstrip("#").upper()
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name
        xml = paragraph._p.xml
        if style_name == "IR Quote" and "<w:left" not in xml:
            items.append(_item("E004", "quote 段落缺少左竖线。", "为 quote 段落添加左边框。"))
        if style_name == "IR Note" and f'w:fill="{note_fill}"' not in xml:
            items.append(_item("E004", "note 段落缺少主题浅底。", "为 note 段落添加主题浅底提示框。"))
    return _dedupe(items)


def _code_block_items(document: Document, theme: dict) -> list[DocxLintItem]:
    token = theme["word_code_block"]
    code_fonts = set(theme["fonts"]["code"])
    fill = theme["colors"][token["fill_color_key"]].lstrip("#").upper()
    border = theme["colors"][token["border_color_key"]].lstrip("#").upper()
    expected_size = float(token["font_size_pt"])
    for paragraph in document.paragraphs:
        if paragraph.style.name != "IR Code":
            continue
        visible_runs = [run for run in paragraph.runs if run.text]
        if any(not _effective_font_names(paragraph, run) or any(name not in code_fonts for name in _effective_font_names(paragraph, run)) for run in visible_runs):
            return [_item("E004", "code_block 未使用主题等宽字体。", "使用 hw_theme.json 的 word_code_block 与 fonts.code token。")]
        if any(_effective_font_size(paragraph, run) != expected_size for run in visible_runs):
            return [_item("E004", "code_block 字号不符合主题 token。", "使用 word_code_block.font_size_pt。")]
        xml = paragraph._p.xml
        if f'w:fill="{fill}"' not in xml or border not in xml:
            return [_item("E004", "code_block 缺少主题底纹或边框。", "使用 word_code_block 的 fill_color_key 与 border_color_key。")]
        if 'xml:space="preserve"' not in xml:
            return [_item("E004", "code_block 未保留 XML 空格语义。", "代码文本节点设置 xml:space=preserve，保留行首缩进。")]
    return []


def _footer_format_items(document: Document, package_xml: str, classification: str | None) -> list[DocxLintItem]:
    if not classification:
        return []
    footer_text = "\n".join(paragraph.text for section in document.sections for paragraph in section.footer.paragraphs)
    if classification in footer_text and "PAGE" not in package_xml:
        return [_item("HW-E01", "页脚缺少页码字段。", "页脚应包含密级文案和 PAGE 页码字段。")]
    return []


def _iter_all_paragraphs(document: Document):
    yield from document.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    # Paragraphs inside drawing text boxes (w:txbxContent) are not exposed by
    # document.paragraphs; lint them too so boxed content cannot bypass rules.
    for container in document.element.body.iter(qn("w:txbxContent")):
        for element in container.findall(qn("w:p")):
            yield Paragraph(element, document)


def _iter_paragraph_runs(paragraph):
    """Direct runs plus runs nested in w:hyperlink (python-docx omits them)."""
    yield from paragraph.runs
    for hyperlink in paragraph._p.findall(qn("w:hyperlink")):
        for element in hyperlink.findall(qn("w:r")):
            yield Run(element, paragraph)


def _effective_font_names(paragraph, run) -> list[str]:
    """latin (w:ascii) + east-asian (w:eastAsia) typefaces, run then style."""
    names = _rfonts_typefaces(run._r.rPr) if run is not None and run._r.rPr is not None else []
    if not names:
        style_rpr = paragraph.style._element.rPr
        names = _rfonts_typefaces(style_rpr) if style_rpr is not None else []
    return names


def _rfonts_typefaces(rpr) -> list[str]:
    rfonts = rpr.rFonts
    if rfonts is None:
        return []
    names: list[str] = []
    for attribute in ("ascii", "eastAsia"):
        value = rfonts.get(qn(f"w:{attribute}"))
        if value:
            names.append(value)
    return names


def _effective_font_size(paragraph, run) -> float | None:
    if run is not None and run.font.size is not None:
        return run.font.size.pt
    if paragraph.style.font.size is not None:
        return paragraph.style.font.size.pt
    return None


def _is_heading(paragraph) -> bool:
    return paragraph.text.strip() and paragraph.style.name.startswith("Heading")


def _contains_text(document: Document, needle: str) -> bool:
    return any(needle in paragraph.text for paragraph in _iter_all_paragraphs(document))


def _is_image_placeholder_table(table) -> bool:
    cells = [cell for row in table.rows for cell in row.cells]
    return len(cells) == 1 and "图片占位" in cells[0].text


def _document_control_caption(table) -> str | None:
    caption = table._tbl.tblPr.find(qn("w:tblCaption"))
    if caption is None:
        return None
    value = caption.get(qn("w:val"))
    return value if value in {"HW_DOCUMENT_CONTROL_INFO", "HW_DOCUMENT_CONTROL_APPROVAL"} else None


def _dedupe(items: list[DocxLintItem]) -> list[DocxLintItem]:
    seen: set[str] = set()
    deduped: list[DocxLintItem] = []
    for item in items:
        if item.code in seen:
            continue
        seen.add(item.code)
        deduped.append(item)
    return deduped


def _item(code: str, message: str, suggestion: str) -> DocxLintItem:
    return DocxLintItem(code=code, level="Error", message=message, suggestion=suggestion)


def write_docx_reports(report: DocxLintReport, output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "report.json"
    md_path = output_dir / "report.md"
    payload = report.to_dict()
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_lines = [
        "# DOCX 复检报告",
        "",
        f"- Errors: {payload['summary']['errors']}",
        f"- Warnings: {payload['summary']['warnings']}",
        f"- Infos: {payload['summary']['infos']}",
        f"- Pass: {payload['summary']['pass']}",
        "",
    ]
    for item in report.items:
        md_lines.append(f"## {item.code} ({item.level})")
        md_lines.append(f"- Message: {item.message}")
        md_lines.append(f"- Suggestion: {item.suggestion}")
        md_lines.append("")
    md_path.write_text("\n".join(md_lines), encoding="utf-8")
    return json_path, md_path
