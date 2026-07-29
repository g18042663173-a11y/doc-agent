from __future__ import annotations

from copy import deepcopy
from functools import lru_cache
from pathlib import Path
import sys

from PIL import ImageFont
from pptx.util import Pt


EMU_PER_INCH = 914400
PX_PER_INCH = 96
FONT_FILE_ALIASES = {
    "arial": ("arial",),
    "aptos": ("aptos",),
    "calibri": ("calibri",),
    "cambria": ("cambria",),
    "microsoftyahei": ("msyh", "microsoftyahei"),
    "微软雅黑": ("msyh",),
    "simhei": ("simhei",),
    "黑体": ("simhei",),
    "simsun": ("simsun",),
    "宋体": ("simsun",),
    "timesnewroman": ("times", "timesnewroman"),
}


def replace_text_preserving_style(
    shape,
    text: str,
    *,
    role: str,
    fallback_font_name: str | None = None,
) -> bool:
    if not getattr(shape, "has_text_frame", False):
        return False
    frame = shape.text_frame
    source_paragraph = frame.paragraphs[0] if frame.paragraphs else None
    source_ppr = (
        deepcopy(source_paragraph._p.pPr)
        if source_paragraph is not None and source_paragraph._p.pPr is not None
        else None
    )
    source_run = next((run for paragraph in frame.paragraphs for run in paragraph.runs), None)
    source_rpr = deepcopy(source_run._r.rPr) if source_run is not None and source_run._r.rPr is not None else None
    source_size = (
        source_run.font.size.pt
        if source_run is not None and source_run.font.size is not None
        else (28 if role == "title" else 16)
    )
    font_name = source_run.font.name if source_run is not None else None
    rendered_font_name = font_name if font_is_available(font_name) else fallback_font_name

    lines = text.splitlines() or [""]
    frame.clear()
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        if source_ppr is not None:
            existing = paragraph._p.pPr
            if existing is not None:
                paragraph._p.remove(existing)
            paragraph._p.insert(0, deepcopy(source_ppr))
        for run in paragraph.runs:
            if source_rpr is not None:
                existing_rpr = run._r.rPr
                if existing_rpr is not None:
                    run._r.remove(existing_rpr)
                run._r.insert(0, deepcopy(source_rpr))

    fitted_size = _fitted_font_size(shape, text, font_name, float(source_size), role)
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(fitted_size)
            if rendered_font_name:
                run.font.name = rendered_font_name
    return _text_fits(shape, text, font_name, fitted_size)


def font_is_available(font_name: str | None) -> bool:
    if not font_name or font_name.startswith("+"):
        return False
    return _font_path(font_name) is not None


def _fitted_font_size(shape, text: str, font_name: str | None, start_size: float, role: str) -> float:
    minimum = 18.0 if role == "title" else 10.5
    size = max(start_size, minimum)
    while size > minimum and not _text_fits(shape, text, font_name, size):
        size = max(minimum, size - 0.5)
    return size


def _text_fits(shape, text: str, font_name: str | None, size_pt: float) -> bool:
    frame = shape.text_frame
    width_px = max(
        1,
        (shape.width - frame.margin_left - frame.margin_right) / EMU_PER_INCH * PX_PER_INCH,
    )
    height_px = max(
        1,
        (shape.height - frame.margin_top - frame.margin_bottom) / EMU_PER_INCH * PX_PER_INCH,
    )
    font = _load_font(font_name, max(1, round(size_pt * PX_PER_INCH / 72)))
    line_height = max(font.getbbox("国Ag")[3] - font.getbbox("国Ag")[1], 1) * 1.2
    line_count = sum(_wrapped_line_count(line, font, width_px) for line in (text.splitlines() or [""]))
    return line_count * line_height <= height_px + 1


def _wrapped_line_count(text: str, font, width_px: float) -> int:
    if not text:
        return 1
    lines = 1
    current = ""
    for character in text:
        candidate = current + character
        if current and font.getlength(candidate) > width_px:
            lines += 1
            current = character
        else:
            current = candidate
    return lines


@lru_cache(maxsize=64)
def _load_font(font_name: str | None, size_px: int):
    path = _font_path(font_name)
    if path is not None:
        try:
            return ImageFont.truetype(str(path), size_px)
        except OSError:
            pass
    for fallback in _fallback_fonts():
        try:
            return ImageFont.truetype(str(fallback), size_px)
        except OSError:
            continue
    return ImageFont.load_default(size=max(10, size_px))


@lru_cache(maxsize=32)
def _font_path(font_name: str | None) -> Path | None:
    if not font_name:
        return None
    folded = font_name.casefold().replace(" ", "")
    aliases = FONT_FILE_ALIASES.get(folded, (folded,))
    for path in _font_files():
        stem = path.stem.casefold().replace(" ", "")
        if any(alias in stem for alias in aliases):
            return path
    return None


@lru_cache(maxsize=1)
def _font_files() -> tuple[Path, ...]:
    roots: list[Path] = []
    if sys.platform == "win32":
        roots.append(Path("C:/Windows/Fonts"))
    elif sys.platform == "darwin":
        roots.extend([Path("/System/Library/Fonts"), Path("/Library/Fonts")])
    else:
        roots.extend([Path("/usr/share/fonts"), Path("/usr/local/share/fonts")])
    return tuple(
        path
        for root in roots
        if root.is_dir()
        for path in root.rglob("*")
        if path.suffix.lower() in {".ttf", ".ttc", ".otf"}
    )


def _fallback_fonts() -> tuple[Path, ...]:
    preferred = ("msyh", "simhei", "arial", "noto", "dejavu")
    files = _font_files()
    return tuple(path for token in preferred for path in files if token in path.stem.casefold())
