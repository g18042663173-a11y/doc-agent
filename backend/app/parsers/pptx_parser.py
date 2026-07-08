from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation

from app.ir.document_ir import DocumentIR


def parse_pptx(path: Path) -> DocumentIR:
    presentation = Presentation(str(path))
    slide_summaries: list[dict] = []
    warnings: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        shape_warnings = _slide_warnings(slide)
        if shape_warnings:
            warnings.extend(f"slide {index}: {warning}" for warning in shape_warnings)
        slide_summaries.append(
            {
                "index": index,
                "layout_name": slide.slide_layout.name,
                "title": _slide_title(slide),
                "bodies": _slide_bodies(slide),
                "tables": _slide_tables(slide),
                "notes": _slide_notes(slide),
                "shape_warnings": shape_warnings,
            }
        )

    return DocumentIR.model_validate(
        {
            "ir_type": "document",
            "ir_version": "1.0",
            "source": {
                "filename": path.name,
                "format": "pptx",
                "size_kb": round(path.stat().st_size / 1024, 2),
                "parsed_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            },
            "stats": {
                "headings": sum(1 for slide in slide_summaries if slide.get("title")),
                "paragraphs": sum(len(slide.get("bodies", [])) for slide in slide_summaries),
                "tables": sum(len(slide.get("tables", [])) for slide in slide_summaries),
                "images": _image_count(presentation),
            },
            "warnings": warnings,
            "content": {"slides": slide_summaries},
        }
    )


def _slide_title(slide) -> str | None:
    title_shape = slide.shapes.title
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        text = title_shape.text.strip()
        return text or None
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                return text
    return None


def _slide_bodies(slide) -> list[str]:
    title_shape = slide.shapes.title
    bodies: list[str] = []
    for shape in slide.shapes:
        if shape is title_shape:
            continue
        if getattr(shape, "has_table", False):
            continue
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                bodies.extend(part.strip() for part in text.splitlines() if part.strip())
        elif getattr(shape, "has_chart", False):
            bodies.append(f"[chart] {shape.chart.chart_type}")
    return bodies


def _slide_tables(slide) -> list[dict]:
    tables: list[dict] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_table", False):
            continue
        table = shape.table
        if not table.rows or not table.columns:
            continue
        header = [_cell_text(table.cell(0, col)) for col in range(len(table.columns))]
        rows: list[list[str]] = []
        for row in range(1, len(table.rows)):
            rows.append([_cell_text(table.cell(row, col)) for col in range(len(table.columns))])
        tables.append({"header": header, "rows": rows})
    return tables


def _cell_text(cell) -> str:
    return "\n".join(paragraph.text.strip() for paragraph in cell.text_frame.paragraphs if paragraph.text.strip())


def _slide_notes(slide) -> str | None:
    try:
        text = slide.notes_slide.notes_text_frame.text.strip()
    except (AttributeError, ValueError):
        return None
    return text or None


def _slide_warnings(slide) -> list[str]:
    xml = slide.element.xml
    warnings: list[str] = []
    if "<p:transition" in xml:
        warnings.append("transition unsupported")
    if "<p:timing" in xml:
        warnings.append("animation timing unsupported")
    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) is not None and str(shape.shape_type) == "GROUP":
            warnings.append("group shape unsupported")
    return warnings


def _image_count(presentation) -> int:
    count = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if "PICTURE" in str(getattr(shape, "shape_type", "")):
                count += 1
    return count
