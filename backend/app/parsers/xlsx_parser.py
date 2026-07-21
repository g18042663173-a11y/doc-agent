from __future__ import annotations

from datetime import date, datetime
from dataclasses import dataclass
import multiprocessing
from pathlib import Path
import posixpath
import time
from typing import Any
from xml.etree import ElementTree
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from app.ir.document_ir import DocumentIR
from app.parsers.errors import ParseFailure, parser_error_boundary
from app.parsers.source_metadata import deterministic_parsed_at
from app.parsers.text_limits import TextLimiter


MAX_PREVIEW_ROWS = 20
MAX_PREVIEW_COLS = 15
MAX_COLUMN_STATS_ROWS = 1000
MAX_FORMULA_SCAN_ROWS = 1000
MAX_PREVIEW_SOURCE_ROWS = 2000
MAX_XLSX_FILE_BYTES = 100 * 1024 * 1024
MAX_PACKAGE_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MIN_RATIO_GUARD_BYTES = 5 * 1024 * 1024
MAX_PARSE_SECONDS = 5.0
HARD_GUARD_FILE_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 100
MAX_TOTAL_SCAN_CELLS = 200_000


@dataclass(frozen=True)
class WorksheetMetadata:
    merge_ranges: list[tuple[int, int, int, int]]
    hidden_rows: set[int]
    hidden_cols: set[int]


@parser_error_boundary
def parse_xlsx(path: Path) -> DocumentIR:
    if path.stat().st_size >= HARD_GUARD_FILE_BYTES:
        return _parse_xlsx_with_hard_deadline(path)
    return _parse_xlsx_impl(path)


def _parse_xlsx_impl(path: Path) -> DocumentIR:
    started = time.monotonic()
    warnings = _preflight_warnings(path)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    metadata_by_sheet = _worksheet_metadata_by_sheet(path)
    package_warnings, image_count = _package_warnings(path)
    warnings.extend(package_warnings)
    limiter = TextLimiter(warnings)
    sheets: list[dict] = []
    scanned_cells = 0

    try:
        sheet_names = value_workbook.sheetnames
        if len(sheet_names) > MAX_SHEETS:
            warnings.append(f"W103: xlsx sheet processing limited to first {MAX_SHEETS} of {len(sheet_names)} sheets")
        for sheet_name in sheet_names[:MAX_SHEETS]:
            sheet = value_workbook[sheet_name]
            formula_sheet = formula_workbook[sheet_name]
            nrows, ncols = _sheet_dimensions(sheet)
            metadata = metadata_by_sheet.get(sheet_name, WorksheetMetadata([], set(), set()))
            visible_cols = _visible_columns(ncols, metadata.hidden_cols)
            truncated = nrows > MAX_PREVIEW_ROWS or ncols > MAX_PREVIEW_COLS
            if truncated:
                warnings.append(f"W103: xlsx preview truncated for sheet {sheet_name} to 20 rows x 15 columns")

            _append_hidden_warning(warnings, sheet_name, metadata)
            preview_rows, preview_source_rows, preview_limited = _preview_rows(
                sheet,
                nrows,
                visible_cols,
                metadata.hidden_rows,
                limiter,
                sheet_name,
            )
            if preview_limited:
                warnings.append(
                    f"W103: xlsx preview source scan limited for sheet {sheet_name} to first {MAX_PREVIEW_SOURCE_ROWS} rows after hidden-row filtering"
                )
            _fill_merged_preview(
                preview_rows,
                preview_source_rows,
                visible_cols,
                metadata.merge_ranges,
                metadata.hidden_rows,
                metadata.hidden_cols,
            )
            if nrows > MAX_COLUMN_STATS_ROWS:
                warnings.append(
                    f"W103: xlsx column statistics sampled from first {MAX_COLUMN_STATS_ROWS} rows for sheet {sheet_name}; hidden rows/columns skipped"
                )
            if nrows > MAX_FORMULA_SCAN_ROWS:
                warnings.append(
                    f"W103: xlsx formula count sampled from first {MAX_FORMULA_SCAN_ROWS} rows for sheet {sheet_name}; count is a lower bound"
                )
            header_guess = preview_rows[0] if preview_rows else []
            optional_cells = (
                min(nrows, MAX_COLUMN_STATS_ROWS) * len(visible_cols)
                + min(nrows, MAX_FORMULA_SCAN_ROWS) * len(visible_cols)
            )
            optional_allowed = (
                time.monotonic() - started <= MAX_PARSE_SECONDS
                and scanned_cells + optional_cells <= MAX_TOTAL_SCAN_CELLS
            )
            if optional_allowed:
                formula_count, type_warnings = _scan_formula_and_cell_types(
                    formula_sheet,
                    nrows,
                    visible_cols,
                    metadata.hidden_rows,
                )
                warnings.extend(f"xlsx {warning} in sheet {sheet_name}" for warning in type_warnings)
                col_stats = _column_stats(
                    sheet,
                    nrows,
                    visible_cols,
                    metadata.hidden_rows,
                    header_guess,
                    limiter,
                    sheet_name,
                )
                scanned_cells += optional_cells
            else:
                formula_count = 0
                col_stats = []
                reason = (
                    f"{MAX_PARSE_SECONDS:.1f}s time"
                    if time.monotonic() - started > MAX_PARSE_SECONDS
                    else f"{MAX_TOTAL_SCAN_CELLS} cell"
                )
                warnings.append(
                    f"W103: xlsx resource budget reached ({reason} budget); optional formula/column scans skipped "
                    f"for sheet {sheet_name}; preview processed through {len(preview_source_rows)} visible source rows"
                )
            sheets.append(
                {
                    "name": sheet_name,
                    "nrows": nrows,
                    "ncols": ncols,
                    "header_guess": header_guess,
                    "preview_rows": preview_rows,
                    "col_stats": col_stats,
                    "formula_count": formula_count,
                    "merged_count": len(metadata.merge_ranges),
                    "truncated": truncated,
                }
            )
    finally:
        value_workbook.close()
        formula_workbook.close()

    return DocumentIR.model_validate(
        {
            "ir_type": "document",
            "ir_version": "1.2",
            "source": {
                "filename": path.name,
                "format": "xlsx",
                "size_kb": round(path.stat().st_size / 1024, 2),
                "parsed_at": deterministic_parsed_at(path),
            },
            "stats": {
                "headings": 0,
                "paragraphs": 0,
                "tables": len(sheets),
                "images": image_count,
            },
            "warnings": warnings,
            "content": {"sheets": sheets},
        }
    )


def _parse_xlsx_with_hard_deadline(path: Path) -> DocumentIR:
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_xlsx_worker, args=(str(path), send), daemon=True)
    process.start()
    send.close()
    timeout = max(float(MAX_PARSE_SECONDS), 0.0)
    try:
        if not receive.poll(timeout):
            process.terminate()
            process.join()
            raise _hard_deadline_failure(path, timeout)
        try:
            status, payload = receive.recv()
        except EOFError as exc:
            raise ParseFailure(
                code="E001",
                loc="source.resource",
                message="xlsx isolated parser exited without a result",
                suggestion="确认工作簿未损坏且未触发操作系统资源限制后重试。",
            ) from exc
    finally:
        receive.close()
        process.join(timeout=1.0)
        if process.is_alive():
            process.terminate()
            process.join()
    if status == "ok":
        return DocumentIR.model_validate(payload)
    if status == "parse_failure":
        raise ParseFailure(**payload)
    raise ParseFailure(
        code="E001",
        loc="source.resource",
        message=f"xlsx isolated parser failed: {payload}",
        suggestion="确认文件存在、格式正确且未损坏后重试。",
    )


def _xlsx_worker(path_value: str, send) -> None:
    try:
        result = _parse_xlsx_impl(Path(path_value))
        send.send(("ok", result.model_dump(mode="json")))
    except ParseFailure as exc:
        send.send(
            (
                "parse_failure",
                {"code": exc.code, "loc": exc.loc, "message": exc.message, "suggestion": exc.suggestion},
            )
        )
    except Exception as exc:
        send.send(("error", f"{type(exc).__name__}: {exc}"))
    finally:
        send.close()


def _hard_deadline_failure(path: Path, timeout: float) -> ParseFailure:
    return ParseFailure(
        code="E001",
        loc="source.resource",
        message=(
            f"xlsx hard resource deadline reached after {timeout:.3f}s; "
            "no worksheet content was committed"
        ),
        suggestion="拆分工作簿、减少工作表中的数据量或在资源更充足的环境中重试。",
    )


def _visible_columns(ncols: int, hidden_cols: set[int]) -> list[int]:
    return [column for column in range(1, ncols + 1) if column not in hidden_cols][:MAX_PREVIEW_COLS]


def _append_hidden_warning(warnings: list[str], sheet_name: str, metadata: WorksheetMetadata) -> None:
    hidden_rows = len(metadata.hidden_rows)
    hidden_cols = len(metadata.hidden_cols)
    if hidden_rows or hidden_cols:
        warnings.append(
            f"xlsx hidden rows/columns detected in sheet {sheet_name}: {hidden_rows} rows, {hidden_cols} columns; skipped from preview and summaries"
        )


def _preview_rows(
    sheet,
    nrows: int,
    visible_cols: list[int],
    hidden_rows: set[int],
    limiter: TextLimiter,
    sheet_name: str,
) -> tuple[list[list[str]], list[int], bool]:
    rows: list[list[str]] = []
    source_rows: list[int] = []
    if not visible_cols:
        return rows, source_rows, False

    max_source_rows = min(nrows, MAX_PREVIEW_SOURCE_ROWS)
    for row_number, row in enumerate(
        sheet.iter_rows(
        min_row=1,
        max_row=max_source_rows,
        min_col=1,
        max_col=max(visible_cols),
        values_only=True,
        ),
        start=1,
    ):
        if row_number in hidden_rows:
            continue
        rows.append(
            [
                _stringify(
                    row[column - 1],
                    limiter=limiter,
                    loc=f"xlsx sheet {sheet_name} row {row_number} column {column}",
                )
                for column in visible_cols
            ]
        )
        source_rows.append(row_number)
        if len(rows) >= MAX_PREVIEW_ROWS:
            break
    return rows, source_rows, nrows > max_source_rows and len(rows) < MAX_PREVIEW_ROWS


def _sheet_dimensions(sheet) -> tuple[int, int]:
    if sheet.max_row and sheet.max_column:
        return int(sheet.max_row), int(sheet.max_column)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(sheet.calculate_dimension(force=True))
    except ValueError:
        return 0, 0
    if min_row == max_row == 1 and min_col == max_col == 1 and sheet["A1"].value is None:
        return 0, 0
    return max_row, max_col


def _fill_merged_preview(
    preview_rows: list[list[str]],
    preview_source_rows: list[int],
    visible_cols: list[int],
    merge_ranges: list[tuple[int, int, int, int]],
    hidden_rows: set[int],
    hidden_cols: set[int],
) -> None:
    row_indexes = {source_row: index for index, source_row in enumerate(preview_source_rows)}
    col_indexes = {source_col: index for index, source_col in enumerate(visible_cols)}
    for min_col, min_row, max_col, max_row in merge_ranges:
        if min_row in hidden_rows or min_col in hidden_cols:
            continue
        source_row_index = row_indexes.get(min_row)
        source_col_index = col_indexes.get(min_col)
        if source_row_index is None or source_col_index is None:
            continue
        source_value = preview_rows[source_row_index][source_col_index]
        for source_row in range(min_row, max_row + 1):
            row_index = row_indexes.get(source_row)
            if row_index is None:
                continue
            for source_col in range(min_col, max_col + 1):
                col_index = col_indexes.get(source_col)
                if col_index is not None and not preview_rows[row_index][col_index]:
                    preview_rows[row_index][col_index] = source_value


def _column_stats(
    sheet,
    nrows: int,
    visible_cols: list[int],
    hidden_rows: set[int],
    header_guess: list[str],
    limiter: TextLimiter,
    sheet_name: str,
) -> list[dict]:
    values_by_col: list[list[Any]] = [[] for _ in visible_cols]
    samples_by_col: list[list[str]] = [[] for _ in visible_cols]
    scanned_rows = 0
    if not visible_cols:
        return []
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(nrows, MAX_COLUMN_STATS_ROWS),
            min_col=1,
            max_col=max(visible_cols),
            values_only=True,
        ),
        start=1,
    ):
        if row_number in hidden_rows:
            continue
        scanned_rows += 1
        for index, column in enumerate(visible_cols):
            value = row[column - 1]
            if value is None:
                continue
            values_by_col[index].append(value)
            if len(samples_by_col[index]) < 3:
                samples_by_col[index].append(
                    _stringify(
                        value,
                        limiter=limiter,
                        loc=f"xlsx sheet {sheet_name} column {column} sample row {row_number}",
                    )
                )

    stats: list[dict] = []
    for index, col in enumerate(visible_cols):
        values = values_by_col[index]
        stats.append(
            {
                "name": header_guess[index] if index < len(header_guess) and header_guess[index] else f"Column {col}",
                "type_guess": _type_guess(values),
                "non_empty_ratio": round(len(values) / scanned_rows, 4) if scanned_rows else 0,
                "samples": samples_by_col[index],
            }
        )
    return stats


def _type_guess(values: list[Any]) -> str:
    if not values:
        return "empty"
    non_header = values[1:] if len(values) > 1 else values
    if non_header and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in non_header):
        return "number"
    if non_header and all(isinstance(value, (datetime, date)) for value in non_header):
        return "date"
    if all(isinstance(value, str) for value in values):
        return "string"
    return "mixed"


def _scan_formula_and_cell_types(
    sheet,
    nrows: int,
    visible_cols: list[int],
    hidden_rows: set[int],
) -> tuple[int, list[str]]:
    if not visible_cols:
        return 0, []
    formula_count = 0
    error_count = 0
    boolean_count = 0
    date_count = 0
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(nrows, MAX_FORMULA_SCAN_ROWS),
            min_col=1,
            max_col=max(visible_cols),
        ),
        start=1,
    ):
        if row_number in hidden_rows:
            continue
        for column in visible_cols:
            cell = row[column - 1]
            if cell.data_type == "f":
                formula_count += 1
            elif cell.data_type == "e":
                error_count += 1
            elif cell.data_type == "b":
                boolean_count += 1
            if cell.is_date:
                date_count += 1
    warnings: list[str] = []
    if error_count:
        warnings.append(f"error cells detected: {error_count}; preserved as text")
    if boolean_count:
        warnings.append(f"boolean cells detected: {boolean_count}; preserved as text")
    if date_count:
        warnings.append(f"date cells detected: {date_count}; preserved as ISO text")
    return formula_count, warnings


def _worksheet_metadata_by_sheet(path: Path) -> dict[str, WorksheetMetadata]:
    with zipfile.ZipFile(path) as package:
        workbook_xml = package.read("xl/workbook.xml")
        rels_xml = package.read("xl/_rels/workbook.xml.rels")
        names = set(package.namelist())

        workbook_root = ElementTree.fromstring(workbook_xml)
        rels_root = ElementTree.fromstring(rels_xml)
        relationship_targets = {
            relationship.attrib["Id"]: relationship.attrib["Target"].lstrip("/")
            for relationship in rels_root
            if relationship.attrib.get("Id") and relationship.attrib.get("Target")
        }

        result: dict[str, WorksheetMetadata] = {}
        workbook_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for sheet in workbook_root.findall(f".//{workbook_ns}sheet"):
            sheet_name = sheet.attrib.get("name")
            rel_id = sheet.attrib.get(f"{rel_ns}id")
            target = relationship_targets.get(rel_id or "")
            if not sheet_name or not target:
                continue
            target_name = target if target.startswith("xl/") else f"xl/{target}"
            if target_name not in names:
                continue
            with package.open(target_name) as sheet_stream:
                result[sheet_name] = _parse_worksheet_metadata_stream(sheet_stream)
    return result


def _parse_worksheet_metadata_stream(sheet_stream) -> WorksheetMetadata:
    spreadsheet_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    ranges: list[tuple[int, int, int, int]] = []
    hidden_rows: set[int] = set()
    hidden_cols: set[int] = set()
    for _, element in ElementTree.iterparse(sheet_stream, events=("end",)):
        if element.tag == f"{spreadsheet_ns}mergeCell":
            ref = element.attrib.get("ref")
            if ref:
                ranges.append(range_boundaries(ref))
        elif element.tag == f"{spreadsheet_ns}row":
            row_number = element.attrib.get("r")
            if _is_hidden(element.attrib.get("hidden")) and row_number and row_number.isdigit():
                hidden_rows.add(int(row_number))
        elif element.tag == f"{spreadsheet_ns}col" and _is_hidden(element.attrib.get("hidden")):
            try:
                start = int(element.attrib["min"])
                end = int(element.attrib["max"])
            except (KeyError, ValueError):
                element.clear()
                continue
            hidden_cols.update(range(start, end + 1))
        element.clear()
    return WorksheetMetadata(ranges, hidden_rows, hidden_cols)


def _is_hidden(value: str | None) -> bool:
    return value in {"1", "true", "True"}


def _package_warnings(path: Path) -> tuple[list[str], int]:
    warnings: list[str] = []
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        sizes = {name: package.getinfo(name).file_size for name in names}
        image_warnings, referenced_images = _xlsx_image_warnings(package, names, sizes)
    chart_count = sum(1 for name in names if name.startswith("xl/charts/"))
    pivot_count = sum(1 for name in names if name.startswith("xl/pivotTables/"))
    vba_count = sum(1 for name in names if name == "xl/vbaProject.bin")
    image_parts = {name for name in names if name.startswith("xl/media/")}
    image_count = len(image_parts)
    ole_parts = [name for name in names if name.startswith("xl/embeddings/")]
    if chart_count:
        warnings.append(f"xlsx chart unsupported: {chart_count}")
    if pivot_count:
        warnings.append(f"xlsx pivot table unsupported: {pivot_count}")
    if vba_count:
        warnings.append("xlsx VBA unsupported: 1")
    if image_count:
        warnings.extend(image_warnings)
        for part in sorted(image_parts - referenced_images):
            warnings.append(
                f"xlsx embedded image unsupported: sheet=unknown, anchor=unknown, width=unknown, height=unknown, "
                f"part={part}, bytes={sizes[part]}; binary not extracted"
            )
    if ole_parts:
        ole_size = sum(sizes[name] for name in ole_parts)
        warnings.append(
            f"xlsx OLE object unsupported: {len(ole_parts)} at parts {', '.join(ole_parts)}; "
            f"binary not extracted, total_bytes={ole_size}"
        )
    return warnings, image_count


def _xlsx_image_warnings(
    package: zipfile.ZipFile,
    names: set[str],
    sizes: dict[str, int],
) -> tuple[list[str], set[str]]:
    warnings: list[str] = []
    referenced: set[str] = set()
    for sheet_name, sheet_part in _worksheet_parts(package, names).items():
        rels_part = _rels_part(sheet_part)
        if rels_part not in names:
            continue
        for relationship in ElementTree.fromstring(package.read(rels_part)):
            if not relationship.attrib.get("Type", "").endswith("/drawing"):
                continue
            drawing_part = _resolve_part(sheet_part, relationship.attrib.get("Target", ""))
            if drawing_part not in names:
                continue
            drawing_rels_part = _rels_part(drawing_part)
            drawing_targets = _relationship_targets(package, drawing_rels_part, drawing_part, names)
            drawing_root = ElementTree.fromstring(package.read(drawing_part))
            for anchor in list(drawing_root):
                image_rel = anchor.find(
                    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
                )
                if image_rel is None:
                    continue
                rel_id = image_rel.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
                image_part = drawing_targets.get(rel_id or "")
                if not image_part or image_part not in names:
                    continue
                referenced.add(image_part)
                anchor_label = _drawing_anchor_label(anchor)
                width, height = _drawing_dimensions(anchor)
                warnings.append(
                    f"xlsx embedded image unsupported: sheet={sheet_name}, anchor={anchor_label}, "
                    f"width={width}, height={height}, part={image_part}, bytes={sizes[image_part]}; "
                    "binary not extracted"
                )
    return warnings, referenced


def _worksheet_parts(package: zipfile.ZipFile, names: set[str]) -> dict[str, str]:
    workbook_root = ElementTree.fromstring(package.read("xl/workbook.xml"))
    relationship_targets = _relationship_targets(
        package,
        "xl/_rels/workbook.xml.rels",
        "xl/workbook.xml",
        names,
    )
    sheet_tag = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    return {
        sheet.attrib["name"]: relationship_targets[sheet.attrib[rel_id]]
        for sheet in workbook_root.findall(f".//{sheet_tag}")
        if sheet.attrib.get("name") and sheet.attrib.get(rel_id) in relationship_targets
    }


def _relationship_targets(
    package: zipfile.ZipFile,
    rels_part: str,
    source_part: str,
    names: set[str],
) -> dict[str, str]:
    if rels_part not in names:
        return {}
    return {
        relationship.attrib["Id"]: _resolve_part(source_part, relationship.attrib["Target"])
        for relationship in ElementTree.fromstring(package.read(rels_part))
        if relationship.attrib.get("Id") and relationship.attrib.get("Target")
    }


def _resolve_part(source_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _rels_part(source_part: str) -> str:
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _drawing_anchor_label(anchor) -> str:
    marker = anchor.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}from"
    )
    if marker is None:
        return "absolute"
    column = marker.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}col"
    )
    row = marker.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row"
    )
    if column is None or row is None or column.text is None or row.text is None:
        return "unknown"
    return f"{get_column_letter(int(column.text) + 1)}{int(row.text) + 1}"


def _drawing_dimensions(anchor) -> tuple[str, str]:
    ext = anchor.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}ext"
    )
    if ext is None:
        ext = anchor.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ext")
    try:
        width = int(ext.attrib["cx"]) / 914400
        height = int(ext.attrib["cy"]) / 914400
    except (AttributeError, KeyError, TypeError, ValueError):
        return "unknown", "unknown"
    return f"{width:.2f}in", f"{height:.2f}in"


def _preflight_warnings(path: Path) -> list[str]:
    file_size = path.stat().st_size
    if file_size > MAX_XLSX_FILE_BYTES:
        raise ParseFailure(
            code="E001",
            loc="source.resource",
            message=f"xlsx package exceeds {MAX_XLSX_FILE_BYTES} byte safety limit: {file_size}",
            suggestion="拆分工作簿或移除大体积嵌入对象后重试。",
        )
    with zipfile.ZipFile(path) as package:
        infos = package.infolist()
        uncompressed = sum(info.file_size for info in infos)
        if uncompressed > MAX_PACKAGE_UNCOMPRESSED_BYTES:
            raise ParseFailure(
                code="E001",
                loc="source.resource",
                message=(
                    f"xlsx uncompressed package exceeds {MAX_PACKAGE_UNCOMPRESSED_BYTES} byte safety limit: "
                    f"{uncompressed}"
                ),
                suggestion="确认文件不是压缩炸弹，并拆分超大工作簿后重试。",
            )
        suspicious = [
            info.filename
            for info in infos
            if info.file_size >= MIN_RATIO_GUARD_BYTES
            and info.compress_size > 0
            and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO
        ]
        if suspicious:
            raise ParseFailure(
                code="E001",
                loc="source.resource",
                message=f"xlsx suspicious compression ratio in parts: {', '.join(suspicious[:5])}",
                suggestion="确认文件来源可信并重新导出为正常 XLSX 后重试。",
            )
    if file_size >= 10 * 1024 * 1024:
        return [
            f"W103: xlsx resource guard active for large package: compressed_bytes={file_size}, "
            f"uncompressed_bytes={uncompressed}"
        ]
    return []


def _stringify(
    value: Any,
    *,
    limiter: TextLimiter | None = None,
    loc: str = "xlsx cell",
) -> str:
    if value is None:
        text = ""
    elif isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, bool):
        text = "true" if value else "false"
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    return limiter.limit(text, loc=loc) if limiter is not None else text
