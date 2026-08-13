from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Iterator
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ir.document_ir import DocumentIR
from app.parsers.errors import parser_error_boundary
from app.parsers.office_preflight import preflight_source_office
from app.parsers.source_metadata import deterministic_parsed_at
from app.parsers.text_limits import TextLimiter


HEADING_STYLE_RE = re.compile(r"^(Heading|标题)\s*([1-6])$")
CODE_STYLE_NAMES = {"ir code", "code", "source code", "preformatted"}
MONOSPACE_FONT_NAMES = {"consolas", "courier new", "courier", "menlo", "monaco"}


@parser_error_boundary
def parse_docx(path: Path) -> DocumentIR:
    warnings = preflight_source_office(path)
    document = Document(str(path))
    numbering = _numbering_num_fmt(path)
    blocks: list[dict] = []
    outline: list[dict] = []
    warnings.extend(_unsupported_warnings(path))
    warnings.extend(_image_warnings(document))
    limiter = TextLimiter(warnings)
    current_list: dict | None = None
    table_index = 0

    def flush_list() -> None:
        nonlocal current_list
        if current_list is not None:
            blocks.append(current_list)
            current_list = None

    for item_index, item in enumerate(_iter_block_items(document), start=1):
        if isinstance(item, Paragraph):
            if _is_code_paragraph(item):
                flush_list()
                code = item.text
                if code.strip():
                    blocks.append({"type": "code_block", "code": code})
                continue
            text = limiter.limit(item.text.strip(), loc=f"docx body paragraph {item_index}")
            if not text:
                continue
            heading_level = _heading_level(item)
            if heading_level is not None:
                flush_list()
                if heading_level > 4:
                    warnings.append(f"heading level {heading_level} capped to 4: {text}")
                    heading_level = 4
                blocks.append({"type": "heading", "level": heading_level, "text": text})
                outline.append({"level": heading_level, "text": text})
                continue

            list_kind = _list_kind(item, numbering)
            if list_kind is not None:
                level = _list_level(item)
                block_type = "numbered_list" if list_kind == "numbered" else "bullet_list"
                if current_list is None or current_list["type"] != block_type:
                    flush_list()
                    current_list = {"type": block_type, "items": []}
                current_list["items"].append({"text": text, "level": level})
                continue

            flush_list()
            blocks.append({"type": "paragraph", "text": text})
        elif isinstance(item, Table):
            flush_list()
            table_index += 1
            table_block, truncated, nested_tables = _parse_table(item, limiter=limiter, table_index=table_index)
            if table_block is not None:
                blocks.append(table_block)
            if truncated:
                warnings.append("W103: docx table preview truncated to 20 rows")
            if nested_tables:
                warnings.append(f"docx nested tables unsupported: {nested_tables}; nested content skipped")

    flush_list()

    return DocumentIR.model_validate(
        {
            "ir_type": "document",
            "ir_version": "1.2",
            "source": {
                "filename": path.name,
                "format": "docx",
                "size_kb": round(path.stat().st_size / 1024, 2),
                "parsed_at": deterministic_parsed_at(path),
            },
            "stats": {
                "headings": len(outline),
                "paragraphs": sum(1 for block in blocks if block["type"] == "paragraph"),
                "tables": sum(1 for block in blocks if block["type"] == "table"),
                "images": len(document.inline_shapes),
            },
            "warnings": warnings,
            "content": {"blocks": blocks, "outline": outline},
        }
    )


def _iter_block_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    match = HEADING_STYLE_RE.match(paragraph.style.name)
    if match is not None:
        return int(match.group(2))
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        return None
    value = outline.get(qn("w:val"))
    if value is None:
        return None
    try:
        # Some writers emit out-of-spec negative outline levels (e.g. -1);
        # clamp to the valid heading range instead of crashing the parse.
        return max(1, int(value) + 1)
    except ValueError:
        return None


def _is_code_paragraph(paragraph: Paragraph) -> bool:
    if paragraph.style.name.strip().casefold() in CODE_STYLE_NAMES:
        return True
    return any(
        run.font.name is not None and run.font.name.strip().casefold() in MONOSPACE_FONT_NAMES
        for run in paragraph.runs
    )


def _list_kind(paragraph: Paragraph, numbering: dict[str, str] | None = None) -> str | None:
    style_name = paragraph.style.name
    if "List Bullet" in style_name or "项目符号" in style_name:
        return "bullet"
    if "List Number" in style_name or "编号" in style_name:
        return "numbered"
    num_pr = _num_pr(paragraph)
    if num_pr is not None:
        num_id_el = num_pr.find(qn("w:numId"))
        if num_id_el is not None:
            num_id = num_id_el.get(qn("w:val"))
            if num_id is not None and num_id not in ("0",):
                num_fmt = (numbering or {}).get(num_id, "")
                return "bullet" if num_fmt in _BULLET_NUM_FMTS else "numbered"
    return None


def _num_pr(paragraph: Paragraph):
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    return ppr.find(qn("w:numPr"))


def _list_level(paragraph: Paragraph) -> int:
    num_pr = _num_pr(paragraph)
    if num_pr is not None:
        ilvl = num_pr.find(qn("w:ilvl"))
        if ilvl is not None:
            try:
                return int(ilvl.get(qn("w:val"))) + 1
            except (ValueError, TypeError):
                pass
    style_name = paragraph.style.name
    return 2 if re.search(r"(^|\D)2($|\D)", style_name) else 1


_BULLET_NUM_FMTS = {"bullet", "square", "circle", "chicago", "hollowSquare", "hollowsquare", "diamond", "disc"}


def _numbering_num_fmt(path: Path) -> dict[str, str]:
    """Map numId -> numFmt for the document's first list level, so real Word
    lists (List Paragraph style + w:numPr) are classified bullet vs numbered."""
    num_fmts: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as package:
            if "word/numbering.xml" not in package.namelist():
                return num_fmts
            root = ElementTree.fromstring(package.read("word/numbering.xml"))
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError):
        return num_fmts
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    abstract_fmts: dict[str, str] = {}
    for abstract_num in root.findall(f"{{{ns}}}abstractNum"):
        abstract_id = abstract_num.get(f"{{{ns}}}abstractNumId")
        for lvl in abstract_num.findall(f"{{{ns}}}lvl"):
            num_fmt = lvl.find(f"{{{ns}}}numFmt")
            if num_fmt is not None and num_fmt.get(f"{{{ns}}}val"):
                abstract_fmts[abstract_id] = num_fmt.get(f"{{{ns}}}val")
                break
    for num in root.findall(f"{{{ns}}}num"):
        num_id = num.get(f"{{{ns}}}numId")
        abstract_id_el = num.find(f"{{{ns}}}abstractNumId")
        if abstract_id_el is not None:
            abstract_id = abstract_id_el.get(f"{{{ns}}}val")
            if abstract_id in abstract_fmts:
                num_fmts[num_id] = abstract_fmts[abstract_id]
    return num_fmts


def _parse_table(
    table: Table,
    *,
    limiter: TextLimiter | None = None,
    table_index: int = 1,
) -> tuple[dict | None, bool, int]:
    if not table.rows or not table.columns:
        return None, False, 0
    column_count = min(len(table.columns), 12)
    if len(table.columns) > column_count and limiter is not None:
        limiter.warnings.append(f"W103: docx table {table_index} preview truncated to 12 columns")
    header = [
        _cell_text(
            cell,
            limiter=limiter,
            loc=f"docx table {table_index} header column {index + 1}",
        )
        or f"Column {index + 1}"
        for index, cell in enumerate(table.rows[0].cells[:column_count])
    ]
    rows: list[list[str]] = []
    truncated = False
    for row_index, row in enumerate(table.rows[1:], start=1):
        values = [
            _cell_text(
                cell,
                limiter=limiter,
                loc=f"docx table {table_index} row {row_index} column {column_index + 1}",
            )
            for column_index, cell in enumerate(row.cells[:column_count])
        ]
        if len(rows) < 20:
            rows.append(values)
        else:
            truncated = True
    return {"type": "table", "header": header, "rows": rows}, truncated, _nested_table_count(table)


def _cell_text(cell, *, limiter: TextLimiter | None = None, loc: str = "docx table cell") -> str:
    value = "\n".join(paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip())
    return limiter.limit(value, loc=loc) if limiter is not None else value


def _nested_table_count(table: Table) -> int:
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for nested_table in cell.tables:
                count += 1 + _nested_table_count(nested_table)
    return count


def _unsupported_warnings(path: Path) -> list[str]:
    warnings: list[str] = []
    patterns = [
        (r"<w:commentRangeStart\b", "comments"),
        (r"<w:ins\b", "revisions"),
        (r"<w:del\b", "revisions"),
        (r"<w:rPrChange\b", "format revisions"),
        (r"<w:pPrChange\b", "format revisions"),
        (r"<w:txbxContent\b", "text boxes"),
        (r"<wps:txbx\b", "text boxes"),
        (r"<dgm:", "SmartArt"),
    ]
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        for name in names:
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            xml = package.read(name)
            if xml[:2] in (b"\xff\xfe", b"\xfe\xff"):
                xml_text = xml.decode("utf-16", errors="replace")
            else:
                xml_text = xml.decode("utf-8", errors="replace")
            counts: dict[str, int] = {}
            for pattern, label in patterns:
                count = len(re.findall(pattern, xml_text))
                if count:
                    counts[label] = counts.get(label, 0) + count
            for label, count in counts.items():
                warnings.append(f"unsupported {label}: {count} at part {name}; content skipped")
        embedding_parts = [name for name in names if name.startswith("word/embeddings/")]
        if embedding_parts:
            total_bytes = sum(package.getinfo(name).file_size for name in embedding_parts)
            warnings.append(
                "unsupported embedded/OLE objects: "
                f"{len(embedding_parts)} at parts {', '.join(embedding_parts)}; binary skipped, total_bytes={total_bytes}"
            )
    return warnings


def _image_warnings(document: DocxDocument) -> list[str]:
    warnings: list[str] = []
    for index, shape in enumerate(document.inline_shapes, start=1):
        try:
            width_in = shape.width / 914400
            height_in = shape.height / 914400
        except Exception as exc:
            # A malformed inline shape (missing <wp:extent>, dangling rId) must
            # not kill the whole document parse.
            warnings.append(
                f"W103: docx image present #{index} with unreadable geometry ({type(exc).__name__}); geometry skipped"
            )
            continue
        warnings.append(f"docx image present #{index} width={width_in:.2f}in height={height_in:.2f}in")
    return warnings
