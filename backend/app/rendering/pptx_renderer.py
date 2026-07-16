from __future__ import annotations

from pathlib import Path

from pptx.chart.data import CategoryChartData
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.ir.deck_ir import (
    AgendaSlide,
    ArchitectureDiagramSlide,
    ArchitectureEdge,
    ArchitectureNode,
    CardsSlide,
    ChartSlide,
    ChartSpec,
    ColumnContent,
    ConclusionSlide,
    CoverSlide,
    DeckTableCell,
    DeckIR,
    ImageSlide,
    ProcessFlowSlide,
    SectionSlide,
    TableSlide,
    TimelineSlide,
    TitleBulletsSlide,
    TwoColumnSlide,
)
from app.rendering.theme import load_theme
from app.rendering.typography import constrained_stack_heights, fit_text_stack


def render_deck_ir(deck: DeckIR, output_path: Path) -> Path:
    theme = load_theme(deck.meta.theme)
    prs = Presentation()
    prs.slide_width = Inches(theme["slide"]["width_in"])
    prs.slide_height = Inches(theme["slide"]["height_in"])
    total_slides = len(deck.slides)

    for slide_number, slide_ir in enumerate(deck.slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if isinstance(slide_ir, CoverSlide):
            _render_cover(slide, slide_ir, theme)
        elif isinstance(slide_ir, AgendaSlide):
            _render_agenda(slide, slide_ir, theme)
        elif isinstance(slide_ir, SectionSlide):
            _render_section(slide, slide_ir, theme)
        elif isinstance(slide_ir, TitleBulletsSlide):
            _render_title_bullets(slide, slide_ir, theme)
        elif isinstance(slide_ir, TableSlide):
            _render_table_slide(slide, slide_ir, theme)
        elif isinstance(slide_ir, TwoColumnSlide):
            _render_two_column(slide, slide_ir, theme)
        elif isinstance(slide_ir, CardsSlide):
            _render_cards(slide, slide_ir, theme)
        elif isinstance(slide_ir, ConclusionSlide):
            _render_conclusion(slide, slide_ir, theme)
        elif isinstance(slide_ir, ChartSlide):
            _render_chart_slide(slide, slide_ir, theme)
        elif isinstance(slide_ir, ArchitectureDiagramSlide):
            _render_architecture_diagram(slide, slide_ir, theme)
        elif isinstance(slide_ir, ProcessFlowSlide):
            _render_process_flow(slide, slide_ir, theme)
        elif isinstance(slide_ir, TimelineSlide):
            _render_timeline(slide, slide_ir, theme)
        elif isinstance(slide_ir, ImageSlide):
            _render_image_placeholder(slide, slide_ir, theme)
        else:
            _render_placeholder(slide, slide_ir.layout, theme)
        _add_footer(slide, deck.meta.classification, slide_number, total_slides, theme)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path


def _render_cover(slide, slide_ir: CoverSlide, theme: dict) -> None:
    layout = theme["layouts"]["cover"]
    typography = theme["ppt_typography"]
    _red_bar(slide, theme, layout["red_bar"])
    title_size = _fit_ppt_text(
        slide_ir.title,
        typography["cover_title_candidates_pt"],
        layout["title"],
        theme,
    )
    _add_text_box(
        slide,
        slide_ir.title,
        layout["title"],
        theme,
        size=title_size,
        bold=True,
        name="HW_RENDERED_TEXT:COVER_TITLE",
    )
    if slide_ir.subtitle:
        _add_text_box(
            slide,
            slide_ir.subtitle,
            layout["subtitle"],
            theme,
            size=typography["cover_subtitle_pt"],
            name="HW_RENDERED_TEXT:COVER_SUBTITLE",
        )
    meta = "  ".join(part for part in [slide_ir.presenter, slide_ir.date] if part)
    if meta:
        _add_text_box(
            slide,
            meta,
            layout["meta"],
            theme,
            size=typography["secondary_pt"],
            color_key="secondary",
            name="HW_RENDERED_TEXT:NOTE",
        )


def _render_agenda(slide, slide_ir: AgendaSlide, theme: dict) -> None:
    _title(slide, "目录", theme)
    layout = theme["layouts"]["agenda"]
    for index, item in enumerate(slide_ir.items, start=1):
        row = (index - 1) % layout["rows_per_column"]
        col = (index - 1) // layout["rows_per_column"]
        left = layout["left_in"] + col * layout["column_gap_in"]
        top = layout["top_in"] + row * layout["row_gap_in"]
        item_box = {"left_in": left, "top_in": top, "width_in": layout["item_width_in"], "height_in": layout["height_in"]}
        _add_agenda_item(slide, index, item, item_box, layout, theme)


def _add_agenda_item(slide, index: int, item: str, box: dict, layout: dict, theme: dict) -> None:
    shape = slide.shapes.add_textbox(
        Inches(box["left_in"]), Inches(box["top_in"]), Inches(box["width_in"]), Inches(box["height_in"])
    )
    shape.name = "HW_RENDERED_TEXT:AGENDA"
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    paragraph = text_frame.paragraphs[0]
    _format_paragraph(paragraph, theme)
    number = paragraph.add_run()
    number.text = f"{index:02d}  "
    _format_run(number, theme, size=layout["font_size_pt"], bold=True, color_key="hw_red", font_role="number")
    label = paragraph.add_run()
    label.text = item
    _format_run(label, theme, size=layout["font_size_pt"], color_key="body")


def _render_section(slide, slide_ir: SectionSlide, theme: dict) -> None:
    layout = theme["layouts"]["section"]
    _add_text_box(
        slide,
        f"{slide_ir.index:02d}",
        layout["index"],
        theme,
        size=layout["index"]["font_size_pt"],
        bold=True,
        color_key="hw_red",
        font_role="number",
        name="HW_RENDERED_TEXT:SECTION_NUMBER",
    )
    _add_text_box(
        slide,
        slide_ir.title,
        layout["title"],
        theme,
        size=_fit_ppt_text(
            slide_ir.title,
            theme["ppt_typography"]["slide_title_candidates_pt"],
            layout["title"],
            theme,
        ),
        bold=True,
        name="HW_RENDERED_TEXT:SECTION_TITLE",
    )
    if slide_ir.subtitle:
        _add_text_box(
            slide,
            slide_ir.subtitle,
            layout["subtitle"],
            theme,
            size=layout["subtitle"]["font_size_pt"],
            color_key="secondary",
            name="HW_RENDERED_TEXT:SUBHEADING",
        )
    _red_bar(slide, theme, layout["red_bar"])


def _render_title_bullets(slide, slide_ir: TitleBulletsSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["title_bullets"]
    typography = theme["ppt_typography"]
    top = layout["top_in"]
    texts = [f"{'    –' if bullet.level == 2 else '•'} {bullet.text}" for bullet in slide_ir.bullets]
    available_height = typography["content_bottom_in"] - top
    fit = _fit_body_stack(texts, layout["width_in"], available_height, theme)
    heights = constrained_stack_heights(
        fit,
        height_in=available_height,
        item_gap_in=typography["body_item_gap_in"],
    )
    for text, height in zip(texts, heights):
        box = {
            "left_in": layout["left_in"],
            "top_in": top,
            "width_in": layout["width_in"],
            "height_in": height,
        }
        _add_text_box(
            slide,
            text,
            box,
            theme,
            size=fit.font_size_pt,
            name="HW_RENDERED_TEXT:BODY",
        )
        top += height + typography["body_item_gap_in"]


def _render_table_slide(slide, slide_ir: TableSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["table"]
    table_ir = slide_ir.table
    has_column_groups = bool(table_ir.column_groups)
    row_group_starts = {group.start_row: group for group in table_ir.row_groups}
    header_row_index = 1 if has_column_groups else 0
    data_start_index = header_row_index + 1
    rows = data_start_index + len(table_ir.rows) + len(table_ir.row_groups)
    cols = len(slide_ir.table.header)
    col_widths = _table_col_widths(table_ir.col_widths, cols, layout, theme)
    row_heights = _table_row_heights(table_ir, rows, header_row_index, data_start_index, layout)
    row_heights = _fit_table_row_heights(row_heights, layout, theme)
    shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(_table_left_in(col_widths, layout, theme)),
        Inches(layout["top_in"]),
        Inches(sum(col_widths)),
        Inches(sum(row_heights)),
    )
    table = shape.table
    for col, width in enumerate(col_widths):
        table.columns[col].width = Inches(width)
    for row, height in enumerate(row_heights):
        table.rows[row].height = Inches(height)

    if has_column_groups:
        _render_column_group_row(table, table_ir, theme)

    for col, value in enumerate(table_ir.header):
        cell = table.cell(0, col)
        if has_column_groups:
            cell = table.cell(header_row_index, col)
        _fill_cell(cell, theme["colors"]["hw_red"])
        _write_cell_text(cell, [value], theme, size=layout["header_font_size_pt"], bold=True, color_key="background")
        _set_cell_border(cell, theme)

    data_row_map: dict[int, int] = {}
    visual_row = data_start_index
    for row_index, row in enumerate(table_ir.rows):
        group = row_group_starts.get(row_index)
        if group is not None:
            _render_row_group(table, visual_row, group.label, cols, theme, layout)
            visual_row += 1
        data_row_map[row_index] = visual_row
        fill_color = theme["colors"]["background"] if row_index % 2 == 0 else theme["colors"]["table_stripe"]
        for col, value in enumerate(row):
            cell = table.cell(visual_row, col)
            _fill_cell(cell, _cell_fill_color(value, fill_color, theme))
            _write_cell_text(
                cell,
                _cell_lines(value),
                theme,
                size=layout["data_font_size_pt"],
                bold=table_ir.conclusion_col == col,
                color_key="body",
            )
            _set_cell_border(cell, theme)
        visual_row += 1

    _apply_table_spans(table, table_ir, header_row_index, data_row_map, theme)


def _table_col_widths(custom_widths: list[float] | None, cols: int, layout: dict, theme: dict) -> list[float]:
    reference = layout.get("decision_matrix_col_widths_in", [])
    weights = list(custom_widths) if custom_widths is not None else list(reference) if cols == len(reference) else [1.0] * cols
    slide = theme["slide"]
    available_width = min(
        layout["width_in"],
        slide["width_in"] - slide["margin_left_in"] - slide["margin_right_in"],
    )
    min_width = available_width / (2 * cols)
    widths = [0.0] * cols
    remaining = set(range(cols))
    remaining_width = available_width

    while remaining:
        remaining_weight = sum(weights[index] for index in remaining)
        too_narrow = {
            index
            for index in remaining
            if remaining_width * weights[index] / remaining_weight < min_width
        }
        if not too_narrow:
            for index in remaining:
                widths[index] = remaining_width * weights[index] / remaining_weight
            break
        for index in too_narrow:
            widths[index] = min_width
        remaining_width -= min_width * len(too_narrow)
        remaining -= too_narrow

    return widths


def _table_left_in(col_widths: list[float], layout: dict, theme: dict) -> float:
    table_width = sum(col_widths)
    centered = (theme["slide"]["width_in"] - table_width) / 2
    return max(layout["left_in"], centered)


def _table_row_heights(table_ir, rows: int, header_row_index: int, data_start_index: int, layout: dict) -> list[float]:
    row_group_starts = {group.start_row for group in table_ir.row_groups}
    heights: list[float] = []
    if header_row_index:
        heights.append(layout["group_header_row_height_in"])
    heights.append(layout["header_row_height_in"])
    for row_index, _row in enumerate(table_ir.rows):
        if row_index in row_group_starts:
            heights.append(layout["row_group_height_in"])
        heights.append(layout["data_row_height_in"])
    if len(heights) != rows:
        return [layout["data_row_height_in"] for _ in range(rows)]
    return heights


def _fit_table_row_heights(row_heights: list[float], layout: dict, theme: dict) -> list[float]:
    max_height = theme["slide"]["footer_top_in"] - theme["grid"]["min_gap_in"] - layout["top_in"]
    current_height = sum(row_heights)
    if current_height <= max_height:
        return row_heights
    scale = max_height / current_height
    return [height * scale for height in row_heights]


def _render_column_group_row(table, table_ir, theme: dict) -> None:
    layout = theme["layouts"]["table"]
    grouped_cols = {
        col
        for group in table_ir.column_groups
        for col in range(group.start_col, group.start_col + group.span)
    }
    for col in range(len(table_ir.header)):
        cell = table.cell(0, col)
        _fill_cell(cell, theme["colors"]["hw_red"] if col in grouped_cols else theme["colors"]["background"])
        _write_cell_text(cell, [""], theme, size=layout["header_font_size_pt"], bold=True, color_key="background")
        _set_cell_border(cell, theme)
    for group in table_ir.column_groups:
        cell = table.cell(0, group.start_col)
        if group.span > 1:
            cell.merge(table.cell(0, group.start_col + group.span - 1))
        _fill_cell(cell, theme["colors"]["hw_red"])
        _write_cell_text(cell, [group.label], theme, size=layout["header_font_size_pt"], bold=True, color_key="background")
        _set_cell_border(cell, theme)


def _render_row_group(table, visual_row: int, label: str, cols: int, theme: dict, layout: dict) -> None:
    cell = table.cell(visual_row, 0)
    if cols > 1:
        cell.merge(table.cell(visual_row, cols - 1))
    _fill_cell(cell, theme["colors"]["table_stripe"])
    _write_cell_text(cell, [label], theme, size=layout["data_font_size_pt"], bold=True, color_key="body")
    _set_cell_border(cell, theme)


def _apply_table_spans(table, table_ir, header_row_index: int, data_row_map: dict[int, int], theme: dict) -> None:
    for span in table_ir.cell_spans:
        start_row = header_row_index if span.area == "header" else data_row_map[span.row]
        start_row += span.row if span.area == "header" else 0
        end_row = start_row + span.rowspan - 1
        end_col = span.col + span.colspan - 1
        if end_row == start_row and end_col == span.col:
            continue
        cell = table.cell(start_row, span.col)
        cell.merge(table.cell(end_row, end_col))
        _set_cell_border(cell, theme)


def _cell_lines(value: DeckTableCell | str) -> list[str]:
    if isinstance(value, str):
        return [value]
    lines: list[str] = []
    if value.text:
        lines.append(value.text)
    lines.extend(f"• {item}" for item in value.items)
    return lines or [""]


def _cell_fill_color(value: DeckTableCell | str, default_color: str, theme: dict) -> str:
    if isinstance(value, DeckTableCell) and value.emphasis:
        return theme["colors"]["accent4"] if value.emphasis == "yellow" else theme["colors"]["accent6"]
    return default_color


def _fill_cell(cell, hex_color: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(hex_color)


def _write_cell_text(cell, lines: list[str], theme: dict, *, size: float, bold: bool, color_key: str) -> None:
    text_frame = cell.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        _format_paragraph(paragraph, theme)
        run = paragraph.add_run()
        run.text = line
        _format_run(run, theme, size=size, bold=bold, color_key=color_key)


def _set_cell_border(cell, theme: dict) -> None:
    border_width = str(int(Pt(theme["strokes"]["table_border_pt"])))
    color = theme["colors"]["border"].lstrip("#")
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        line = tc_pr.find(qn(f"a:{edge}"))
        if line is None:
            line = OxmlElement(f"a:{edge}")
            tc_pr.append(line)
        line.set("w", border_width)
        for child in list(line):
            line.remove(child)
        solid_fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color)
        solid_fill.append(srgb)
        line.append(solid_fill)


def _render_two_column(slide, slide_ir: TwoColumnSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["two_column"]
    typography = theme["ppt_typography"]
    column_fits = []
    for content in (slide_ir.left, slide_ir.right):
        texts = _column_texts(content)
        body_top = layout["top_in"] + (layout["heading_gap_in"] if content.heading else 0)
        column_fits.append(
            _fit_body_stack(
                texts,
                layout["column_width_in"],
                typography["content_bottom_in"] - body_top,
                theme,
            )
        )
    body_size = min(fit.font_size_pt for fit in column_fits)
    _render_column(
        slide,
        slide_ir.left,
        layout["left_column_left_in"],
        layout["top_in"],
        layout,
        theme,
        body_size,
    )
    _render_column(
        slide,
        slide_ir.right,
        layout["right_column_left_in"],
        layout["top_in"],
        layout,
        theme,
        body_size,
    )


def _render_column(
    slide,
    content: ColumnContent,
    left: float,
    top: float,
    layout: dict,
    theme: dict,
    body_size: float,
) -> None:
    if content.heading:
        box = {"left_in": left, "top_in": top, "width_in": layout["column_width_in"], "height_in": layout["heading_height_in"]}
        _add_text_box(
            slide,
            content.heading,
            box,
            theme,
            size=layout["heading_font_size_pt"],
            bold=True,
            color_key="hw_red",
            name="HW_RENDERED_TEXT:SUBHEADING",
        )
        top += layout["heading_gap_in"]
    texts = _column_texts(content)
    typography = theme["ppt_typography"]
    available_height = typography["content_bottom_in"] - top
    fit = _fit_body_stack(texts, layout["column_width_in"], available_height, theme, candidates=[body_size])
    heights = constrained_stack_heights(
        fit,
        height_in=available_height,
        item_gap_in=typography["body_item_gap_in"],
    )
    for text, height in zip(texts, heights):
        box = {
            "left_in": left,
            "top_in": top,
            "width_in": layout["column_width_in"],
            "height_in": height,
        }
        _add_text_box(
            slide,
            text,
            box,
            theme,
            size=body_size,
            color_key="body",
            name="HW_RENDERED_TEXT:BODY",
        )
        top += height + typography["body_item_gap_in"]


def _column_texts(content: ColumnContent) -> list[str]:
    texts = [content.text] if content.text else []
    texts.extend(f"{'    –' if bullet.level == 2 else '•'} {bullet.text}" for bullet in content.bullets)
    return texts


def _render_cards(slide, slide_ir: CardsSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["cards"]
    count = len(slide_ir.cards)
    span = int(layout["grid_span_by_count"][str(count)])
    gap = int(layout["grid_gap_columns"])
    used_columns = count * span + (count - 1) * gap
    start_column = max(0, (theme["grid"]["columns"] - used_columns) // 2)
    for index, card in enumerate(slide_ir.cards):
        left, card_width = _grid_horizontal_box(theme, start_column + index * (span + gap), span)
        card_shape = _card_box(
            slide,
            left,
            layout["top_in"],
            card_width,
            layout["height_in"],
            layout,
            theme,
            fill_color_key="background" if slide_ir.variant == "kpi" else "surface",
        )
        if slide_ir.variant == "kpi":
            _kpi_card_content(card_shape, card, index, layout, theme)
        else:
            _card_content(card_shape, card, layout, theme)


def _grid_horizontal_box(theme: dict, start_column: int, span: int) -> tuple[float, float]:
    margin = theme["slide"]["margin_in"]
    column_width = (theme["slide"]["width_in"] - margin * 2) / theme["grid"]["columns"]
    return margin + start_column * column_width, span * column_width


def _card_content(shape, card, layout: dict, theme: dict) -> None:
    shape.name = "HW_RENDERED_TEXT:CARD"
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    text_frame.margin_left = Inches(layout["padding_in"])
    text_frame.margin_right = Inches(layout["padding_in"])
    text_frame.margin_top = Inches(layout["padding_in"])
    text_frame.margin_bottom = Inches(layout["padding_in"])
    parts = [
        (card.title, layout["title_font_size_pt"], True, "title"),
        (card.desc, layout["desc_font_size_pt"], False, "body"),
    ]
    if card.tag:
        parts.append((card.tag, layout["tag_font_size_pt"], False, "secondary"))
    for part_index, (text, size, bold, color_key) in enumerate(parts):
        paragraph = text_frame.paragraphs[0] if part_index == 0 else text_frame.add_paragraph()
        _format_paragraph(paragraph, theme)
        if part_index < len(parts) - 1:
            paragraph.space_after = Pt(layout["paragraph_gap_pt"])
        run = paragraph.add_run()
        run.text = text
        _format_run(run, theme, size=size, bold=bold, color_key=color_key)


def _kpi_card_content(shape, card, index: int, layout: dict, theme: dict) -> None:
    shape.name = f"HW_RENDERED_TEXT:KPI_CARD:{index + 1}"
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for part_index, (text, size, bold, color_key, font_role) in enumerate(
        [
            (card.desc, layout["kpi_value_font_size_pt"], True, "hw_red", "number"),
            (card.title, layout["kpi_title_font_size_pt"], True, "title", "body"),
            *(([(card.tag, layout["kpi_tag_font_size_pt"], False, "secondary", "body")]) if card.tag else []),
        ]
    ):
        paragraph = text_frame.paragraphs[0] if part_index == 0 else text_frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        _format_paragraph(paragraph, theme)
        if part_index < 2:
            paragraph.space_after = Pt(layout["kpi_paragraph_gap_pt"])
        run = paragraph.add_run()
        run.text = text
        _format_run(run, theme, size=size, bold=bold, color_key=color_key, font_role=font_role)


def _card_box(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    layout: dict,
    theme: dict,
    *,
    fill_color_key: str,
):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = "HW_LAYOUT_CONTAINER:CARD"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["colors"][fill_color_key])
    shape.line.color.rgb = _rgb(theme["colors"]["border"])
    shape.line.width = Pt(theme["strokes"]["card_border_pt"])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(layout["red_bar_height_in"]))
    bar.name = "HW_DECORATION:CARD_BAR"
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(theme["colors"]["hw_red"])
    bar.line.fill.background()
    return shape


def _render_conclusion(slide, slide_ir: ConclusionSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["conclusion"]
    typography = theme["ppt_typography"]
    top = layout["top_in"]
    texts = [f"• {bullet}" for bullet in slide_ir.bullets]
    body_bottom = min(typography["content_bottom_in"], layout["cta_top_in"] - typography["body_item_gap_in"])
    available_height = body_bottom - top
    fit = _fit_body_stack(texts, layout["width_in"], available_height, theme)
    heights = constrained_stack_heights(
        fit,
        height_in=available_height,
        item_gap_in=typography["body_item_gap_in"],
    )
    for text, height in zip(texts, heights):
        box = {"left_in": layout["left_in"], "top_in": top, "width_in": layout["width_in"], "height_in": height}
        _add_text_box(
            slide,
            text,
            box,
            theme,
            size=fit.font_size_pt,
            color_key="body",
            name="HW_RENDERED_TEXT:BODY",
        )
        top += height + typography["body_item_gap_in"]
    if slide_ir.cta:
        cta_box = {
            "left_in": layout["left_in"],
            "top_in": layout["cta_top_in"],
            "width_in": layout["cta_width_in"],
            "height_in": layout["cta_height_in"],
        }
        _add_text_box(
            slide,
            slide_ir.cta,
            cta_box,
            theme,
            size=layout["cta_font_size_pt"],
            bold=True,
            color_key="hw_red",
            name="HW_RENDERED_TEXT:SUBHEADING",
        )


def _render_chart_slide(slide, slide_ir: ChartSlide, theme: dict) -> None:
    layout = theme["layouts"]["chart"]
    _chart_title(slide, slide_ir.title, layout, theme)
    chart_layout = layout["plot"]
    chart_data = _chart_data(slide_ir.chart)
    chart_shape = slide.shapes.add_chart(
        _chart_type(slide_ir.chart),
        Inches(chart_layout["left_in"]),
        Inches(chart_layout["top_in"]),
        Inches(chart_layout["width_in"]),
        Inches(chart_layout["height_in"]),
        chart_data,
    )
    chart = chart_shape.chart
    _format_chart(chart, slide_ir.chart, theme)
    _add_threshold_lines(slide, slide_ir.chart, chart_layout, chart_shape, theme)
    _add_chart_side_content(slide, slide_ir.chart, layout, theme)


def _chart_data(chart_ir: ChartSpec) -> CategoryChartData:
    chart_data = CategoryChartData()
    chart_data.categories = chart_ir.categories
    for series in chart_ir.series:
        chart_data.add_series(series.name, series.values)
    return chart_data


def _chart_type(chart_ir: ChartSpec):
    if chart_ir.kind == "bar" and chart_ir.orientation == "horizontal":
        return XL_CHART_TYPE.BAR_CLUSTERED
    if chart_ir.kind == "bar":
        return XL_CHART_TYPE.COLUMN_CLUSTERED
    if chart_ir.kind == "line":
        return XL_CHART_TYPE.LINE_MARKERS
    return XL_CHART_TYPE.PIE


def _chart_title(slide, text: str, layout: dict, theme: dict) -> None:
    box = layout.get("title", theme["layouts"]["title"])
    size = _fit_ppt_text(text, theme["ppt_typography"]["slide_title_candidates_pt"], box, theme)
    _add_text_box(slide, text, box, theme, size=size, bold=True, name="HW_RENDERED_TEXT:TITLE")
    _red_bar(slide, theme, theme["layouts"]["title_bar"])


def _format_chart(chart, chart_ir: ChartSpec, theme: dict) -> None:
    chart.has_title = False
    chart.has_legend = chart_ir.legend_position != "none"
    if chart.has_legend:
        chart.legend.position = _legend_position(chart_ir.legend_position)
        chart.legend.include_in_layout = False
        _format_chart_font(chart.legend.font, theme, size=theme["layouts"]["chart"]["legend_font_size_pt"], color_key="secondary")
    plot = chart.plots[0]
    if chart_ir.kind == "bar" and hasattr(plot, "gap_width"):
        plot.gap_width = theme["layouts"]["chart"]["bar_gap_width"]
    _format_chart_series(chart, chart_ir, theme)
    if chart_ir.show_data_labels:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END if chart_ir.kind == "bar" else XL_DATA_LABEL_POSITION.ABOVE
        _format_chart_font(labels.font, theme, size=theme["layouts"]["chart"]["data_label_font_size_pt"], color_key="secondary")
    if chart_ir.kind != "pie":
        _format_chart_axes(chart, chart_ir, theme)


def _legend_position(position: str):
    if position == "right":
        return XL_LEGEND_POSITION.RIGHT
    return XL_LEGEND_POSITION.TOP


def _format_chart_series(chart, chart_ir: ChartSpec, theme: dict) -> None:
    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]
    for index, series in enumerate(chart.series):
        source = chart_ir.series[index]
        color_key = "hw_red" if source.emphasis else accent_keys[index % len(accent_keys)]
        rgb = _rgb(theme["colors"][color_key])
        if chart_ir.kind == "line":
            series.format.line.color.rgb = rgb
            series.format.line.width = Pt(theme["layouts"]["chart"]["series_line_width_pt"])
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb
            series.format.line.color.rgb = rgb


def _format_chart_axes(chart, chart_ir: ChartSpec, theme: dict) -> None:
    axis_size = theme["layouts"]["chart"]["axis_font_size_pt"]
    for axis in (chart.category_axis, chart.value_axis):
        _format_chart_font(axis.tick_labels.font, theme, size=axis_size, color_key="secondary")
    minimum, maximum = _chart_value_range(chart_ir)
    chart.value_axis.minimum_scale = minimum
    chart.value_axis.maximum_scale = maximum
    if chart_ir.kind == "bar" and chart_ir.orientation == "horizontal":
        chart.category_axis.reverse_order = True
    chart.value_axis.has_major_gridlines = True
    gridline = chart.value_axis.major_gridlines.format.line
    gridline.color.rgb = _rgb(theme["colors"]["border"])
    gridline.dash_style = MSO_LINE_DASH_STYLE.DASH
    if chart_ir.unit:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = chart_ir.unit
        _format_chart_text_frame(chart.value_axis.axis_title.text_frame, theme, size=axis_size, color_key="secondary")


def _format_chart_font(font, theme: dict, *, size: float, color_key: str) -> None:
    font.name = theme["fonts"]["east_asia"][0]
    font.size = Pt(size)
    font.color.rgb = _rgb(theme["colors"][color_key])


def _format_chart_text_frame(text_frame, theme: dict, *, size: float, color_key: str) -> None:
    for paragraph in text_frame.paragraphs:
        _format_paragraph(paragraph, theme)
        for run in paragraph.runs:
            _format_run(run, theme, size=size, color_key=color_key)


def _add_threshold_lines(slide, chart_ir: ChartSpec, chart_layout: dict, chart_shape, theme: dict) -> None:
    if chart_ir.kind == "pie" or not chart_ir.thresholds:
        return
    minimum, maximum = _chart_value_range(chart_ir)
    if maximum <= minimum:
        return
    threshold_layout = theme["layouts"]["chart"]["threshold"]
    plot_bounds = _threshold_plot_bounds(chart_layout, threshold_layout)
    left = plot_bounds["left_in"]
    right = plot_bounds["right_in"]
    top = plot_bounds["top_in"]
    bottom = plot_bounds["bottom_in"]
    horizontal_bar = chart_ir.kind == "bar" and chart_ir.orientation == "horizontal"
    label_top = _snap_to_baseline_in(
        chart_layout["top_in"] - threshold_layout["label_height_in"] - threshold_layout["label_gap_in"],
        theme,
    )
    for index, threshold in enumerate(chart_ir.thresholds, start=1):
        ratio = max(0, min(1, (threshold.value - minimum) / (maximum - minimum)))
        if horizontal_bar:
            x = left + plot_bounds["width_in"] * ratio
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(top), Inches(x), Inches(bottom))
        else:
            y = top + plot_bounds["height_in"] * (1 - ratio)
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(left), Inches(y), Inches(right), Inches(y))
        line.name = f"HW_THRESHOLD_LINE_{index}"
        line.line.color.rgb = _rgb(theme["colors"]["hw_red"])
        line.line.width = Pt(threshold_layout["line_width_pt"])
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        _move_shape_behind(line, chart_shape)
        if horizontal_bar:
            label_left = _clamp(
                x - threshold_layout["label_width_in"] / 2,
                left,
                right - threshold_layout["label_width_in"],
            )
            label = {
                "left_in": label_left,
                "top_in": label_top,
                "width_in": threshold_layout["label_width_in"],
                "height_in": threshold_layout["label_height_in"],
            }
        else:
            fallback_left = min(
                right + threshold_layout["label_gap_in"],
                theme["slide"]["width_in"] - threshold_layout["label_width_in"] - theme["slide"]["margin_right_in"],
            )
            label = {
                "left_in": threshold_layout.get("label_left_in", fallback_left),
                "top_in": label_top,
                "width_in": threshold_layout["label_width_in"],
                "height_in": threshold_layout["label_height_in"],
            }
        label_shape = _add_text_box(
            slide,
            _threshold_label_text(threshold, chart_ir),
            label,
            theme,
            size=threshold_layout["label_font_size_pt"],
            bold=True,
            color_key="hw_red",
        )
        label_shape.name = f"HW_THRESHOLD_LABEL:{index}"
        label_margin = Inches(threshold_layout["label_margin_in"])
        label_shape.text_frame.margin_left = label_margin
        label_shape.text_frame.margin_right = label_margin
        label_shape.text_frame.margin_top = label_margin
        label_shape.text_frame.margin_bottom = label_margin


def _threshold_plot_bounds(chart_layout: dict, threshold_layout: dict) -> dict[str, float]:
    left = chart_layout["left_in"] + threshold_layout.get("plot_left_offset_in", 0)
    right = chart_layout["left_in"] + chart_layout["width_in"] - threshold_layout.get("plot_right_offset_in", 0)
    top = chart_layout["top_in"] + threshold_layout.get("plot_top_offset_in", 0)
    bottom = chart_layout["top_in"] + chart_layout["height_in"] - threshold_layout.get("plot_bottom_offset_in", 0)
    return {
        "left_in": left,
        "right_in": right,
        "top_in": top,
        "bottom_in": bottom,
        "width_in": max(right - left, 0.1),
        "height_in": max(bottom - top, 0.1),
    }


def _move_shape_behind(shape, reference_shape) -> None:
    parent = shape.element.getparent()
    parent.remove(shape.element)
    parent.insert(parent.index(reference_shape.element), shape.element)


def _threshold_label_text(threshold, chart_ir: ChartSpec) -> str:
    label = threshold.label.strip() or "阈值"
    if label.endswith("线"):
        label = label.removesuffix("线")
    if "阈值" not in label:
        label = f"{label}阈值"
    target = _threshold_target_series(chart_ir)
    if target and not label.startswith(target):
        label = f"{target}{label}"
    value = f"{threshold.value:g}{chart_ir.unit or ''}"
    return f"{label} {value}"


def _threshold_target_series(chart_ir: ChartSpec) -> str | None:
    emphasized = [series.name for series in chart_ir.series if series.emphasis]
    if len(emphasized) == 1:
        return emphasized[0]
    if len(chart_ir.series) == 1:
        return chart_ir.series[0].name
    return None


def _snap_to_baseline_in(value: float, theme: dict) -> float:
    step = theme["grid"]["baseline_pt"] / 72
    return round(value / step) * step


def _chart_value_range(chart_ir: ChartSpec) -> tuple[float, float]:
    values = [value for series in chart_ir.series for value in series.values]
    values.extend(threshold.value for threshold in chart_ir.thresholds)
    minimum = min(values, default=0)
    maximum = max(values, default=1)
    if minimum > 0:
        minimum = 0
    padding = (maximum - minimum) * 0.1 or 1
    return minimum, maximum + padding


def _add_chart_side_content(slide, chart_ir: ChartSpec, layout: dict, theme: dict) -> None:
    side = layout["side"]
    top = side["top_in"]
    if chart_ir.side_conclusion:
        box = {
            "left_in": side["left_in"],
            "top_in": top,
            "width_in": side["width_in"],
            "height_in": side["conclusion_height_in"],
        }
        _add_text_box(
            slide,
            chart_ir.side_conclusion,
            box,
            theme,
            size=side["font_size_pt"],
            bold=True,
            color_key="body",
            name="HW_RENDERED_TEXT:COMPACT",
        )
        top += side["conclusion_height_in"] + side["gap_in"]
    if chart_ir.side_table:
        rows = len(chart_ir.side_table.rows) + 1
        cols = len(chart_ir.side_table.header)
        shape = slide.shapes.add_table(
            rows,
            cols,
            Inches(side["left_in"]),
            Inches(top),
            Inches(side["width_in"]),
            Inches(side["row_height_in"] * rows),
        )
        table = shape.table
        for col, value in enumerate(chart_ir.side_table.header):
            cell = table.cell(0, col)
            _fill_cell(cell, theme["colors"]["hw_red"])
            _write_cell_text(cell, [value], theme, size=side["table_font_size_pt"], bold=True, color_key="background")
            _set_cell_border(cell, theme)
        for row_index, row in enumerate(chart_ir.side_table.rows, start=1):
            for col, value in enumerate(row):
                cell = table.cell(row_index, col)
                _fill_cell(cell, theme["colors"]["background"] if row_index % 2 else theme["colors"]["table_stripe"])
                _write_cell_text(cell, [value], theme, size=side["table_font_size_pt"], bold=False, color_key="body")
                _set_cell_border(cell, theme)


def _render_architecture_diagram(slide, slide_ir: ArchitectureDiagramSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["architecture_diagram"]
    boxes, layers = _architecture_node_boxes(slide_ir, layout)
    _render_architecture_groups(slide, slide_ir, layers, layout, theme)
    _render_architecture_edges(slide, slide_ir.edges, boxes, layers, layout, theme)
    for node in slide_ir.nodes:
        _render_architecture_node(slide, node, boxes[node.id], layout, theme)


def _render_process_flow(slide, slide_ir: ProcessFlowSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["process_flow"]
    boxes = _sequence_boxes(
        len(slide_ir.steps),
        slide_ir.orientation,
        layout,
        horizontal_height_key="horizontal_step_height_in",
        vertical_width_key="vertical_step_width_in",
    )
    for index, (first, second) in enumerate(zip(boxes, boxes[1:]), start=1):
        _render_sequence_connector(slide, first, second, slide_ir.orientation, index, layout, theme)
    for step, box in zip(slide_ir.steps, boxes):
        shape = _sequence_card(
            slide,
            box,
            name=f"HW_PROCESS_STEP:{step.id}",
            layout=layout,
            theme=theme,
            line_color_key="accent6",
        )
        _write_sequence_text(
            shape,
            heading=step.title,
            detail=step.description,
            label=None,
            layout=layout,
            theme=theme,
        )


def _render_timeline(slide, slide_ir: TimelineSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["timeline"]
    boxes = _sequence_boxes(
        len(slide_ir.milestones),
        slide_ir.orientation,
        layout,
        horizontal_height_key="horizontal_card_height_in",
        vertical_width_key="vertical_card_width_in",
    )
    marker_centers = _timeline_marker_centers(boxes, slide_ir.orientation, layout)
    _render_timeline_axis(slide, marker_centers, slide_ir.orientation, layout, theme)
    for index, (milestone, box, center) in enumerate(zip(slide_ir.milestones, boxes, marker_centers), start=1):
        color_key = layout["status_colors"][milestone.status]
        _render_timeline_marker(slide, center, index, color_key, layout, theme)
        shape = _sequence_card(
            slide,
            box,
            name=f"HW_TIMELINE_MILESTONE:{index}",
            layout=layout,
            theme=theme,
            line_color_key=color_key,
        )
        _write_sequence_text(
            shape,
            heading=milestone.title,
            detail=milestone.description,
            label=milestone.label,
            layout=layout,
            theme=theme,
            label_color_key="secondary",
        )


def _sequence_boxes(
    count: int,
    orientation: str,
    layout: dict,
    *,
    horizontal_height_key: str,
    vertical_width_key: str,
) -> list[dict[str, float]]:
    content = layout["content"]
    gap = layout["gap_in"]
    boxes: list[dict[str, float]] = []
    if orientation == "horizontal":
        width = (content["width_in"] - gap * (count - 1)) / count
        height = layout[horizontal_height_key]
        top = content["top_in"] + (content["height_in"] - height) / 2
        for index in range(count):
            boxes.append(
                {
                    "left_in": content["left_in"] + index * (width + gap),
                    "top_in": top,
                    "width_in": width,
                    "height_in": height,
                }
            )
        return boxes
    height = (content["height_in"] - gap * (count - 1)) / count
    width = layout[vertical_width_key]
    left = content["left_in"] + (content["width_in"] - width) / 2
    for index in range(count):
        boxes.append(
            {
                "left_in": left,
                "top_in": content["top_in"] + index * (height + gap),
                "width_in": width,
                "height_in": height,
            }
        )
    return boxes


def _render_sequence_connector(
    slide,
    first: dict[str, float],
    second: dict[str, float],
    orientation: str,
    index: int,
    layout: dict,
    theme: dict,
) -> None:
    if orientation == "horizontal":
        start = (first["left_in"] + first["width_in"], first["top_in"] + first["height_in"] / 2)
        end = (second["left_in"], second["top_in"] + second["height_in"] / 2)
    else:
        start = (first["left_in"] + first["width_in"] / 2, first["top_in"] + first["height_in"])
        end = (second["left_in"] + second["width_in"] / 2, second["top_in"])
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]),
        Inches(start[1]),
        Inches(end[0]),
        Inches(end[1]),
    )
    connector.name = f"HW_PROCESS_CONNECTOR:{index}"
    connector.line.color.rgb = _rgb(theme["colors"]["secondary"])
    connector.line.width = Pt(layout["connector_width_pt"])
    _set_connector_arrowheads(connector, "forward")


def _sequence_card(
    slide,
    box: dict[str, float],
    *,
    name: str,
    layout: dict,
    theme: dict,
    line_color_key: str,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(box["left_in"]),
        Inches(box["top_in"]),
        Inches(box["width_in"]),
        Inches(box["height_in"]),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["colors"]["background"])
    shape.line.color.rgb = _rgb(theme["colors"][line_color_key])
    shape.line.width = Pt(layout["border_pt"])
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    padding = Inches(layout["step_padding_in"] if "step_padding_in" in layout else 0.1)
    text_frame.margin_left = padding
    text_frame.margin_right = padding
    text_frame.margin_top = padding
    text_frame.margin_bottom = padding
    return shape


def _write_sequence_text(
    shape,
    *,
    heading: str,
    detail: str | None,
    label: str | None,
    layout: dict,
    theme: dict,
    label_color_key: str = "secondary",
) -> None:
    text_frame = shape.text_frame
    lines = []
    if label:
        lines.append((label, layout["label_font_size_pt"], True, label_color_key))
    lines.append((heading, layout["title_font_size_pt"], True, "title"))
    if detail:
        lines.append((detail, layout["description_font_size_pt"], False, "body"))
    for index, (text, size, bold, color_key) in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        _format_paragraph(paragraph, theme)
        run = paragraph.add_run()
        run.text = text
        _format_run(run, theme, size=size, bold=bold, color_key=color_key)


def _timeline_marker_centers(
    boxes: list[dict[str, float]],
    orientation: str,
    layout: dict,
) -> list[tuple[float, float]]:
    marker_gap = layout["marker_size_in"]
    if orientation == "horizontal":
        y = boxes[0]["top_in"] + boxes[0]["height_in"] + marker_gap * 1.5
        return [(box["left_in"] + box["width_in"] / 2, y) for box in boxes]
    x = boxes[0]["left_in"] - marker_gap * 1.5
    return [(x, box["top_in"] + box["height_in"] / 2) for box in boxes]


def _render_timeline_axis(
    slide,
    centers: list[tuple[float, float]],
    orientation: str,
    layout: dict,
    theme: dict,
) -> None:
    first, last = centers[0], centers[-1]
    axis = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(first[0]),
        Inches(first[1]),
        Inches(last[0]),
        Inches(last[1]),
    )
    axis.name = "HW_TIMELINE_AXIS"
    axis.line.color.rgb = _rgb(theme["colors"]["border"])
    axis.line.width = Pt(layout["axis_width_pt"])


def _render_timeline_marker(
    slide,
    center: tuple[float, float],
    index: int,
    color_key: str,
    layout: dict,
    theme: dict,
) -> None:
    size = layout["marker_size_in"]
    marker = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(center[0] - size / 2),
        Inches(center[1] - size / 2),
        Inches(size),
        Inches(size),
    )
    marker.name = f"HW_TIMELINE_MARKER:{index}"
    marker.fill.solid()
    marker.fill.fore_color.rgb = _rgb(theme["colors"][color_key])
    marker.line.color.rgb = _rgb(theme["colors"][color_key])


def _architecture_node_boxes(
    slide_ir: ArchitectureDiagramSlide,
    layout: dict,
) -> tuple[dict[str, dict[str, float]], list[tuple[str | None, str | None, list[str], dict[str, float]]]]:
    content = layout["content"]
    group_by_node = {
        node_id: group.id
        for group in slide_ir.groups
        for node_id in group.node_ids
    }
    layers: list[tuple[str | None, str | None, list[str], dict[str, float]]] = [
        (group.id, group.label, list(group.node_ids), {}) for group in slide_ir.groups
    ]
    ungrouped = [node.id for node in slide_ir.nodes if node.id not in group_by_node]
    if ungrouped:
        layers.append((None, None, ungrouped, {}))

    layer_count = max(len(layers), 1)
    available_height = content["height_in"] - layout["group_gap_in"] * (layer_count - 1)
    layer_height = available_height / layer_count
    node_by_id = {node.id: node for node in slide_ir.nodes}
    boxes: dict[str, dict[str, float]] = {}
    positioned_layers: list[tuple[str | None, str | None, list[str], dict[str, float]]] = []

    for layer_index, (group_id, label, node_ids, _bounds) in enumerate(layers):
        top = content["top_in"] + layer_index * (layer_height + layout["group_gap_in"])
        bounds = {
            "left_in": content["left_in"],
            "top_in": top,
            "width_in": content["width_in"],
            "height_in": layer_height,
        }
        positioned_layers.append((group_id, label, node_ids, bounds))
        horizontal_space = bounds["width_in"] - layout["group_padding_in"] * 2
        slot_width = horizontal_space / max(len(node_ids), 1)
        max_node_width = max(slot_width - layout["node_gap_in"], layout["node_min_width_in"])
        max_node_height = max(
            bounds["height_in"] - layout["group_label_height_in"] - layout["group_padding_in"] * 2,
            layout["node_min_height_in"],
        )
        for node_index, node_id in enumerate(node_ids):
            node = node_by_id[node_id]
            width = min(_architecture_node_width(node, slide_ir, content, layout), max_node_width)
            height = min(_architecture_node_height(node, slide_ir, content, layout), max_node_height)
            center_x = bounds["left_in"] + layout["group_padding_in"] + slot_width * (node_index + 0.5)
            node_top = (
                bounds["top_in"]
                + layout["group_label_height_in"]
                + (bounds["height_in"] - layout["group_label_height_in"] - height) / 2
            )
            boxes[node_id] = {
                "left_in": center_x - width / 2,
                "top_in": node_top,
                "width_in": width,
                "height_in": height,
            }

    for node in slide_ir.nodes:
        position = _architecture_node_position(node, slide_ir)
        if position is None:
            continue
        width = _architecture_node_width(node, slide_ir, content, layout)
        height = _architecture_node_height(node, slide_ir, content, layout)
        left = content["left_in"] + position.x * content["width_in"] - width / 2
        top = content["top_in"] + position.y * content["height_in"] - height / 2
        boxes[node.id] = {
            "left_in": _clamp(left, content["left_in"], content["left_in"] + content["width_in"] - width),
            "top_in": _clamp(top, content["top_in"], content["top_in"] + content["height_in"] - height),
            "width_in": width,
            "height_in": height,
        }
    return boxes, positioned_layers


def _architecture_node_position(node: ArchitectureNode, slide_ir: ArchitectureDiagramSlide):
    if slide_ir.manual_hints and node.id in slide_ir.manual_hints.node_positions:
        return slide_ir.manual_hints.node_positions[node.id]
    return node.position


def _architecture_node_size(node: ArchitectureNode, slide_ir: ArchitectureDiagramSlide):
    if slide_ir.manual_hints and node.id in slide_ir.manual_hints.node_sizes:
        return slide_ir.manual_hints.node_sizes[node.id]
    return node.size


def _architecture_node_width(node: ArchitectureNode, slide_ir: ArchitectureDiagramSlide, content: dict, layout: dict) -> float:
    size = _architecture_node_size(node, slide_ir)
    return size.width * content["width_in"] if size else layout["node_width_in"]


def _architecture_node_height(node: ArchitectureNode, slide_ir: ArchitectureDiagramSlide, content: dict, layout: dict) -> float:
    size = _architecture_node_size(node, slide_ir)
    return size.height * content["height_in"] if size else layout["node_height_in"]


def _render_architecture_groups(slide, slide_ir, layers, layout: dict, theme: dict) -> None:
    for group_id, label, _node_ids, bounds in layers:
        if group_id is None:
            continue
        frame = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(bounds["left_in"]),
            Inches(bounds["top_in"]),
            Inches(bounds["width_in"]),
            Inches(bounds["height_in"]),
        )
        frame.name = f"HW_ARCH_GROUP:{group_id}"
        frame.fill.background()
        frame.line.color.rgb = _rgb(theme["colors"]["border"])
        frame.line.width = Pt(layout["group_border_pt"])
        frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        text_frame = frame.text_frame
        text_frame.clear()
        _fit_text_frame(text_frame)
        text_frame.vertical_anchor = MSO_ANCHOR.TOP
        text_frame.margin_left = Inches(layout["group_label_margin_in"])
        text_frame.margin_top = Inches(layout["group_label_margin_in"])
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        _format_paragraph(paragraph, theme)
        run = paragraph.add_run()
        run.text = label or ""
        _format_run(run, theme, size=layout["group_label_font_size_pt"], bold=True, color_key="secondary")


def _render_architecture_edges(
    slide,
    edges: list[ArchitectureEdge],
    boxes: dict[str, dict[str, float]],
    layers: list[tuple[str | None, str | None, list[str], dict[str, float]]],
    layout: dict,
    theme: dict,
) -> None:
    occupied_boxes = list(boxes.values())
    for index, edge in enumerate(edges, start=1):
        route = _architecture_edge_route(
            edge,
            index,
            boxes,
            layers,
            layout,
        )
        segments = list(zip(route, route[1:]))
        for segment_index, (start, end) in enumerate(segments, start=1):
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(start[0]),
                Inches(start[1]),
                Inches(end[0]),
                Inches(end[1]),
            )
            if segment_index == len(segments):
                connector.name = f"HW_ARCH_EDGE:{index}:{segment_index}:{edge.from_node}->{edge.to}"
            else:
                connector.name = f"HW_ARCH_EDGE_SEGMENT:{index}:{segment_index}:{edge.from_node}->{edge.to}"
            connector.line.color.rgb = _rgb(theme["colors"]["secondary"])
            connector.line.width = Pt(layout["edge_width_pt"])
            if edge.style == "dashed":
                connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            _set_connector_arrowheads(
                connector,
                _architecture_segment_direction(edge.direction, segment_index, len(segments)),
            )
        if edge.label:
            label_box = _render_architecture_edge_label(
                slide,
                edge.label,
                index,
                segments,
                occupied_boxes,
                layout,
                theme,
            )
            occupied_boxes.append(label_box)


def _architecture_edge_route(
    edge: ArchitectureEdge,
    edge_index: int,
    boxes: dict[str, dict[str, float]],
    layers: list[tuple[str | None, str | None, list[str], dict[str, float]]],
    layout: dict,
) -> list[tuple[float, float]]:
    source = boxes[edge.from_node]
    target = boxes[edge.to]
    layer_by_node = {
        node_id: layer_index
        for layer_index, (_group_id, _label, node_ids, _bounds) in enumerate(layers)
        for node_id in node_ids
    }
    source_layer = layer_by_node[edge.from_node]
    target_layer = layer_by_node[edge.to]
    if source_layer == target_layer:
        points = _architecture_same_layer_route(
            source,
            target,
            edge.from_node,
            edge.to,
            edge_index,
            boxes,
            layers[source_layer],
            layout,
        )
    else:
        points = _architecture_cross_layer_route(
            source,
            target,
            source_layer,
            target_layer,
            edge_index,
            layers,
            layout,
        )
    return _simplify_architecture_route(points, layout["edge_route_alignment_tolerance_in"])


def _architecture_same_layer_route(
    source: dict[str, float],
    target: dict[str, float],
    source_id: str,
    target_id: str,
    edge_index: int,
    boxes: dict[str, dict[str, float]],
    layer: tuple[str | None, str | None, list[str], dict[str, float]],
    layout: dict,
) -> list[tuple[float, float]]:
    source_x = source["left_in"] + source["width_in"] / 2
    source_y = source["top_in"] + source["height_in"] / 2
    target_x = target["left_in"] + target["width_in"] / 2
    target_y = target["top_in"] + target["height_in"] / 2
    aligned = abs(target_y - source_y) <= layout["edge_route_alignment_tolerance_in"]
    if target_x >= source_x:
        start = (source["left_in"] + source["width_in"], source_y)
        end = (target["left_in"], target_y)
    else:
        start = (source["left_in"], source_y)
        end = (target["left_in"] + target["width_in"], target_y)
    if aligned and _architecture_horizontal_route_is_clear(
        start,
        end,
        boxes,
        {source_id, target_id},
        layout["edge_route_clearance_in"],
    ):
        return [start, end]

    group_id, _label, node_ids, bounds = layer
    layer_boxes = [boxes[node_id] for node_id in node_ids]
    top_limit = bounds["top_in"] + (
        layout["group_label_height_in"] if group_id is not None else layout["edge_route_clearance_in"]
    )
    top_limit += layout["edge_route_clearance_in"]
    top_near_nodes = min(box["top_in"] for box in layer_boxes) - layout["edge_route_clearance_in"]
    bottom_near_nodes = max(box["top_in"] + box["height_in"] for box in layer_boxes) + layout["edge_route_clearance_in"]
    bottom_limit = bounds["top_in"] + bounds["height_in"] - layout["edge_route_clearance_in"]
    use_top = edge_index % 2 == 1
    if top_near_nodes < top_limit:
        use_top = False
    if bottom_near_nodes > bottom_limit:
        use_top = True
    lane_slot = ((edge_index - 1) // 2) % layout["edge_route_lane_count"]
    lane_offset = lane_slot * layout["edge_route_lane_gap_in"]
    if use_top:
        lane_y = _clamp(top_near_nodes - lane_offset, top_limit, top_near_nodes)
        source_anchor = (source_x, source["top_in"])
        target_anchor = (target_x, target["top_in"])
    else:
        lane_y = _clamp(bottom_near_nodes + lane_offset, bottom_near_nodes, bottom_limit)
        source_anchor = (source_x, source["top_in"] + source["height_in"])
        target_anchor = (target_x, target["top_in"] + target["height_in"])
    return [source_anchor, (source_x, lane_y), (target_x, lane_y), target_anchor]


def _architecture_horizontal_route_is_clear(
    start: tuple[float, float],
    end: tuple[float, float],
    boxes: dict[str, dict[str, float]],
    endpoint_ids: set[str],
    clearance: float,
) -> bool:
    left, right = sorted((start[0], end[0]))
    y = (start[1] + end[1]) / 2
    for node_id, box in boxes.items():
        if node_id in endpoint_ids:
            continue
        box_left = box["left_in"] - clearance
        box_right = box["left_in"] + box["width_in"] + clearance
        box_top = box["top_in"] - clearance
        box_bottom = box["top_in"] + box["height_in"] + clearance
        if box_left < right and left < box_right and box_top < y < box_bottom:
            return False
    return True


def _architecture_cross_layer_route(
    source: dict[str, float],
    target: dict[str, float],
    source_layer: int,
    target_layer: int,
    edge_index: int,
    layers: list[tuple[str | None, str | None, list[str], dict[str, float]]],
    layout: dict,
) -> list[tuple[float, float]]:
    source_x = source["left_in"] + source["width_in"] / 2
    target_x = target["left_in"] + target["width_in"] / 2
    moving_down = target_layer > source_layer
    source_y = source["top_in"] + source["height_in"] if moving_down else source["top_in"]
    target_y = target["top_in"] if moving_down else target["top_in"] + target["height_in"]
    source_gap = _architecture_layer_gap_y(
        source_layer if moving_down else source_layer - 1,
        edge_index,
        layers,
        layout,
    )
    target_gap = _architecture_layer_gap_y(
        target_layer - 1 if moving_down else target_layer,
        edge_index,
        layers,
        layout,
    )
    if abs(source_layer - target_layer) == 1:
        return [(source_x, source_y), (source_x, source_gap), (target_x, target_gap), (target_x, target_y)]

    content = layout["content"]
    left_lane = content["left_in"] + layout["edge_route_frame_margin_in"]
    right_lane = content["left_in"] + content["width_in"] - layout["edge_route_frame_margin_in"]
    left_cost = abs(source_x - left_lane) + abs(target_x - left_lane)
    right_cost = abs(source_x - right_lane) + abs(target_x - right_lane)
    side_lane = right_lane if right_cost <= left_cost else left_lane
    return [
        (source_x, source_y),
        (source_x, source_gap),
        (side_lane, source_gap),
        (side_lane, target_gap),
        (target_x, target_gap),
        (target_x, target_y),
    ]


def _architecture_layer_gap_y(
    upper_layer: int,
    edge_index: int,
    layers: list[tuple[str | None, str | None, list[str], dict[str, float]]],
    layout: dict,
) -> float:
    upper_bounds = layers[upper_layer][3]
    lower_bounds = layers[upper_layer + 1][3]
    low = upper_bounds["top_in"] + upper_bounds["height_in"] + layout["edge_route_gap_margin_in"]
    high = lower_bounds["top_in"] - layout["edge_route_gap_margin_in"]
    lane_count = layout["edge_route_lane_count"]
    lane_slot = (edge_index - 1) % lane_count
    centered_slot = lane_slot - (lane_count - 1) / 2
    return _clamp(
        (low + high) / 2 + centered_slot * layout["edge_route_lane_gap_in"],
        low,
        high,
    )


def _simplify_architecture_route(
    points: list[tuple[float, float]],
    tolerance: float,
) -> list[tuple[float, float]]:
    simplified: list[tuple[float, float]] = []
    for point in points:
        if simplified and abs(point[0] - simplified[-1][0]) <= tolerance and abs(point[1] - simplified[-1][1]) <= tolerance:
            continue
        simplified.append(point)
    index = 1
    while index < len(simplified) - 1:
        previous, current, following = simplified[index - 1 : index + 2]
        same_x = abs(previous[0] - current[0]) <= tolerance and abs(current[0] - following[0]) <= tolerance
        same_y = abs(previous[1] - current[1]) <= tolerance and abs(current[1] - following[1]) <= tolerance
        if same_x or same_y:
            simplified.pop(index)
        else:
            index += 1
    return simplified


def _architecture_segment_direction(direction: str, segment_index: int, segment_count: int) -> str:
    if segment_count == 1:
        return direction
    if segment_index == 1 and direction in {"backward", "both"}:
        return "backward"
    if segment_index == segment_count and direction in {"forward", "both"}:
        return "forward"
    return "none"


def _set_connector_arrowheads(connector, direction: str) -> None:
    line = connector.element.spPr.get_or_add_ln()
    for tag in ("a:headEnd", "a:tailEnd"):
        child = line.find(qn(tag))
        if child is not None:
            line.remove(child)
    if direction in {"forward", "both"}:
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        tail.set("w", "sm")
        tail.set("len", "sm")
        line.append(tail)
    if direction in {"backward", "both"}:
        head = OxmlElement("a:headEnd")
        head.set("type", "triangle")
        head.set("w", "sm")
        head.set("len", "sm")
        line.append(head)


def _render_architecture_edge_label(
    slide,
    text: str,
    index: int,
    segments: list[tuple[tuple[float, float], tuple[float, float]]],
    occupied_boxes: list[dict[str, float]],
    layout: dict,
    theme: dict,
) -> dict[str, float]:
    content = layout["content"]
    candidates = _architecture_label_candidates(segments, content, layout)
    preferred_box = candidates[0]
    box = next(
        (
            candidate
            for candidate in candidates
            if not any(_boxes_overlap(candidate, occupied, theme["grid"]["min_gap_in"]) for occupied in occupied_boxes)
        ),
        None,
    )
    if box is None:
        box = _place_architecture_label(preferred_box, occupied_boxes, content, theme["grid"]["min_gap_in"])
    label = _add_text_box(slide, text, box, theme, size=layout["edge_label_font_size_pt"], color_key="secondary")
    label.name = f"HW_ARCH_EDGE_LABEL:{index}"
    label.fill.solid()
    label.fill.fore_color.rgb = _rgb(theme["colors"]["background"])
    label.line.fill.background()
    return box


def _architecture_label_candidates(
    segments: list[tuple[tuple[float, float], tuple[float, float]]],
    content: dict[str, float],
    layout: dict,
) -> list[dict[str, float]]:
    width = layout["edge_label_width_in"]
    height = layout["edge_label_height_in"]
    offset = layout["edge_label_offset_in"]
    ordered = sorted(
        segments,
        key=lambda segment: (
            segment[0][1] == segment[1][1],
            abs(segment[1][0] - segment[0][0]) + abs(segment[1][1] - segment[0][1]),
        ),
        reverse=True,
    )
    steps = layout["edge_label_local_steps"]
    fractions = [0.5]
    for step in range(1, steps + 1):
        delta = step / (2 * (steps + 1))
        fractions.extend((0.5 - delta, 0.5 + delta))

    candidates: list[dict[str, float]] = []
    seen: set[tuple[float, float]] = set()
    for start, end in ordered:
        horizontal = start[1] == end[1]
        for fraction in fractions:
            x = start[0] + (end[0] - start[0]) * fraction
            y = start[1] + (end[1] - start[1]) * fraction
            if horizontal:
                positions = ((x - width / 2, y - offset - height), (x - width / 2, y + offset))
            else:
                positions = ((x + offset, y - height / 2), (x - offset - width, y - height / 2))
            for left, top in positions:
                box = {
                    "left_in": _clamp(left, content["left_in"], content["left_in"] + content["width_in"] - width),
                    "top_in": _clamp(top, content["top_in"], content["top_in"] + content["height_in"] - height),
                    "width_in": width,
                    "height_in": height,
                }
                position = (round(box["left_in"], 6), round(box["top_in"], 6))
                if position not in seen:
                    seen.add(position)
                    candidates.append(box)
    return candidates


def _place_architecture_label(
    preferred: dict[str, float],
    occupied_boxes: list[dict[str, float]],
    content: dict[str, float],
    gap: float,
) -> dict[str, float]:
    step_x = preferred["width_in"] + gap
    step_y = preferred["height_in"] + gap
    max_dx = int(content["width_in"] / step_x) + 1
    max_dy = int(content["height_in"] / step_y) + 1
    offsets = [
        (dx, dy)
        for dx in range(-max_dx, max_dx + 1)
        for dy in range(-max_dy, max_dy + 1)
    ]
    offsets.sort(key=lambda offset: (abs(offset[0]) + abs(offset[1]), abs(offset[0]), abs(offset[1]), offset))
    seen: set[tuple[float, float]] = set()
    for dx, dy in offsets:
        left = _clamp(
            preferred["left_in"] + dx * step_x,
            content["left_in"],
            content["left_in"] + content["width_in"] - preferred["width_in"],
        )
        top = _clamp(
            preferred["top_in"] + dy * step_y,
            content["top_in"],
            content["top_in"] + content["height_in"] - preferred["height_in"],
        )
        position = (round(left, 6), round(top, 6))
        if position in seen:
            continue
        seen.add(position)
        candidate = {**preferred, "left_in": left, "top_in": top}
        if not any(_boxes_overlap(candidate, occupied, gap) for occupied in occupied_boxes):
            return candidate
    return preferred


def _boxes_overlap(first: dict[str, float], second: dict[str, float], gap: float = 0) -> bool:
    return not (
        first["left_in"] + first["width_in"] + gap <= second["left_in"]
        or second["left_in"] + second["width_in"] + gap <= first["left_in"]
        or first["top_in"] + first["height_in"] + gap <= second["top_in"]
        or second["top_in"] + second["height_in"] + gap <= first["top_in"]
    )


def _render_architecture_node(slide, node: ArchitectureNode, box: dict[str, float], layout: dict, theme: dict) -> None:
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if node.type in {"primary", "emphasis"} else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(box["left_in"]),
        Inches(box["top_in"]),
        Inches(box["width_in"]),
        Inches(box["height_in"]),
    )
    shape.name = f"HW_ARCH_NODE:{node.id}"
    color_key = {
        "primary": "accent6",
        "secondary": "accent4",
        "emphasis": "hw_red",
        "data": "accent5",
    }[node.type]
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["colors"][color_key])
    shape.line.color.rgb = _rgb(theme["colors"][color_key])
    shape.line.width = Pt(layout["node_border_pt"])
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = Inches(layout["node_text_margin_in"])
    text_frame.margin_right = Inches(layout["node_text_margin_in"])
    text_frame.margin_top = Inches(layout["node_text_margin_in"])
    text_frame.margin_bottom = Inches(layout["node_text_margin_in"])
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    _format_paragraph(paragraph, theme)
    run = paragraph.add_run()
    run.text = node.text
    text_color = "background" if node.type == "emphasis" else "body"
    _format_run(run, theme, size=layout["node_font_size_pt"], bold=True, color_key=text_color)


def _clamp(value: float, minimum: float, maximum: float) -> float:
    return max(minimum, min(value, maximum))


def _render_image_placeholder(slide, slide_ir: ImageSlide, theme: dict) -> None:
    _title(slide, slide_ir.title, theme)
    layout = theme["layouts"]["image"]
    box = layout["box"]
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(box["left_in"]),
        Inches(box["top_in"]),
        Inches(box["width_in"]),
        Inches(box["height_in"]),
    )
    shape.name = "HW_IMAGE_PLACEHOLDER"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["colors"]["surface"])
    shape.line.color.rgb = _rgb(theme["colors"]["muted"])
    reference = slide_ir.placeholder or slide_ir.image_ref or ""
    shape.text = f"16:9 图片占位\n{reference}\n等比放入，禁止随意裁切".strip()
    _fit_text_frame(shape.text_frame)
    for paragraph in shape.text_frame.paragraphs:
        _format_paragraph(paragraph, theme)
        for run in paragraph.runs:
            _format_run(run, theme, size=layout["placeholder_font_size_pt"], bold=True, color_key="body")
    if slide_ir.caption:
        _add_text_box(
            slide,
            slide_ir.caption,
            layout["caption"],
            theme,
            size=layout["caption"]["font_size_pt"],
            color_key="secondary",
            name="HW_RENDERED_TEXT:NOTE",
        )


def _render_placeholder(slide, layout: str, theme: dict) -> None:
    _title(slide, f"{layout} 待渲染", theme)
    placeholder = theme["layouts"]["placeholder"]
    _add_text_box(
        slide,
        "该版式将在后续任务卡实现。",
        placeholder,
        theme,
        size=placeholder["font_size_pt"],
        color_key="secondary",
        name="HW_RENDERED_TEXT:BODY",
    )


def _title(slide, text: str, theme: dict) -> None:
    box = theme["layouts"]["title"]
    size = _fit_ppt_text(text, theme["ppt_typography"]["slide_title_candidates_pt"], box, theme)
    _add_text_box(slide, text, box, theme, size=size, bold=True, name="HW_RENDERED_TEXT:TITLE")
    _red_bar(slide, theme, theme["layouts"]["title_bar"])


def _fit_ppt_text(text: str, candidates: list[float], box: dict, theme: dict) -> float:
    typography = theme["ppt_typography"]
    return fit_text_stack(
        [text],
        candidates,
        width_in=box["width_in"],
        height_in=box["height_in"],
        line_spacing=theme["typography"]["line_spacing"],
        item_gap_in=0,
        baseline_pt=theme["grid"]["baseline_pt"],
        horizontal_margin_in=typography["text_horizontal_margin_in"],
        vertical_margin_in=typography["text_vertical_margin_in"],
    ).font_size_pt


def _fit_body_stack(
    texts: list[str],
    width_in: float,
    height_in: float,
    theme: dict,
    *,
    candidates: list[float] | None = None,
):
    typography = theme["ppt_typography"]
    return fit_text_stack(
        texts,
        candidates or typography["body_candidates_pt"],
        width_in=width_in,
        height_in=max(height_in, 0),
        line_spacing=theme["typography"]["line_spacing"],
        item_gap_in=typography["body_item_gap_in"],
        baseline_pt=theme["grid"]["baseline_pt"],
        horizontal_margin_in=typography["text_horizontal_margin_in"],
        vertical_margin_in=typography["text_vertical_margin_in"],
    )


def _red_bar(slide, theme: dict, box: dict) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(box["left_in"]),
        Inches(box["top_in"]),
        Inches(box["width_in"]),
        Inches(box["height_in"]),
    )
    shape.name = "HW_DECORATION:RED_BAR"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["colors"]["hw_red"])
    shape.line.fill.background()


def _add_footer(slide, classification: str, page_number: int, total_pages: int, theme: dict) -> None:
    top = theme["slide"]["footer_top_in"]
    layout = theme["layouts"]["footer"]
    security_text = f"{theme['footer'].get('security_prefix', '')}{classification}"
    page_text = theme["footer"].get("page_number_format", "{current}").format(current=page_number, total=total_pages)
    classification_box = {
        "left_in": layout["classification_left_in"],
        "top_in": top,
        "width_in": layout["classification_width_in"],
        "height_in": layout["height_in"],
    }
    copyright_box = {
        "left_in": layout["copyright_left_in"],
        "top_in": top,
        "width_in": layout["copyright_width_in"],
        "height_in": layout["height_in"],
    }
    page_box = {
        "left_in": layout["page_left_in"],
        "top_in": top,
        "width_in": layout["page_width_in"],
        "height_in": layout["height_in"],
    }
    _add_text_box(slide, security_text, classification_box, theme, size=theme["font_sizes_pt"]["footer"], color_key="muted")
    _add_text_box(slide, theme["footer"]["copyright"], copyright_box, theme, size=theme["font_sizes_pt"]["copyright"], color_key="muted")
    _add_text_box(slide, page_text, page_box, theme, size=theme["font_sizes_pt"]["footer"], color_key="muted")


def _add_text_box(
    slide,
    text: str,
    box: dict,
    theme: dict,
    *,
    size: float,
    bold: bool = False,
    color_key: str = "title",
    font_role: str = "body",
    name: str = "HW_RENDERED_TEXT",
):
    return _add_text(
        slide,
        text,
        box["left_in"],
        box["top_in"],
        box["width_in"],
        box["height_in"],
        theme,
        size=size,
        bold=bold,
        color_key=color_key,
        font_role=font_role,
        name=name,
    )


def _add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict,
    *,
    size: float,
    bold: bool = False,
    color_key: str = "title",
    font_role: str = "body",
    name: str = "HW_RENDERED_TEXT",
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    text_frame = shape.text_frame
    text_frame.clear()
    _fit_text_frame(text_frame)
    paragraph = text_frame.paragraphs[0]
    _format_paragraph(paragraph, theme)
    run = paragraph.add_run()
    run.text = text
    _format_run(run, theme, size=size, bold=bold, color_key=color_key, font_role=font_role)
    return shape


def _format_cell(cell, theme: dict, *, bold: bool = False, size: float | None = None, color_key: str = "body") -> None:
    for paragraph in cell.text_frame.paragraphs:
        _format_paragraph(paragraph, theme)
        for run in paragraph.runs:
            _format_run(run, theme, size=size or theme["font_sizes_pt"]["note"], bold=bold, color_key=color_key)


def _format_paragraph(paragraph, theme: dict) -> None:
    paragraph.line_spacing = theme["typography"]["line_spacing"]


def _fit_text_frame(text_frame) -> None:
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _format_run(run, theme: dict, *, size: float, bold: bool = False, color_key: str = "title", font_role: str = "body") -> None:
    _set_run_font(run, theme, font_role=font_role)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(theme["colors"][color_key])


def _set_run_font(run, theme: dict, *, font_role: str) -> None:
    latin_font = theme["fonts"]["latin"][0]
    if font_role == "number":
        latin_font = theme["fonts"].get("number", theme["fonts"]["latin"])[0]
    east_asia_font = theme["fonts"]["east_asia"][0]
    run.font.name = latin_font
    rpr = run._r.get_or_add_rPr()
    rpr.get_or_add_latin().typeface = latin_font
    east_asia = _get_or_add_font_node(rpr, "ea")
    east_asia.set("typeface", east_asia_font)


def _get_or_add_font_node(rpr, tag: str):
    node_tag = qn(f"a:{tag}")
    for child in rpr:
        if child.tag == node_tag:
            return child
    node = OxmlElement(f"a:{tag}")
    rpr.append(node)
    return node


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
