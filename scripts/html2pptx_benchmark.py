"""Run the isolated DeckIR HTML/PptxGenJS experiment beside the production renderer."""

from __future__ import annotations

import argparse
from hashlib import sha256
import importlib.util
import json
from importlib.metadata import version as distribution_version
from pathlib import Path
import shutil
import subprocess
import sys
from typing import Any

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.ir.shell import validate_deck_ir_text
from app.lint.pptx_lint import check_pptx, write_reports
from app.rendering.pptx_renderer import render_deck_ir
from app.template.package import validate_pptx_package
from app.template.profile import extract_template_profile
from app.template.renderer import render_deck_ir_with_template


EXPERIMENT = ROOT / "experiments" / "html2pptx"
THEME_PATH = ROOT / "backend" / "app" / "rendering" / "themes" / "hw_theme.json"


def _assessment_module():
    path = EXPERIMENT / "assessment.py"
    spec = importlib.util.spec_from_file_location("html2pptx_assessment", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("HTML experiment assessment module is unavailable")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Compare validated DeckIR rendering engines without changing production defaults.")
    parser.add_argument("--deck-ir", type=Path, action="append", help="DeckIR JSON fixture; repeat to select individual cases.")
    parser.add_argument("--suite", action="store_true", help="Run the fixed core/long-Chinese/dense-table fixtures.")
    parser.add_argument("--template", type=Path, help="Optional HIT or business template; Python uses it, HTML is explicitly marked unsupported.")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--skip-office", action="store_true", help="Do not invoke optional Windows Office visual exports.")
    parser.add_argument("--visual-review", type=Path, help="Approved per-case agent/human review JSON to replace provisional scores.")
    return parser


def _fixture_paths(args: argparse.Namespace) -> list[Path]:
    paths = list(args.deck_ir or [])
    if args.suite:
        manifest = json.loads((EXPERIMENT / "fixtures.json").read_text(encoding="utf-8"))
        paths.extend(ROOT / item["deck_ir"] for item in manifest["fixtures"])
    unique = list(dict.fromkeys(path.resolve() for path in paths))
    if not unique:
        raise ValueError("provide --deck-ir or --suite")
    return unique


def _validated_deck(path: Path):
    result = validate_deck_ir_text(path.read_text(encoding="utf-8"))
    if not result.ok or result.value is None:
        raise ValueError(f"DeckIR fixture is invalid: {path.name}")
    return result.value


def _render_python(deck, case_dir: Path, template: Path | None, template_profile=None):
    output = case_dir / "python" / "deck.pptx"
    if template is None:
        return render_deck_ir(deck, output)
    return render_deck_ir_with_template(
        deck,
        template,
        output,
        audit_dir=output.parent,
        profile=template_profile,
    ).artifact_path


def _render_html(deck, case_dir: Path) -> tuple[Path, dict[str, Any]]:
    engine_dir = case_dir / "html"
    engine_dir.mkdir(parents=True, exist_ok=True)
    input_path = engine_dir / "deck_ir.json"
    output = engine_dir / "deck.pptx"
    report_path = engine_dir / "html_render_report.json"
    input_path.write_text(deck.model_dump_json(indent=2, by_alias=True) + "\n", encoding="utf-8")
    node = shutil.which("node")
    if node is None:
        raise RuntimeError("Node is unavailable; install the isolated experiment dependencies before running this benchmark")
    if not (EXPERIMENT / "node_modules" / "pptxgenjs").exists():
        raise RuntimeError("HTML experiment dependencies are missing; run npm ci in experiments/html2pptx")
    command = [
        node,
        str(EXPERIMENT / "render.mjs"),
        "--input",
        str(input_path),
        "--output",
        str(output),
        "--report",
        str(report_path),
        "--theme",
        str(THEME_PATH),
    ]
    completed = subprocess.run(command, cwd=EXPERIMENT, capture_output=True, text=True, encoding="utf-8")
    if completed.returncode != 0:
        raise RuntimeError(f"HTML experiment renderer failed: {completed.stderr.strip() or completed.stdout.strip()}")
    return output, json.loads(report_path.read_text(encoding="utf-8"))


def _engine_evidence(
    engine: str,
    artifact: Path | None,
    deck,
    engine_dir: Path,
    *,
    template_case: bool,
    skip_office: bool,
    html_render: dict[str, Any] | None = None,
    template_profile=None,
) -> dict[str, Any]:
    if artifact is None or not artifact.is_file():
        return _failed_evidence("artifact_missing")
    package = validate_pptx_package(artifact)
    lint_report = check_pptx(
        artifact,
        classification=deck.meta.classification,
        template_profile=template_profile,
    )
    write_reports(lint_report, engine_dir)
    lint_payload = lint_report.to_dict()
    objects = _editable_objects(artifact)
    requires_table = any(slide.layout == "table" for slide in deck.slides)
    requires_chart = any(slide.layout == "chart" for slide in deck.slides)
    native_editability = (
        objects["text_shapes"] >= len(deck.slides)
        and (not requires_table or objects["tables"] >= 1)
        and (not requires_chart or objects["charts"] >= 1)
    )
    layout_issues = [item for item in lint_report.items if item.code in {"HW-W03", "HW-W07"}]
    office = {"status": "not_run"}
    if not skip_office:
        office = _run_office_visual_qa(artifact, engine_dir / "visual")
    hard_gates = {
        "schema_valid": True,
        "package_valid": bool(package["pass"]),
        "lint_clean": lint_payload["summary"]["errors"] == 0,
        "no_layout_issues": not layout_issues,
        "office_opened": skip_office or bool(office.get("office_export_pass")),
        "native_editability": native_editability and not (engine == "html" and template_case),
    }
    provisional = {
        "visual_hierarchy": 3 if hard_gates["no_layout_issues"] else 0,
        "chinese_typography": 3 if hard_gates["lint_clean"] else 0,
        "template_fidelity": 5 if template_case and engine == "python" else (0 if template_case else 3),
        "deployment_reliability": 5 if engine == "python" else 2,
    }
    return {
        "artifact": artifact.name,
        "artifact_sha256": _file_sha256(artifact),
        "hard_gates": hard_gates,
        "scores": provisional,
        "visual_review": "pending",
        "package_report": package,
        "lint_summary": lint_payload["summary"],
        "layout_warning_count": len(layout_issues),
        "editable_objects": objects,
        "office": office,
        "html_render": html_render,
    }


def _failed_evidence(reason: str) -> dict[str, Any]:
    return {
        "hard_gates": {
            "schema_valid": False,
            "package_valid": False,
            "lint_clean": False,
            "no_layout_issues": False,
            "office_opened": False,
            "native_editability": False,
        },
        "scores": {"visual_hierarchy": 0, "chinese_typography": 0, "template_fidelity": 0, "deployment_reliability": 0},
        "visual_review": "pending",
        "failure": reason,
    }


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _engine_versions() -> dict[str, str]:
    package = json.loads((EXPERIMENT / "package.json").read_text(encoding="utf-8"))
    return {
        "python": sys.version.split()[0],
        "python_pptx": distribution_version("python-pptx"),
        "html_node_requirement": package["engines"]["node"],
        "html_pptxgenjs": package["dependencies"]["pptxgenjs"],
        "html_playwright": package["dependencies"]["playwright"],
        "html_sharp": package["dependencies"]["sharp"],
    }


def _editable_objects(path: Path) -> dict[str, int]:
    presentation = Presentation(path)
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    return {
        "slides": len(presentation.slides),
        "text_shapes": sum(bool(getattr(shape, "has_text_frame", False)) for shape in shapes),
        "tables": sum(bool(getattr(shape, "has_table", False)) for shape in shapes),
        "charts": sum(bool(getattr(shape, "has_chart", False)) for shape in shapes),
    }


def _run_office_visual_qa(artifact: Path, output_dir: Path) -> dict[str, Any]:
    command = [
        sys.executable,
        str(ROOT / "scripts" / "ppt_visual_qa.py"),
        "--kind",
        "deck",
        "--input",
        str(artifact),
        "--output-dir",
        str(output_dir),
    ]
    subprocess.run(command, cwd=ROOT, check=False, capture_output=True, text=True, encoding="utf-8")
    path = output_dir / "visual_qa_report.json"
    report = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {"status": "failed"}
    export_path = output_dir / "visual_export_report.json"
    export = json.loads(export_path.read_text(encoding="utf-8-sig")) if export_path.is_file() else {"pass": False}
    report["office_export_pass"] = bool(export.get("pass"))
    return report


def _apply_visual_reviews(cases: list[dict[str, Any]], path: Path | None) -> None:
    if path is None:
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    for case in cases:
        review = payload.get("cases", {}).get(case["name"], {})
        for engine in ("python", "html"):
            update = review.get(engine)
            if not isinstance(update, dict):
                continue
            if update.get("visual_review") in {"approved", "pending", "rejected"}:
                case[engine]["visual_review"] = update["visual_review"]
            if isinstance(update.get("scores"), dict):
                case[engine]["scores"].update(update["scores"])


def run_benchmark(args: argparse.Namespace) -> dict[str, Any]:
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    template_profile = extract_template_profile(args.template) if args.template is not None else None
    cases: list[dict[str, Any]] = []
    for source in _fixture_paths(args):
        deck = _validated_deck(source)
        case_dir = output_dir / source.stem
        case_dir.mkdir(parents=True, exist_ok=True)
        template_case = args.template is not None
        python_artifact = _render_python(deck, case_dir, args.template, template_profile)
        html_artifact = None
        html_render = None
        html_failure = None
        try:
            html_artifact, html_render = _render_html(deck, case_dir)
        except RuntimeError as exc:
            html_failure = str(exc)
        python_evidence = _engine_evidence(
            "python",
            python_artifact,
            deck,
            case_dir / "python",
            template_case=template_case,
            skip_office=args.skip_office,
            template_profile=template_profile,
        )
        html_evidence = (
            _engine_evidence(
                "html", html_artifact, deck, case_dir / "html", template_case=template_case,
                skip_office=args.skip_office, html_render=html_render,
            )
            if html_artifact is not None
            else _failed_evidence(html_failure or "html_renderer_failed")
        )
        cases.append({"name": source.stem, "template_case": template_case, "python": python_evidence, "html": html_evidence})
    _apply_visual_reviews(cases, args.visual_review)
    assessment = _assessment_module().assess_engines(cases)
    assessment.update(
        {
            "production_engine": "python-pptx",
            "html_engine": "isolated_experiment",
            "engine_versions": _engine_versions(),
            "template": args.template.name if args.template else None,
            "visual_review_source": str(args.visual_review) if args.visual_review else None,
        }
    )
    (output_dir / "engine_assessment.json").write_text(
        json.dumps(assessment, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return assessment


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    assessment = run_benchmark(args)
    print(f"engine assessment: {args.output_dir / 'engine_assessment.json'}")
    print(f"recommendation: {assessment['recommendation']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
