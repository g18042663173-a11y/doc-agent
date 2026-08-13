from __future__ import annotations
# ruff: noqa: E402

import json
import os
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.generators.stub import StubGenerator
from app.ir.deck_ir import DeckIR
from app.ir.schema_export import SchemaSnapshotError, verify_schema_snapshots
from app.ir.word_ir import WordIR
from app.lint.pptx_lint import check_pptx, write_reports
from app.rendering.docx_renderer import render_word_ir
from app.rendering.pptx_renderer import render_deck_ir


CORE_PACKAGE_THRESHOLD = 80.0
OVERALL_THRESHOLD = 70.0


def main(output_dir: Path | None = None) -> int:
    try:
        verify_schema_snapshots(ROOT / "backend" / "schemas")
    except SchemaSnapshotError as exc:
        print(f"schema verification failed: {exc}", file=sys.stderr)
        return 1

    generator = StubGenerator()
    word = WordIR.model_validate_json(generator.generate("", target="word_ir"))
    deck = DeckIR.model_validate_json(generator.generate("", target="deck_ir"))

    output_dir = output_dir or (ROOT / "output")
    word_path = render_word_ir(word, output_dir / "c0_word.docx")
    deck_path = render_deck_ir(deck, output_dir / "c0_deck.pptx")
    report = check_pptx(deck_path, classification=deck.meta.classification)
    report_json, _ = write_reports(report, output_dir)
    e2e_ok = _run_four_format_e2e(output_dir)
    coverage_ok = _run_pytest_with_coverage(output_dir)

    if report.summary["pass"] and e2e_ok and coverage_ok:
        print("C0 verify passed")
    else:
        print("C0 verify failed")
    print(f"schemas: {ROOT / 'backend' / 'schemas'}")
    print(f"artifacts: {word_path.name}, {deck_path.name}")
    print(f"report: {report_json.name}")
    return 0 if report.summary["pass"] and e2e_ok and coverage_ok else 1


def _run_four_format_e2e(output_dir: Path) -> bool:
    samples = {
        "md": ROOT / "samples" / "input" / "parser_samples" / "md_sample_01.md",
        "docx": ROOT / "samples" / "input" / "parser_samples" / "docx_sample_01.docx",
        "xlsx": ROOT / "samples" / "input" / "parser_samples" / "xlsx_sample_01.xlsx",
        "pptx": ROOT / "samples" / "input" / "parser_samples" / "pptx_sample_01.pptx",
    }
    e2e_root = output_dir / f"verify-e2e-{os.getpid()}"
    if e2e_root.exists():
        shutil.rmtree(e2e_root)
    for suffix, sample in samples.items():
        for target in ("word", "deck"):
            target_dir = e2e_root / f"{suffix}-{target}"
            command = [
                sys.executable,
                "scripts/demo_e2e.py",
                str(sample),
                "--target",
                target,
                "--generator",
                "stub",
                "--output-dir",
                str(target_dir),
            ]
            if target == "deck":
                command.append("--lint")
            result = subprocess.run(
                command,
                cwd=ROOT,
                text=True,
                encoding="utf-8",
                errors="replace",
                capture_output=True,
                check=False,
            )
            if result.returncode != 0:
                print(f"e2e: {suffix}->{target} failed")
                print(result.stdout)
                print(result.stderr, file=sys.stderr)
                return False
            report_path = target_dir / "report.json"
            if not report_path.exists():
                print(f"e2e: {suffix}->{target} did not write report.json")
                return False
            report = json.loads(report_path.read_text(encoding="utf-8"))
            if not report.get("summary", {}).get("pass"):
                print(f"e2e: {suffix}->{target} report did not pass")
                return False
            document = json.loads((target_dir / "document_ir.json").read_text(encoding="utf-8"))
            facts = _semantic_facts(document)
            if len(facts) < 3:
                print(f"e2e: {suffix}->{target} exposed only {len(facts)} source facts; fixture is not meaningful")
                return False
            generated_path = target_dir / f"{target}_ir.json"
            if not generated_path.exists():
                print(f"e2e: {suffix}->{target} did not write {generated_path.name}")
                return False
            generated_text = generated_path.read_text(encoding="utf-8")
            missing_ir = _missing_semantic_facts(facts, generated_text)
            if missing_ir:
                fact = missing_ir[0]
                print(f"e2e: {suffix}->{target} generated IR lost {fact.loc}={fact.value!r}")
                return False
            artifact_path = target_dir / ("word.docx" if target == "word" else "deck.pptx")
            if not artifact_path.exists():
                print(f"e2e: {suffix}->{target} did not write {artifact_path.name}")
                return False
            missing_artifact = _missing_semantic_facts(facts, _artifact_text(artifact_path, target))
            if missing_artifact:
                fact = missing_artifact[0]
                print(f"e2e: {suffix}->{target} final artifact lost {fact.loc}={fact.value!r}")
                return False
    print("e2e: four input formats passed word/deck stub chains")
    return True


@dataclass(frozen=True)
class SemanticFact:
    loc: str
    value: str


def _semantic_facts(document: dict) -> list[SemanticFact]:
    facts: list[SemanticFact] = []
    seen: set[str] = set()

    def add(loc: str, value) -> None:
        normalized = _normalize_semantic_text(value)
        if not normalized or normalized in seen:
            return
        seen.add(normalized)
        facts.append(SemanticFact(loc=loc, value=normalized))

    content = document.get("content", {})
    for block_index, block in enumerate(content.get("blocks", [])):
        block_type = block.get("type")
        if block_type in {"heading", "paragraph"}:
            add(f"content.blocks[{block_index}].text", block.get("text"))
        elif block_type in {"bullet_list", "numbered_list"}:
            for item_index, item in enumerate(block.get("items", [])):
                add(f"content.blocks[{block_index}].items[{item_index}].text", item.get("text"))
        elif block_type == "table":
            for column, value in enumerate(block.get("header", [])):
                add(f"content.blocks[{block_index}].header[{column}]", value)
            for row_index, row in enumerate(block.get("rows", [])):
                for column, value in enumerate(row):
                    add(f"content.blocks[{block_index}].rows[{row_index}][{column}]", value)
        elif block_type == "image_placeholder":
            add(f"content.blocks[{block_index}].ref", block.get("ref"))
            add(f"content.blocks[{block_index}].caption", block.get("caption"))

    for sheet_index, sheet in enumerate(content.get("sheets", [])):
        add(f"content.sheets[{sheet_index}].name", sheet.get("name"))
        for column, value in enumerate(sheet.get("header_guess", [])):
            add(f"content.sheets[{sheet_index}].header_guess[{column}]", value)
        for row_index, row in enumerate(sheet.get("preview_rows", [])):
            for column, value in enumerate(row):
                add(f"content.sheets[{sheet_index}].preview_rows[{row_index}][{column}]", value)

    for slide_index, slide in enumerate(content.get("slides", [])):
        add(f"content.slides[{slide_index}].title", slide.get("title"))
        for body_index, value in enumerate(slide.get("bodies", [])):
            add(f"content.slides[{slide_index}].bodies[{body_index}]", value)
        for table_index, table in enumerate(slide.get("tables", [])):
            for column, value in enumerate(table.get("header", [])):
                add(f"content.slides[{slide_index}].tables[{table_index}].header[{column}]", value)
            for row_index, row in enumerate(table.get("rows", [])):
                for column, value in enumerate(row):
                    add(f"content.slides[{slide_index}].tables[{table_index}].rows[{row_index}][{column}]", value)
        add(f"content.slides[{slide_index}].notes", slide.get("notes"))

    if not facts:
        add("source.filename", Path(document["source"]["filename"]).stem)
    return facts


def _missing_semantic_facts(facts: list[SemanticFact], value: str) -> list[SemanticFact]:
    corpus = _semantic_corpus(value)
    return [fact for fact in facts if fact.value not in corpus]


def _semantic_corpus(value: str) -> str:
    try:
        parsed = json.loads(value)
    except (json.JSONDecodeError, TypeError):
        return _normalize_semantic_text(value)
    return "\n".join(_iter_json_strings(parsed))


def _iter_json_strings(value) -> list[str]:
    if isinstance(value, str):
        return [_normalize_semantic_text(value)]
    if isinstance(value, dict):
        return [text for child in value.values() for text in _iter_json_strings(child)]
    if isinstance(value, list):
        return [text for child in value for text in _iter_json_strings(child)]
    if value is None:
        return []
    return [_normalize_semantic_text(value)]


def _artifact_text(path: Path, target: str) -> str:
    values: list[str] = []
    if target == "word":
        document = Document(str(path))
        values.extend(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            values.extend(cell.text for row in table.rows for cell in row.cells)
    else:
        presentation = Presentation(str(path))
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    values.append(shape.text)
                if getattr(shape, "has_table", False):
                    values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(_normalize_semantic_text(value) for value in values if _normalize_semantic_text(value))


def _normalize_semantic_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _run_pytest_with_coverage(output_dir: Path) -> bool:
    output_dir.mkdir(parents=True, exist_ok=True)
    coverage_json = output_dir / "coverage.json"
    env = os.environ.copy()
    existing_pythonpath = env.get("PYTHONPATH")
    env["PYTHONPATH"] = str(BACKEND) if not existing_pythonpath else f"{BACKEND}{os.pathsep}{existing_pythonpath}"
    env.pop("VERIFY_RUNNING", None)
    command = [
        sys.executable,
        "-m",
        "pytest",
        "backend/tests",
        "-q",
        "--cov=app",
        "--cov-report=term-missing",
        f"--cov-report=json:{coverage_json}",
    ]
    result = subprocess.run(
        command,
        cwd=ROOT,
        env=env,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        print("coverage: pytest failed")
        print(result.stdout)
        print(result.stderr, file=sys.stderr)
        return False

    coverage = json.loads(coverage_json.read_text(encoding="utf-8"))
    package_results = {
        package: _package_coverage(coverage, package)
        for package in ("parsers", "ir", "lint")
    }
    overall = float(coverage["totals"]["percent_covered"])
    failures = [
        f"app/{package} {percent:.2f}% < {CORE_PACKAGE_THRESHOLD:.0f}%"
        for package, percent in package_results.items()
        if percent < CORE_PACKAGE_THRESHOLD
    ]
    if overall < OVERALL_THRESHOLD:
        failures.append(f"overall {overall:.2f}% < {OVERALL_THRESHOLD:.0f}%")
    if failures:
        print("coverage: threshold failed")
        for failure in failures:
            print(f"- {failure}")
        return False

    packages = ", ".join(f"app/{package}={percent:.2f}%" for package, percent in package_results.items())
    print(f"coverage: {packages}, overall={overall:.2f}%")
    return True


def _package_coverage(coverage: dict, package: str) -> float:
    total = 0
    covered = 0
    needle = f"app/{package}/"
    for filename, data in coverage["files"].items():
        normalized = filename.replace("\\", "/")
        if needle not in normalized:
            continue
        summary = data["summary"]
        statements = int(summary["num_statements"])
        if statements == 0:
            continue
        total += statements
        covered += int(summary["covered_lines"])
    if total == 0:
        return 0.0
    return (covered / total) * 100


if __name__ == "__main__":
    raise SystemExit(main())
