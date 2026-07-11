from __future__ import annotations
# ruff: noqa: E402

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.ir.deck_ir import DeckIR
from app.lint.pptx_lint import check_pptx, write_reports
from app.rendering.pptx_renderer import render_deck_ir


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Render a legal DeckIR and inject stable PPTX lint violations.")
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "samples" / "ir" / "deck_lint_violation.json",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "samples" / "output" / "deck" / "deck_lint_violation.pptx",
    )
    parser.add_argument(
        "--report-dir",
        type=Path,
        default=ROOT / "samples" / "output" / "deck" / "deck_lint_violation_report",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    deck = DeckIR.model_validate_json(args.source.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    render_deck_ir(deck, args.output)
    _inject_violations(args.output, deck.meta.classification)
    report = check_pptx(args.output, classification=deck.meta.classification)
    write_reports(report, args.report_dir)
    print(json.dumps(report.summary, ensure_ascii=False, sort_keys=True))
    return 0


def _inject_violations(path: Path, classification: str) -> None:
    presentation = Presentation(path)
    slide = presentation.slides[0]
    body = None
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if classification in shape.text:
            shape.text_frame.clear()
        if "关键结论" in shape.text:
            body = shape
    if body is None:
        raise RuntimeError("body injection target not found")
    body.left = Inches(0.05)
    for paragraph in body.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Comic Sans MS"
            run.font.size = Pt(6)
            run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    presentation.save(path)


if __name__ == "__main__":
    raise SystemExit(main())
