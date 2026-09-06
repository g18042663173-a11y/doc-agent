"""Create deterministic synthetic samples for rhetoric-deck-workflow."""

from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "skills/rhetoric-deck-workflow/samples"
RED = RGBColor(0xC7, 0x00, 0x0B)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF5, 0xF5, 0xF5)


def main() -> int:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    expected = SAMPLES / "expected"
    expected.mkdir(exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    _remove_default_slides(presentation)
    cover_refs = _add_cover(presentation)
    card_refs = _add_cards(presentation)
    table_refs = _add_table(presentation)
    node_refs = _add_nodes(presentation)
    _add_master_footer(presentation, "源件密级内部资料不可进入产物")
    presentation.save(SAMPLES / "source_sample.pptx")

    document = Document()
    document.add_heading("用户材料：统一校验方案", 0)
    document.add_paragraph("密级：HUAWEI CONFIDENTIAL")
    document.add_paragraph("背景：当前流程存在重复录入和问题定位慢的痛点，目标是缩短处理链路。")
    document.add_paragraph("封面副标题：把校验前移到契约层。")
    document.add_paragraph("汇报人：交付组")
    document.add_paragraph("卡片一：现状重复录入；卡片二：定位链路过长；卡片三：缺少统一契约；卡片四：回归证据不足。")
    document.add_paragraph("方案 A 保留现有接口，改造成本低；方案 B 增加统一校验层，问题提前暴露。", style="List Bullet")
    document.add_paragraph("推荐方案 B，因为校验职责集中且回归边界清晰。", style="List Bullet")
    document.add_paragraph("问题是字段不一致；措施是 Schema 前置；负责人是契约组；状态是进行中。")
    document.add_paragraph("表体：接口字段缺少必填标记。统一校验拦截非法输入。契约组跟进。本周完成。")
    document.add_paragraph("表体：下游失败无法定位。增加审计清单。交付组跟进。下周验证。")
    document.add_paragraph("节点：解析输入、校验契约、生成产物、执行审计。")
    document.add_paragraph("测试覆盖正常输入、缺失字段和异常包，非法输入不生成产物。")
    document.save(SAMPLES / "material_sample.docx")

    pages = [
        _page(
            "p01",
            "struct_cover",
            "cover_title_subtitle",
            [
                _slot("s1", "page.title", "title", cover_refs["title"], chars=36, lines=2),
                _slot("s2", "page.subtitle", "subtitle", cover_refs["subtitle"], chars=40, lines=2),
                _slot("s3", "page.byline", "caption", cover_refs["byline"], chars=24, lines=1),
            ],
            items=1,
        ),
        _page(
            "p02",
            "context_pain",
            "context_pain_goal",
            [
                _slot("s1", "page.title", "title", card_refs["title"], chars=36, lines=2),
                *[
                    _slot(f"s{index + 2}", f"context.card_{index + 1}", "caption", ref, chars=24, lines=2)
                    for index, ref in enumerate(card_refs["cards"])
                ],
            ],
            items=4,
        ),
        _page(
            "p03",
            "issue_retro",
            "problem_symptom_analysis_action",
            [
                _slot("s1", "page.title", "title", table_refs["title"], chars=36, lines=2),
                *[
                    _slot(
                        f"s{index + 2}",
                        f"issue.cell_{row}_{column}",
                        "table_cell",
                        f"{table_refs['table']}.cell_{row}_{column}",
                        chars=24,
                        lines=2,
                    )
                    for index, (row, column) in enumerate((r, c) for r in range(3) for c in range(4))
                ],
            ],
            items=12,
        ),
        _page(
            "p04",
            "method_walkthrough",
            "steps_compare_benefit",
            [
                _slot("s1", "page.title", "title", node_refs["title"], chars=36, lines=2),
                *[
                    _slot(f"s{index + 2}", f"method.node_{index + 1}", "node_label", ref, chars=16, lines=2)
                    for index, ref in enumerate(node_refs["nodes"])
                ],
            ],
            items=4,
            diagram=True,
        ),
    ]
    (expected / "skeleton.json").write_text(
        json.dumps(
            {
                "format": "deck_skeleton",
                "version": "1.0",
                "source_kind": "extracted",
                "deck_pattern": "review_solution",
                "page_count": 4,
                "pages": pages,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    (expected / "README.md").write_text(
        "# Expected sample\n\n"
        "Four 16:9 pages: cover, independent card boxes, a native 3x4 table, and rounded node labels.\n"
        "The skeleton binds every visible-text candidate. It retains no source sentence.\n"
        "Master footer chrome is kept and rewritten to the user classification.\n",
        encoding="utf-8",
    )
    (expected / "content_points.md").write_text(
        "# Content points\n\n"
        "- Cover title, subtitle, and byline come from user material.\n"
        "- Card page titles are filled; no empty title bars.\n"
        "- Every table cell has user text; source cell sentences are absent.\n"
        "- Node labels have user text.\n"
        "- Footer classification is `HUAWEI CONFIDENTIAL` from the material line.\n"
        "- Source phrases such as `源件高度敏感` / `源件表格机密` / `源件节点框` do not appear.\n",
        encoding="utf-8",
    )
    (expected / "forbidden_source_phrases.txt").write_text(
        "\n".join(
            [
                "源件高度敏感",
                "源件封面机密标题不可进入产物",
                "源件卡片标题",
                "源件表格机密",
                "源件节点框",
                "源件密级内部资料不可进入产物",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    return 0


def _remove_default_slides(presentation) -> None:
    slide_id_list = presentation.slides._sldIdLst
    for slide_id in list(slide_id_list):
        r_id = getattr(slide_id, "rId", None) or slide_id.get(qn("r:id"))
        if r_id:
            presentation.part.drop_rel(r_id)
        slide_id_list.remove(slide_id)


def _add_master_footer(presentation, text: str) -> None:
    donor = presentation.slides[0]
    shape = donor.shapes.add_textbox(Inches(0.4), Inches(7.15), Inches(6.4), Inches(0.28))
    _set_text(shape, text, size=10, color=GRAY)
    element = deepcopy(shape._element)
    donor.shapes._spTree.remove(shape._element)
    nvpr = element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    if nvpr is not None:
        nvpr.set("id", "4096")
        nvpr.set("name", "FooterClassification")
    tree = presentation.slide_masters[0]._element.find(qn("p:cSld")).find(qn("p:spTree"))
    tree.append(element)


def _add_cover(presentation) -> dict[str, str]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    _fill(bar, RED)
    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.1))
    _set_text(title, "源件封面机密标题不可进入产物", size=32, bold=True)
    subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.6))
    _set_text(subtitle, "源件封面机密副标题不可进入产物", size=18)
    byline = slide.shapes.add_textbox(Inches(0.8), Inches(4.4), Inches(11.7), Inches(0.4))
    _set_text(byline, "源件封面机密署名不可进入产物", size=14, color=GRAY)
    return {
        "title": f"sp_{title.shape_id}",
        "subtitle": f"sp_{subtitle.shape_id}",
        "byline": f"sp_{byline.shape_id}",
    }


def _add_cards(presentation) -> dict[str, str | list[str]]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.8), Inches(0.6))
    _set_text(title, "源件卡片页标题不可进入产物", size=24, bold=True)
    cards = []
    labels = [
        "源件卡片标题甲不可进入产物",
        "源件卡片标题乙不可进入产物",
        "源件卡片标题丙不可进入产物",
        "源件卡片标题丁不可进入产物",
    ]
    bodies = [
        "源件卡片正文甲高度敏感说明不可进入产物",
        "源件卡片正文乙高度敏感说明不可进入产物",
        "源件卡片正文丙高度敏感说明不可进入产物",
        "源件卡片正文丁高度敏感说明不可进入产物",
    ]
    positions = [(0.7, 1.2), (7.0, 1.2), (0.7, 4.2), (7.0, 4.2)]
    for (left, top), label, body in zip(positions, labels, bodies):
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(5.6), Inches(2.6))
        _fill(panel, LIGHT)
        heading = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(5.2), Inches(0.45))
        _set_text(heading, label, size=16, bold=True, color=RED)
        text = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.7), Inches(5.2), Inches(1.6))
        _set_text(text, body, size=14)
        cards.extend([f"sp_{heading.shape_id}", f"sp_{text.shape_id}"])
    return {"title": f"sp_{title.shape_id}", "cards": cards}


def _add_table(presentation) -> dict[str, str]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.8), Inches(0.55))
    _set_text(title, "源件表格页标题不可进入产物", size=24, bold=True)
    table_shape = slide.shapes.add_table(3, 4, Inches(0.7), Inches(1.2), Inches(11.9), Inches(4.4))
    headers = ["问题", "措施", "负责人", "状态"]
    bodies = [
        ["源件表格机密单元格甲不可进入产物", "源件表格机密单元格乙不可进入产物", "源件表格机密单元格丙不可进入产物", "源件表格机密单元格丁不可进入产物"],
        ["源件表格机密单元格戊不可进入产物", "源件表格机密单元格己不可进入产物", "源件表格机密单元格庚不可进入产物", "源件表格机密单元格辛不可进入产物"],
    ]
    table = table_shape.table
    for column, header in enumerate(headers):
        table.cell(0, column).text = header
    for row, values in enumerate(bodies, start=1):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    return {"title": f"sp_{title.shape_id}", "table": f"sp_{table_shape.shape_id}"}


def _add_nodes(presentation) -> dict[str, str | list[str]]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.8), Inches(0.55))
    _set_text(title, "源件节点页标题不可进入产物", size=24, bold=True)
    labels = [
        "源件节点框甲不可进入产物",
        "源件节点框乙不可进入产物",
        "源件节点框丙不可进入产物",
        "源件节点框丁不可进入产物",
    ]
    nodes = []
    for index, label in enumerate(labels):
        left = Inches(0.7 + index * 3.15)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.6), Inches(2.9), Inches(1.5))
        _fill(shape, RGBColor(0xFF, 0xF2, 0xF2))
        shape.line.color.rgb = RED
        _set_text(shape, label, size=14, align=PP_ALIGN.CENTER)
        nodes.append(f"sp_{shape.shape_id}")
    return {"title": f"sp_{title.shape_id}", "nodes": nodes}


def _page(page_id: str, pattern: str, flow: str, slots: list[dict], *, items: int, diagram: bool = False) -> dict:
    return {
        "page_id": page_id,
        "page_pattern": pattern,
        "argument_flow": flow,
        "confidence": 0.95,
        "structure": {"items": {"count": items, "expandable": False, "max": items}},
        "slots": slots,
        "emphasis": [],
        "diagram": {
            "type": "linear_flow" if diagram else "none",
            "groups": 1 if diagram else 0,
            "nodes_per_group": items if diagram else 0,
            "flow": "left_right" if diagram else "none",
            "annotations": 0,
            "labels": None,
        },
        "structural_labels": {"_display": []},
    }


def _slot(slot_id: str, semantic: str, slot_type: str, shape_ref: str, *, chars: int, lines: int) -> dict:
    return {
        "slot_id": slot_id,
        "semantic": semantic,
        "type": slot_type,
        "cardinality": {"min": 1, "max": 1},
        "capacity": {"chars_cjk": chars, "lines": lines},
        "required": True,
        "shape_ref": shape_ref,
    }


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _set_text(shape, text: str, *, size: int, bold: bool = False, color: RGBColor | None = None, align=None) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or DARK
    run.font.name = "Microsoft YaHei"


if __name__ == "__main__":
    raise SystemExit(main())
