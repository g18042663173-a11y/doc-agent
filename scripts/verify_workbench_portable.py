"""Exercise the extracted workbench API using only its bundled Python runtime."""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import sys
import time


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    root, output = args.root.resolve(), args.out.resolve()
    Path(sys.executable).resolve().relative_to(root / "runtime/python")
    output.mkdir(parents=True, exist_ok=False)
    os.environ["PATH"] = str(root / "app/tools/graphviz/bin") + os.pathsep + str(Path(os.environ.get("SystemRoot", "C:/Windows")) / "System32")
    def deny_network(event, arguments):
        if event in {"socket.connect", "socket.connect_ex", "socket.getaddrinfo", "socket.bind"}:
            raise PermissionError("Network disabled for portable acceptance")
    sys.addaudithook(deny_network)
    try:
        socket.getaddrinfo("portable-acceptance.invalid", 443)
        raise RuntimeError("Network guard was not activated")
    except PermissionError:
        pass
    from app.generators.stub import StubGenerator
    from app.web_api import create_api_app
    from docx import Document
    from pptx import Presentation
    app = create_api_app(work_dir=output / "jobs", generator=StubGenerator())
    client = app.test_client()
    cases = []
    material = "# 中期工作报告\n\n## 工作进展\n- 已完成需求分析与系统设计。\n- 已完成核心模块实现和测试。\n\n## 后续计划\n- 完成真实材料验证。\n- 汇总验收报告。\n"
    for target in ("word", "deck"):
        response = client.post("/api/generate", data={"type": target,
            "input_file": (BytesIO(material.encode("utf-8")), "中期材料.md")}, content_type="multipart/form-data")
        assert response.status_code == 202, response.get_json()
        job_id = response.get_json()["job_id"]
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            job = client.get(f"/api/status/{job_id}").get_json()
            if job["status"] in {"done", "failed", "cancelled"}:
                break
            time.sleep(0.1)
        assert job["status"] == "done", job
        assert job["report"]["summary"]["pass"] is True, job["report"]
        artifact = client.get(job["artifact"]["download_url"])
        assert artifact.status_code == 200
        path = output / ("word.docx" if target == "word" else "deck.pptx")
        path.write_bytes(artifact.data)
        if target == "word":
            assert Document(path).paragraphs
        else:
            assert len(Presentation(path).slides) >= 2
        downloaded = []
        for name, asset in job.get("assets", {}).items():
            result = client.get(asset["download_url"])
            assert result.status_code == 200, name
            downloaded.append(name)
        cases.append({"target": target, "job_id": job_id, "status": job["status"], "artifact": path.name,
                      "audit_downloads": downloaded})
    # Opening the job store again verifies persisted task recovery/listing.
    restored = create_api_app(work_dir=output / "jobs", generator=StubGenerator()).test_client()
    jobs = restored.get("/api/jobs").get_json()["jobs"]
    assert {case["job_id"] for case in cases} <= {job["job_id"] for job in jobs}
    report = {"pass": True, "python": sys.executable, "network_blocked": True, "cases": cases,
              "job_store_reopened": True, "ui_review": "separate_native_window_acceptance_required"}
    (output / "acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
