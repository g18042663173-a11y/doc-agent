from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from threading import Event, Thread
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.util import Inches, Pt
from werkzeug.serving import BaseWSGIServer, make_server


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.generators.stub import StubGenerator
from app.web_api import APP_VERSION, create_api_app


VIEWPORTS = (("desktop", 1280, 900), ("mobile", 390, 844))
NATIVE_THEME_TOKENS = {
    "light": {"accent": "#0078d4", "selected_background": "rgb(220, 235, 250)"},
    "dark": {"accent": "#4cc2ff", "selected_background": "rgb(40, 66, 92)"},
}


class FailOnceGenerator:
    name = "workbench-fail-once"

    def __init__(self) -> None:
        self.calls = 0
        self.stub = StubGenerator()
        self.blocking_started = Event()
        self.blocking_release = Event()

    def generate(self, prompt: str, *, target: str) -> str:
        self.calls += 1
        if self.calls == 1:
            raise RuntimeError("synthetic provider failure")
        if self.calls == 3:
            self.blocking_started.set()
            self.blocking_release.wait(timeout=10)
        return self.stub.generate(prompt, target=target)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run development-only browser checks for the document workbench.")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "output" / "qa" / "workbench-ui")
    parser.add_argument("--channel", default="chrome", help="Playwright browser channel; defaults to installed Google Chrome.")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        from playwright.sync_api import sync_playwright
    except ModuleNotFoundError:
        print("Playwright 未安装。开发机执行 .venv\\Scripts\\python.exe -m pip install -r requirements-dev-ui.txt 后重试。")
        return 2

    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    with sync_playwright() as playwright:
        for name, width, height in VIEWPORTS:
            results.append(_run_viewport(playwright, output_dir, args.channel, name, width, height))
    report = {"workbench_ui_version": "1.0", "pass": all(item["pass"] for item in results), "viewports": results}
    report_path = output_dir / "workbench_ui_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"workbench ui report: {report_path}")
    return 0 if report["pass"] else 1


def _run_viewport(playwright, output_dir: Path, channel: str, name: str, width: int, height: int) -> dict[str, Any]:
    workspace = output_dir / f"{name}-jobs"
    generator = FailOnceGenerator()
    app = create_api_app(work_dir=workspace, generator=generator)
    server: BaseWSGIServer = make_server("127.0.0.1", 0, app)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    browser = None
    try:
        browser = playwright.chromium.launch(headless=True, channel=channel)
        context = browser.new_context(viewport={"width": width, "height": height}, accept_downloads=True)
        page = context.new_page()
        console_errors: list[str] = []
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        page.on("pageerror", lambda error: console_errors.append(str(error)))

        with TemporaryDirectory(dir=output_dir, prefix=f"{name}-fixtures-") as fixture_dir:
            fixture_root = Path(fixture_dir)
            source = fixture_root / "report.md"
            source.write_text("# Workbench reliability\n\n- first run fails\n- second run succeeds\n", encoding="utf-8")
            template = _write_template(fixture_root / "template.pptx")
            unsafe_template = _write_ole_template(fixture_root / "template-with-ole.pptx")
            page.goto(f"http://127.0.0.1:{server.server_port}/static/index.html", wait_until="networkidle")
            page.get_by_test_id("service-status").get_by_text(f"本地服务 v{APP_VERSION}").wait_for(timeout=5_000)

            # Windows Settings visual shell: fresh browser state follows the system,
            # preserves manual overrides, and exposes all four workbench views.
            theme_mode = page.get_by_test_id("theme-toggle")
            if theme_mode.input_value() != "system":
                raise AssertionError("fresh workbench theme must default to system")
            _assert_native_theme_tokens(page, page.locator("html").get_attribute("data-theme"))
            page.screenshot(path=str(output_dir / f"workbench-{name}-system.png"), full_page=True)
            theme_mode.select_option("dark")
            if page.locator("html").get_attribute("data-theme-mode") != "dark":
                raise AssertionError("manual dark appearance was not applied")
            _assert_native_theme_tokens(page, "dark")
            page.screenshot(path=str(output_dir / f"workbench-{name}-dark.png"), full_page=True)
            theme_mode.select_option("light")
            if page.locator("html").get_attribute("data-theme-mode") != "light":
                raise AssertionError("manual light appearance was not applied")
            _assert_native_theme_tokens(page, "light")
            page.screenshot(path=str(output_dir / f"workbench-{name}-light.png"), full_page=True)
            theme_mode.select_option("dark")
            if page.locator("html").get_attribute("data-theme-mode") != "dark":
                raise AssertionError("manual dark appearance was not restored")
            page.get_by_test_id("nav-settings").click()
            page.locator("#settingsView").wait_for(state="visible")
            if page.locator("#settingsThemeMode").input_value() != "dark":
                raise AssertionError("settings appearance control did not reflect the header override")
            page.locator("#settingsThemeMode").select_option("system")
            if page.locator("html").get_attribute("data-theme-mode") != "system":
                raise AssertionError("settings appearance control did not restore system mode")
            page.screenshot(path=str(output_dir / f"workbench-{name}-settings.png"), full_page=True)
            page.get_by_test_id("nav-diagnostics").click()
            page.get_by_test_id("diagnostics-list").locator("dt").first.wait_for(timeout=5_000)
            refresh_box = page.get_by_test_id("refresh-diagnostics-web").bounding_box()
            if refresh_box is None or refresh_box["height"] > 42:
                raise AssertionError("diagnostics refresh button must remain a single compact row")
            page.screenshot(path=str(output_dir / f"workbench-{name}-diagnostics.png"), full_page=True)
            page.get_by_test_id("nav-tasks").click()
            page.get_by_test_id("task-list").wait_for(state="visible")
            if name == "mobile":
                nav_box = page.locator(".app-nav").bounding_box()
                task_box = page.get_by_test_id("task-list").bounding_box()
                if nav_box is None or task_box is None or task_box["y"] - (nav_box["y"] + nav_box["height"]) > 220:
                    raise AssertionError("mobile task view must not leave a stretched blank navigation row")
            page.screenshot(path=str(output_dir / f"workbench-{name}-tasks.png"), full_page=True)
            page.get_by_test_id("nav-generate").click()
            page.locator("#generateView").wait_for(state="visible")

            # 生成设置回归：CLI 模式字段显隐与按传输方式校验（曾误要求 base_url 导致 CLI 无法保存）
            page.get_by_test_id("generator-settings").locator("summary").click()
            if page.locator("#ngaBaseUrlField").is_visible():
                raise AssertionError("CLI transport must hide HTTP-only fields on load")
            page.get_by_test_id("save-generator-settings").click()
            if page.locator("#generatorOperationStatus").inner_text() != "请填写模型名。":
                raise AssertionError("CLI mode must ask for the model name, not base URL")
            page.locator("#ngaModel").fill("w3/GLM-5.1-WX-Auto")
            page.get_by_test_id("save-generator-settings").click()
            page.locator("#generatorOperationStatus").get_by_text("配置已保存").wait_for(timeout=5_000)
            page.get_by_test_id("nga-transport").select_option("http")
            if not page.locator("#ngaBaseUrlField").is_visible():
                raise AssertionError("HTTP transport must show the base URL field")
            page.get_by_test_id("save-generator-settings").click()
            if page.locator("#generatorOperationStatus").inner_text() != "请填写服务地址和模型。":
                raise AssertionError("HTTP mode must require base URL and model")
            page.get_by_test_id("nga-transport").select_option("cli")

            page.get_by_test_id("input-file").set_input_files(str(source))
            page.get_by_test_id("type-word").click()
            page.get_by_test_id("generate-button").click()
            page.get_by_test_id("failure-diagnostic").wait_for(state="visible", timeout=15_000)
            if page.get_by_test_id("failure-code").inner_text() != "E001":
                raise AssertionError("synthetic generator failure did not expose E001")
            failure_href = page.get_by_test_id("failure-report-link").get_attribute("href")
            if not failure_href or not failure_href.endswith("/failure-report"):
                raise AssertionError("failure report link was not rendered")

            page.get_by_test_id("type-deck").click()
            page.get_by_test_id("template-file").set_input_files(str(template))
            page.get_by_test_id("remove-template").wait_for(state="visible")
            page.get_by_test_id("remove-template").click()
            page.get_by_test_id("template-file").set_input_files(str(template))
            page.get_by_test_id("generate-button").click()
            page.locator('[data-testid="result-panel"].result--success').wait_for(state="visible", timeout=30_000)
            page.screenshot(path=str(output_dir / f"workbench-{name}-success.png"), full_page=True)
            artifact_href = page.get_by_test_id("artifact-download").get_attribute("href")
            if not artifact_href or not artifact_href.startswith("/api/download/"):
                raise AssertionError("artifact download link was not rendered")
            if not page.locator("text=模板 Profile").count():
                raise AssertionError("template audit links were not rendered")

            page.reload(wait_until="networkidle")
            page.locator('[data-testid="result-panel"].result--success').wait_for(state="visible", timeout=10_000)
            restored_href = page.get_by_test_id("artifact-download").get_attribute("href")
            if restored_href != artifact_href:
                raise AssertionError("completed job was not restored after refresh")
            if not page.locator("text=模板 Profile").count():
                raise AssertionError("template audit links were not restored after refresh")

            page.get_by_test_id("input-file").set_input_files(str(source))
            page.get_by_test_id("type-deck").click()

            console_count_before_ole = len(console_errors)
            page.get_by_test_id("template-file").set_input_files(str(unsafe_template))
            page.get_by_test_id("generate-button").click()
            page.get_by_test_id("failure-diagnostic").wait_for(state="visible", timeout=15_000)
            if page.get_by_test_id("failure-code").inner_text() != "E003":
                raise AssertionError("unsafe OLE template did not expose E003")
            if page.locator("#failureStage").inner_text() != "定位：template_file":
                raise AssertionError("unsafe OLE template did not expose template_file location")
            suggestion = page.locator("#failureSuggestion").inner_text()
            if "删除" not in suggestion or "PNG" not in suggestion:
                raise AssertionError("unsafe OLE template did not expose an actionable suggestion")
            if "第 1 页" not in page.get_by_test_id("result-detail").inner_text():
                raise AssertionError("unsafe OLE template did not expose the referencing slide")
            if page.get_by_test_id("failure-report-link").is_visible():
                raise AssertionError("pre-job template rejection must not expose a stale failure report")
            ole_console_errors = console_errors[console_count_before_ole:]
            unexpected_ole_errors = [
                message for message in ole_console_errors if "status of 400 (BAD REQUEST)" not in message
            ]
            if unexpected_ole_errors:
                raise AssertionError(f"unexpected OLE rejection console errors: {unexpected_ole_errors}")
            del console_errors[console_count_before_ole:]

            page.get_by_test_id("remove-template").click()
            page.get_by_test_id("generate-button").click()
            if not generator.blocking_started.wait(timeout=5):
                raise AssertionError("blocking generator did not start")
            page.get_by_test_id("cancel-job").wait_for(state="visible", timeout=5_000)
            page.screenshot(path=str(output_dir / f"workbench-{name}-running.png"), full_page=True)
            page.get_by_test_id("cancel-job").click()
            page.get_by_test_id("failure-diagnostic").wait_for(state="visible", timeout=5_000)
            if page.get_by_test_id("failure-code").inner_text() != "E009":
                raise AssertionError("canceled job did not expose E009")
            if "is-visible" in (page.locator("#toast").get_attribute("class") or ""):
                raise AssertionError("terminal failure toast must not cover the persistent error diagnostic")
            generator.blocking_release.set()
            page.screenshot(path=str(output_dir / f"workbench-{name}-failure.png"), full_page=True)
            overflow = page.evaluate("document.documentElement.scrollWidth > window.innerWidth")
            screenshot = output_dir / f"workbench-{name}.png"
            page.screenshot(path=str(screenshot), full_page=True)
            context.close()

        if console_errors:
            raise AssertionError(f"browser console errors: {console_errors}")
        if overflow:
            raise AssertionError(f"{name} viewport has horizontal overflow")
        return {"name": name, "width": width, "height": height, "pass": True, "screenshot": screenshot.name}
    except Exception as exc:
        return {"name": name, "width": width, "height": height, "pass": False, "error": str(exc)}
    finally:
        generator.blocking_release.set()
        if browser is not None:
            browser.close()
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()


def _write_template(path: Path) -> Path:
    presentation = Presentation()
    for title, body in (("TEMPLATE COVER", "TEMPLATE SUBTITLE"), ("TEMPLATE TITLE", "TEMPLATE BODY")):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _textbox(slide, title, 0.8, 0.6, 11.0, 0.8, 26)
        _textbox(slide, body, 0.8, 1.8, 11.0, 4.6, 18)
    presentation.save(path)
    return path


def _assert_native_theme_tokens(page, theme: str | None) -> None:
    if theme not in NATIVE_THEME_TOKENS:
        raise AssertionError(f"workbench did not resolve a supported theme: {theme}")
    expected = NATIVE_THEME_TOKENS[theme]
    actual = page.evaluate(
        """() => {
          const root = getComputedStyle(document.documentElement);
          const selected = document.querySelector('.app-nav__button.is-selected');
          return {
            accent: root.getPropertyValue('--accent').trim().toLowerCase(),
            productMark: root.getPropertyValue('--product-mark').trim().toLowerCase(),
            selectedBackground: selected ? getComputedStyle(selected).backgroundColor : null,
          };
        }"""
    )
    if actual["accent"] != expected["accent"]:
        raise AssertionError(f"{theme} theme must use the Fluent accent, got {actual['accent']}")
    if actual["selectedBackground"] != expected["selected_background"]:
        raise AssertionError(
            f"{theme} selected navigation must use the semantic accent surface, got {actual['selectedBackground']}"
        )
    if actual["productMark"] != "#c7000b":
        raise AssertionError("Huawei red must remain limited to the product mark")


def _write_ole_template(path: Path) -> Path:
    source = _write_template(path).read_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(source)) as source_package, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source_package.infolist():
            data = source_package.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/oleObject" '
                    b'Target="../embeddings/oleObject1.bin"/></Relationships>',
                )
            target.writestr(info, data)
        target.writestr("ppt/embeddings/oleObject1.bin", b"unsafe")
    path.write_bytes(output.getvalue())
    return path


def _textbox(slide, text: str, left: float, top: float, width: float, height: float, size: float) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)


if __name__ == "__main__":
    raise SystemExit(main())
