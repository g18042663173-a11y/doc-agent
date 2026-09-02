from pathlib import Path

from pptx import Presentation

from doc_agent.ir.schemas import CardIR, ChartIR, ChartSeriesIR, DeckIR, SlideIR, VisualIR
from doc_agent.renderers.pptx_python_renderer import PythonPptxRenderer


def test_python_pptx_renderer_writes_editable_presentation(tmp_path: Path) -> None:
    deck = DeckIR(
        deck_title="企业文档生成 Agent",
        slides=[
            SlideIR(layout="cover", title="企业文档生成 Agent", subtitle="内部汇报"),
            SlideIR(layout="agenda", title="目录", bullets=["背景", "方案", "结论"]),
            SlideIR(layout="title_bullets", title="方案重点", bullets=["结构化生成", "模板化渲染"]),
            SlideIR(
                layout="two_column",
                title="能力对比",
                left_title="输入",
                left_bullets=["md", "docx"],
                right_title="输出",
                right_bullets=["pptx", "docx"],
            ),
            SlideIR(
                layout="table",
                title="模块状态",
                table_headers=["模块", "状态"],
                table_rows=[["Parser", "ready"]],
            ),
            SlideIR(layout="conclusion", title="结论", bullets=["默认本地运行", "可接入 GLM"]),
        ],
    )
    output = tmp_path / "demo.pptx"

    result = PythonPptxRenderer().render(deck.model_dump(), output)

    prs = Presentation(str(result))
    assert len(prs.slides) == 6
    assert result.stat().st_size > 0


def test_python_pptx_renderer_supports_v11_layouts(tmp_path: Path) -> None:
    deck = DeckIR(
        deck_title="DeckIR v1.1",
        slides=[
            SlideIR(layout="cover", title="DeckIR v1.1"),
            SlideIR(layout="cards", title="卡片页", cards=[CardIR(title="重点", body="基础可编辑渲染")]),
            SlideIR(
                layout="chart",
                title="图表页",
                chart=ChartIR(labels=["A", "B"], series=[ChartSeriesIR(name="指标", values=[1, 3])], summary="占位图表"),
            ),
            SlideIR(layout="image", title="图片页", visuals=[VisualIR(kind="placeholder", alt_text="图片占位")]),
            SlideIR(layout="conclusion", title="结论", bullets=["完成"]),
        ],
    )

    result = PythonPptxRenderer().render(deck.model_dump(), tmp_path / "v11.pptx")

    prs = Presentation(str(result))
    assert len(prs.slides) == 5
    assert "卡片页" in "\n".join(shape.text for shape in prs.slides[1].shapes if hasattr(shape, "text"))
