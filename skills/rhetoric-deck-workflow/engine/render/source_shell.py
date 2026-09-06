from __future__ import annotations

from pathlib import Path
import shutil

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION

from engine.extract.recover import write_diagram_node
from engine.fit.text_fit import replace_text_preserving_style
from engine.security.office_package import OfficePackageError, preflight_office_package
from engine.shared.chrome import DEFAULT_CLASSIFICATION, rewrite_classification_chrome
from engine.shared.common import RdwError, sha256, read_json
from engine.native.package import inspect_package, update_smartart_file, update_shared_text_file, audit_output


class _CellProxy:
    def __init__(self, table_shape, row: int, column: int):
        cell = table_shape.table.cell(row, column)
        self.text_frame = cell.text_frame
        self.has_text_frame = True
        span_width = cell.span_width if cell.is_merge_origin else 1
        span_height = cell.span_height if cell.is_merge_origin else 1
        self.width = sum(table_shape.table.columns[i].width for i in range(column, column + span_width))
        self.height = sum(table_shape.table.rows[i].height for i in range(row, row + span_height))


def fill_source_shell(
    shell: Path,
    skeleton: dict,
    content: dict,
    output: Path,
    *,
    classification: str | None = None,
    skip_pages: list[dict] | None = None,
) -> dict:
    output = output.resolve()
    if output.exists():
        raise RdwError("RD-E050", "out", "目标 deck.pptx 已存在。", "使用新的输出目录。")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.building{output.suffix}")
    shutil.copy2(shell, temporary)
    warnings: list[dict] = []
    inventory = inspect_package(shell)
    native_pages = {page["page_id"]: {obj["shape_ref"]: obj for obj in page["objects"]} for page in inventory["pages"]}
    source_inventory_path = shell.parent / "native_inventory.json"
    source_inventory = read_json(source_inventory_path, code="RD-E050", loc="native_inventory") if source_inventory_path.is_file() else inventory
    source_pages = {page["page_id"]: {obj["shape_ref"]: obj for obj in page["objects"]} for page in source_inventory["pages"]}
    native_replacements = []
    shared_replacements = []
    shared_objects = {obj["shape_ref"]: obj for obj in inventory["shared_objects"]}
    filled_charts = set()
    try:
        presentation = Presentation(temporary)
        # Normalize retained template chrome before author-supplied slots.
        # A business label beginning with "classification:" is still authored text.
        rewrite_classification_chrome(presentation, classification or DEFAULT_CLASSIFICATION)
        pages = {page["page_id"]: page for page in skeleton["pages"]}
        for page_content in content["pages"]:
            page = pages[page_content["page_id"]]
            slide_index = int(page["page_id"][1:])
            if slide_index > len(presentation.slides):
                raise RdwError("RD-E050", page["page_id"], "骨架页面超出 shell 页数。", "修正 skeleton 后重新 seal。")
            slide = presentation.slides[slide_index - 1]
            slots = {slot["slot_id"]: slot for slot in page["slots"]}
            for item in page_content["slots"]:
                if "value" not in item:
                    continue
                slot = slots[item["slot_id"]]
                if not slot["shape_ref"]:
                    raise RdwError("RD-E050", f"{page['page_id']}.{slot['slot_id']}", "source-shell 槽位缺少 shape_ref。", "抽取骨架时为模式 A 槽位绑定源形状。")
                value = item["value"]
                if slot["shape_ref"].endswith(".chart"):
                    target = _resolve_ref(slide, slot["shape_ref"].removesuffix(".chart"))
                    if target is None or not getattr(target, "has_chart", False) or not isinstance(value, dict):
                        raise RdwError("RD-E050", slot["shape_ref"], "图表需要结构化 categories/series 数据。", "依据材料提供原生图表数据。")
                    data = CategoryChartData()
                    data.categories = value["categories"]
                    for series in value["series"]:
                        data.add_series(series["name"], series["values"])
                    target.chart.replace_data(data)
                    # Old source axis bounds are data, not reusable geometry.
                    try:
                        axis = target.chart.value_axis
                        axis.minimum_scale = None
                        axis.maximum_scale = None
                        axis.tick_labels.number_format = value.get("number_format", "0.000")
                        axis.tick_labels.number_format_is_linked = False
                    except ValueError:
                        pass  # Pie charts have no value axis.
                    _format_replaced_chart(target.chart, value)
                    filled_charts.add((page["page_id"], slot["shape_ref"]))
                    continue
                text = "\n".join(value) if isinstance(value, list) else value
                if slot["shape_ref"].startswith("part_"):
                    obj = shared_objects.get(slot["shape_ref"])
                    if obj is None:
                        raise RdwError("RD-E050", slot["shape_ref"], "找不到共享文字对象。", "重新抽取完整对象清单。")
                    shared_replacements.append({"part": obj["part"], "shape_id": obj["shape_id"], "text": text})
                    continue
                role = "title" if slot["type"] == "title" else "body"
                if ".dgm_" in slot["shape_ref"]:
                    native = native_pages[page["page_id"]].get(slot["shape_ref"])
                    if native is not None:
                        native_replacements.append({"diagram_part": native["diagram_part"], "node_id": native["node_id"], "text": text})
                        continue
                    shape_part, _, node_part = slot["shape_ref"].partition(".dgm_")
                    frame_id = int(shape_part.removeprefix("sp_"))
                    node_index = int(node_part)
                    if not write_diagram_node(slide, frame_id, node_index, text):
                        raise RdwError("RD-E050", slot["shape_ref"], "shell 中找不到 SmartArt 节点。", "修正 shape_ref 并重新 seal。")
                    continue
                target = _resolve_ref(slide, slot["shape_ref"])
                if target is None:
                    raise RdwError("RD-E050", slot["shape_ref"], "shell 中找不到槽位形状。", "修正 shape_ref 并重新 seal。")
                if not getattr(target, "has_text_frame", False):
                    warnings.append({"code": "RD-W004", "page_id": page["page_id"], "slot_id": slot["slot_id"], "message": "shape has no text frame; source visual was kept"})
                    continue
                original = source_pages.get(page["page_id"], {}).get(slot["shape_ref"], {})
                fits = replace_text_preserving_style(target, text, role=role, preserve_source_size=True, source_run_lengths=original.get("paragraph_run_lengths"))
                if not fits:
                    warnings.append({"code": "RD-W003", "page_id": page["page_id"], "slot_id": slot["slot_id"], "message": "minimum readable font size still overflows; split or shorten manually"})
        expected_charts = {(page_id, ref) for page_id, objects in native_pages.items() for ref, obj in objects.items() if obj["kind"] == "chart"}
        if expected_charts != filled_charts:
            raise RdwError("RD-E050", "charts", "全部原生图表必须提供新材料数据。", "补齐图表；内部 shell 源缓存不可交付。")
        presentation.save(temporary)
        if native_replacements:
            update_smartart_file(temporary, native_replacements)
        if shared_replacements:
            update_shared_text_file(temporary, shared_replacements)
        try:
            preflight_office_package(temporary, purpose="output")
        except OfficePackageError as exc:
            raise RdwError("RD-E050", exc.loc, exc.message, "输出包安全检查失败，未发布产物。") from exc
        audit = audit_output(shell, temporary, inventory)
        if not audit["pass"]:
            raise RdwError("RD-E050", "native_audit", "源页、图片或对象几何发生变化。", "修复原生保留失败后再输出。")
        temporary.replace(output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return {"artifact": str(output), "sha256": sha256(output), "warnings": warnings, "native_audit": audit}


def _format_replaced_chart(chart, value):
    """Keep the native chart object, but discard source-data-dependent scales."""
    pie_types = {XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED, XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED}
    pie = chart.chart_type in pie_types
    for node in chart._chartSpace.xpath('.//*[local-name()="majorUnit" or local-name()="minorUnit" or local-name()="crossesAt"]'):
        node.getparent().remove(node)
    if not pie:
        for node in chart._chartSpace.xpath('.//*[local-name()="varyColors"]'):
            node.set("val", "0")
        chart.has_legend = len(value["series"]) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        values = [number for series in value["series"] for number in series["values"] if number is not None]
        if values and all(0 <= number <= 1 for number in values):
            chart.value_axis.minimum_scale = 0.0
            chart.value_axis.maximum_scale = 1.0
        for index, series in enumerate(chart.series):
            # Per-point source colors override a series fill. They described
            # old categories and would make cloned new series all look red.
            for point_style in series._element.xpath('./*[local-name()="dPt"]'):
                point_style.getparent().remove(point_style)
            color = ("C7000A", "626262", "A6A6A6", "D9D9D9")[index % 4]
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.format.line.color.rgb = RGBColor.from_string(color)
            if hasattr(series, "marker"):
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
                series.marker.format.line.color.rgb = RGBColor.from_string(color)
                for effect in series._element.xpath('./*[local-name()="marker"]/*[local-name()="spPr"]/*[local-name()="effectLst" or local-name()="effectDag"]'):
                    effect.getparent().remove(effect)
        for plot in chart.plots:
            if plot.has_data_labels:
                plot.data_labels.number_format = value.get("number_format", "0.000")
                plot.data_labels.number_format_is_linked = False
    else:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        # python-pptx omits the default enum value; Office needs an explicit
        # value on this element in existing template charts.
        for position in chart._chartSpace.xpath('.//*[local-name()="legendPos"]'):
            position.set("val", "r")
        chart.legend.include_in_layout = False
        for plot in chart.plots:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.show_value = False
            labels.show_percentage = True
            labels.show_category_name = False
            labels.number_format = "0.00%"
            labels.number_format_is_linked = False
            labels.position = XL_DATA_LABEL_POSITION.BEST_FIT
        values = value["series"][0]["values"]
        total = sum(number for number in values if number is not None)
        for series in chart.series:
            for index, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(("C7000A", "626262", "A6A6A6", "D9D9D9")[index % 4])
                if total and (values[index] or 0) / total < .05:
                    label = point.data_label._get_or_add_dLbl()
                    delete = label.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}delete")
                    if delete is None:
                        from pptx.oxml.xmlchemy import OxmlElement
                        delete = OxmlElement("c:delete")
                        label.insert(1, delete)
                    delete.set("val", "1")
                    # CT_DLbl chooses either delete OR visible-label
                    # properties. Keeping python-pptx's default properties
                    # beside delete makes PowerPoint refuse the package.
                    for child in list(label):
                        if child.tag.rsplit("}", 1)[-1] not in {"idx", "delete"}:
                            label.remove(child)


def _resolve_ref(slide, ref: str):
    shape_part, _, cell_part = ref.partition(".cell_")
    shape_id = int(shape_part.removeprefix("sp_"))
    shape = next((item for item in _walk_shapes(slide.shapes) if item.shape_id == shape_id), None)
    if shape is None:
        return None
    if not cell_part:
        return shape
    if not getattr(shape, "has_table", False):
        return None
    row, column = map(int, cell_part.split("_"))
    if row >= len(shape.table.rows) or column >= len(shape.table.columns):
        return None
    return _CellProxy(shape, row, column)


def _walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6:
            yield from _walk_shapes(shape.shapes)
