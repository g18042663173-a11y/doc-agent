from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import json
from pathlib import Path

from pptx import Presentation

from app.assets.pipeline import AssetRegistry
from app.ir.deck_ir import DeckIR
from app.rendering.pptx_renderer import _add_footer, _render_slide_content
from app.template.contracts import TemplatePlan, TemplateProfile, TemplateSlidePlan, TemplateWarning
from app.template.package import TemplateInputError, validate_pptx_package
from app.template.planner import build_template_plan
from app.template.profile import extract_template_profile, write_template_profile
from app.template.ppt_template_structure import (
    extract_ppt_template_structure,
    write_ppt_template_structure,
)
from app.template.replacement_audit import (
    build_template_replacement_audit,
    write_template_replacement_audit,
)
from app.template.text_fit import replace_text_preserving_style
from app.template.theme import build_template_render_theme


@dataclass(frozen=True)
class TemplateRenderResult:
    artifact_path: Path
    profile_path: Path
    plan_path: Path
    package_report_path: Path
    structure_json_path: Path
    structure_markdown_path: Path
    replacement_audit_path: Path
    profile: TemplateProfile
    plan: TemplatePlan
    package_report: dict


def render_deck_ir_with_template(
    deck: DeckIR,
    template_path: Path,
    output_path: Path,
    *,
    audit_dir: Path | None = None,
    profile: TemplateProfile | None = None,
    plan: TemplatePlan | None = None,
    asset_registry: AssetRegistry | None = None,
) -> TemplateRenderResult:
    template_path = template_path.resolve()
    output_path = output_path.resolve()
    audit_dir = (audit_dir or output_path.parent).resolve()
    audit_dir.mkdir(parents=True, exist_ok=True)

    profile = profile or extract_template_profile(template_path)
    plan = plan or build_template_plan(deck, profile)
    if plan.template_sha256 != profile.source.sha256:
        raise TemplateInputError("E003", "template_plan.template_sha256", "模板计划与 Profile 来源不一致。")
    if len(plan.slides) != len(deck.slides):
        raise TemplateInputError(
            "E003",
            "template_plan.slides",
            f"模板计划 {len(plan.slides)} 页与 Deck {len(deck.slides)} 页不一致。",
        )
    if plan.deck_ir_version != deck.ir_version:
        raise TemplateInputError(
            "E003",
            "template_plan.deck_ir_version",
            f"模板计划面向 DeckIR {plan.deck_ir_version}，当前 Deck 为 {deck.ir_version}。",
        )
    profile_path = write_template_profile(profile, audit_dir / "template_profile.json")
    structure = extract_ppt_template_structure(template_path)
    structure_json_path = write_ppt_template_structure(structure, audit_dir / "template_structure.json")
    structure_markdown_path = write_ppt_template_structure(structure, audit_dir / "template_structure.md")
    presentation = Presentation(template_path)
    original_slides = list(presentation.slides)
    original_slide_ids = list(presentation.slides._sldIdLst)
    theme = build_template_render_theme(deck.meta.theme, profile)
    for slide_plan in plan.slides:
        if (slide_plan.prototype_index or 1) > len(original_slides):
            raise TemplateInputError(
                "E003",
                "template_plan.slides[].prototype_index",
                f"模板计划引用原型页 {slide_plan.prototype_index}，模板只有 {len(original_slides)} 页。",
            )

    for index, (slide_ir, slide_plan) in enumerate(zip(deck.slides, plan.slides), start=1):
        source_slide = original_slides[(slide_plan.prototype_index or 1) - 1]
        target_slide = presentation.slides.add_slide(source_slide.slide_layout)
        _clear_slide_shapes(target_slide)
        if slide_plan.strategy == "prototype_replace":
            _clone_safe_shapes(source_slide, target_slide)
            _mark_template_decorations(
                target_slide,
                profile.slides[source_slide_index(original_slides, source_slide)],
            )
            replacements_fit = _apply_replacements(
                target_slide,
                slide_ir,
                profile.slides[source_slide_index(original_slides, source_slide)],
                slide_plan,
                fallback_font_name=theme["fonts"]["east_asia"][0],
                allowed_font_names=set(theme["fonts"]["whitelist"]),
            )
            if not replacements_fit:
                _clear_slide_shapes(target_slide)
                _render_slide_content(
                    target_slide, slide_ir, theme, asset_registry=asset_registry, slide_number=index
                )
                _record_runtime_fallback(plan, slide_plan, index)
        else:
            _render_slide_content(target_slide, slide_ir, theme, asset_registry=asset_registry, slide_number=index)
        _add_footer(target_slide, deck.meta.classification, index, len(deck.slides), theme)

    _remove_original_slides(presentation, original_slide_ids)
    plan_path = audit_dir / "template_plan.json"
    plan_path.write_text(plan.model_dump_json(indent=2) + "\n", encoding="utf-8")
    replacement_audit = build_template_replacement_audit(
        deck, profile, plan, presentation, asset_registry=asset_registry
    )
    replacement_audit_path = write_template_replacement_audit(
        replacement_audit, audit_dir / "template_replacement_audit.json"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    if asset_registry is not None:
        asset_registry.write_usage_audit(audit_dir / "asset_usage_audit.json")
    package_report = validate_pptx_package(output_path)
    package_report_path = audit_dir / "pptx_package_report.json"
    package_report_path.write_text(
        json.dumps(package_report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    if not package_report["pass"]:
        raise TemplateInputError("E003", "output", f"生成的 PPTX 包关系不完整: {package_report['errors'][0]}")
    return TemplateRenderResult(
        artifact_path=output_path,
        profile_path=profile_path,
        plan_path=plan_path,
        package_report_path=package_report_path,
        structure_json_path=structure_json_path,
        structure_markdown_path=structure_markdown_path,
        replacement_audit_path=replacement_audit_path,
        profile=profile,
        plan=plan,
        package_report=package_report,
    )


def source_slide_index(slides: list, source_slide) -> int:
    return next(index for index, slide in enumerate(slides) if slide.part is source_slide.part)


def _clear_slide_shapes(slide) -> None:
    tree = slide.shapes._spTree
    for element in list(tree)[2:]:
        if element.tag.endswith("}extLst"):
            continue
        tree.remove(element)


def _clone_safe_shapes(source_slide, target_slide) -> None:
    relationship_map: dict[str, str] = {}
    for shape in source_slide.shapes:
        element = deepcopy(shape.element)
        for node in element.iter():
            for attribute, value in list(node.attrib.items()):
                if not isinstance(value, str) or not value.startswith("rId"):
                    continue
                if value not in relationship_map:
                    relationship = source_slide.part.rels.get(value)
                    if relationship is None:
                        raise TemplateInputError(
                            "E003",
                            f"template.slide[{source_slide.slide_id}]",
                            f"原型形状引用了不存在的关系 {value}。",
                        )
                    if not _relationship_is_clone_safe(relationship):
                        _strip_relationship_reference(node, attribute)
                        continue
                    relationship_map[value] = target_slide.part.relate_to(
                        relationship.target_part,
                        relationship.reltype,
                    )
                node.attrib[attribute] = relationship_map[value]
        target_slide.shapes._spTree.insert_element_before(element, "p:extLst")


def _mark_template_decorations(target_slide, slide_profile) -> None:
    """Keep Profile-approved visual decoration out of generated-content geometry checks."""

    decorative_ids = {
        shape.shape_id
        for shape in slide_profile.shapes
        if shape.role == "decorative"
    }
    for shape in _iter_shapes(target_slide.shapes):
        if shape.shape_id in decorative_ids:
            shape.name = f"HW_TEMPLATE_DECORATION:{shape.shape_id}"


def _apply_replacements(
    target_slide,
    slide_ir,
    slide_profile,
    slide_plan: TemplateSlidePlan,
    *,
    fallback_font_name: str,
    allowed_font_names: set[str],
) -> bool:
    shape_by_id = {shape.shape_id: shape for shape in _iter_shapes(target_slide.shapes)}
    planned_ids = {replacement.shape_id for replacement in slide_plan.replacements}
    for shape_profile in slide_profile.shapes:
        if shape_profile.role == "brand":
            continue
        shape = shape_by_id.get(shape_profile.shape_id)
        if shape is None or not getattr(shape, "has_text_frame", False):
            continue
        if shape_profile.shape_id not in planned_ids:
            replace_text_preserving_style(
                shape,
                "",
                role=shape_profile.role,
                fallback_font_name=fallback_font_name,
                allowed_font_names=allowed_font_names,
            )

    replacements_fit = True
    for replacement in slide_plan.replacements:
        shape = shape_by_id.get(replacement.shape_id)
        if shape is None:
            raise TemplateInputError(
                "E003",
                f"template.{slide_profile.prototype_id}.shape[{replacement.shape_id}]",
                "计划引用的模板形状不存在。",
            )
        text = _source_text(slide_ir, replacement.source_path)
        replacements_fit = (
            replace_text_preserving_style(
                shape,
                text,
                role=replacement.role,
                fallback_font_name=fallback_font_name,
                allowed_font_names=allowed_font_names,
            )
            and replacements_fit
        )
    return replacements_fit


def _record_runtime_fallback(plan: TemplatePlan, slide_plan: TemplateSlidePlan, output_index: int) -> None:
    warning = TemplateWarning(
        code="W201",
        loc=f"slides[{output_index - 1}].template",
        message="模板原型在最小字号下仍预计溢出，已在同一模板母版/主题下重绘。",
    )
    slide_plan.strategy = "master_redraw"
    slide_plan.replacements = []
    slide_plan.warnings.append(warning)
    slide_plan.selection_reason += "; runtime_text_fit=overflow"
    plan.warnings.append(warning)


def _source_text(slide, source_path: str) -> str:
    if source_path == "title":
        return slide.title
    if source_path == "subtitle":
        return slide.subtitle or ""
    if source_path in {"presenter", "date", "index"}:
        return str(getattr(slide, source_path, "") or "")
    if source_path.startswith("items["):
        index = int(source_path.removeprefix("items[").removesuffix("]"))
        return slide.items[index]
    if source_path == "bullets":
        values = slide.bullets
        return "\n".join(f"• {getattr(item, 'text', item)}" for item in values)
    if source_path == "left":
        return _column_text(slide.left)
    if source_path == "right":
        return _column_text(slide.right)
    if source_path == "cta":
        return slide.cta or ""
    raise TemplateInputError("E003", f"template.source.{source_path}", "模板计划引用了未知内容路径。")


def _column_text(column) -> str:
    values: list[str] = []
    if column.heading:
        values.append(column.heading)
    if column.text:
        values.append(column.text)
    values.extend(f"• {item.text}" for item in column.bullets)
    return "\n".join(values)


def _remove_original_slides(presentation, slide_ids: list) -> None:
    slide_id_list = presentation.slides._sldIdLst
    for slide_id in slide_ids:
        presentation.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def _iter_shapes(shapes):
    for shape in shapes:
        if "GROUP" in str(getattr(shape, "shape_type", "")) and hasattr(shape, "shapes"):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _relationship_is_clone_safe(relationship) -> bool:
    if relationship.is_external:
        return False
    if relationship.reltype.endswith("/image"):
        return True
    return str(getattr(getattr(relationship, "target_part", None), "content_type", "")).startswith("image/")


def _strip_relationship_reference(node, attribute: str) -> None:
    local_name = node.tag.rsplit("}", 1)[-1]
    parent = node.getparent()
    if parent is not None and local_name == "tags":
        grandparent = parent.getparent()
        parent.remove(node)
        if grandparent is not None and len(parent) == 0 and parent.tag.rsplit("}", 1)[-1] == "custDataLst":
            grandparent.remove(parent)
        return
    if parent is not None and local_name in {
        "audioFile",
        "videoFile",
        "hlinkClick",
        "hlinkHover",
        "snd",
        "ext",
        "custData",
    }:
        parent.remove(node)
    else:
        node.attrib.pop(attribute, None)
